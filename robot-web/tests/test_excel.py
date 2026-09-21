import datetime as dt

import pytest
from openpyxl import Workbook, load_workbook

from autoweb.erreurs import ErreurExcel
from autoweb.excel import STATUT_ERREUR, STATUT_OK, ClasseurSuivi, creer_classeur


def _classeur_test(chemin):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plans"
    ws.append(["Numéro plan", "Titre", "Quantité", "Total", "Statut", "Message", "Horodatage"])
    ws.append(["PL-1", "Un", 2, "=C2*10", None, None, None])
    ws.append(["PL-2", "Deux", 3, "=C3*10", "OK", "déjà fait", "01/01/2026"])
    ws.append([None, None, None, None, None, None, None])  # ligne vide : ignorée
    ws.append(["PL-4", "Quatre", 4.0, "=C5*10", None, None, None])
    wb.save(chemin)
    return chemin


def test_lecture_lignes_et_valeurs(tmp_path):
    chemin = _classeur_test(tmp_path / "suivi.xlsx")
    classeur = ClasseurSuivi(chemin, feuille="Plans").ouvrir()
    try:
        lignes = classeur.lignes()
        assert [l.numero for l in lignes] == [2, 3, 5]
        assert lignes[0].valeur("numéro plan") == "PL-1"
        assert lignes[0].valeur("NUMERO PLAN") == "PL-1"  # tolérant
        assert lignes[2].valeurs["Quantité"] == 4.0
        assert classeur.entetes[:2] == ["Numéro plan", "Titre"]
        assert classeur.chemin_sauvegarde is not None and classeur.chemin_sauvegarde.exists()
    finally:
        classeur.fermer()


def test_ecriture_statut_conserve_formules_et_cree_colonne(tmp_path):
    chemin = _classeur_test(tmp_path / "suivi.xlsx")
    classeur = ClasseurSuivi(chemin, feuille="Plans").ouvrir()
    classeur.marquer(2, STATUT_OK, "")
    classeur.marquer(5, STATUT_ERREUR, "  message   sur plusieurs\nlignes  " + "x" * 600)
    classeur.ecrire(2, "Référence outil", "REF-0001")
    classeur.sauvegarder()
    classeur.fermer()

    wb = load_workbook(chemin)
    ws = wb["Plans"]
    assert ws["E2"].value == "OK"
    assert ws["E5"].value == "ERREUR"
    assert ws["F5"].value.startswith("message sur plusieurs lignes")
    assert len(ws["F5"].value) <= 500
    assert ws["G2"].value  # horodatage
    assert ws["D2"].value == "=C2*10"  # formule intacte
    assert ws["H1"].value == "Référence outil" and ws["H2"].value == "REF-0001"
    # la valeur écrite est relue immédiatement
    classeur2 = ClasseurSuivi(chemin, feuille="Plans", sauvegarde=False).ouvrir()
    assert classeur2.valeur(2, "référence outil") == "REF-0001"
    assert classeur2.lignes()[0].valeur("Statut") == "OK"
    classeur2.fermer()


def test_erreurs_explicites(tmp_path):
    with pytest.raises(ErreurExcel, match="introuvable"):
        ClasseurSuivi(tmp_path / "absent.xlsx").ouvrir()
    chemin = _classeur_test(tmp_path / "suivi.xlsx")
    with pytest.raises(ErreurExcel, match="Feuille"):
        ClasseurSuivi(chemin, feuille="Inexistante").ouvrir()
    (tmp_path / "x.csv").write_text("a;b")
    with pytest.raises(ErreurExcel, match="xlsx"):
        ClasseurSuivi(tmp_path / "x.csv").ouvrir()


def test_entete_en_double_refuse(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B", "A"])
    ws.append([1, 2, 3])
    wb.save(tmp_path / "double.xlsx")
    with pytest.raises(ErreurExcel, match="double"):
        ClasseurSuivi(tmp_path / "double.xlsx").ouvrir()


def test_creer_classeur(tmp_path):
    chemin = creer_classeur(tmp_path / "nouveau.xlsx", ["Numéro", "Date"], [["N1", dt.datetime(2026, 1, 2)]], feuille="Suivi")
    ws = load_workbook(chemin)["Suivi"]
    assert [c.value for c in ws[1]] == ["Numéro", "Date", "Statut", "Message", "Horodatage"]
    assert ws["B2"].number_format == "DD/MM/YYYY"
    classeur = ClasseurSuivi(chemin, sauvegarde=False).ouvrir()
    assert classeur.lignes()[0].valeur("Numéro") == "N1"
    classeur.fermer()
