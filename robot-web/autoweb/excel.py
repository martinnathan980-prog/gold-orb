"""Lecture / écriture du fichier Excel de suivi (openpyxl).

Le fichier Excel est la source de vérité : une ligne = un élément à traiter.
Le robot y écrit, pour chaque ligne, Statut / Message / Horodatage (et
éventuellement d'autres colonnes via l'étape « lire »), puis sauvegarde après
CHAQUE ligne : on peut interrompre et reprendre sans rien perdre.

Détails importants :
- Les valeurs sont lues avec `data_only=True` (résultat des formules tel que
  calculé par Excel lors du dernier enregistrement), mais l'écriture se fait sur
  un second exemplaire du classeur ouvert normalement, pour ne PAS écraser les
  formules existantes.
- Une copie de sauvegarde est faite dans `sauvegardes/` à l'ouverture.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .erreurs import ErreurExcel
from .gabarit import normaliser_cle

STATUT_OK = "OK"
STATUT_ERREUR = "ERREUR"
STATUT_IGNORE = "IGNORE"
STATUT_A_FAIRE = "A faire"

COULEURS_STATUT = {
    STATUT_OK: "C6EFCE",
    STATUT_ERREUR: "FFC7CE",
    STATUT_IGNORE: "FFEB9C",
}

NB_SAUVEGARDES_CONSERVEES = 10


@dataclass
class Ligne:
    """Une ligne de données : `numero` est le numéro de ligne Excel (1 = première)."""

    numero: int
    valeurs: Dict[str, Any] = field(default_factory=dict)

    def valeur(self, colonne: str) -> Any:
        if colonne in self.valeurs:
            return self.valeurs[colonne]
        cible = normaliser_cle(colonne)
        for cle, val in self.valeurs.items():
            if normaliser_cle(cle) == cible:
                return val
        return None

    def est_vide(self) -> bool:
        return all(v is None or (isinstance(v, str) and not v.strip()) for v in self.valeurs.values())


class ClasseurSuivi:
    """Le fichier Excel de suivi, ouvert en lecture (valeurs) et en écriture (statuts)."""

    def __init__(
        self,
        chemin: Path,
        feuille: Optional[str] = None,
        ligne_entete: int = 1,
        colonne_statut: str = "Statut",
        colonne_message: str = "Message",
        colonne_horodatage: str = "Horodatage",
        sauvegarde: bool = True,
    ) -> None:
        self.chemin = Path(chemin)
        self.nom_feuille = feuille
        self.ligne_entete = int(ligne_entete)
        self.colonne_statut = colonne_statut
        self.colonne_message = colonne_message
        self.colonne_horodatage = colonne_horodatage
        self.sauvegarde = sauvegarde
        self._wb_valeurs = None
        self._wb_ecriture = None
        self._ws_valeurs = None
        self._ws_ecriture = None
        self._colonnes: Dict[str, int] = {}
        self.modifie = False
        self.chemin_sauvegarde: Optional[Path] = None

    # ------------------------------------------------------------------ ouverture
    def ouvrir(self) -> "ClasseurSuivi":
        if not self.chemin.exists():
            raise ErreurExcel(f"Fichier Excel introuvable : {self.chemin}")
        if self.chemin.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ErreurExcel(
                f"Format non pris en charge : {self.chemin.suffix}. "
                "Enregistrez le fichier au format .xlsx (Excel > Enregistrer sous)."
            )
        try:
            self._wb_valeurs = load_workbook(self.chemin, data_only=True)
            self._wb_ecriture = load_workbook(self.chemin, keep_vba=self.chemin.suffix.lower() == ".xlsm")
        except PermissionError:
            raise ErreurExcel(
                f"Impossible d'ouvrir {self.chemin.name} : il est probablement ouvert dans Excel. "
                "Fermez-le puis relancez."
            )
        except Exception as e:  # fichier corrompu, mauvais format...
            raise ErreurExcel(f"Impossible de lire {self.chemin.name} : {e}")

        if self.nom_feuille:
            if self.nom_feuille not in self._wb_ecriture.sheetnames:
                raise ErreurExcel(
                    f"Feuille « {self.nom_feuille} » absente de {self.chemin.name}. "
                    f"Feuilles disponibles : {', '.join(self._wb_ecriture.sheetnames)}."
                )
            self._ws_valeurs = self._wb_valeurs[self.nom_feuille]
            self._ws_ecriture = self._wb_ecriture[self.nom_feuille]
        else:
            self._ws_valeurs = self._wb_valeurs.active
            self._ws_ecriture = self._wb_ecriture.active
            self.nom_feuille = self._ws_ecriture.title

        self._lire_entetes()
        if self.sauvegarde:
            self._faire_sauvegarde()
        return self

    def _lire_entetes(self) -> None:
        self._colonnes = {}
        ws = self._ws_ecriture
        if ws.max_row < self.ligne_entete:
            raise ErreurExcel(
                f"La ligne d'en-tête {self.ligne_entete} n'existe pas dans la feuille « {self.nom_feuille} »."
            )
        for cellule in ws[self.ligne_entete]:
            if cellule.value is None:
                continue
            nom = str(cellule.value).strip()
            if not nom:
                continue
            if nom in self._colonnes:
                raise ErreurExcel(
                    f"En-tête en double : « {nom} » (colonnes "
                    f"{get_column_letter(self._colonnes[nom])} et {get_column_letter(cellule.column)}). "
                    "Renommez l'une des deux colonnes."
                )
            self._colonnes[nom] = cellule.column
        if not self._colonnes:
            raise ErreurExcel(
                f"Aucun en-tête trouvé en ligne {self.ligne_entete} de la feuille « {self.nom_feuille} ». "
                "La première ligne doit contenir les noms de colonnes."
            )

    def _faire_sauvegarde(self) -> None:
        dossier = self.chemin.parent / "sauvegardes"
        try:
            dossier.mkdir(exist_ok=True)
            horodatage = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
            cible = dossier / f"{self.chemin.stem}.{horodatage}{self.chemin.suffix}"
            shutil.copy2(self.chemin, cible)
            self.chemin_sauvegarde = cible
            anciennes = sorted(dossier.glob(f"{self.chemin.stem}.*{self.chemin.suffix}"))
            for vieille in anciennes[:-NB_SAUVEGARDES_CONSERVEES]:
                vieille.unlink(missing_ok=True)
        except OSError:
            self.chemin_sauvegarde = None  # pas bloquant

    # ------------------------------------------------------------------ lecture
    @property
    def entetes(self) -> List[str]:
        return list(self._colonnes.keys())

    def colonne_existe(self, nom: str) -> bool:
        return self._resoudre_colonne(nom) is not None

    def _resoudre_colonne(self, nom: str) -> Optional[str]:
        if nom in self._colonnes:
            return nom
        cible = normaliser_cle(nom)
        for entete in self._colonnes:
            if normaliser_cle(entete) == cible:
                return entete
        return None

    def lignes(self) -> List[Ligne]:
        """Toutes les lignes non vides sous l'en-tête, avec leurs valeurs par nom de colonne."""
        resultat: List[Ligne] = []
        ws = self._ws_valeurs
        for numero in range(self.ligne_entete + 1, ws.max_row + 1):
            valeurs = {nom: ws.cell(row=numero, column=col).value for nom, col in self._colonnes.items()}
            ligne = Ligne(numero=numero, valeurs=valeurs)
            if ligne.est_vide():
                continue
            resultat.append(ligne)
        return resultat

    def valeur(self, numero: int, colonne: str) -> Any:
        entete = self._resoudre_colonne(colonne)
        if entete is None:
            return None
        return self._ws_valeurs.cell(row=numero, column=self._colonnes[entete]).value

    # ------------------------------------------------------------------ écriture
    def _assurer_colonne(self, nom: str) -> int:
        entete = self._resoudre_colonne(nom)
        if entete is not None:
            return self._colonnes[entete]
        col = (max(self._colonnes.values()) + 1) if self._colonnes else 1
        for ws in (self._ws_ecriture, self._ws_valeurs):
            cellule = ws.cell(row=self.ligne_entete, column=col, value=nom)
            cellule.font = Font(bold=True)
        self._colonnes[nom] = col
        self.modifie = True
        return col

    def ecrire(self, numero: int, colonne: str, valeur: Any) -> None:
        """Écrit une valeur dans la cellule (crée la colonne en fin de tableau si besoin)."""
        col = self._assurer_colonne(colonne)
        self._ws_ecriture.cell(row=numero, column=col, value=valeur)
        self._ws_valeurs.cell(row=numero, column=col, value=valeur)
        self.modifie = True

    def marquer(self, numero: int, statut: str, message: str = "") -> None:
        """Renseigne Statut / Message / Horodatage pour une ligne (avec couleur)."""
        message = " ".join(str(message).split())
        if len(message) > 500:
            message = message[:497] + "..."
        self.ecrire(numero, self.colonne_statut, statut)
        self.ecrire(numero, self.colonne_message, message)
        self.ecrire(numero, self.colonne_horodatage, dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
        couleur = COULEURS_STATUT.get(statut)
        cellule = self._ws_ecriture.cell(row=numero, column=self._colonnes[self._resoudre_colonne(self.colonne_statut)])
        if couleur:
            cellule.fill = PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")
        else:
            cellule.fill = PatternFill(fill_type=None)

    def sauvegarder(self) -> None:
        if not self.modifie:
            return
        try:
            self._wb_ecriture.save(self.chemin)
        except PermissionError:
            raise ErreurExcel(
                f"Impossible d'enregistrer {self.chemin.name} : le fichier est ouvert dans Excel. "
                "Fermez-le pour que le robot puisse écrire le suivi."
            )
        self.modifie = False

    def fermer(self) -> None:
        for wb in (self._wb_valeurs, self._wb_ecriture):
            try:
                if wb is not None:
                    wb.close()
            except Exception:
                pass


# ---------------------------------------------------------------------- création
def creer_classeur(
    chemin: Path,
    colonnes: Sequence[str],
    lignes: Iterable[Sequence[Any]] = (),
    feuille: str = "Suivi",
    colonnes_suivi: Sequence[str] = ("Statut", "Message", "Horodatage"),
) -> Path:
    """Crée un fichier Excel de suivi propre : en-têtes en gras, colonnes de suivi à droite."""
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = feuille
    entetes = list(colonnes) + [c for c in colonnes_suivi if c not in colonnes]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
        cellule.alignment = Alignment(vertical="center")
    for ligne in lignes:
        ws.append(list(ligne))
    for idx, nom in enumerate(entetes, start=1):
        largeur = max(12, min(45, len(str(nom)) + 4))
        if nom in ("Message",):
            largeur = 60
        ws.column_dimensions[get_column_letter(idx)].width = largeur
    ws.freeze_panes = "A2"
    for ligne_cellules in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cellule in ligne_cellules:
            if isinstance(cellule.value, (dt.datetime, dt.date)):
                cellule.number_format = "DD/MM/YYYY"
    try:
        wb.save(chemin)
    except PermissionError:
        raise ErreurExcel(f"Impossible de créer {chemin.name} : fichier ouvert dans Excel ou dossier protégé.")
    return chemin


def lire_csv(chemin: Path) -> List[List[str]]:
    """Lit un CSV « à la française » ou autre : encodage et séparateur détectés."""
    brut = Path(chemin).read_bytes()
    texte = None
    for encodage in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            texte = brut.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    if texte is None:
        texte = brut.decode("utf-8", errors="replace")
    echantillon = texte[:8192]
    try:
        separateur = csv.Sniffer().sniff(echantillon, delimiters=";,\t|").delimiter
    except csv.Error:
        separateur = ";" if echantillon.count(";") >= echantillon.count(",") else ","
    return [ligne for ligne in csv.reader(io.StringIO(texte), delimiter=separateur)]


def csv_vers_excel(chemin_csv: Path, chemin_xlsx: Optional[Path] = None, feuille: str = "Export") -> Path:
    """Convertit un export CSV en classeur Excel (toutes les valeurs en texte, en-tête en gras)."""
    chemin_csv = Path(chemin_csv)
    cible = Path(chemin_xlsx) if chemin_xlsx else chemin_csv.with_suffix(".xlsx")
    lignes = lire_csv(chemin_csv)
    wb = Workbook()
    ws = wb.active
    ws.title = feuille[:31]
    for ligne in lignes:
        ws.append([cellule.strip() for cellule in ligne])
    if ws.max_row >= 1:
        for cellule in ws[1]:
            cellule.font = Font(bold=True)
        ws.freeze_panes = "A2"
    largeurs: Dict[int, int] = {}
    for ligne_cellules in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
        for cellule in ligne_cellules:
            if cellule.value is not None:
                largeurs[cellule.column] = max(largeurs.get(cellule.column, 8), min(60, len(str(cellule.value)) + 2))
    for col, largeur in largeurs.items():
        ws.column_dimensions[get_column_letter(col)].width = largeur
    try:
        wb.save(cible)
    except PermissionError:
        raise ErreurExcel(f"Impossible d'écrire {cible.name} : fichier ouvert dans Excel ?")
    return cible
