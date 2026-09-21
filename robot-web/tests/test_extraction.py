import pytest
from openpyxl import load_workbook

from autoweb.erreurs import ErreurAutoweb
from autoweb.extraction import charger_regles, extraire, extraire_champs, lire_document

REGLES = """
fichiers: ["*.txt", "*.docx"]
champs:
  Numéro plan:
    regex: 'N[°o]\\s*plan\\s*:\\s*(\\S+)'
    obligatoire: true
  Titre: 'Titre\\s*:\\s*(.+)'
  Dates:
    regex: '\\d{2}/\\d{2}/\\d{4}'
    tous: true
  Absent:
    regex: 'jamais présent (x)'
    defaut: "?"
"""


def test_extraction_champs():
    regles_path = None
    texte = "N° plan : PL-77\nTitre :  Façade   nord \nRévisé le 01/02/2026 puis le 03/04/2026 et le 01/02/2026"
    from autoweb.extraction import Champ  # noqa
    import re
    champs = [
        Champ("Numéro plan", re.compile(r"N[°o]\s*plan\s*:\s*(\S+)", re.I)),
        Champ("Titre", re.compile(r"Titre\s*:\s*(.+)", re.I)),
        Champ("Dates", re.compile(r"\d{2}/\d{2}/\d{4}"), groupe=0, tous=True),
        Champ("Absent", re.compile(r"zzz(x)"), defaut="?"),
    ]
    resultat = extraire_champs(texte, champs)
    assert resultat == {"Numéro plan": "PL-77", "Titre": "Façade nord", "Dates": "01/02/2026; 03/04/2026", "Absent": "?"}


def test_extraction_vers_excel(tmp_path):
    docs = tmp_path / "docs"
    (docs / "sous").mkdir(parents=True)
    (docs / "a.txt").write_text("N° plan : PL-1\nTitre : Plan A\nle 01/01/2026", encoding="utf-8")
    (docs / "sous" / "b.txt").write_text("Titre : sans numéro", encoding="cp1252")
    try:
        import docx
        d = docx.Document()
        d.add_paragraph("No plan : PL-3")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Titre : Plan C"
        table.rows[0].cells[1].text = "02/02/2026"
        d.save(docs / "c.docx")
        avec_docx = True
    except ImportError:
        avec_docx = False
    regles_path = tmp_path / "regles.yaml"
    regles_path.write_text(REGLES, encoding="utf-8")
    regles = charger_regles(regles_path)
    sortie = tmp_path / "extraction.xlsx"
    bilan = extraire(regles, docs, sortie)
    assert bilan.fichiers == (3 if avec_docx else 2)
    assert bilan.reussis == (2 if avec_docx else 1) and bilan.incomplets == 1 and bilan.illisibles == 0

    ws = load_workbook(sortie)["Suivi"]
    entetes = [c.value for c in ws[1]]
    assert entetes == ["Fichier", "Numéro plan", "Titre", "Dates", "Absent", "Statut", "Message", "Horodatage"]
    lignes = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(1, 9)] for r in range(2, ws.max_row + 1)}
    assert lignes["a.txt"][1:6] == ["PL-1", "Plan A", "01/01/2026", "?", None]
    b = lignes[str(__import__("pathlib").Path("sous") / "b.txt")]
    assert b[5] == "A verifier" and "Numéro plan" in b[6]
    if avec_docx:
        assert lignes["c.docx"][1:4] == ["PL-3", "Plan C", "02/02/2026"]

    with pytest.raises(ErreurAutoweb, match="existe déjà"):
        extraire(regles, docs, sortie)
    extraire(regles, docs, sortie, ecraser=True)


def test_regles_invalides(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text("champs:\n  X: '(['\n", encoding="utf-8")
    with pytest.raises(ErreurAutoweb, match="régulière"):
        charger_regles(p)
    p.write_text("champs:\n  X: {regex: 'a(b)', groupe: 2}\n", encoding="utf-8")
    with pytest.raises(ErreurAutoweb, match="groupe"):
        charger_regles(p)
    p.write_text("fichiers: '*.txt'\n", encoding="utf-8")
    with pytest.raises(ErreurAutoweb, match="champs"):
        charger_regles(p)


def test_lire_document_extensions(tmp_path):
    (tmp_path / "x.bin").write_bytes(b"\x00\x01")
    with pytest.raises(ErreurAutoweb, match="extension"):
        lire_document(tmp_path / "x.bin")
    (tmp_path / "x.doc").write_bytes(b"\x00")
    with pytest.raises(ErreurAutoweb, match="docx"):
        lire_document(tmp_path / "x.doc")
