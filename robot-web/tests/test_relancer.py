"""Relancer une tâche déjà faite : le robot ne refait pas une ligne OK, mais il doit
le dire clairement, et le menu doit proposer de la refaire."""

import builtins
from pathlib import Path

from openpyxl import load_workbook

from autoweb import cli
from autoweb.excel import creer_classeur


def _tache(dossier: Path, statut: str) -> Path:
    creer_classeur(dossier / "essai.xlsx", ["Tache"], [["essai"]])
    wb = load_workbook(dossier / "essai.xlsx")
    ws = wb["Suivi"]
    entetes = [c.value for c in ws[1]]
    ws.cell(row=2, column=entetes.index("Statut") + 1, value=statut)
    wb.save(dossier / "essai.xlsx")
    chemin = dossier / "t.yaml"
    chemin.write_text(
        "nom: essai\n"
        "navigateur: {canal: auto, visible: false}\n"
        "excel: {fichier: essai.xlsx, feuille: Suivi}\n"
        "variables: {url: \"https://exemple.invalid\"}\n"
        "etapes:\n"
        "  - aller: \"{{url}}\"\n",
        encoding="utf-8",
    )
    return chemin


def test_lancer_signale_que_tout_est_deja_fait(tmp_path, capsys):
    chemin = _tache(tmp_path, "OK")
    code = cli.main(["lancer", str(chemin), "--sans-pause"])
    sortie = capsys.readouterr().out
    assert code == cli.CODE_TOUT_DEJA_FAIT
    assert "déjà traitées" in sortie  # la cause est dite, pas seulement « 0 ligne »


def test_un_fichier_vide_n_est_pas_confondu_avec_tout_fait(tmp_path, capsys):
    creer_classeur(tmp_path / "essai.xlsx", ["Tache"], [])
    chemin = tmp_path / "t.yaml"
    chemin.write_text(
        "nom: vide\nnavigateur: {canal: auto, visible: false}\n"
        "excel: {fichier: essai.xlsx, feuille: Suivi}\netapes:\n  - aller: \"https://exemple.invalid\"\n",
        encoding="utf-8",
    )
    code = cli.main(["lancer", str(chemin), "--sans-pause"])
    sortie = capsys.readouterr().out
    assert code == 0
    assert "aucune ligne de données" in sortie


def test_le_menu_propose_de_refaire_une_tache_deja_faite(tmp_path, monkeypatch, capsys):
    chemin = _tache(tmp_path, "OK")
    monkeypatch.setattr(cli, "lister_scenarios", lambda racine=None: [chemin])
    lancements = []

    def faux_lancer(args):
        lancements.append(dict(tout=args.tout, limite=args.limite))
        return cli.CODE_TOUT_DEJA_FAIT if not args.tout else 0

    monkeypatch.setattr(cli, "cmd_lancer", faux_lancer)
    reponses = iter(["2", "1", "1", "o", "0"])  # lancer, tâche 1, une ligne, refaire : oui, quitter
    monkeypatch.setattr(builtins, "input", lambda *a: next(reponses))
    assert cli.main([]) == 0
    sortie = capsys.readouterr().out
    assert "deja faites" in sortie and "Les refaire quand meme" in sortie
    assert lancements == [dict(tout=False, limite=1), dict(tout=True, limite=1)]


def test_le_menu_offre_tout_refaire_directement(tmp_path, monkeypatch, capsys):
    chemin = _tache(tmp_path, "OK")
    monkeypatch.setattr(cli, "lister_scenarios", lambda racine=None: [chemin])
    lancements = []
    monkeypatch.setattr(cli, "cmd_lancer", lambda args: lancements.append(args.tout) or 0)
    reponses = iter(["2", "1", "4", "0"])  # lancer, tâche 1, « tout refaire », quitter
    monkeypatch.setattr(builtins, "input", lambda *a: next(reponses))
    assert cli.main([]) == 0
    assert "Tout refaire" in capsys.readouterr().out
    assert lancements == [True]
