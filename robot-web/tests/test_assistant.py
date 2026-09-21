"""L'assistant construit, à partir d'un relevé de la page de démo (connectée) et
des colonnes de l'Excel de démo, un scénario qui fonctionne de bout en bout."""

import datetime as dt
import shutil
from pathlib import Path

from openpyxl import load_workbook

from autoweb.assistant import Dialogue, construire
from autoweb.excel import creer_classeur
from autoweb.navigateur import Navigateur
from autoweb.releve import charger_releve, relever
from autoweb.runner import Options, lancer
from autoweb.scenario import ConfigNavigateur, charger


def _releve_connecte(dossier: Path) -> Path:
    """Ouvre la démo, se connecte, relève l'écran du formulaire."""
    nav = Navigateur(ConfigNavigateur(canal="auto", visible=False), dossier, visible=False)
    page = nav.ouvrir()
    try:
        page.goto((dossier / "formulaire_demo.html").resolve().as_uri())
        page.fill("#utilisateur", "demo")
        page.fill("#mot-de-passe", "demo")
        page.click("#btn-connexion")
        page.wait_for_selector("#formulaire-plan")
        return relever(page, dossier / "releves", nom="nouveau plan")
    finally:
        nav.fermer()


def test_assistant_bout_en_bout(tmp_path, dossier_modeles, navigateur_ok):
    shutil.copyfile(dossier_modeles / "formulaire_demo.html", tmp_path / "formulaire_demo.html")
    date = dt.datetime(2026, 9, 21)
    excel = creer_classeur(
        tmp_path / "suivi.xlsx",
        ["Numéro plan", "Titre", "Type", "Date", "Urgent", "Commentaire"],
        [
            ["PL-1", "Rez-de-chaussée", "Architecture", date, "oui", "Premier"],
            [None, "Sans numéro", "STR", None, None, None],
            ["PL-3", "Toiture", "PLB", date, "x", None],
        ],
        feuille="Plans",
    )
    dossier = _releve_connecte(tmp_path)
    releve = charger_releve(dossier)
    colonnes = ["Numéro plan", "Titre", "Type", "Date", "Urgent", "Commentaire"]

    # réponses scriptées : les propositions par défaut sont acceptées (Entrée = "")
    # champs visibles dans l'ordre : numero, titre, type, date, urgent, commentaire
    reponses = [
        "",            # type de scénario : 1 (formulaire)
        "",            # adresse : celle du relevé
        "",            # Numéro de plan -> proposition « Numéro plan »
        "",            # Titre -> Titre
        "",            # Type -> Type
        "",            # Date du plan -> Date
        "",            # Urgent -> Urgent
        "",            # Commentaire -> Commentaire
        "",            # bouton : proposition « Enregistrer »
        "n",           # pas de pause (test automatique)
        "enregistré",  # texte de succès
        "#reference",  # référence à relever
        "",            # colonne « Référence outil »
        "n",           # pas de connexion manuelle (on l'ajoute nous-mêmes ci-dessous)
        "o",           # capture
    ]
    texte = construire(releve, colonnes, Dialogue(reponses), nom="nouveau plan", fichier_excel="suivi.xlsx", feuille="Plans")
    assert 'remplir: {selecteur: "#numero", valeur: "{{Numéro plan}}"}' in texte
    assert 'choisir: {selecteur: "#type", valeur: "{{Type}}"}' in texte
    assert 'cocher: {selecteur: "#urgent", valeur: "{{Urgent}}"}' in texte
    assert 'cliquer: "role=button:Enregistrer"' in texte
    assert 'verifier: {texte_page: "enregistré"}' in texte
    assert 'lire: {selecteur: "#reference", vers: "Référence outil"}' in texte
    assert "pause" not in texte and "avant:" not in texte

    # connexion à la démo (dans la vraie vie : pause pour se connecter à la main)
    texte = texte.replace(
        "etapes:\n",
        "avant:\n  - aller: \"{{url}}\"\n  - remplir: {\"#utilisateur\": demo, \"#mot-de-passe\": demo}\n"
        "  - cliquer: \"#btn-connexion\"\n  - attendre: \"#formulaire-plan\"\netapes:\n",
    )
    chemin = tmp_path / "nouveau_plan.yaml"
    chemin.write_text(texte, encoding="utf-8")
    scenario = charger(chemin)
    assert scenario.excel.feuille == "Plans"

    bilan = lancer(scenario, Options(visible=False, interactif=False))
    assert (bilan.ok, bilan.erreurs) == (2, 1)
    ws = load_workbook(excel)["Plans"]
    entetes = [c.value for c in ws[1]]
    col = entetes.index("Référence outil") + 1
    assert ws.cell(row=2, column=col).value == "REF-0001"
    assert ws.cell(row=3, column=entetes.index("Statut") + 1).value == "ERREUR"
    assert ws.cell(row=4, column=col).value == "REF-0002"


def test_assistant_sans_excel_et_valeur_fixe(tmp_path, dossier_modeles, navigateur_ok):
    shutil.copyfile(dossier_modeles / "formulaire_demo.html", tmp_path / "formulaire_demo.html")
    releve = charger_releve(_releve_connecte(tmp_path))
    reponses = [
        "1",           # formulaire
        "",            # adresse
        "Numéro",      # numero -> {{Numéro}}
        "",            # titre ignoré
        "v", "STR",    # type : valeur fixe
        "", "", "",    # date, urgent, commentaire ignorés
        "0",           # aucun bouton
        "o",           # pause
        "", "",        # pas de succès, pas de référence
        "o",           # connexion manuelle
        "n",           # pas de capture
    ]
    texte = construire(releve, [], Dialogue(reponses), nom="test")
    assert 'remplir: {selecteur: "#numero", valeur: "{{Numéro}}"}' in texte
    assert 'choisir: {selecteur: "#type", valeur: "STR"}' in texte
    assert "#titre" not in texte and "cliquer" not in texte
    assert "avant:" in texte and "pause:" in texte and "capture" not in texte
    (tmp_path / "t.yaml").write_text(texte, encoding="utf-8")
    assert charger(tmp_path / "t.yaml").etapes[0].action == "aller"
