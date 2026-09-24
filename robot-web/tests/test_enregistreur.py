"""Le mode « montre-moi » : le robot regarde faire la tâche, puis écrit le scénario.

Les actions de l'utilisateur sont simulées avec Playwright (clics et saisies réels
dans la page), ce qui exerce le script d'enregistrement exactement comme un humain.
"""

import time
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from autoweb.assistant import Dialogue, construire_depuis_enregistrement, deviner_colonne
from autoweb.cli import _servir_dossier, lister_scenarios, preparer_bac_a_sable
from autoweb.enregistreur import Enregistreur
from autoweb.navigateur import Navigateur
from autoweb.runner import Options, lancer
from autoweb.scenario import ConfigNavigateur, charger


@pytest.fixture
def bac(tmp_path):
    preparer_bac_a_sable(tmp_path)
    serveur, url = _servir_dossier(tmp_path, "base_demo.html", port=0)
    try:
        yield tmp_path, url
    finally:
        serveur.shutdown()
        serveur.server_close()


def _enregistrer(dossier: Path, url: str, actions):
    """Ouvre le navigateur, enregistre pendant que `actions(page)` joue la tâche."""
    nav = Navigateur(ConfigNavigateur(canal="auto", visible=False), dossier, visible=False)
    nav.ouvrir()
    enregistreur = Enregistreur(nav)
    try:
        page = enregistreur.demarrer(url)
        page.wait_for_selector("#__autoweb_badge")  # le bandeau d'enregistrement est là
        actions(page)
        page.wait_for_timeout(800)  # laisse arriver les derniers événements
        return enregistreur.arreter()
    finally:
        nav.fermer()


def test_enregistrement_export_par_contrat(bac, navigateur_ok):
    dossier, url = bac

    def tache(page):
        page.fill("#utilisateur", "demo")
        page.fill("#mot-de-passe", "demo")
        page.click("#btn-connexion")
        page.wait_for_selector("#liste")
        page.select_option("#filtre-contrat", "HDK")
        page.click("#btn-rechercher")
        page.wait_for_selector("text=15 plan(s)")
        page.click("#btn-exporter")

    etapes = _enregistrer(dossier, url, tache)
    actions = [(e.action, e.args.get("selecteur") or e.args.get("cliquer") or e.args.get("url")) for e in etapes]
    assert ("remplir", "#utilisateur") in actions
    assert ("remplir", "#mot-de-passe") in actions
    assert ("cliquer", "#btn-connexion") in actions
    assert ("choisir", "#filtre-contrat") in actions
    assert ("cliquer", "#btn-rechercher") in actions
    assert ("telecharger", "#btn-exporter") in actions  # le clic devient un téléchargement
    assert etapes[0].action == "aller"
    mdp = [e for e in etapes if e.type_champ == "password"][0]
    assert "demo" not in mdp.resume()  # le mot de passe n'est jamais réaffiché

    # ---- transformation en scénario, avec les colonnes de l'Excel
    colonnes = ["Contrat"]
    lignes_excel = [{"Contrat": "HDK"}, {"Contrat": "QRS"}]
    assert deviner_colonne("HDK", colonnes, lignes_excel) == 0
    assert deviner_colonne("Exporter CSV", colonnes, lignes_excel) is None
    reponses = [
        "",     # ne rien retirer
        "",     # connexion une seule fois au début : oui
        "",     # « HDK » -> colonne Contrat (proposée)
        "",     # dossier exports
        "",     # nom export_{{Contrat}}
        "",     # convertir en Excel
        "",     # colonne Fichier export
        "n",    # pas de pause supplémentaire
        "n",    # pas de capture
    ]
    texte = construire_depuis_enregistrement(
        etapes, colonnes, lignes_excel, Dialogue(reponses), nom="export contrat",
        fichier_excel="contrats.xlsx", feuille="Contrats", canal="chromium", url_depart=url,
    )
    donnees = yaml.safe_load(texte)
    assert donnees["excel"] == {"fichier": "contrats.xlsx", "feuille": "Contrats", "colonne_libelle": "Contrat",
                                "refaire": "toujours"}  # tâche répétitive : refaite à chaque lancement
    # la connexion est dans « avant », conditionnée à l'écran de connexion
    avant = donnees["avant"]
    assert avant[1]["si"]["visible"] == "#mot-de-passe"
    connexion = avant[1]["si"]["alors"]
    assert connexion[0]["remplir"] == {"selecteur": "#utilisateur", "valeur": "demo"}
    assert connexion[1]["remplir"]["valeur"] == "{{mot_de_passe}}"   # jamais le vrai mot de passe
    ligne_mdp = [l for l in texte.splitlines() if "#mot-de-passe" in l and "remplir" in l][0]
    assert "demo" not in ligne_mdp and "{{mot_de_passe}}" in ligne_mdp
    # la variable est déclarée (sinon « verifier » la signalerait comme inconnue)
    assert 'mot_de_passe: ""' in texte and "--var mot_de_passe=" in texte
    assert donnees["variables"]["mot_de_passe"] == ""
    # la tâche répétée par ligne : filtre, recherche, export
    etapes_yaml = donnees["etapes"]
    assert etapes_yaml[0] == {"aller": "{{url}}"}
    assert {"choisir": {"selecteur": "#filtre-contrat", "valeur": "{{Contrat}}"}} in etapes_yaml
    assert {"cliquer": "#btn-rechercher"} in etapes_yaml
    telecharger = [e["telecharger"] for e in etapes_yaml if isinstance(e, dict) and "telecharger" in e][0]
    assert telecharger["cliquer"] == "#btn-exporter"
    assert telecharger["renommer"] == "export_{{Contrat}}" and telecharger["convertir_excel"] is True

    # ---- le scénario enregistré fonctionne vraiment
    chemin = dossier / "export_contrat.yaml"
    chemin.write_text(texte, encoding="utf-8")
    scenario = charger(chemin)
    bilan = lancer(scenario, Options(visible=False, interactif=False, variables={"mot_de_passe": "demo"}))
    assert (bilan.ok, bilan.erreurs) == (2, 0), bilan.resume()
    for contrat in ("HDK", "QRS"):
        assert (dossier / "exports" / f"export_{contrat}.xlsx").exists()
    ws = load_workbook(dossier / "exports" / "export_HDK.xlsx").active
    lignes = list(ws.iter_rows(values_only=True))
    assert len(lignes) == 16 and all(l[1] == "HDK" for l in lignes[1:])
    suivi = load_workbook(dossier / "contrats.xlsx")["Contrats"]
    entetes = [c.value for c in suivi[1]]
    assert suivi.cell(row=2, column=entetes.index("Statut") + 1).value == "OK"
    assert chemin in lister_scenarios(dossier)
    # tâche répétitive : relancée, elle refait tout, sans « déjà fait »
    bilan = lancer(charger(chemin), Options(visible=False, interactif=False, variables={"mot_de_passe": "demo"}))
    assert (bilan.ok, bilan.erreurs) == (2, 0), bilan.resume()


def test_tache_sans_excel_relancee_autant_de_fois_qu_on_veut(bac, navigateur_ok):
    """Pas d'Excel : la tâche est rejouée en entier à chaque lancement, sans « une seule fois »."""
    dossier, url = bac

    def tache(page):
        page.fill("#utilisateur", "demo")
        page.fill("#mot-de-passe", "demo")
        page.click("#btn-connexion")
        page.wait_for_selector("#liste")
        page.select_option("#filtre-contrat", "HDK")
        page.click("#btn-rechercher")
        page.wait_for_selector("text=15 plan(s)")
        page.click("#btn-exporter")

    etapes = _enregistrer(dossier, url, tache)
    reponses = [
        "",     # ne rien retirer
        "",     # connexion une seule fois au début : oui
        "",     # dossier exports
        "",     # nom proposé : date et heure, pour ne jamais écraser l'export précédent
        "",     # convertir en Excel
        "n",    # pas de pause
        "n",    # pas de capture
    ]
    texte = construire_depuis_enregistrement(
        etapes, [], [], Dialogue(reponses), nom="export hdk",
        fichier_excel=None, canal="chromium", url_depart=url,
    )
    donnees = yaml.safe_load(texte)
    assert "excel" not in donnees
    assert {"choisir": {"selecteur": "#filtre-contrat", "valeur": "HDK"}} in donnees["etapes"]
    telecharger = [e["telecharger"] for e in donnees["etapes"] if isinstance(e, dict) and "telecharger" in e][0]
    assert telecharger["renommer"] == "export_hdk_{{horodatage}}" and "vers_colonne" not in telecharger

    chemin = dossier / "export_hdk.yaml"
    chemin.write_text(texte, encoding="utf-8")
    for fois in range(2):
        if fois:
            time.sleep(1.1)  # l'horodatage est à la seconde
        bilan = lancer(charger(chemin), Options(visible=False, interactif=False, variables={"mot_de_passe": "demo"}))
        assert (bilan.ok, bilan.erreurs) == (1, 0), bilan.resume()
        assert bilan.resume().startswith("Tâche réussie")
        assert not bilan.tout_deja_fait
    assert len(list((dossier / "exports").glob("export_hdk_*.xlsx"))) == 2


def test_enregistrement_modification_de_fiche(bac, navigateur_ok):
    """Plusieurs écrans : liste, clic sur une fiche, saisie, enregistrement."""
    dossier, url = bac

    def tache(page):
        page.fill("#utilisateur", "demo")
        page.fill("#mot-de-passe", "demo")
        page.click("#btn-connexion")
        page.wait_for_selector("#liste")
        page.select_option("#filtre-contrat", "HDK")
        page.click("#btn-rechercher")
        page.wait_for_selector("text=15 plan(s)")
        page.click("text=HDK-ARC-001")          # ouvre la fiche
        page.wait_for_selector("#formulaire-fiche")
        page.fill("#fiche-indice", "D")
        page.select_option("#fiche-statut", "Archivé")
        page.click("#btn-enregistrer")
        page.wait_for_selector("#message-fiche")

    etapes = _enregistrer(dossier, url, tache)
    clics = [e for e in etapes if e.action == "cliquer"]
    # le clic dans le tableau est ancré sur le texte de la ligne, pas sur sa position
    assert any('tr:has-text("HDK-ARC-001")' in e.args["selecteur"] for e in clics), [e.args for e in clics]
    assert any(e.action == "remplir" and e.args["selecteur"] == "#fiche-indice" for e in etapes)

    colonnes = ["Numéro", "Nouveau statut", "Indice"]
    lignes_excel = [
        {"Numéro": "HDK-ARC-001", "Nouveau statut": "Archivé", "Indice": "D"},
        {"Numéro": "HDK-ARC-002", "Nouveau statut": "Diffusé", "Indice": "B"},
    ]
    reponses = [
        "",     # ne rien retirer
        "",     # connexion une seule fois au début
        "",     # filtre « HDK » : aucune colonne ne correspond, reste figé
        "",     # clic « HDK-ARC-001 » -> colonne Numéro (proposée)
        "",     # statut « Archivé » -> colonne Nouveau statut (proposée)
        "",     # indice « D » -> colonne Indice (proposée)
        "enregistrée",  # texte de succès
        "n",    # pas de pause supplémentaire
        "n",    # pas de capture
    ]
    texte = construire_depuis_enregistrement(
        etapes, colonnes, lignes_excel, Dialogue(reponses), nom="modifier fiche",
        fichier_excel="fiches.xlsx", feuille="Fiches", canal="chromium", url_depart=url,
    )
    donnees = yaml.safe_load(texte)
    assert {"cliquer": 'tr:has-text("{{Numéro}}") a'} in donnees["etapes"]
    assert '{selecteur: "#fiche-indice", valeur: "{{Indice}}"}' in texte
    assert '{selecteur: "#fiche-statut", valeur: "{{Nouveau statut}}"}' in texte
    assert '- verifier: {texte_page: "enregistrée"}' in texte

    chemin = dossier / "modifier_fiche.yaml"
    chemin.write_text(texte, encoding="utf-8")
    bilan = lancer(charger(chemin), Options(visible=False, interactif=False, variables={"mot_de_passe": "demo"}))
    assert (bilan.ok, bilan.erreurs) == (2, 1), bilan.resume()  # 3e ligne de l'Excel : indice « 12 » refusé
    ws = load_workbook(dossier / "fiches.xlsx")["Fiches"]
    entetes = [c.value for c in ws[1]]
    statut = entetes.index("Statut") + 1
    assert [ws.cell(row=r, column=statut).value for r in (2, 3, 4)] == ["OK", "OK", "ERREUR"]
