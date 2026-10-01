"""La procédure de suppression d'un plan, telle que l'utilisateur la décrit : ouvrir le plan, vérifier
qu'aucun client n'y est associé et qu'aucune référence n'est validée ou utilisée ailleurs ; sinon NE
RIEN TOUCHER, prévenir, passer au plan suivant. Sinon retirer les références une à une, puis
supprimer avec la corbeille. Toute fenêtre du portail qui n'était pas attendue est refusée : la ligne
est laissée de côté, rien n'est validé à l'aveugle."""

import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

from autoweb.excel import creer_classeur
from autoweb.runner import Options, lancer
from autoweb.scenario import charger


def _etat_initial():
    return {
        "A": {"refs": [("R1", "Brouillon", "non"), ("R2", "En cours", "non")], "clients": []},
        "B": {"refs": [("R3", "Validé", "oui")], "clients": []},
        "C": {"refs": [], "clients": ["Client du plan C"]},
        "D": {"refs": [], "clients": [], "dossier": True},  # règle du serveur inconnue de la tâche
    }


class _Base(BaseHTTPRequestHandler):
    plans = {}
    envois = []

    def log_message(self, *args):
        pass

    def _html(self, corps):
        donnees = corps.encode()
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(donnees)))
        self.end_headers()
        self.wfile.write(donnees)

    def do_GET(self):
        u = urlsplit(self.path)
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        n, onglet = q.get("n", ""), q.get("onglet", "general")
        plan = type(self).plans.get(n)
        if u.path == "/supprime":
            return self._html("<!doctype html><meta charset=utf-8><h1>Plan supprimé</h1>")
        if plan is None:
            return self._html("<!doctype html><meta charset=utf-8><h1>Plan introuvable</h1>")
        onglets = (f"<nav class=tabs><a href='/plan?n={n}&onglet=general'>Général</a> "
                   f"<a href='/plan?n={n}&onglet=references'>Références</a> "
                   f"<a href='/plan?n={n}&onglet=clients'>Clients</a></nav>")
        if onglet == "references":
            lignes = "".join(f"<tr><td>{r}</td><td>{s}</td><td>{a}</td><td><button class=retirer "
                             f"onclick=\"fetch('/retirer?n={n}&r={r}',{{method:'POST'}}).then(() => location.reload())\">Retirer"
                             f"</button></td></tr>" for r, s, a in plan["refs"])
            lignes = lignes or "<tr><td colspan=4>Aucune référence</td></tr>"
            contenu = ("<table id=refs><thead><tr><th>Référence</th><th>Statut</th><th>Utilisé ailleurs</th><th></th></tr>"
                       f"</thead><tbody>{lignes}</tbody></table>")
        elif onglet == "clients":
            lignes = "".join(f"<tr><td>{c}</td></tr>" for c in plan["clients"]) or "<tr><td>Aucun client</td></tr>"
            contenu = f"<table id=clients><thead><tr><th>Client</th></tr></thead><tbody>{lignes}</tbody></table>"
        else:
            contenu = f"<p>Plan {n}</p>"
        corbeille = ("<button class=corbeille title=Supprimer onclick=\"if (confirm('Voulez-vous vraiment supprimer ce plan ?')) "
                     f"fetch('/supprimer?n={n}',{{method:'POST'}}).then(r => r.text()).then(t => {{ if (t !== 'ok') alert(t); "
                     "else location = '/supprime'; })\">🗑</button>")
        return self._html(f"<!doctype html><meta charset=utf-8><title>Plan {n}</title><h1>Plan {n}</h1>{onglets}{contenu}{corbeille}")

    def do_POST(self):
        u = urlsplit(self.path)
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        type(self).envois.append(self.path)
        plan = type(self).plans.get(q.get("n", ""))
        reponse = "ok"
        if u.path == "/retirer" and plan is not None:
            plan["refs"] = [r for r in plan["refs"] if r[0] != q.get("r")]
        elif u.path == "/supprimer" and plan is not None:
            if plan["refs"] or plan["clients"]:
                reponse = "Attention : il reste des références ou des clients associés"
            elif plan.get("dossier"):
                reponse = "Attention : ce plan est utilisé dans un dossier d'affaire"
            else:
                del type(self).plans[q["n"]]
        donnees = reponse.encode()
        self.send_response(200)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.send_header("Content-Length", str(len(donnees)))
        self.end_headers()
        self.wfile.write(donnees)


@pytest.fixture
def base():
    _Base.plans = _etat_initial()
    _Base.envois = []
    serveur = ThreadingHTTPServer(("127.0.0.1", 0), _Base)
    threading.Thread(target=serveur.serve_forever, daemon=True).start()
    try:
        yield f"http://127.0.0.1:{serveur.server_address[1]}"
    finally:
        serveur.shutdown()
        serveur.server_close()


TACHE = """nom: supprimer des plans
navigateur: {canal: auto, visible: false, profil: profil, dialogues: prudent,
             dialogues_ok: ["Voulez-vous vraiment supprimer ce plan"]}
excel: {fichier: plans.xlsx, feuille: Suivi, colonne_libelle: Numéro}
variables: {url: "%s"}
confirmer: true
etapes:
  # 1. vérifications, AVANT de toucher quoi que ce soit
  - aller: "{{url}}/plan?n={{Numéro}}&onglet=clients"
  - si:
      tableau: "#clients"
      vide: false
      alors:
        - prevenir: "Plan {{Numéro}} : un client est associé. Je n'y touche pas, je passe au suivant."
        - ignorer: "client associé : rien n'a été touché"
  - aller: "{{url}}/plan?n={{Numéro}}&onglet=references"
  - si:
      tableau: "#refs"
      colonne: "Utilisé ailleurs"
      ligne_contient: "oui"
      alors:
        - prevenir: "Plan {{Numéro}} : une référence est utilisée ailleurs. Je n'y touche pas."
        - ignorer: "référence utilisée ailleurs : rien n'a été touché"
  - si:
      tableau: "#refs"
      colonne: "Statut"
      ligne_contient: "Validé"
      alors:
        - prevenir: "Plan {{Numéro}} : une référence est validée. Je n'y touche pas."
        - ignorer: "référence validée : rien n'a été touché"
  # 2. retirer les références une à une, puis la corbeille
  - repeter:
      tant_que: {tableau: "#refs", vide: false}
      max: 20
      etapes:
        - cliquer: {selecteur: "#refs tbody tr button.retirer", premier: true}
        - attendre: {chargement: reseau, delai: 5000}
  - attendre: {chargement: reseau, delai: 5000}
  - attendre: "button.corbeille"
  - cliquer: "button.corbeille"
  - attendre: {chargement: reseau, delai: 5000}
  - verifier: {texte_page: "Plan supprimé"}
"""


def test_suppression_prudente_plan_par_plan(base, tmp_path, navigateur_ok):
    creer_classeur(tmp_path / "plans.xlsx", ["Numéro"], [["A"], ["B"], ["C"], ["D"]])
    chemin = tmp_path / "supprimer.yaml"
    chemin.write_text(TACHE % base, encoding="utf-8")
    bilan = lancer(charger(chemin), Options(visible=False, interactif=False))

    # A : références retirées puis plan supprimé ; B, C, D : rien n'a été touché
    assert "A" not in _Base.plans
    assert _Base.plans["B"]["refs"] == [("R3", "Validé", "oui")]
    assert _Base.plans["C"]["clients"] == ["Client du plan C"]
    assert "D" in _Base.plans
    assert not any("n=B" in e or "n=C" in e for e in _Base.envois), _Base.envois
    assert (bilan.ok, bilan.ignorees, bilan.erreurs) == (1, 3, 0), bilan.resume()

    resume = bilan.resume()
    assert "Laissé(s) de côté" in resume
    assert "référence utilisée ailleurs" in resume and "client associé" in resume
    assert "fenêtre inattendue du portail" in resume and "dossier d'affaire" in resume  # D : règle du serveur

    ws = load_workbook(tmp_path / "plans.xlsx")["Suivi"]
    entetes = [c.value for c in ws[1]]
    statuts = [ws.cell(row=i, column=entetes.index("Statut") + 1).value for i in range(2, 6)]
    assert statuts == ["OK", "IGNORE", "IGNORE", "IGNORE"]


def test_tableau_ou_colonne_introuvable_n_est_jamais_rien_trouve(base, tmp_path, navigateur_ok):
    """Une colonne mal nommée (le portail a changé) : ERREUR pour la ligne, jamais « rien trouvé, on supprime »."""
    creer_classeur(tmp_path / "plans.xlsx", ["Numéro"], [["A"]])
    chemin = tmp_path / "supprimer.yaml"
    chemin.write_text((TACHE % base).replace('colonne: "Statut"', 'colonne: "État"'), encoding="utf-8")
    bilan = lancer(charger(chemin), Options(visible=False, interactif=False))
    assert bilan.erreurs == 1 and "A" in _Base.plans
    assert _Base.plans["A"]["refs"]  # aucune référence retirée
