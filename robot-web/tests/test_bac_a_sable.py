"""La fausse base documentaire (base_demo.html) : export CSV par contrat, modification
de fiches, et construction du scénario d'export par l'assistant."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from autoweb.assistant import Dialogue, construire
from autoweb.cli import _servir_dossier, preparer_bac_a_sable
from autoweb.excel import creer_classeur
from autoweb.navigateur import Navigateur
from autoweb.releve import charger_releve, relever
from autoweb.runner import Options, lancer
from autoweb.scenario import ConfigNavigateur, charger

AVANT = """avant:
  - aller: "{{url}}"
  - si:
      visible: "#btn-connexion"
      alors:
        - remplir: {"#utilisateur": demo, "#mot-de-passe": demo}
        - cliquer: "#btn-connexion"
  - attendre: "#liste"
"""


def _plans_attendus():
    """Même génération que base_demo.html : (numero, contrat) des 60 plans."""
    contrats = ["HDK", "LMN", "QRS", "TUV"]
    types = ["ARC", "STR", "ELE", "PLB"]
    compteurs = {}
    plans = []
    for i in range(60):
        c, t = contrats[i % 4], types[(i // 4) % 4]
        cle = f"{c}-{t}"
        compteurs[cle] = compteurs.get(cle, 0) + 1
        plans.append((f"{cle}-{compteurs[cle]:03d}", c))
    return plans


@pytest.fixture
def bac(tmp_path):
    preparer_bac_a_sable(tmp_path)
    serveur, url = _servir_dossier(tmp_path, "base_demo.html", port=0)
    try:
        yield tmp_path, url
    finally:
        serveur.shutdown()
        serveur.server_close()


def test_export_par_contrat(bac, navigateur_ok):
    dossier, url = bac
    (dossier / "export.yaml").write_text(f"""
nom: export contrat
navigateur: {{canal: auto, visible: false}}
excel: {{fichier: contrats.xlsx, feuille: Contrats}}
variables: {{url: "{url}"}}
{AVANT}
etapes:
  - aller: "{{{{url}}}}"
  - choisir: {{selecteur: "#filtre-contrat", valeur: "{{{{Contrat}}}}"}}
  - cliquer: "#btn-rechercher"
  - attendre: {{chargement: reseau, delai: 5000}}
    optionnel: true
  - telecharger: {{cliquer: "#btn-exporter", vers: "exports/", renommer: "plans_{{{{Contrat}}}}", convertir_excel: true, vers_colonne: "Fichier export"}}
""", encoding="utf-8")
    bilan = lancer(charger(dossier / "export.yaml"), Options(interactif=False))
    assert (bilan.ok, bilan.erreurs) == (2, 0)

    csv_hdk = dossier / "exports" / "plans_HDK.csv"
    xlsx_hdk = dossier / "exports" / "plans_HDK.xlsx"
    assert csv_hdk.exists() and xlsx_hdk.exists() and (dossier / "exports" / "plans_QRS.xlsx").exists()
    ws = load_workbook(xlsx_hdk).active
    lignes = list(ws.iter_rows(values_only=True))
    assert lignes[0] == ("Numéro", "Contrat", "Titre", "Type", "Statut", "Indice", "Date")
    attendus = [n for n, c in _plans_attendus() if c == "HDK"]
    assert [l[0] for l in lignes[1:]] == attendus and len(attendus) == 15
    assert all(l[1] == "HDK" for l in lignes[1:])
    assert "Électricité" in {l[3] for l in lignes[1:]}  # accents conservés (BOM utf-8)

    suivi = load_workbook(dossier / "contrats.xlsx")["Contrats"]
    entetes = [c.value for c in suivi[1]]
    assert suivi.cell(row=2, column=entetes.index("Fichier export") + 1).value.endswith("plans_HDK.xlsx")
    assert suivi.cell(row=2, column=entetes.index("Statut") + 1).value == "OK"


def test_modification_de_fiches(bac, navigateur_ok):
    dossier, url = bac
    (dossier / "fiches.yaml").write_text(f"""
nom: fiches
navigateur: {{canal: auto, visible: false}}
excel: {{fichier: fiches.xlsx, feuille: Fiches, colonne_libelle: Numéro}}
variables: {{url: "{url}"}}
{AVANT}
etapes:
  - aller: "{{{{url}}}}#fiche={{{{Numéro}}}}"
  - attendre: "#formulaire-fiche"
  - verifier: {{selecteur: "#fiche-numero", contient: "{{{{Numéro}}}}"}}
  - choisir: {{selecteur: "#fiche-statut", valeur: "{{{{Nouveau statut}}}}"}}
  - remplir:
      "#fiche-indice": "{{{{Indice}}}}"
      "#fiche-commentaire": "{{{{Commentaire}}}}"
  - cliquer: "#btn-enregistrer"
  - verifier: {{selecteur: "#message-fiche", contient: "enregistrée", delai: 3000}}
""", encoding="utf-8")
    bilan = lancer(charger(dossier / "fiches.yaml"), Options(interactif=False))
    assert (bilan.ok, bilan.erreurs) == (2, 1)
    ws = load_workbook(dossier / "fiches.xlsx")["Fiches"]
    entetes = [c.value for c in ws[1]]
    statut = entetes.index("Statut") + 1
    assert [ws.cell(row=r, column=statut).value for r in (2, 3, 4)] == ["OK", "OK", "ERREUR"]
    assert "enregistrée" in ws.cell(row=4, column=entetes.index("Message") + 1).value


def test_assistant_construit_l_export(bac, navigateur_ok):
    dossier, url = bac
    nav = Navigateur(ConfigNavigateur(canal="auto", visible=False), dossier, visible=False)
    page = nav.ouvrir()
    try:
        page.goto(url)
        page.fill("#utilisateur", "demo")
        page.fill("#mot-de-passe", "demo")
        page.click("#btn-connexion")
        page.wait_for_selector("#liste")
        releve = charger_releve(relever(page, dossier / "releves", nom="export contrat"))
    finally:
        nav.fermer()

    reponses = [
        "2",     # filtres + export
        "",      # adresse
        "",      # Contrat -> colonne Contrat proposée
        "",      # Type de plan -> 0
        "",      # Statut -> 0
        "",      # Numéro contient -> 0
        "",      # bouton recherche : Rechercher proposé
        "",      # bouton export : Exporter CSV proposé (désactivé au moment du relevé, mais listé)
        "",      # dossier exports
        "",      # nom : export_{{Contrat}}
        "",      # convertir en Excel : oui
        "",      # colonne Fichier export
        "n",     # pas de connexion manuelle (ajoutée ci-dessous)
    ]
    texte = construire(releve, ["Contrat"], Dialogue(reponses), nom="export contrat", fichier_excel="contrats.xlsx", feuille="Contrats")
    assert 'choisir: {selecteur: "#filtre-contrat", valeur: "{{Contrat}}"}' in texte
    assert "#filtre-type" not in texte and "#recherche" not in texte
    assert 'cliquer: "#btn-rechercher"' in texte
    assert 'telecharger: {cliquer: "#btn-exporter", vers: "exports/", renommer: "export_{{Contrat}}", convertir_excel: true, vers_colonne: "Fichier export"}' in texte
    assert "pause" not in texte and "capture" not in texte

    texte = texte.replace("etapes:\n", AVANT + "etapes:\n")
    (dossier / "export_assistant.yaml").write_text(texte, encoding="utf-8")
    bilan = lancer(charger(dossier / "export_assistant.yaml"), Options(visible=False, interactif=False))
    assert (bilan.ok, bilan.erreurs) == (2, 0)
    assert (dossier / "exports" / "export_HDK.xlsx").exists() and (dossier / "exports" / "export_QRS.xlsx").exists()
