# -*- coding: utf-8 -*-
# ROBOT - lancement :  python robot.txt   (1 = test, 2 = releve des pages, 3 = lancer une tache, 4 = voir,
#                                          5 = explorer le portail (carte), 6 = aller a une page)
# Il ne modifie rien sur le portail sans votre accord ; tout ce qu'il ecrit est dans Bureau > Robot.

import ctypes
import hashlib
import json
import os
import platform
import re
import secrets
import subprocess
import sys
import tempfile
import time
import unicodedata
import webbrowser
from datetime import datetime
from pathlib import Path
from collections import deque
from urllib.parse import parse_qsl, unquote, urljoin, urlsplit

VERSION = "5.1"
WINDOWS = os.name == "nt"
# pour les essais automatiques du programme seulement (navigateur caché)
INVISIBLE = os.environ.get("PRISE_EN_MAIN_INVISIBLE") == "1"
FICHIER_CLAVIER = "essai_clavier.txt"

# Sur le portail, le robot ne clique jamais sur ce qui pourrait modifier quelque chose.
# Comparé à mots() : texte sans accents, en minuscules, « camelCase » et ponctuation découpés en mots
# (« btnSupprimer », « lnk_Suppr », « Supprimé »... sont tous pris).
MOTS_INTERDITS = re.compile(
    r"suppr|\bdel\b|delet|remov|retir|retrait|effac|erase|destro|detrui|vider|\bclear|nettoy|corbeil|trash|purg|"
    r"dupliq|duplic|\bcopi|\bcopy|\bclon|valid|enregistr|sauv|\bsave|\bcreer|\bcree\b|\bcreat|nouveau|nouvelle|"
    r"\bnew\b|ajout|\badd\b|modif|\bedit|envoy|envoi|\bsend\b|soumet|soumis|submit|confirm|annul|cancel|"
    r"\barchiv|\bpublier|\bpublish|transfer|transmet|transmis|forward|\bstatut|\bstatus|\betat\b|rejet|reject|"
    r"approuv|approb|approv|\bsigner|\bsign\b|\bsignature|\bassocier|\bassociation|rattach|attach|joindre|detach|\bimport|"
    r"\breviser\b|\bliberer|\brelease|verrou|\block|unlock|bascul|clotur|\bclore|\bclos\b|ferm|\bclose|"
    r"mettre a jour|mise a jour|\bmaj\b|update|upgrade|reinitialis|reset|\braz\b|restaur|restore|\blancer|execut|"
    r"\brun\b|\bstart|demarr|arret|\bstop\b|activer|desactiv|activate|deactivat|\benable|\bdisable|affect|assign|"
    r"attribu|\baccept|refus|declin|\bterminer|finish|\bcompleter|\bcomplete\b|resolu|resou|resolv|appliqu|"
    r"\bapply|upload|televers|\bdepos|deplac|\bmove\b|renomm|rename|rempla|replace|\bgenerer|regener|generat|"
    r"\bcommander|\border\b|\bpayer|paiement|\bpay\b|factur|repond|reply|\bcommenter|relanc|prise en charge|"
    r"prendre en charge|\btraiter|\btraite\b|synchro|recalcul|fusion|merge|deploy|deploi|abandon|revoq|revok|"
    r"autoris|inscri|mise en|mettre en|\bdissocier|unlink|enlev|resili|\bradier|radiation|suspend|\bbloqu|debloqu|"
    r"ecras|\bexclu|\bdelier|declass|rebut|remettre a zero|\bcontinuer\b|poursuiv|"
    r"partag|\bshare|\bcalcul|planifi|schedul|prolong|renouvel|acquitt|acknowledg|rouvr|reouvr|reopen|"
    r"^oui\b|^ok\b|^yes\b|^non$|^no$|^go$",
    re.IGNORECASE,
)

# Adresse copiée juste après un clic : celle d'une ACTION si l'un de ses mots est un de ces verbes.
VERBES_ACTION = set("""supprimer suppression suppr delete del remove retirer effacer erase destroy detruire purge purger
vider clear valider validate validation enregistrer save sauvegarder dupliquer duplicate copier copy clone creer create
new nouveau nouvelle ajouter add edit editer modifier modification update maj envoyer send submit soumettre approuver
approve rejeter reject annuler cancel archiver archive publier publish importer import upload deplacer move renommer
rename remplacer replace activer desactiver activate deactivate enable disable lock unlock verrouiller deverrouiller
signer sign cloturer close reset restaurer restore transferer transfer assigner assign affecter attribuer appliquer
apply confirmer confirm executer execute diffuser viser emettre notifier indicer figer reserver""".split())


def mots(texte):
    t = sans_accents(texte)
    t = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", t)
    t = re.sub(r"[^A-Za-z0-9]+", " ", t)
    return " ".join(t.split()).lower()


# expressions qui contiennent un mot d'action mais ne modifient rien
NEUTRES = ("nouvelle recherche", "new search", "afficher les nouveautes", "nouveautes")


def interdit(texte):
    t = " " + mots(texte) + " "
    for phrase in NEUTRES:
        t = t.replace(" " + phrase + " ", " ")
    return bool(MOTS_INTERDITS.search(t.strip()))


# ---------------------------------------------------------------------------- console
def preparer_console():
    for flux in (sys.stdout, sys.stderr):
        try:
            flux.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


def ecrire(texte=""):
    print(texte, flush=True)


def vider_clavier():
    """Une touche tapée pendant que le robot travaillait ne doit pas répondre à la question suivante."""
    if not WINDOWS:
        return
    try:
        import msvcrt
        while msvcrt.kbhit():
            msvcrt.getwch()
    except Exception:
        pass


def demander(question, defaut=""):
    vider_clavier()
    try:
        reponse = input(question + " ")
    except EOFError:
        reponse = ""
    return reponse.strip() or defaut


def oui(question, defaut=""):
    """o ou n. Sans défaut (questions « avez-vous vu ? »), Entrée seule redemande : on ne note jamais OK par erreur."""
    while True:
        reponse = demander(f"{question} (o/n)" + (f" [{defaut}]" if defaut else "") + " :", defaut).lower()
        if reponse[:1] in ("o", "y", "n"):
            return reponse[:1] != "n"
        ecrire("   Repondez o ou n.")


def pause(texte="Appuyez sur Entree pour commencer..."):
    demander("   >>> " + texte)


def titre_etape(numero, texte):
    ecrire()
    ecrire("=" * 70)
    ecrire(f" ETAPE {numero} : {texte}")
    ecrire("=" * 70)


def revenez_ici():
    revenir_console()
    ecrire("   Revenez dans CETTE fenetre noire (cliquez dedans) pour repondre.")


def sans_accents(texte):
    return "".join(c for c in unicodedata.normalize("NFD", str(texte)) if unicodedata.category(c) != "Mn")


def masquer(texte, url=""):
    """Rien du portail ni de votre compte dans le résultat : ni adresse, ni chemin, ni nom d'utilisateur."""
    texte = str(texte or "")
    if url:
        texte = texte.replace(url, "<portail>")
        hote = urlsplit(url).hostname or ""
        if len(hote) >= 3:
            texte = re.sub(r"(?<![\w.-])" + re.escape(hote) + r"(?![\w-])", "<portail>", texte, flags=re.IGNORECASE)
    texte = re.sub(r"(https?|file)://\S+", "<adresse>", texte)
    texte = texte.replace("\\\\", "\\")  # chemins affichés avec des barres doublées (repr)
    for chemin in sorted({str(Path.home()), os.environ.get("USERPROFILE", ""), os.environ.get("OneDrive", ""),
                          os.environ.get("LOCALAPPDATA", ""), os.environ.get("APPDATA", "")}, key=len, reverse=True):
        if len(chemin) > 3:
            texte = texte.replace(chemin, "~")
    texte = re.sub(r"(?i)\b[a-z]:\\[^'\"\n]*", "<chemin>", texte)
    texte = re.sub(r"/(home|Users)/[^/\s'\"]+", r"/\1/<moi>", texte)
    utilisateur = os.environ.get("USERNAME") or os.environ.get("USER") or ""
    if len(utilisateur) >= 3:
        texte = re.sub(r"(?<![\w])" + re.escape(utilisateur) + r"(?![\w])", "<moi>", texte, flags=re.IGNORECASE)
    return " ".join(texte.split())[:100]


def erreur_courte(e):
    """Une erreur de fichier sans son chemin (qui contient votre nom et celui de l'entreprise)."""
    code = getattr(e, "winerror", None) or getattr(e, "errno", None)
    return e.__class__.__name__ + (f" ({code})" if code else "")


def premiere_ligne(erreur):
    lignes = [l.strip() for l in str(erreur).splitlines() if l.strip()]
    return lignes[0] if lignes else erreur.__class__.__name__


def expliquer(erreur):
    """Une erreur du navigateur, en quelques mots compréhensibles (sans adresse ni texte de page)."""
    t = str(erreur).lower()
    connues = (
        (("disallowed by the system admin", "remotedebuggingallowed", "remote debugging is disallowed"),
         "pilotage du navigateur interdit par une strategie de l'entreprise"),
        (("executable doesn't exist", "is not found at", "looks like playwright", "please run the following command"),
         "navigateur absent de ce poste"),
        (("err_blocked_by_administrator",), "page bloquee par une strategie de l'entreprise"),
        (("err_invalid_auth_credentials", "http authentication"),
         "le portail demande une connexion par une petite fenetre du navigateur"),
        (("err_name_not_resolved",), "adresse introuvable depuis ce poste"),
        (("err_cert", "certificate"), "certificat de securite refuse"),
        (("err_tunnel_connection_failed", "err_proxy", "proxy"), "probleme de proxy"),
        (("err_connection_refused", "err_connection_timed_out", "err_timed_out", "err_internet_disconnected"),
         "le portail ne repond pas depuis ce poste"),
        (("target page, context or browser has been closed", "browser has been closed", "target closed"),
         "la fenetre du robot a ete fermee"),
        (("timeout",), "delai depasse"),
    )
    for indices, explication in connues:
        if any(i in t for i in indices):
            return explication
    return masquer(premiere_ligne(erreur))


# ---------------------------------------------------------------------------- bilan
CODES = {"OK": "1", "ECHEC": "0", "SAUTE": "-", "PAS VU": "?", "NON DISPONIBLE": "x"}
ORDRE = "PFKSWNDRCV"


class Bilan:
    def __init__(self):
        self.lignes = {}  # lettre -> (intitule, statut, detail)

    def noter(self, lettre, intitule, statut, detail="", afficher=True):
        self.lignes[lettre] = (intitule, statut, masquer(detail))
        if afficher:
            ecrire(f"   --> {intitule} : {statut}" + (f"  ({masquer(detail)})" if detail else ""))
            if statut == "ECHEC":
                ecrire("       Ce n'est pas grave : c'est justement ce qu'on cherche a savoir. On continue.")

    def tries(self):
        return sorted(self.lignes.items(), key=lambda kv: ORDRE.index(kv[0]) if kv[0] in ORDRE else 99)

    def code(self):
        return " ".join(lettre + CODES.get(v[1], "?") for lettre, v in self.tries())

    def texte(self):
        lignes = [f"RESULTAT DE LA PRISE EN MAIN (version {VERSION}) - {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                  "-" * 70]
        for lettre, (intitule, statut, detail) in self.tries():
            lignes.append(f" {lettre}  {(intitule + ' ').ljust(40, '.')} {statut}")
            if detail:
                lignes.append(f"      {detail}")
        lignes += ["-" * 70, "CODE : " + self.code()]
        return "\n".join(lignes)


# ---------------------------------------------------------------------------- Windows : clavier, souris, fenêtres
if WINDOWS:
    from ctypes import wintypes

    user32 = ctypes.WinDLL("user32", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    ULONG_PTR = ctypes.c_size_t

    class MOUSEINPUT(ctypes.Structure):
        _fields_ = [("dx", wintypes.LONG), ("dy", wintypes.LONG), ("mouseData", wintypes.DWORD),
                    ("dwFlags", wintypes.DWORD), ("time", wintypes.DWORD), ("dwExtraInfo", ULONG_PTR)]

    class KEYBDINPUT(ctypes.Structure):
        _fields_ = [("wVk", wintypes.WORD), ("wScan", wintypes.WORD), ("dwFlags", wintypes.DWORD),
                    ("time", wintypes.DWORD), ("dwExtraInfo", ULONG_PTR)]

    class HARDWAREINPUT(ctypes.Structure):
        _fields_ = [("uMsg", wintypes.DWORD), ("wParamL", wintypes.WORD), ("wParamH", wintypes.WORD)]

    class _UNION_ENTREE(ctypes.Union):
        _fields_ = [("mi", MOUSEINPUT), ("ki", KEYBDINPUT), ("hi", HARDWAREINPUT)]

    class INPUT(ctypes.Structure):
        _anonymous_ = ("u",)
        _fields_ = [("type", wintypes.DWORD), ("u", _UNION_ENTREE)]

    user32.SendInput.argtypes = [wintypes.UINT, ctypes.POINTER(INPUT), ctypes.c_int]
    user32.SendInput.restype = wintypes.UINT
    user32.GetForegroundWindow.restype = wintypes.HWND
    user32.SetForegroundWindow.argtypes = [wintypes.HWND]
    user32.SetForegroundWindow.restype = wintypes.BOOL
    user32.GetClassNameW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
    user32.GetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
    user32.GetCursorPos.argtypes = [ctypes.POINTER(wintypes.POINT)]
    user32.SetCursorPos.argtypes = [ctypes.c_int, ctypes.c_int]
    user32.GetWindowRect.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.RECT)]
    user32.WindowFromPoint.argtypes = [wintypes.POINT]
    user32.WindowFromPoint.restype = wintypes.HWND
    user32.GetAncestor.argtypes = [wintypes.HWND, wintypes.UINT]
    user32.GetAncestor.restype = wintypes.HWND
    user32.IsWindowVisible.argtypes = [wintypes.HWND]
    user32.IsIconic.argtypes = [wintypes.HWND]
    user32.ShowWindow.argtypes = [wintypes.HWND, ctypes.c_int]
    user32.GetSystemMetrics.argtypes = [ctypes.c_int]
    kernel32.GetConsoleWindow.restype = wintypes.HWND
    kernel32.GetStdHandle.argtypes = [wintypes.DWORD]
    kernel32.GetStdHandle.restype = wintypes.HANDLE
    kernel32.GetConsoleMode.argtypes = [wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD)]
    kernel32.SetConsoleMode.argtypes = [wintypes.HANDLE, wintypes.DWORD]

    INPUT_MOUSE, INPUT_KEYBOARD = 0, 1
    KEYEVENTF_KEYUP, KEYEVENTF_UNICODE = 0x0002, 0x0004
    MOUSEEVENTF_MOVE, MOUSEEVENTF_ABSOLUTE, MOUSEEVENTF_VIRTUALDESK = 0x0001, 0x8000, 0x4000
    MOUSEEVENTF_LEFTDOWN, MOUSEEVENTF_LEFTUP = 0x0002, 0x0004
    MOUSEEVENTF_RIGHTDOWN, MOUSEEVENTF_RIGHTUP = 0x0008, 0x0010
    VK_RETURN = 0x0D
    GA_ROOT, GA_ROOTOWNER = 2, 3
    SM_SWAPBUTTON = 23
    SM_XVIRTUALSCREEN, SM_YVIRTUALSCREEN, SM_CXVIRTUALSCREEN, SM_CYVIRTUALSCREEN = 76, 77, 78, 79


def rendre_net():
    """Coordonnées de la souris en vrais pixels, même avec un affichage à 125 % ou 150 %."""
    if not WINDOWS:
        return
    try:
        if ctypes.windll.shcore.SetProcessDpiAwareness(2) == 0:  # renvoie un code, ne lève pas d'erreur
            return
    except Exception:
        pass
    try:
        user32.SetProcessDPIAware()
    except Exception:
        pass


def edition_rapide_coupee():
    """Console classique : un clic dedans la met en mode « Sélectionner » et Entrée ne répond plus.
    On coupe ce mode le temps du test ; renvoie l'ancien réglage, à remettre à la fin."""
    if not WINDOWS:
        return None
    try:
        h = kernel32.GetStdHandle(wintypes.DWORD(-10 & 0xFFFFFFFF))
        mode = wintypes.DWORD()
        if not kernel32.GetConsoleMode(h, ctypes.byref(mode)):
            return None
        # sans « Sélectionner » (0x40) ni souris captée (0x10) : le clic droit colle à nouveau
        kernel32.SetConsoleMode(h, (mode.value | 0x0080) & ~0x0040 & ~0x0010)
        return mode.value
    except Exception:
        return None


def remettre_console(ancien):
    if WINDOWS and ancien is not None:
        try:
            kernel32.SetConsoleMode(kernel32.GetStdHandle(wintypes.DWORD(-10 & 0xFFFFFFFF)), ancien)
        except Exception:
            pass


def _envoyer(*entrees):
    tableau = (INPUT * len(entrees))(*entrees)
    n = user32.SendInput(len(entrees), tableau, ctypes.sizeof(INPUT))
    if n != len(entrees):
        raise OSError(f"Windows a refuse le geste simule (SendInput, code {ctypes.get_last_error()})")


def _touche(vk=0, unite=0, relache=False):
    e = INPUT()
    e.type = INPUT_KEYBOARD
    e.ki = KEYBDINPUT(vk, unite, (KEYEVENTF_UNICODE if not vk else 0) | (KEYEVENTF_KEYUP if relache else 0), 0, 0)
    return e


def _souris(drapeaux, dx=0, dy=0):
    e = INPUT()
    e.type = INPUT_MOUSE
    e.mi = MOUSEINPUT(dx, dy, 0, drapeaux, 0, 0)
    return e


def cliquer_ici():
    """Appui et relâchement du bouton, là où est la souris (vérifiée juste avant), en un seul envoi."""
    gaucher = bool(user32.GetSystemMetrics(SM_SWAPBUTTON))
    appui, relache = (MOUSEEVENTF_RIGHTDOWN, MOUSEEVENTF_RIGHTUP) if gaucher else (MOUSEEVENTF_LEFTDOWN, MOUSEEVENTF_LEFTUP)
    _envoyer(_souris(appui), _souris(relache))


def zone_de_travail(x, y):
    """L'écran (sans la barre des tâches) qui contient ce point : (gauche, haut, droite, bas)."""
    class MONITORINFO(ctypes.Structure):
        _fields_ = [("cbSize", wintypes.DWORD), ("rcMonitor", wintypes.RECT), ("rcWork", wintypes.RECT),
                    ("dwFlags", wintypes.DWORD)]
    try:
        user32.MonitorFromPoint.argtypes = [wintypes.POINT, wintypes.DWORD]
        user32.MonitorFromPoint.restype = wintypes.HANDLE
        user32.GetMonitorInfoW.argtypes = [wintypes.HANDLE, ctypes.POINTER(MONITORINFO)]
        ecran = user32.MonitorFromPoint(wintypes.POINT(int(x), int(y)), 2)  # MONITOR_DEFAULTTONEAREST
        info = MONITORINFO()
        info.cbSize = ctypes.sizeof(MONITORINFO)
        if ecran and user32.GetMonitorInfoW(ecran, ctypes.byref(info)):
            r = info.rcWork
            return r.left, r.top, r.right, r.bottom
    except Exception:
        pass
    return 0, 0, user32.GetSystemMetrics(0), user32.GetSystemMetrics(1)


def fenetre_active():
    return user32.GetForegroundWindow()


def _classe(hwnd):
    tampon = ctypes.create_unicode_buffer(256)
    user32.GetClassNameW(hwnd, tampon, 256)
    return tampon.value


def _titre(hwnd):
    tampon = ctypes.create_unicode_buffer(512)
    user32.GetWindowTextW(hwnd, tampon, 512)
    return tampon.value


def est_notre_fichier(hwnd):
    """La fenêtre qui affiche NOTRE fichier essai_clavier.txt (Bloc-notes, ou l'éditeur choisi par
    l'entreprise) : jamais un document de l'utilisateur, jamais un navigateur ni l'Explorateur."""
    if not hwnd:
        return False
    classe = _classe(hwnd)
    if classe in ("CabinetWClass", "ExploreWClass", "Chrome_WidgetWin_1", "MozillaWindowClass", "ConsoleWindowClass",
                  "CASCADIA_HOSTING_WINDOW_CLASS"):
        return False
    return FICHIER_CLAVIER.split(".")[0] in _titre(hwnd).lower()


def attendre_fenetre(test, secondes):
    fin = time.time() + secondes
    while time.time() < fin:
        hwnd = fenetre_active()
        if test(hwnd):
            return hwnd
        time.sleep(0.2)
    return None


def taper(texte, hwnd, delai=0.04):
    """Tape le texte, caractère par caractère, SEULEMENT tant que notre fichier est au premier plan :
    si une autre fenêtre passe devant, le robot s'arrête (il n'écrit jamais ailleurs)."""
    ecrits = 0
    for caractere in texte:
        active = fenetre_active()
        if active != hwnd or not est_notre_fichier(active):
            return ecrits, False
        if caractere == "\n":
            _envoyer(_touche(vk=VK_RETURN), _touche(vk=VK_RETURN, relache=True))
        else:
            donnees = caractere.encode("utf-16-le")
            for i in range(0, len(donnees), 2):
                unite = int.from_bytes(donnees[i:i + 2], "little")
                _envoyer(_touche(unite=unite), _touche(unite=unite, relache=True))
        ecrits += 1
        time.sleep(delai)
    return ecrits, True


def position_souris():
    p = wintypes.POINT()
    user32.GetCursorPos(ctypes.byref(p))
    return p.x, p.y


def deplacer_doucement(x, y, duree=0.5):
    x0, y0 = position_souris()
    pas = 25
    for i in range(1, pas + 1):
        user32.SetCursorPos(int(x0 + (x - x0) * i / pas), int(y0 + (y - y0) * i / pas))
        time.sleep(duree / pas)


def revenir_console():
    """Remet la fenêtre noire devant (cmd classique ; dans Windows Terminal, sa fenêtre visible)."""
    if not WINDOWS:
        return
    try:
        h = kernel32.GetConsoleWindow()
        if h and not user32.IsWindowVisible(h):
            h = user32.GetAncestor(h, GA_ROOTOWNER)
        if h and user32.IsWindowVisible(h):
            user32.SetForegroundWindow(h)
    except Exception:
        pass


def dossier_bureau():
    if WINDOWS:
        tampon = ctypes.create_unicode_buffer(260)
        try:
            if ctypes.windll.shell32.SHGetFolderPathW(None, 0x10, None, 0, tampon) == 0 and tampon.value:
                return Path(tampon.value)
        except Exception:
            pass
    for p in (Path.home() / "Desktop", Path.home() / "Bureau"):
        if p.is_dir():
            return p
    return Path.home()


def effacer_dossier(chemin):
    """Efface un dossier temporaire du robot, en réessayant : sous Windows, le navigateur garde parfois
    des fichiers ouverts quelques instants après sa fermeture. Jamais d'erreur : au pire, il reste."""
    import shutil
    for _ in range(6):
        shutil.rmtree(chemin, ignore_errors=True)
        if not os.path.exists(chemin):
            return
        time.sleep(0.5)


def ouvrir_dans_windows(chemin):
    if WINDOWS:
        os.startfile(str(chemin))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(chemin)])
    else:
        subprocess.Popen(["xdg-open", str(chemin)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


# ---------------------------------------------------------------------------- étapes 1 à 5
def etape_python(bilan):
    titre_etape(1, "Python fonctionne")
    detail = f"Python {platform.python_version()} {platform.architecture()[0]}, {platform.system()} {platform.release()}"
    ecrire("   " + detail)
    bilan.noter("P", "Python fonctionne", "OK", detail)


def etape_fichiers(bilan):
    titre_etape(2, "Fichiers et dossiers")
    dossier = dossier_bureau() / "Robot"
    ecrire("   Le robot va creer le dossier « Robot » sur votre Bureau, y ecrire")
    ecrire("   le fichier bonjour.txt, puis ouvrir ce dossier.")
    pause()
    try:
        dossier.mkdir(parents=True, exist_ok=True)
        (dossier / "bonjour.txt").write_text(
            "Bonjour ! Ce fichier a ete cree par le robot le %s.\n"
            "Le robot range ici tout ce qu'il ecrit (releves, resultats des taches).\n"
            % datetime.now().strftime("%d/%m/%Y a %H:%M"), encoding="utf-8")
    except Exception as e:
        bilan.noter("F", "Fichiers et dossiers", "ECHEC", erreur_courte(e))
        return dossier if dossier.is_dir() else None
    try:
        ouvrir_dans_windows(dossier)
    except Exception as e:
        bilan.noter("F", "Fichiers et dossiers", "PAS VU", "fichier cree, dossier pas ouvert : " + erreur_courte(e))
        return dossier
    time.sleep(1.5)
    revenez_ici()
    if oui("   Voyez-vous le dossier « Robot » avec le fichier bonjour.txt ?"):
        bilan.noter("F", "Fichiers et dossiers", "OK")
    else:
        bilan.noter("F", "Fichiers et dossiers", "PAS VU", "fichier cree mais dossier pas vu a l'ecran")
    return dossier


def etape_clavier(bilan, dossier):
    titre_etape(3, "Le clavier : ecrire tout seul dans le Bloc-notes")
    if not WINDOWS:
        bilan.noter("K", "Clavier (Bloc-notes)", "NON DISPONIBLE", "pas Windows")
        return None
    ecrire("   Le robot va ouvrir SON fichier essai_clavier.txt dans le Bloc-notes et y ecrire")
    ecrire("   une phrase tout seul. Pendant qu'il ecrit (10 secondes) : NE TOUCHEZ NI AU CLAVIER")
    ecrire("   NI A LA SOURIS. (Si une autre fenetre passe devant, il s'arrete : il n'ecrit jamais")
    ecrire("   ailleurs que dans son fichier.)")
    pause()
    try:
        dossier = dossier if dossier is not None and dossier.is_dir() else Path(tempfile.gettempdir())
        fichier = dossier / FICHIER_CLAVIER
        fichier.write_text("", encoding="utf-8")
        subprocess.Popen(["notepad.exe", str(fichier)])
    except Exception as e:
        bilan.noter("K", "Clavier (Bloc-notes)", "ECHEC", "Bloc-notes impossible a ouvrir : " + erreur_courte(e))
        return None
    hwnd = attendre_fenetre(est_notre_fichier, 10)
    if not hwnd:
        ecrire("   Cliquez sur l'icone du Bloc-notes « essai_clavier » en bas, dans la barre des taches,")
        ecrire("   puis LACHEZ LA SOURIS. Le robot attend 20 secondes...")
        hwnd = attendre_fenetre(est_notre_fichier, 20)
        if hwnd:
            time.sleep(2)  # le temps de lâcher la souris
    if not hwnd:
        devant = fenetre_active()
        bilan.noter("K", "Clavier (Bloc-notes)", "ECHEC",
                    f"le fichier essai_clavier n'est jamais passe devant (fenetre devant : {_classe(devant) if devant else 'aucune'})")
        revenez_ici()
        return None
    time.sleep(1.2)  # le temps que la zone de texte soit prête à recevoir les touches
    try:
        ecrits, complet = taper("Bonjour ! Je suis le robot. J'ecris tout seul, sans toucher au clavier.\n"
                                "Avec les accents : é è à ç ù ê.\n", hwnd)
    except OSError as e:
        bilan.noter("K", "Clavier (Bloc-notes)", "ECHEC", str(e))
        revenez_ici()
        return hwnd
    revenez_ici()
    if not complet:
        ecrire(f"   Arrete apres {ecrits} caractere(s) : une autre fenetre est passee devant.")
    if oui("   Avez-vous vu la phrase s'ecrire toute seule dans le Bloc-notes ?"):
        bilan.noter("K", "Clavier (Bloc-notes)", "OK" if complet else "ECHEC",
                    "" if complet else f"interrompu apres {ecrits} caracteres (autre fenetre devant)")
    else:
        bilan.noter("K", "Clavier (Bloc-notes)", "PAS VU", f"{ecrits} caracteres envoyes, rien vu a l'ecran")
    return hwnd


def _ramener(hwnd):
    """Notre fichier au premier plan (restauré s'il a été réduit). Faux si ce n'est pas possible."""
    if user32.IsIconic(hwnd):
        user32.ShowWindow(hwnd, 9)  # SW_RESTORE
        time.sleep(0.5)
    user32.SetForegroundWindow(hwnd)
    if attendre_fenetre(lambda h: h == hwnd, 2):
        return True
    ecrire("   Cliquez sur l'icone du Bloc-notes « essai_clavier » en bas, dans la barre des taches,")
    ecrire("   puis LACHEZ LA SOURIS. Le robot attend 20 secondes...")
    if not attendre_fenetre(lambda h: h == hwnd and est_notre_fichier(h), 20):
        return False
    time.sleep(2)  # le temps de lâcher la souris
    return True


def etape_souris(bilan, hwnd):
    titre_etape(4, "La souris : bouger et cliquer toute seule")
    if not WINDOWS:
        bilan.noter("S", "Souris", "NON DISPONIBLE", "pas Windows")
        return
    ecrire("   Le robot va faire un carre avec la souris" + (", puis cliquer DANS son fichier" if hwnd else "")
           + (" et y ecrire une ligne." if hwnd else "."))
    ecrire("   NE TOUCHEZ PAS A LA SOURIS pendant 5 secondes. Il la remet ensuite ou elle etait.")
    pause()
    depart = position_souris()
    clic = "pas de Bloc-notes : mouvement seulement"
    try:
        if hwnd and not _ramener(hwnd):
            hwnd = None
        if hwnd:
            r = wintypes.RECT()
            user32.GetWindowRect(hwnd, ctypes.byref(r))
            cx, cy = (r.left + r.right) // 2, (r.top + r.bottom) // 2
        else:
            cx, cy = position_souris()
        gauche, haut, droite, bas = zone_de_travail(cx, cy)
        cx = min(max(cx, gauche + 130), droite - 130)
        cy = min(max(cy, haut + 90), bas - 90)
        carre = [(cx - 120, cy - 80), (cx + 120, cy - 80), (cx + 120, cy + 80), (cx - 120, cy + 80), (cx, cy)]
        bouge = True
        for x, y in carre:
            deplacer_doucement(x, y)
            px, py = position_souris()
            bouge = bouge and abs(px - x) <= 2 and abs(py - y) <= 2
        if hwnd:
            # on ne clique QUE si la souris est bien là où le robot l'a mise, sur NOTRE fichier, au premier plan
            px, py = position_souris()
            dessous = user32.WindowFromPoint(wintypes.POINT(px, py))
            sous = user32.GetAncestor(dessous, GA_ROOT) if dessous else None
            r2 = wintypes.RECT()
            user32.GetWindowRect(hwnd, ctypes.byref(r2))
            dans_notre_fenetre = r2.left <= px <= r2.right and r2.top <= py <= r2.bottom
            # Windows 11 : la zone de texte du Bloc-notes est une fenêtre interne d'une autre classe
            # (Windows.UI...) ; elle est à nous si elle est dans notre rectangle et que nous sommes devant
            classe_sous = _classe(sous) if sous else ""
            notre = sous == hwnd or (sous and est_notre_fichier(sous)) or (
                dans_notre_fenetre and fenetre_active() == hwnd and
                classe_sous.startswith(("Windows.UI", "ApplicationFrame", "Microsoft.UI")))
            if not bouge or abs(px - cx) > 3 or abs(py - cy) > 3:
                clic = "clic annule par prudence : la souris n'est pas la ou le robot l'a mise"
            elif not (notre and fenetre_active() == hwnd and est_notre_fichier(hwnd)):
                clic = ("clic annule par prudence : sous la souris, " + (classe_sous or "rien")
                        + ("" if fenetre_active() == hwnd else " ; le Bloc-notes n'est plus devant"))
            else:
                cliquer_ici()
                time.sleep(0.4)
                taper("\nLe robot a aussi bouge la souris et clique ici.\n", hwnd)
                clic = "clic fait dans le Bloc-notes"
        deplacer_doucement(*depart, duree=0.3)
    except OSError as e:
        bilan.noter("S", "Souris", "ECHEC", str(e))
        revenez_ici()
        return
    revenez_ici()
    if not bouge:
        bilan.noter("S", "Souris", "ECHEC", "la souris n'est pas allee ou le robot voulait")
    elif oui("   Avez-vous vu la souris bouger toute seule" + (" et la ligne s'ecrire apres le clic ?" if hwnd else " ?")):
        bilan.noter("S", "Souris", "OK", clic)
    else:
        bilan.noter("S", "Souris", "PAS VU", clic)


def adresse_prudente(adresse):
    """Une adresse copiée juste après un clic peut être celle d'une ACTION (…/Plans/Suppression?id=5,
    …#/plans/5/supprimer) : l'ouvrir la referait. Dans ce cas, on ne garde que l'accueil du portail."""
    morceaux = urlsplit(adresse)
    reste = mots(unquote(" ".join((morceaux.path, morceaux.query, morceaux.fragment))))
    if VERBES_ACTION & set(reste.split()):
        return f"{morceaux.scheme}://{morceaux.netloc}/", True
    return adresse, False


def demander_adresse():
    ecrire()
    ecrire("=" * 70)
    ecrire(" L'ADRESSE DE VOTRE PORTAIL")
    ecrire("=" * 70)
    ecrire("   Dans votre Chrome, allez sur la page d'ACCUEIL du portail. Cliquez dans la barre")
    ecrire("   d'adresse tout en haut, faites Ctrl+C pour la copier. Revenez dans cette fenetre noire")
    ecrire("   et faites Ctrl+V pour la coller (ou clic droit, puis Coller), puis Entree.")
    precedente = adresse_retenue()
    if precedente:
        ecrire("   Entree sans rien = la meme adresse que la derniere fois.")
    else:
        ecrire("   Entree sans rien = sauter les etapes du portail.")
    ecrire("   (L'adresse reste sur ce poste, dans Bureau > Robot ; elle n'est jamais dans le RESULTAT.)")
    adresse = demander("   Adresse du portail :").strip().strip('"') or precedente
    if not adresse:
        return ""
    if adresse.count("://") > 1 or " " in adresse:
        ecrire("   Ce n'est pas une adresse (collee deux fois ?). Recommencez : Ctrl+V une seule fois.")
        return demander_adresse()
    if not re.match(r"^[a-z][a-z0-9+.-]*://", adresse, re.I):
        if "." not in adresse:
            ecrire("   Ce n'est pas une adresse. Copiez-la dans la barre d'adresse de Chrome (Ctrl+C), puis Ctrl+V ici.")
            return demander_adresse()
        adresse = "https://" + adresse
    try:
        valable = bool(urlsplit(adresse).hostname)
        adresse, coupee = adresse_prudente(adresse)
    except ValueError:
        valable = False
    if not valable:
        ecrire("   Ce n'est pas une adresse de site : les etapes du portail sont sautees.")
        return ""
    if coupee:
        ecrire("   Cette adresse ressemble a celle d'une ACTION (supprimer, valider...). Par prudence,")
        ecrire("   le robot ouvrira seulement l'ACCUEIL du portail.")
    retenir_adresse(adresse)
    return adresse


def etape_navigateur_habituel(bilan, url):
    titre_etape(5, "Votre navigateur habituel")
    if not url:
        bilan.noter("W", "Portail dans votre navigateur", "SAUTE", "pas d'adresse donnee")
        return
    ecrire("   Le robot va ouvrir le portail dans VOTRE navigateur habituel (comme un clic sur un lien).")
    ecrire("   Ne tapez rien dans la page du portail : regardez seulement si elle s'ouvre.")
    pause()
    try:
        ouvert = webbrowser.open(url)
    except Exception as e:
        bilan.noter("W", "Portail dans votre navigateur", "ECHEC", masquer(premiere_ligne(e), url))
        return
    time.sleep(2.5)
    revenez_ici()
    if oui("   La page du portail s'est-elle ouverte dans votre navigateur ?"):
        bilan.noter("W", "Portail dans votre navigateur", "OK")
    else:
        bilan.noter("W", "Portail dans votre navigateur", "PAS VU" if ouvert else "ECHEC",
                    "" if ouvert else "Windows n'a pas trouve de navigateur par defaut")
    ecrire("   Vous pouvez fermer cet onglet du portail dans votre Chrome.")


# ---------------------------------------------------------------------------- étapes 6 à 8 : navigateur piloté
def _python_avec_playwright(chemin):
    try:
        return subprocess.call([str(chemin), "-c", "import playwright.sync_api"],
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60) == 0
    except Exception:
        return False


def trouver_python_playwright():
    """Le Python qui a le pilote de navigateur : celui-ci, ou celui du dossier robot-web déjà installé."""
    try:
        import playwright.sync_api  # noqa: F401
        return sys.executable
    except ImportError:
        pass
    ici = Path(__file__).resolve().parent
    bases = [ici, ici.parent, Path.home() / "Downloads", Path.home() / "Telechargements", Path.home() / "Téléchargements",
             dossier_bureau(), Path.home() / "Documents", Path.home()]
    for base in bases:
        for dossier in (base, base / "robot-web"):
            for morceaux in (("venv", "Scripts", "python.exe"), ("venv", "bin", "python")):
                candidat = dossier.joinpath(*morceaux)
                if candidat.is_file() and _python_avec_playwright(candidat):
                    return str(candidat)
    return None


RAISON_PILOTE = ["pilote de navigateur absent"]


def obtenir_python_playwright():
    ecrire("   Recherche du pilote de navigateur sur ce poste...")
    python = trouver_python_playwright()
    if python:
        return python
    ecrire("   Le pilote de navigateur (Playwright) n'est pas trouve sur ce poste.")
    ecrire("   Si vous avez un dossier robot-web : faites-le GLISSER dans cette fenetre noire,")
    ecrire("   puis Entree. Sinon, Entree sans rien.")
    chemin = demander("   Dossier robot-web :").strip().strip('"')
    if chemin:
        for morceaux in (("venv", "Scripts", "python.exe"), ("venv", "bin", "python")):
            candidat = Path(chemin).joinpath(*morceaux)
            if candidat.is_file() and _python_avec_playwright(candidat):
                return str(candidat)
        ecrire("   Pas de pilote de navigateur dans ce dossier.")
    commande = [sys.executable, "-m", "pip", "install", "--disable-pip-version-check", "--no-warn-script-location"]
    if sys.prefix == sys.base_prefix:  # pas dans un environnement à part : installation dans votre compte
        commande.append("--user")
    commande.append("playwright")
    ecrire("   Le robot peut l'installer depuis Internet (2 a 5 minutes, dans votre compte Windows,")
    ecrire("   sans droits d'administrateur) avec la commande :")
    ecrire("      python " + " ".join(commande[1:]) + "   (rien a taper : le robot la lance lui-meme)")
    if not oui("   Tapez o pour l'installer, n pour sauter les etapes du navigateur.", "n"):
        return None
    if subprocess.call(commande) != 0 or not _python_avec_playwright(sys.executable):
        ecrire("   L'installation n'a pas abouti.")
        pause("Prenez en photo les lignes ci-dessus, puis Entree...")
        RAISON_PILOTE[0] = "installation du pilote echouee"
        return None
    return sys.executable


def etape_navigateur_robot(bilan, url):
    titre_etape(6, "Le navigateur du robot")
    ecrire("   Le robot va ouvrir SON navigateur (une fenetre a part, vide, qui ne touche pas a")
    ecrire("   votre Chrome ni a vos mots de passe), puis remplir et cliquer une page d'essai.")
    pause()
    python = obtenir_python_playwright()
    if not python:
        for lettre, intitule in (("N", "Navigateur du robot (page d'essai)"), ("R", "Navigateur du robot sur le portail"),
                                 ("C", "Clic choisi sur le portail")):
            bilan.noter(lettre, intitule, "ECHEC" if "echouee" in RAISON_PILOTE[0] else "SAUTE", RAISON_PILOTE[0],
                        afficher=False)
        ecrire(f"   --> Etapes du navigateur sautees : {RAISON_PILOTE[0]}.")
        return
    echange = Path(tempfile.mkdtemp(prefix="robot_essai_"))
    try:
        (echange / "adresse.txt").write_text(url or "", encoding="utf-8")  # jamais sur la ligne de commande
        interrompu = None
        try:
            enfant = subprocess.Popen([python, str(Path(__file__).resolve()), "--navigateur", str(echange)])
            try:
                enfant.wait()
            except KeyboardInterrupt as e:  # Ctrl+C : on laisse l'enfant noter où il en était
                interrompu = e
                try:
                    enfant.wait(timeout=5)
                except Exception:
                    enfant.kill()
        except Exception as e:
            bilan.noter("N", "Navigateur du robot (page d'essai)", "ECHEC", expliquer(e), afficher=False)
            return
        try:
            notes = json.loads((echange / "resultat.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            notes = [["N", "Navigateur du robot (page d'essai)", "ECHEC", "le programme du navigateur s'est arrete"]]
        for lettre, intitule, statut, detail in notes:
            if lettre in ("N", "D", "R", "C", "V"):
                bilan.noter(lettre, intitule, statut, masquer(detail, url), afficher=False)
        if interrompu is not None:
            raise interrompu
    finally:
        effacer_dossier(echange)


PAGE_ESSAI = """<!doctype html><html lang=fr><meta charset=utf-8><title>Page d'essai du robot</title>
<style>body{font:18px Arial,sans-serif;margin:40px;background:#f4f6fb}label,button,input{font-size:20px}
input{padding:6px;margin:0 10px}button{padding:8px 18px;cursor:pointer}#reponse{margin-top:30px;font-size:26px;color:#15803d}</style>
<h1>Page d'essai du robot</h1>
<p>Le robot va écrire son nom dans le champ, puis cliquer sur « Valider ».</p>
<label for=prenom>Votre prénom</label><input id=prenom autocomplete=off><button id=valider>Valider</button>
<div id=reponse></div>
<script>document.getElementById('valider').onclick = () => {
  document.getElementById('reponse').textContent = 'Bonjour ' + document.getElementById('prenom').value + ', le clic a marché !';
};</script></html>"""

# Tout ce qui est CLIQUABLE, de toutes les façons que le robot connaît. « avant » : liens, boutons, champs, rôles,
# main de la souris ; « chrome » : élément qui a un écouteur de clic (signalé par Chrome, data-robot-ecoute) ;
# « attribut » : attributs des applications (Angular, AngularJS, jQuery...) ; « appris » : montré par vous (choix 4).
# Composants web (shadow DOM) compris. Les grands blocs (plus de 30 % de l'écran) ne sont pas des éléments.
JS_DETECTION = r"""
  const SEL_AVANT = 'a[href], button, input:not([type=hidden]), select, textarea, summary, area[href], label[for], ' +
    '[role=button], [role=link], [role=menuitem], [role=menuitemcheckbox], [role=menuitemradio], [role=tab], [role=treeitem], ' +
    '[role=option], [role=checkbox], [role=radio], [role=switch], [onclick], [tabindex]:not([tabindex="-1"])';
  const SEL_ATTRIBUT = '[onmousedown], [onmouseup], [onpointerdown], [ng-click], [data-ng-click], [ui-sref], [routerlink], ' +
    '[ng-reflect-router-link], [jsaction], [data-toggle], [data-bs-toggle]';
  const visDet = e => { const r = e.getBoundingClientRect(); if (r.width < 4 || r.height < 4) return false;
                        const s = getComputedStyle(e); return s.display !== 'none' && s.visibility !== 'hidden' && s.opacity !== '0'; };
  const pointeurDet = e => { const s = getComputedStyle(e).cursor;
                             return s === 'pointer' && !(e.parentElement && getComputedStyle(e.parentElement).cursor === 'pointer'); };
  const detecter = () => {
    const trouves = new Map();
    const ajouter = (e, s) => { if (!trouves.has(e)) trouves.set(e, s); };
    const appris = Array.isArray(window.__robotAppris) ? window.__robotAppris : [];
    const fouiller = racine => {
      for (const e of racine.querySelectorAll(SEL_AVANT)) ajouter(e, 'avant');
      for (const e of racine.querySelectorAll('div, span, li, td, th, i, img, svg, p, label, mat-icon, b, strong, em, h1, h2, h3, h4, h5, h6'))
        if (!trouves.has(e) && pointeurDet(e)) ajouter(e, 'avant');
      for (const e of racine.querySelectorAll('[data-robot-ecoute]')) ajouter(e, 'chrome');
      for (const e of racine.querySelectorAll(SEL_ATTRIBUT)) ajouter(e, 'attribut');
      for (const sel of appris) { try { for (const e of racine.querySelectorAll(sel)) ajouter(e, 'appris'); } catch (x) {} }
      for (const e of racine.querySelectorAll('*')) if (e.shadowRoot) fouiller(e.shadowRoot);
    };
    fouiller(document);
    // une liste dont le clic est géré par la liste entière (ul, tbody...) : ce sont ses lignes qui sont cliquables
    for (const [e, s] of Array.from(trouves)) {
      if (s !== 'chrome') continue;
      let c = e;
      if (c.children.length === 1 && c.firstElementChild.children.length >= 2) c = c.firstElementChild;   // ul > li, table > tbody > tr
      const enfants = Array.from(c.children).filter(k => visDet(k) && (k.textContent || '').trim());
      const estListe = /^(UL|OL|MENU|DL|TBODY|TABLE)$/.test(c.tagName) || /^(list|listbox|menu|menubar|grid|tree|tablist)$/.test(c.getAttribute('role') || '') ||
                    (enfants.length && /^(LI|TR)$/.test(enfants[0].tagName)) ||
                    (enfants.length >= 3 && /^(DIV|A)$/.test(enfants[0].tagName) &&   // (pas « libellé + flèche » d'une entrée,
                     c.getBoundingClientRect().width * c.getBoundingClientRect().height <= innerWidth * innerHeight * 0.3);  // ni la page)
      if (estListe && enfants.length >= 2 && enfants.length <= 200 && enfants.every(k => k.tagName === enfants[0].tagName) &&
          enfants.length >= 0.8 * c.children.length) {
        trouves.delete(e);
        for (const k of enfants) ajouter(k, 'chrome');
      }
    }
    const grand = innerWidth * innerHeight * 0.3, liste = [];
    for (const [e, s] of trouves) {
      if (e === document.body || e === document.documentElement || (e.closest && e.closest('#robot-cadres')) || !visDet(e)) continue;
      const r = e.getBoundingClientRect();
      if (r.width * r.height <= grand) liste.push({ e: e, s: s, r: r });
    }
    // le même élément vu deux fois (lien > texte, bouton > icône, même cadre) : on garde l'extérieur
    const ens = new Set(liste.map(x => x.e));
    const meme = (a, b) => Math.abs(a.left - b.left) <= 4 && Math.abs(a.top - b.top) <= 4 &&
                           Math.abs(a.right - b.right) <= 4 && Math.abs(a.bottom - b.bottom) <= 4;
    return liste.filter(x => { for (let p = x.e.parentElement, n = 0; p && n < 4; p = p.parentElement, n++)
                                 if (ens.has(p) && meme(p.getBoundingClientRect(), x.r)) return false;
                               return true; });
  };
"""

JS_STRUCTURE = r"""() => {""" + JS_DETECTION + r"""
  const vis = e => { const s = getComputedStyle(e); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const q = s => Array.from(document.querySelectorAll(s));
  const liste = detecter();
  let ombres = 0;
  for (const e of document.querySelectorAll('*')) if (e.shadowRoot) ombres++;
  return { cliquables: liste.length, nouveaux: liste.filter(x => x.s !== 'avant').length,
           champs: q('input:not([type=hidden]), select, textarea').filter(vis).length,
           frameset: !!document.querySelector('frameset'), ombres: ombres,
           dessins: q('canvas, embed, object').filter(vis).length };
}"""

JS_MARQUER_CLIQUABLES = """() => Array.from(document.querySelectorAll('a[href], button, input[type=submit], input[type=button], [role=button], [role=menuitem], [role=tab], [onclick], li, span, div'))
  .filter(e => { const r = e.getBoundingClientRect(); if (!(r.width > 4 && r.height > 4 && r.top >= 0 && r.top < innerHeight)) return false;
                 if (['LI', 'SPAN', 'DIV'].includes(e.tagName) && !e.matches('[onclick], [role=button], [role=menuitem], [role=tab]'))
                   return getComputedStyle(e).cursor === 'pointer' && (e.textContent || '').length < 200 &&
                     !(e.parentElement && getComputedStyle(e.parentElement).cursor === 'pointer');
                 return true; })
  .slice(0, 5).map(e => { e.setAttribute('data-robot-essai', '1'); return 1; }).length"""

JS_ENCADRER = """(e, ms) => { e.scrollIntoView({block: 'nearest', inline: 'nearest'});
  if (e.dataset.robotContour === undefined) e.dataset.robotContour = e.style.outline || '';
  e.style.outline = '4px solid #dc2626'; e.style.outlineOffset = '2px';
  if (ms > 0) setTimeout(() => { e.style.outline = e.dataset.robotContour || ''; e.style.outlineOffset = '';
                                 delete e.dataset.robotContour; }, ms); }"""

# Après les essais : le robot retire ses marques et ses cadres de la page.
JS_NETTOYER = """(tout) => { for (const e of document.querySelectorAll('[data-robot-cible], [data-robot-texte], [data-robot-essai], [data-robot-contour], [data-robot-trouve], [data-robot-barre], [data-robot-fort], [data-robot-nouveau], [data-robot-mot], [data-robot-survole]')) {
  if (e.dataset.robotContour !== undefined) { e.style.outline = e.dataset.robotContour; e.style.outlineOffset = ''; }
  e.removeAttribute('data-robot-contour');
  if (tout) for (const a of ['data-robot-cible', 'data-robot-texte', 'data-robot-essai', 'data-robot-trouve', 'data-robot-barre',
                             'data-robot-fort', 'data-robot-nouveau', 'data-robot-mot', 'data-robot-survole']) e.removeAttribute(a);
} }"""

# Les éléments dont le PROPRE texte (sans celui de leurs sous-menus) est ce mot, majuscules ignorées.
JS_TROUVER_PROPRE_TEXTE = r"""([mot, jeton]) => {
  const norme = t => (t || '').replace(/\s+/g, ' ').trim().toLowerCase();
  const cherche = norme(mot);
  for (const e of document.querySelectorAll('body *')) {
    let propre = '';
    for (const n of e.childNodes) if (n.nodeType === 3) propre += n.textContent;
    if (norme(propre) === cherche) e.setAttribute('data-robot-trouve', jeton);
  }
}"""

# Recherche tolérante : accents, majuscules, flèches et symboles (« GATES ▾ », « » Historique »), bulle d'aide,
# texte de remplacement d'une image (menus en images). mode 'propre' = propre texte de l'élément ;
# 'entier' = tout son texte s'il est court ; 'image' = alt / title d'une image ou de son lien.
JS_TROUVER_TOLERANT = r"""([mot, jeton, mode]) => {
  const norme = t => (t || '').normalize('NFD').replace(/[\u0300-\u036f]/g, '').toLowerCase()
                     .replace(/[^a-z0-9]+/g, ' ').trim();
  const cherche = norme(mot);
  if (!cherche) return 0;
  let n = 0;
  if (mode === 'image') {
    for (const img of document.querySelectorAll('img, input[type=image], svg, [class*=icon i]')) {
      const t = [img.getAttribute('alt'), img.getAttribute('title'), img.getAttribute('aria-label')].map(norme);
      const a = img.closest('a, button, [onclick], [role=button], [role=menuitem]');
      const t2 = a ? [a.getAttribute('title'), a.getAttribute('aria-label')].map(norme) : [];
      if (t.includes(cherche) || t2.includes(cherche)) { (a || img).setAttribute('data-robot-trouve', jeton); n++; }
    }
    return n;
  }
  for (const e of document.querySelectorAll('body *')) {
    if (['SCRIPT', 'STYLE', 'OPTION'].includes(e.tagName)) continue;
    let t;
    if (mode === 'propre') { t = ''; for (const c of e.childNodes) if (c.nodeType === 3) t += c.textContent; t = norme(t); }
    else { if ((e.textContent || '').length > 60) continue; t = norme(e.textContent); }
    if (t === cherche) { e.setAttribute('data-robot-trouve', jeton); n++; }
  }
  return n;
}"""

# Zone de MENU d'un élément (par mots ENTIERS du nom de ses parents : « sidenav », « menu »... et jamais « inside »).
# La recherche s'arrête au contenu de la page : « mat-sidenav-content », barre d'actions, fenêtre, formulaire, main.
JS_ZONE_MENU = r"""
  const jetonsDe = x => (x.tagName + ' ' + (x.id || '') + ' ' + (typeof x.className === 'string' ? x.className : '') + ' ' +
                         (x.getAttribute('role') || '')).replace(/([a-z0-9])([A-Z])/g, '$1 $2').toLowerCase()
                         .split(/[^a-z0-9]+/).filter(Boolean);
  const MOTS_MENU = ['menu', 'menus', 'menubar', 'menuitem', 'submenu', 'menulist', 'nav', 'navbar', 'navigation', 'sidenav',
                     'sidebar', 'sidemenu', 'leftmenu', 'mainmenu', 'megamenu', 'topnav', 'mainnav', 'subnav', 'leftnav',
                     'navmenu', 'navlist', 'navitem', 'flyout', 'dropdown', 'drawer', 'rail'];
  const MOTS_ARRET = ['content', 'contents', 'actions', 'action', 'toolbar', 'buttons', 'dialog', 'alertdialog', 'modal'];
  const zoneMenu = n => {
    for (let x = n; x && x !== document.body && x !== document.documentElement; x = x.parentElement) {
      const t = jetonsDe(x);
      if (/^(MAIN|FORM|DIALOG)$/.test(x.tagName) || x.getAttribute('role') === 'main' || t.some(m => MOTS_ARRET.includes(m))) return false;
      if (/^(NAV|ASIDE)$/.test(x.tagName) || t.some(m => MOTS_MENU.includes(m))) return true;
    }
    return false;
  };
  // dans le CONTENU de la page (main, formulaire, fenêtre, « ...-content », barre d'actions) : jamais un menu certain
  const contenu = n => { for (let y = n; y && y !== document.body && y !== document.documentElement; y = y.parentElement)
                           if (/^(MAIN|FORM|DIALOG)$/.test(y.tagName) || y.getAttribute('role') === 'main' ||
                               jetonsDe(y).some(m => MOTS_ARRET.includes(m))) return true;
                         return false; };
"""

# Outils communs pour les MENUS : visibilité, nom d'une entrée (sans ses icônes ni ses flèches), éléments cliquables
# (liens, boutons, entrées de menu, et les « div » ou « span » avec la main de la souris), lien vers un enregistrement.
JS_OUTILS_MENUS = JS_ZONE_MENU + r"""
  const vis = e => { const s = getComputedStyle(e); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     if (e.checkVisibility && !e.checkVisibility({opacityProperty: true})) return false;
                     const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const classe = x => typeof x.className === 'string' ? x.className : ((x.className && x.className.baseVal) || '');
  const SOUS = 'ul, ol, [role=menu], [role=group], [role=listbox], table, select';
  const ICONE = /(^|[\s_-])(ico|icon|icons|icone|material-icons|material-symbols\S*|mat-icon|fa|fas|far|fal|fab|glyphicon|bi|mdi|chevron|arrow|caret|angle|fleche)([\s_-]|$)/i;
  const motIcone = t => !t || /^[a-z0-9_]+$/.test(t) || /^[^A-Za-zÀ-ÿ0-9]{1,3}$/.test(t);
  // une pastille de compteur (« 3 » nouveaux messages) n'est pas un nom : seulement le compteur lui-même (des chiffres),
  // jamais l'élément qui porte le texte (« <span matBadge=4>Plans à valider</span> »)
  const PASTILLE = /badge|pastille|counter|compteur|bubble|count|scroll-number/i;
  const compteur = x => (/^\s*\d{1,4}\+?\s*$/.test(x.textContent || '') &&
                         (x.tagName === 'SUP' || PASTILLE.test(classe(x)) || (x.parentElement && PASTILLE.test(classe(x.parentElement))))) ||
                        // (une étiquette d'un mot, « Nouveau », « Beta », dans une pastille)
                        (!x.children.length && /(^|[\s_-])(badge|pastille|tag|chip-new|label-new)([\s_-]|$)/i.test(classe(x)) &&
                         /^\s*[A-Za-zÀ-ÿ]{2,10}\s*$/.test(x.textContent || ''));
  const texteDuClone = c => {
    c.querySelectorAll(SOUS + ', svg, img, script, style, input, textarea').forEach(x => x.remove());
    for (const x of Array.from(c.querySelectorAll('*'))) if (c.contains(x) && compteur(x)) x.remove();
    for (const x of Array.from(c.querySelectorAll('*')))
      if (!x.children.length && (['I', 'MAT-ICON'].includes(x.tagName) || ICONE.test(classe(x))) && motIcone((x.textContent || '').trim()))
        x.remove();
    return (c.textContent || '').replace(/\s+/g, ' ').trim().replace(/^[\s›»<>▸▶►▾▼⌄+•·|-]+|[\s›»<>▸▶►▾▼⌄+•·|-]+$/g, '').trim();
  };
  const texteDe = e => {
    let t = texteDuClone(e.cloneNode(true)), base = e, elague = false;
    // une entrée de MENU dont le sous-menu FERMÉ est dans la page (caché) : son nom, c'est seulement ce qui se voit.
    // Jamais dans le contenu de la page (une tuile, une fiche dont le détail est caché n'est pas une entrée de menu),
    // et seulement si la partie cachée contient des éléments cliquables (un sous-menu).
    // (aussi une entrée courte dont le sous-menu caché est dedans : une icône seule, ou un nom suivi de son sous-menu)
    const cache = Array.from(e.querySelectorAll('a, button, [role=menuitem], [onclick]')).slice(0, 20).some(x => !vis(x));
    if (cache || (t.length > 30 && !contenu(e) && (e.innerText || '').replace(/\s+/g, ' ').trim().length < t.length)) {
      const c = e.cloneNode(true), o = Array.from(e.querySelectorAll('*')), k = Array.from(c.querySelectorAll('*'));
      if (o.length <= 400) {
        let sousMenu = false;
        o.forEach((x, i) => { const s = getComputedStyle(x);
          if (s.display === 'none' || s.visibility === 'hidden') {
            if (x.matches('a, button, [role=menuitem], [onclick]') || x.querySelector('a, button, [role=menuitem], [onclick]')) sousMenu = true;
            k[i].remove(); } });
        const visible = c.cloneNode(true), vu = texteDuClone(c);
        // (vide pour une icône seule : elle sera nommée par sa bulle ou son icône, jamais par ce qui est caché)
        if (sousMenu) { t = vu; base = visible; elague = true; }
      }
    }
    if (!t) t = (e.getAttribute('aria-label') || e.getAttribute('title') || '').replace(/\s+/g, ' ').trim();
    if (!t) { const i = Array.from(base.querySelectorAll('img[alt], [title], [aria-label]'))
                              .find(x => { const s = x.parentElement && x.parentElement.closest(SOUS);
                                           return (elague || vis(x)) && !(s && base.contains(s) && s !== base); });
              if (i) t = (i.getAttribute('alt') || i.getAttribute('title') || i.getAttribute('aria-label') || '').replace(/\s+/g, ' ').trim(); }
    return t;
  };
  // le nom d'une entrée faite d'une ICÔNE seule (barre d'icônes à gauche) : sa bulle, sinon le nom de son icône
  // (« icone dashboard » : le mot de l'icône Material, la classe « fa-folder », le symbole SVG, le composant d'icône).
  // Jamais le nom d'un fichier image (photo d'une personne, logo d'un client) ; une bulle qui ressemble à une donnée
  // (chiffre, « Nom Prénom », longue phrase) n'est pas un nom.
  // (jamais « alt » : le texte d'une image est souvent un logo ou une photo)
  const BULLES = ['aria-label', 'title', 'data-tooltip', 'mattooltip', 'ng-reflect-message', 'data-original-title',
                  'data-bs-original-title', 'tooltip', 'ptooltip', 'data-title', 'data-tip'];
  // une bulle qui peut être un NOM DE PERSONNE n'est pas un nom d'entrée : « Jean-Pierre Dupont », « J. Dupont »,
  // « Agnieszka K. », « nathan martin », « dupont.j », « JD »
  const MAJ = x => /^[A-ZÀ-Ý][a-zà-ÿ']*(-[A-ZÀ-Ý][a-zà-ÿ']*)*$/.test(x) || /^[A-ZÀ-Ý'-]{2,}$/.test(x);
  const HOTE = (location.hostname || '').toLowerCase().split(/[.-]/).filter(x => x.length >= 4 &&
                 !/^(www|intranet|portail|portal|apps?|plans?|test|prod|recette|local|localhost|corp|com|net|org|gouv)$/.test(x));
  const deLHote = t => HOTE.some(x => t.toLowerCase().replace(/[^a-z0-9]/g, '').includes(x));
  const bulleSure = v => { const t = (v || '').replace(/\s+/g, ' ').trim(), m = t.split(' ');
    if (t.length < 2 || t.length > 30 || /[:\d@_]/.test(t) || m.length > 4 || !/[A-Za-zÀ-ÿ]/.test(t) || deLHote(t)) return '';
    // (« J. Dupont », « Dupont J », « dupont.j », « Jan Kowalski », « nathan martin » ; un sigle « P.L.U. », « SIG »
    // reste un nom ; « Mes Plans » est aussi refusée : l'entrée sera nommée par son icône)
    if (m.length >= 2 && m.some(x => /^[A-ZÀ-Ý]\.?$/.test(x))) return '';
    if (/^[a-z]+\.[a-z]+$/.test(t)) return '';
    if (m.length >= 2 && (m.every(x => /^[a-zà-ÿ'-]+$/.test(x)) || m.some((x, i) => i && MAJ(x) && MAJ(m[i - 1])))) return '';
    return t; };
  // un nom d'icône qui serait celui de l'entreprise ou d'une personne (logo, marque, photo) n'est jamais un nom
  const ICONE_PERSO = /logo|brand|marque|avatar|photo|portrait|profil|picture/i;
  // les noms d'icônes d'un registre PROPRE à l'application (svgIcon, data-icon, symbole « #... », classe « icon-... ») :
  // seulement des mots courants d'icônes (sinon, le préfixe de l'entreprise pourrait s'écrire)
  const MOTS_ICONE = new Set(('home house dashboard tableau bord accueil folder folders dossier dossiers file files fichier ' +
    'document documents description doc docs settings setting parametres reglages gear cog cogs user users person people ' +
    'group groups team equipe chart charts graph bar line pie map maps carte cartes location pin place search recherche list ' +
    'liste grid table calendar calendrier event events clock time history historique book books library archive archives ' +
    'inbox mail envelope message messages chat bell notification notifications star heart flag tag tags label bookmark help ' +
    'aide question info information warning error alert shield lock key link print download upload cloud sync refresh edit ' +
    'pencil plus add minus delete trash close check cross arrow arrows chevron caret left right up down back forward next ' +
    'previous menu bars apps app view eye filter sort task tasks todo assignment work briefcase building city factory tools ' +
    'tool wrench hammer box package truck car cart money euro wallet card bank report reports rapport rapports analytics ' +
    'stats statistiques tree network sitemap layers layer stack database server code bug globe world language image images ' +
    'photo camera video play phone mobile laptop desktop monitor plan plans contract contrat contrats client clients ' +
    'project projet projets admin administration config configuration module modules nav icon ico svg sprite outline ' +
    'filled fill round sharp solid o alt circle square small large').split(' '));
  const registreSur = v => !ICONE_PERSO.test(v) && !deLHote(v) && v.toLowerCase().split(/[-_]+/).every(w => !w || MOTS_ICONE.has(w) || /^\d+$/.test(w));
  // les classes de STYLE d'une police d'icônes (fa-solid, fa-fw, pi-fw, mdi-24px...) ne sont pas le nom de l'icône
  const STYLE_ICONE = /^(fa|fas|far|fal|fab|fad|icon|ico|pi|mdi|bi|ti|ri|la|bx|bxs|lucide|feather|k-i|k-svg-i|e|ph)-(solid|regular|light|thin|duotone|sharp|brands|kit|fw|lg|sm|xs|xl|2xs|2xl|[0-9]+x|[0-9]+px|spin|pulse|border|inverse|stack|stack-[0-9]x|flip|flip-\w+|rotate-\w+|beat|beat-fade|bounce|fade|shake|spin-pulse|spin-reverse|li|ul|dark|inactive|icon|icons|size-\w+|fill)$/i;
  const CLASSE_ICONE = /^(?:fa|fas|far|fal|fab|fad|mdi|bi|icon|ico|glyphicon|pi|k-i|k-svg-i|ti|ri|la|lucide|feather|bx|bxs|bxl|octicon|ion|ph|e)-([a-z][\w-]{1,30})$|^ms-Icon--([A-Za-z]\w{1,30})$/i;
  const COMPTE_ICONE = /^icone (person|user|account|account_circle|account_box|manage_accounts|circle-user|user-circle|face|id-badge|avatar|user-tie|portrait|users|people|contacts)([-_](circle|alt|outline|outlined|solid|fill|filled|round|rounded|o))*$/i;
  const iconeSeule = k => cleIcone(k, true);
  const cleIcone = (k, sansBulle) => {
    // (ce qui se VOIT seulement : jamais la bulle d'un lien d'un sous-menu caché dans l'entrée)
    const tous = [k, ...Array.from(k.querySelectorAll('*')).slice(0, 80)
                    .filter(x => !x.checkVisibility || x.checkVisibility({visibilityProperty: true, opacityProperty: true})).slice(0, 39)];
    if (!sansBulle)
      for (const x of tous)
        for (const a of BULLES) { const v = bulleSure(x.getAttribute && x.getAttribute(a)); if (v) return v; }
    for (const x of tous)
      // un composant d'icône qui porte son nom : <mat-icon svgIcon="plans">, <app-icon name="plans">, data-icon="plans"
      for (const a of ['data-mat-icon-name', 'svgicon', 'data-icon', 'icon', 'iconname', 'data-icon-name'].concat(
                         /ICON/.test(x.tagName) ? ['name', 'type', 'key'] : [])) {
        const v = (x.getAttribute && x.getAttribute(a) || '').split(':').pop().trim();
        if (/^[A-Za-z][\w-]{1,30}$/.test(v) && registreSur(v)) return 'icone ' + v;
      }
    for (const x of tous) {
      if (!['MAT-ICON', 'I', 'SPAN', 'SVG', 'USE', 'EM'].includes(x.tagName.toUpperCase()) && x !== k) continue;
      if (x.tagName.toLowerCase() === 'use') { const m = (x.getAttribute('href') || x.getAttribute('xlink:href') || '').match(/#([A-Za-z][\w-]{1,30})$/);
                                              const v = m ? m[1].replace(/^(icon|ico|i|svg|sprite)[-_]/i, '') : '';
                                              if (v && registreSur(v)) return 'icone ' + v; }
      // la classe d'icône (la DERNIÈRE qui n'est pas une classe de style : « fa-solid fa-house » -> house) ; une classe
      // « icon-... » ou « ico-... » est propre à l'application : mots courants seulement
      const c = classe(x).split(/\s+/).filter(c => CLASSE_ICONE.test(c) && !STYLE_ICONE.test(c) &&
                                                 (!/^e-/i.test(c) || /(^|\s)e-icons(\s|$)/.test(classe(x)))).pop();
      if (c) { const m = c.match(CLASSE_ICONE), v = (m[1] || m[2]).toLowerCase();
               if (!ICONE_PERSO.test(v) && (!/^(icon|ico)-/i.test(c) || registreSur(v))) return 'icone ' + v; }
      let t = ''; for (const n of x.childNodes) if (n.nodeType === 3) t += n.textContent;
      t = t.trim();
      if (['MAT-ICON', 'I', 'SPAN'].includes(x.tagName) && /^[a-z][a-z_]{2,30}$/.test(t) &&
          (x.tagName !== 'SPAN' || ICONE.test(classe(x)))) return 'icone ' + t;   // mot d'une icône Material
    }
    return '';
  };
  const CLIQ = 'a, button, [role=menuitem], [role=menuitemcheckbox], [role=menuitemradio], [role=option], [role=treeitem], ' +
               '[role=link], [role=button], [role=tab], [onclick], li, [tabindex], input[type=submit], input[type=button], ' +
               '[data-robot-ecoute], [ng-click], [data-ng-click], [ui-sref], [routerlink], [ng-reflect-router-link], [jsaction]';
  const pointeur = e => getComputedStyle(e).cursor === 'pointer' && !(e.parentElement && getComputedStyle(e.parentElement).cursor === 'pointer');
  const cliquables = () => {
    const s = new Set(document.querySelectorAll(CLIQ));
    for (const sel of (Array.isArray(window.__robotAppris) ? window.__robotAppris : []))
      try { document.querySelectorAll(sel).forEach(e => s.add(e)); } catch (x) {}
    for (const e of document.querySelectorAll('div, span, td')) if (!s.has(e) && pointeur(e)) s.add(e);
    return Array.from(s).filter(e => (e.textContent || '').length <= 300 && vis(e))
      .sort((a, b) => a.compareDocumentPosition(b) & Node.DOCUMENT_POSITION_FOLLOWING ? -1 : 1);
  };
  const enregistrement = e => { const a = e.closest('a[href]') || e.querySelector('a[href]'); const h = (a && a.getAttribute('href')) || '';
    return /^(tel|mailto|callto|sip):/i.test(h) || /[?&][^=&]+=([^&]*\d{2,}|[A-Za-z]+-\d+)/.test(h) || /\/\d{2,}(\/|$|\?)/.test(h) ||
           /[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}/i.test(h); };
  // le titre de section juste au-dessus d'un élément (texte court non cliquable, ou titre h1-h6 sans ses icônes)
  // (un élément cliquable ou une autre entrée du menu n'est pas un titre ; un titre peut avoir une icône d'action à côté)
  const ACTION = 'a[href], button, [role=menuitem], [role=button], [role=link], [onclick]';
  const titreSection = x => {
    for (let y = x, n = 0; y && y !== document.body && n < 3; y = y.parentElement, n++)
      for (let p = y.previousElementSibling, k = 0; p && k < 40; p = p.previousElementSibling, k++) {
        if (p.matches(ACTION) || getComputedStyle(p).cursor === 'pointer') continue;
        const c = p.cloneNode(true); c.querySelectorAll(ACTION + ', svg, img, i, mat-icon').forEach(z => z.remove());
        const t = (c.textContent || '').replace(/\s+/g, ' ').trim();
        if (t && t.length <= 40) return t;
      }
    return '';
  };
  // la zone de la personne connectée (« user », « utilisateur », « compte », « profil », « avatar »...)
  const COMPTE = ['user', 'username', 'utilisateur', 'account', 'compte', 'profil', 'profile', 'avatar', 'login', 'identite'];
  const compteDe = e => { for (let x = e, n = 0; x && x !== document.body && n < 2; x = x.parentElement, n++)
                            if (jetonsDe(x).some(m => COMPTE.includes(m))) return true;
                          return Array.from(e.querySelectorAll('*')).slice(0, 30).some(x => jetonsDe(x).some(m => COMPTE.includes(m))); };
  // une entrée qui ouvre un sous-menu : attribut, classe, flèche (icône, image ou caractère), ou sous-liste cachée
  const indiceSous = e => {
    for (const x of [e, ...e.querySelectorAll('[aria-haspopup], [aria-expanded]')]) {
      const h = x.getAttribute('aria-haspopup'); if (h && h !== 'false') return true;
      if (x.hasAttribute('aria-expanded')) return true; }
    if (/(^|[\s_-])(has-?sub\w*|sub-?menu\w*|has-?children|parent|dropdown(-toggle)?|expand\w*)([\s_-]|$)/i.test(classe(e))) return true;
    if (e.querySelector('[class*=chevron i], [class*=arrow i], [class*=caret i], [class*=angle i], [data-icon*=chevron i], ' +
                        '[data-icon*=angle i], [data-icon*=arrow i], img[src*=arrow i], img[src*=chevron i], ul, ol, [role=menu]')) return true;
    return /([›»▸▶►]|(chevron|arrow|angle|navigate)[a-z_]*)\s*$/i.test((e.textContent || '').trim());
  };
"""

# Ce que le robot voit comme menus, onglets, liens et boutons (noms courts, structure seulement) : les liens et boutons,
# et ce qui est dans un menu. Pas les « div » cliquables du contenu (cartes, lignes de résultats : ce sont des données).
JS_CE_QUE_JE_VOIS = r"""() => {""" + JS_OUTILS_MENUS + r"""
  const vus = [];
  const NET = 'a, button, li, [role=menuitem], [role=tab], [role=button], input[type=submit], input[type=button]';
  const MENUS = '[data-robot-fort], [data-robot-nouveau], nav, [role=navigation], [role=menu], [role=menubar], header';
  for (const e of cliquables()) {
    if (e.closest('td, [role=gridcell], [role=row], [role=option], [role=listbox], [role=tree]')) continue;
    if (!(e.closest('[data-robot-fort], [data-robot-nouveau]') || (e.matches(NET) && (e.closest(MENUS) || zoneMenu(e)) && !contenu(e))))
      continue;
    const t = e.tagName === 'INPUT' ? (e.value || '').replace(/\s+/g, ' ').trim() :
              (enregistrement(e) ? '(donnee masquee)' : (compteDe(e) ? '(menu du compte)' : texteDe(e)));
    if (t && t.length <= 40 && !vus.some(v => v[0] === t)) vus.push([t, titreSection(e)]);
    if (vus.length >= 40) break;
  }
  return vus;
}"""

# Les ENTRÉES DES MENUS de la page : barre du haut, menu de gauche (même fait de « div »), en ligne ou en colonne.
# Un groupe = 2 à 25 éléments frères, visibles, cliquables, de nom court, alignés, dans une zone de menu ou au bord
# (en haut / à gauche). Jamais : tableau, arbre, liste de choix, formulaire, onglets, fenêtre, « Libellé : valeur »,
# barre de boutons d'action. « fort » : menu certain (nav, rôle menu, ou flèches de sous-menus sur la moitié des entrées).
# Un groupe « faible » (une liste au bord de la page...) n'est gardé que si le survol de ses entrées ouvre des sous-menus.
# Chaque entrée est marquée (data-robot-barre ; data-robot-fort si le menu est certain). Renvoie [{nom, groupe, fort}].
JS_MENUS = r"""() => {""" + JS_OUTILS_MENUS + r"""
  document.querySelectorAll('[data-robot-barre], [data-robot-fort]').forEach(x => { x.removeAttribute('data-robot-barre');
                                                                                   x.removeAttribute('data-robot-fort'); });
  const dansZone = (n, colonne) => { if (!zoneMenu(n)) return false;
    if (!colonne) return true;
    for (let x = n; x && x !== document.body; x = x.parentElement)
      if (/^(NAV|ASIDE)$/.test(x.tagName) || jetonsDe(x).some(m => MOTS_MENU.includes(m)))
        return x.getBoundingClientRect().width <= Math.max(400, innerWidth * 0.35);
    return false; };
  const ONGLETS = /(^|[\s_-])(tabs?|onglets?|tab-?bar|tablist|breadcrumbs?|ariane|fil|pagination|pager)([\s_-]|$)/i;
  // bouton qui envoie un formulaire (Rechercher, Effacer...) : une rangée de ces boutons n'est pas un menu
  const envoi = k => [k, ...k.querySelectorAll('button, input')].some(x => x.form && ['submit', 'reset', 'image'].includes(x.type));
  const B = 'button, [role=button], input, a[href^="javascript:" i]', LIEN = 'a[href]:not([href^="javascript:" i])';
  const estBouton = k => k.matches(B) || (!k.matches(LIEN) && !k.querySelector(LIEN) && !!k.querySelector(B));
  const entreeCliquable = k => k.matches(CLIQ) || !!k.querySelector(CLIQ) || getComputedStyle(k).cursor === 'pointer';
  // le titre du groupe (« Récents » au-dessus d'une liste de fiches) : bulle, titre lié, ou titre juste au-dessus
  const titreDe = g => {
    const l = g.getAttribute('aria-label'); if (l) return l;
    const id = (g.getAttribute('aria-labelledby') || '').split(' ')[0];
    if (id && document.getElementById(id)) return document.getElementById(id).textContent || '';
    for (let p = g, n = 0; p && p !== document.body && n < 3; p = p.parentElement, n++) {
      const prev = p.previousElementSibling;
      if (!prev) continue;
      const t = (prev.textContent || '').replace(/\s+/g, ' ').trim();
      return /^H[1-6]$/.test(prev.tagName) || (t.length <= 30 && !prev.matches(CLIQ) && !prev.querySelector(CLIQ)) ? t : '';
    }
    return '';
  };
  // le nom écrit d'une entrée (une seule fois par élément : la page peut être grande)
  // (un nom qui ne vient que du texte d'une image, « alt » : jamais un logo ni un lien vers l'accueil, et pas un nom de
  // personne ou de marque en deux mots)
  const altSeul = (k, t) => !texteDuClone(k.cloneNode(true)) && !k.getAttribute('aria-label') && !k.getAttribute('title') &&
                            Array.from(k.querySelectorAll('img[alt]')).some(i => (i.getAttribute('alt') || '').replace(/\s+/g, ' ').trim() === t);
  const altRefuse = (k, t) => { const a = k.closest('a[href]') || k.querySelector('a[href]'), h = a ? (a.getAttribute('href') || '').trim() : '';
    const autour = [k, a].filter(Boolean).map(x => classe(x) + ' ' + (x.id || '')).join(' ');
    const logo = ICONE_PERSO.test(t + ' ' + autour + ' ' + Array.from(k.querySelectorAll('img')).map(i => (i.getAttribute('src') || '').slice(0, 200) + ' ' + classe(i)).join(' ')) ||
                 /brand|navbar-header|entete|header-logo/i.test(autour);
    let accueil = false;
    try { const u = new URL(h || '#', location.href), b = new URL(document.baseURI);
          accueil = h !== '#' && u.origin === location.origin && u.search === '' &&
                    (/^\/(index\.\w+|home|accueil|dashboard)?\/?$/i.test(u.pathname) || u.pathname.replace(/\/$/, '') === b.pathname.replace(/\/$/, '')) &&
                    /^(#\/?(home|accueil)?\/?)?$/i.test(u.hash); } catch (x) {}
    return !bulleSure(t) || logo || accueil; };
  // (un nom qui ne se VOIT pas, pris d'une bulle, d'un « title » ou d'un « alt » : contrôlé comme une bulle, jamais
  // un nom de personne)
  const memo = new Map(), T = k => { if (!memo.has(k)) {
      let t = texteDe(k);
      if (t && !texteDuClone(k.cloneNode(true)) && (!bulleSure(t) || (altSeul(k, t) && altRefuse(k, t)))) t = '';
      memo.set(k, t); }
    return memo.get(k); };
  const nomCourt = t => (t && t.length <= 30 && /[A-Za-zÀ-ÿ]/.test(t) && !t.includes(':')) ? t : '';
  // une entrée faite d'une ICÔNE seule (petite, sans texte à elle) : sa bulle ou son icône ; sans aucun nom, elle
  // sera nommée par sa place (« icone 1 », « icone 2 »...)
  const nomIcone = k => { const t = T(k).replace(/^\d+\+?$/, ''), r = k.getBoundingClientRect();   // (« 3 » : une pastille)
    if (r.width > 160 || r.height > 160 || (t && (nomCourt(t) || texteDuClone(k.cloneNode(true)).replace(/^\d+\+?$/, '')))) return '';
    return cleIcone(k) || ((k.matches('svg, img') || k.querySelector('svg, img, i, mat-icon, [class*=icon i]') ||
                            [k, ...k.querySelectorAll('*')].slice(0, 20).some(x => /ICON/i.test(x.tagName))) ? 'icone ?' : ''); };
  // un groupe d'entrées : 2 à 40, même genre d'élément, alignées, cliquables, pas un formulaire ni une liste de fiches
  const essai = (g, v, e, icones) => {
    if (e.length < 2 || e.length > 40 || e.length < 0.7 * v.length || !e.every(k => k.tagName === e[0].tagName)) return null;
    if (e.some(k => k.matches('input, select, textarea') || k.querySelector('input:not([type=hidden]), select, textarea') || envoi(k))) return null;
    if (e.filter(entreeCliquable).length < 0.7 * e.length) return null;
    if (e.filter(enregistrement).length >= e.length / 2) return null;   // une liste de fiches (données), pas un menu
    let fort = !contenu(g) &&
               (!!g.closest('nav, [role=navigation], [role=menu], [role=menubar]') || e.filter(indiceSous).length >= e.length / 2);
    const rs = e.map(k => k.getBoundingClientRect());
    const ecart = f => Math.max(...rs.map(f)) - Math.min(...rs.map(f));
    const enLigne = ecart(r => r.top) <= 6, enColonne = !enLigne && ecart(r => r.left) <= 6;
    if (!enLigne && !enColonne) return null;
    const gauche = enColonne && rs.every(r => r.left < 100);
    // des icônes ailleurs qu'à gauche (barre du haut : replier le menu, langue, aide...) : gardées seulement si leur
    // survol ouvre des sous-menus (jamais un menu « certain » cliqué entrée par entrée)
    if (icones && !gauche) fort = false;
    // un dessin sans aucun nom : seulement dans une colonne au bord gauche (barre d'icônes) ; ailleurs (barre d'outils),
    // on ne sait pas ce qu'il fait
    if (icones && !gauche) { e = e.filter(k => nomIcone(k) !== 'icone ?'); if (e.length < 2) return null; }
    // une colonne à gauche reconnue par ses textes : ses entrées faites d'une icône nommée (sans bulle) en font partie
    if (!icones && gauche) {
      const en = new Set(e), x0 = rs[0].left;
      e = v.filter(k => en.has(k) || (k.tagName === e[0].tagName && Math.abs(k.getBoundingClientRect().left - x0) <= 6 &&
                                      (k.textContent || '').length <= 3000 && !nomCourt(T(k)) && /^icone [a-z]/.test(nomIcone(k))));
    }
    // des boutons hors d'un menu certain et hors d'une zone de menu : barre de boutons d'action (Calculer, Imprimer...)
    if (!fort && e.filter(estBouton).length >= e.length / 2 && !dansZone(g, enColonne)) return null;
    const gr = g.getBoundingClientRect();
    const ok = enLigne ? (gr.top < 260 || !!g.closest('header') || dansZone(g, false))
                       : (dansZone(g, true) || (gr.left < 100 && gr.width <= 400));
    return ok ? { g: g, e: e, fort: fort, icones: icones } : null;
  };
  let anonymes = 0;
  const groupes = [];
  for (const g of document.querySelectorAll('body *')) {
    if (g.children.length < 2 || g.children.length > 40 || /^(SELECT|DATALIST|SVG|OPTGROUP)$/i.test(g.tagName)) continue;
    if (g.closest('[role=grid], [role=row], [role=tablist], [role=tree], [role=treegrid], [role=listbox], [role=dialog], ' +
                  '[role=alertdialog], dialog, [role=tooltip]') ||
        ONGLETS.test(classe(g)) || (/^(TABLE|TBODY|THEAD|TR)$/.test(g.tagName) && !dansZone(g, false)) || !vis(g)) continue;
    const v = Array.from(g.children).filter(vis);
    const court = v.filter(k => (k.textContent || '').length <= 3000);
    // d'abord les entrées qui ont un texte ; les icônes seules seulement si le groupe n'est pas un menu sans elles
    // (une barre d'icônes, ou des entrées surtout faites d'icônes)
    let G = essai(g, v, court.filter(k => nomCourt(T(k))), false);
    if (!G && court.some(k => !nomCourt(T(k))))
      G = essai(g, v, court.filter(k => nomCourt(T(k)) || nomIcone(k)), true);
    if (G) groupes.push(G);
  }
  // un groupe qui n'est qu'une enveloppe d'un autre menu (en-tête = logo + menu + nom) : on garde le menu
  const reste = (k, c) => texteDe(k).toLowerCase().replace(/[^a-z0-9à-ÿ]/g, '').replace(texteDe(c).toLowerCase().replace(/[^a-z0-9à-ÿ]/g, ''), '');
  const garde = groupes.filter(G => !G.e.some(k => groupes.some(H => H !== G && k.contains(H.g) && !/[a-zà-ÿ]/.test(reste(k, H.g)))));
  const entrees = [];
  garde.forEach((G, n) => {
    if (G.e.some(k => k.closest('[data-robot-barre]'))) return;   // sous-menu déjà ouvert d'une entrée retenue
    for (const k of G.e) {
      if (entrees.length >= 60) break;
      k.setAttribute('data-robot-barre', String(entrees.length));
      if (G.fort) k.setAttribute('data-robot-fort', '1');
      const nomK = nomCourt(T(k)) || cleIcone(k) || ('icone ' + (++anonymes));
      // (2 ou 3 majuscules dans une pastille RONDE : les initiales de la personne connectée, pas un sigle « SIG »)
      const rond = [k, ...k.querySelectorAll('*')].slice(0, 6).some(x => { const r = x.getBoundingClientRect();
        return r.width > 0 && parseFloat(getComputedStyle(x).borderTopLeftRadius) >= 0.4 * Math.min(r.width, r.height); });
      entrees.push({ nom: enregistrement(k) ? '(donnee masquee)' : nomK, texte: nomK, groupe: n, fort: G.fort,
                     titre: titreDe(G.g), section: titreSection(k),
                     compte: compteDe(k) || COMPTE_ICONE.test(iconeSeule(k)) ||
                             (/^[A-ZÀ-Ý]{2,3}$/.test(texteDuClone(k.cloneNode(true))) && (rond || G.icones)),
                     // tout le texte, sous-menu caché compris (« Se déconnecter » : menu de la personne) ; jamais écrit
                     cache: (k.textContent || '').replace(/\s+/g, ' ').trim().slice(0, 600) });
    }
  });
  // deux icônes du même nom (« icone folder ») : la deuxième devient « icone folder 2 » (retrouvée par sa place)
  const fois = {};
  for (const x of entrees) if (/^icone /.test(x.texte)) {
    fois[x.texte] = (fois[x.texte] || 0) + 1;
    if (fois[x.texte] > 1) { x.texte = x.texte + ' ' + fois[x.texte]; if (x.nom !== '(donnee masquee)') x.nom = x.texte; }
  }
  return entrees;
}"""

# L'entrée faite d'une icône seule dont le nom (bulle, ou « icone dashboard ») est celui-ci : l'élément le plus
# intérieur (le lien dans « li > a », pas le « li »). numero : ou bien l'entrée de menu marquée data-robot-barre=numero.
JS_TROUVER_ICONE = r"""([cle, jeton, numero]) => {""" + JS_OUTILS_MENUS + r"""
  const LIEN = 'a[href], button, [role=menuitem], [role=link], [role=button], [role=tab], [routerlink], [ng-reflect-router-link]';
  // une entrée qui contient UN seul lien ou bouton : c'est lui qu'on clique
  const interieur = e => { if (e.matches(LIEN)) return e; const d = Array.from(e.querySelectorAll(LIEN)).filter(vis);
                           return d.length === 1 ? d[0] : e; };
  let trouves;
  if (numero !== null && numero !== undefined) {
    trouves = Array.from(document.querySelectorAll('[data-robot-barre="' + numero + '"]'));
  } else {
    // d'abord les entrées des menus (la bulle peut être sur l'entrée, autour du lien : « div[mattooltip] > a > icône ») ;
    // sinon un élément cliquable de la page qui porte cette icône
    trouves = Array.from(document.querySelectorAll('[data-robot-barre]')).filter(e => !texteDe(e) && cleIcone(e) === cle);
    if (!trouves.length) trouves = cliquables().filter(e => !texteDe(e) && cleIcone(e) === cle);
    trouves = trouves.filter(e => !trouves.some(f => f !== e && e.contains(f)));
  }
  const marques = new Set(trouves.map(interieur));
  marques.forEach(e => e.setAttribute('data-robot-trouve', jeton));
  return marques.size;
}"""

# Avant un survol : ce qui est déjà visible (pour voir ensuite ce qui APPARAÎT). La marque du survol précédent
# (data-robot-survole) est retirée.
JS_VUS = r"""() => {""" + JS_OUTILS_MENUS + r"""
  document.querySelectorAll('[data-robot-survole]').forEach(x => x.removeAttribute('data-robot-survole'));
  window.__robotVus = new WeakMap();
  for (const e of cliquables()) window.__robotVus.set(e, texteDe(e));
  return 0;
}"""

# Après un survol : les éléments cliquables qui viennent d'apparaître PRÈS de l'élément survolé (data-robot-survole), ou
# dans un panneau de menu (rôle menu, « flyout », « dropdown », « overlay »...), et ceux de l'entrée survolée elle-même.
# Ce qui apparaît ailleurs dans la page (contenu qui finit de se charger) est ignoré ; trop de nouveautés = la page a changé.
# Ils sont marqués (data-robot-nouveau) pour pouvoir être survolés à leur tour.
JS_NOUVEAUX = r"""([numero, jeton]) => {""" + JS_OUTILS_MENUS + r"""
  const entree = numero >= 0 ? document.querySelector('[data-robot-barre="' + numero + '"]') : null;
  const survole = document.querySelector('[data-robot-survole]');
  const groupe = entree ? entree.parentElement : null;
  const nomEntree = (entree || survole) ? texteDe(entree || survole) : '';
  const vus = window.__robotVus || new WeakMap();
  const PANNEAU = ['menu', 'menus', 'submenu', 'menuitem', 'menubar', 'flyout', 'dropdown', 'popover', 'popup', 'overlay', 'cdk',
                   'megamenu', 'nav', 'navigation', 'listbox'];
  const dansPanneau = n => { for (let x = n; x && x !== document.body; x = x.parentElement)
                               if (jetonsDe(x).some(m => PANNEAU.includes(m))) return true; return false; };
  const ra = survole ? survole.getBoundingClientRect() : null;
  const pres = e => { if (!ra) return false; const r = e.getBoundingClientRect();
    const dx = Math.max(0, ra.left - r.right, r.left - ra.right), dy = Math.max(0, ra.top - r.bottom, r.top - ra.bottom);
    return dx <= 400 && dy <= 400; };
  const nouveaux = [];
  for (const e of cliquables()) {
    if (entree && (e === entree || e.contains(entree))) continue;
    if (survole && (e === survole || e.contains(survole))) continue;
    const hote = entree || survole;
    const dedans = !!(hote && hote.contains(e));
    if (!dedans && ((groupe && groupe.contains(e)) ||
        e.closest('[role=dialog], [role=alertdialog], dialog, [aria-modal=true], [role=tooltip], [data-robot-barre]'))) continue;
    if (dedans) {   // dans l'entrée survolée : un sous-menu (liste imbriquée, ou dessiné hors de l'entrée), pas une icône d'action
      const l = e.closest(SOUS), r = e.getBoundingClientRect(), b = hote.getBoundingClientRect();
      const cx = (r.left + r.right) / 2, cy = (r.top + r.bottom) / 2;
      if (!((l && l !== hote && hote.contains(l)) || cx < b.left || cx > b.right || cy < b.top || cy > b.bottom)) continue;
    }
    const t = e.tagName === 'INPUT' ? '' : texteDe(e);
    if (!t || t.length > 45 || !/[A-Za-zÀ-ÿ0-9]/.test(t) || t === nomEntree) continue;
    if (vus.has(e) && vus.get(e) === t) continue;
    nouveaux.push({ e: e, t: t, ok: dedans || dansPanneau(e) || (pres(e) && !contenu(e)) });
  }
  if (nouveaux.length > 80) return [];   // la page entière a changé : ce n'est pas un sous-menu
  // les voisins (mêmes parents) d'un élément retenu font partie du même sous-menu
  for (let tour = 0; tour < 2; tour++) {
    const parents = new Set(nouveaux.filter(n => n.ok).map(n => n.e.parentElement));
    for (const n of nouveaux) if (!n.ok && parents.has(n.e.parentElement)) n.ok = true;
  }
  const retenus = nouveaux.filter(n => n.ok);
  const ens = new Map(retenus.map(n => [n.e, n.t]));
  // un panneau (role=menu, tabindex...) qui contient directement d'autres éléments retenus n'est pas un élément :
  // seuls ses éléments comptent (une entrée « li » qui a un sous-menu, elle, reste)
  const contientDirect = (x, y) => x !== y && x.contains(y) && !Array.from(x.querySelectorAll(SOUS)).some(z => z.contains(y));
  const garde = retenus.filter(n => { for (let p = n.e.parentElement; p; p = p.parentElement) if (ens.get(p) === n.t) return false;
                                      return !retenus.some(m => m.t !== n.t && contientDirect(n.e, m.e)); });
  return garde.slice(0, 40).map((n, i) => { n.e.setAttribute('data-robot-nouveau', jeton + '-' + i);
    return { texte: enregistrement(n.e) ? '(donnee masquee)' : n.t, sous: indiceSous(n.e), cle: jeton + '-' + i,
             section: titreSection(n.e) }; });
}"""

# Un sous-menu en train de se charger près de l'élément survolé (« Chargement... », roue, barre, aria-busy) : on attend.
JS_EN_ATTENTE = r"""() => {""" + JS_OUTILS_MENUS + r"""
  const survole = document.querySelector('[data-robot-survole]');
  if (!survole) return false;
  const ra = survole.getBoundingClientRect();
  const PANNEAU = ['menu', 'menus', 'submenu', 'flyout', 'dropdown', 'popover', 'popup', 'overlay', 'cdk', 'megamenu', 'listbox'];
  const ici = x => { for (let y = x; y && y !== document.body; y = y.parentElement)
                       if (jetonsDe(y).some(m => PANNEAU.includes(m))) return true;
                     const r = x.getBoundingClientRect();
                     return Math.max(0, ra.left - r.right, r.left - ra.right) <= 400 && Math.max(0, ra.top - r.bottom, r.top - ra.bottom) <= 400; };
  for (const x of document.querySelectorAll('[aria-busy=true], [class*=spinner i], [class*=loading i], [class*=loader i], ' +
                                            '[class*=progress i], mat-spinner, mat-progress-spinner, mat-progress-bar'))
    if (vis(x) && ici(x)) return true;
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  for (let n = w.nextNode(); n; n = w.nextNode())
    if (/chargement|loading|patientez|en cours/i.test(n.textContent) && n.textContent.length < 60 && n.parentElement &&
        vis(n.parentElement) && ici(n.parentElement)) return true;
  return false;
}"""

# L'élément trouvé par son texte -> l'élément que le clic déclencherait vraiment.
# LISTE BLANCHE : un lien, une entrée de menu / d'onglet, ou un élément d'un menu (nav, menu...), un bouton
# qui ouvre un menu. Refusé : bouton de formulaire, case, champ, bascule, zone modifiable, élément de tableau
# qui n'est pas un simple lien, élément qui contient un autre bouton (icône poubelle...).
# Les attributs de l'élément et les gestionnaires de clic de TOUS ses parents sont renvoyés pour le contrôle des mots.
JS_EXAMINER = r"""(e, [jeton, premier]) => {""" + JS_ZONE_MENU + r"""
  const SOUS = 'ul, ol, [role=menu], [role=group], [role=listbox], table, select';
  const CLIQ = 'a, button, input, select, textarea, label, summary, option, [role=button], [role=link], [role=menuitem], ' +
               '[role=menuitemcheckbox], [role=menuitemradio], [role=tab], [role=treeitem], [role=checkbox], [role=switch], ' +
               '[role=radio], [role=option], [onclick]';
  const MENU = 'nav, header, [role=navigation], [role=menu], [role=menubar], [role=tablist], [role=tree], ' +
               '[data-robot-fort], [data-robot-nouveau]';
  let c = e.closest(CLIQ) || e;
  // l'élément est dans un SOUS-MENU du bloc cliquable trouvé (« Données MK1 » sous « GATES ») :
  // la cible reste l'entrée du sous-menu, jamais le menu parent
  if (c !== e && Array.from(c.querySelectorAll(SOUS)).some(x => x.contains(e))) c = e.closest('li, [role=menuitem]') || e;
  c.setAttribute('data-robot-cible', jeton);
  // le mot trouvé : c'est sur LUI que la souris se posera (jamais sur le premier texte d'un bloc plus grand)
  if (premier) {
    document.querySelectorAll('[data-robot-mot]').forEach(x => x.removeAttribute('data-robot-mot'));
    if (c !== e && c.contains(e)) e.setAttribute('data-robot-mot', '1');
  }
  const tag = c.tagName.toLowerCase(), role = (c.getAttribute('role') || '').toLowerCase();
  const href = (c.getAttribute('href') || '').trim();
  const lien = tag === 'a' && c.hasAttribute('href') && !/^javascript:/i.test(href);
  const lienScript = tag === 'a' && /^javascript:/i.test(href);
  const entree = ['menuitem', 'tab', 'link', 'treeitem'].includes(role);
  const bouton = tag === 'button' || role === 'button';
  // dans une fenêtre (« Êtes-vous sûr ? ») : jamais un menu
  const classes = x => typeof x.className === 'string' ? x.className : '';
  let fenetre = !!c.closest('[role=dialog], [role=alertdialog], dialog, [aria-modal=true]');
  for (let x = c; x && !fenetre; x = x.parentElement) if (/modal|dialog|k-window|alert/i.test(classes(x))) fenetre = true;
  // un bouton qui n'envoie aucun formulaire : hors formulaire, « type=button », ou rôle bouton sur un div / span
  const sansFormulaire = !(c.form || c.closest('form')) || (tag === 'button' && c.type === 'button') ||
                         (role === 'button' && !['button', 'input'].includes(tag));
  // un vrai bouton n'est accepté que s'il ouvre un menu, est une entrée de menu (rôle), ou appartient à un menu CERTAIN :
  // entrée d'un menu sûr (data-robot-fort) ou élément apparu en survolant un menu (data-robot-nouveau)
  const menuSur = !!c.closest('[data-robot-fort], [data-robot-nouveau]');
  const menuBouton = bouton && sansFormulaire && !fenetre &&
                     (entree || c.hasAttribute('aria-haspopup') || c.hasAttribute('aria-expanded') || menuSur);
  // un élément d'un menu (lien, li, div, span) : liste, menu repéré, ou parent dont le NOM est un mot de menu
  // (un « li » hors du contenu de la page — main, formulaire, « ...-content » — compte aussi : vieux menus en listes)
  const dansMenu = !!(c.closest(MENU) || zoneMenu(c) || (c.closest('li') && !contenu(c))) &&
                   !fenetre && ['li', 'span', 'div', 'a', 'p'].includes(tag);
  let refus = '';
  if (['input', 'select', 'textarea', 'label', 'summary', 'option'].includes(tag) ||
      ['checkbox', 'switch', 'radio', 'option', 'menuitemcheckbox', 'menuitemradio'].includes(role) ||
      c.hasAttribute('aria-checked') || c.hasAttribute('aria-pressed') || (bouton && !menuBouton) ||
      (tag === 'button' && !sansFormulaire))
    refus = 'bouton, case ou champ';
  else if (fenetre && !lien)
    refus = 'element d une fenetre';
  else if (c.isContentEditable || document.designMode === 'on')
    refus = 'zone modifiable';
  else if (!(lien || entree || menuBouton || dansMenu || (lienScript && (c.closest(MENU) || zoneMenu(c)))))
    refus = 'pas un lien ni un menu';
  else if (!lien && c.closest('table, [role=grid], [role=row], [role=gridcell], [role=treegrid]'))
    refus = 'element de tableau';
  else if (c.hasAttribute('data-confirm') || c.hasAttribute('data-method') || c.hasAttribute('data-turbo-method') ||
           c.hasAttribute('data-ajax-method') || ['hx-post', 'hx-put', 'hx-patch', 'hx-delete'].some(a => c.hasAttribute(a)))
    refus = 'action envoyee au serveur';
  // le texte propre de l'élément (sous-menus à part), et ce qu'il contient
  const copie = c.cloneNode(true);
  copie.querySelectorAll(SOUS).forEach(x => x.remove());
  // le mot d'une icône (« archive » dans « <mat-icon>archive</mat-icon> Archives ») n'est pas le nom de l'élément, sauf
  // s'il n'y a rien d'autre (icône seule : elle est alors jugée par son icône)
  // (un mot d'icône : dans mat-icon, ou dans un i / span qui porte une classe d'icône ; jamais un mot en italique)
  const motDIcone = x => !x.children.length && /^[a-z0-9_]+$/.test((x.textContent || '').trim()) &&
      (x.tagName === 'MAT-ICON' || /icon|symbol|material|^fa|\bfa-|mdi|glyph/i.test(classes(x)));
  const ligatures = Array.from(copie.querySelectorAll('mat-icon, i, span')).filter(motDIcone);
  const sansIcones = copie.cloneNode(true);
  Array.from(sansIcones.querySelectorAll('mat-icon, i, span')).filter(motDIcone).forEach(x => x.remove());
  const propre = (sansIcones.textContent || '').replace(/\s+/g, ' ').trim();
  const motsIcones = ligatures.map(x => (x.textContent || '').trim());
  if (propre && ligatures.length) ligatures.forEach(x => x.remove());
  if (!refus && copie.querySelector('a[href], button, input, select, textarea, label, [role=button], [role=checkbox], [onclick]'))
    refus = 'contient un autre bouton';
  let texte = (copie.textContent || '').replace(/\s+/g, ' ').trim();
  if (!texte)   // lien ou bouton fait d'une image : son texte, c'est le nom de l'image ou sa bulle
    texte = [c.getAttribute('title'), c.getAttribute('aria-label'), ...Array.from(copie.querySelectorAll('img, input[type=image]'))
             .map(i => i.getAttribute('alt') || i.getAttribute('title'))].filter(Boolean).join(' ').replace(/\s+/g, ' ').trim();
  const attributs = [];
  // les attributs d'ÉTAT (aria-disabled="false", aria-expanded...) décrivent l'élément, ils ne sont pas des actions
  const ETAT = /^aria-(disabled|expanded|haspopup|selected|hidden|current|level|setsize|posinset|orientation|busy|live|atomic)$/;
  for (const n of [copie, ...copie.querySelectorAll('*')].slice(0, 80))
    for (const a of Array.from(n.attributes))
      if (a.name !== 'style' && !a.name.startsWith('data-robot') && !ETAT.test(a.name)) attributs.push(a.name + ' ' + a.value);
  // les parents : leurs gestionnaires de clic, et le formulaire qui serait envoyé
  const GEST = /^(on(click|dblclick|mousedown|mouseup|pointerdown|pointerup|touchstart|touchend|submit))$|click|^hx-|^wire:|^x-on|^@|^data-(ajax|turbo|confirm|method|action|url)/i;
  for (let n = c.parentElement; n && n !== document.documentElement; n = n.parentElement)
    for (const a of Array.from(n.attributes))
      if (GEST.test(a.name)) attributs.push(a.name + ' ' + a.value);
  const formulaire = c.form || c.closest('form');
  if (formulaire) attributs.push('form ' + (formulaire.getAttribute('action') || ''));
  // une entrée de MENU qui porte son propre nom (bulle, titre) : jugée par ce nom, pas par le dessin de son icône
  const NOMS_BULLE = ['title', 'aria-label', 'mattooltip', 'data-tooltip', 'ng-reflect-message', 'data-bs-original-title', 'ptooltip'];
  const bulle = NOMS_BULLE.some(a => (c.getAttribute(a) || '').trim() || (c.parentElement && (c.parentElement.getAttribute(a) || '').trim()));
  // (la bulle du parent, « <span mattooltip="Supprimer"><a>... » : contrôlée aussi)
  if (c.parentElement) for (const a of NOMS_BULLE) { const v = (c.parentElement.getAttribute(a) || '').trim(); if (v) attributs.push(a + ' ' + v); }
  const menuNomme = bulle && !!c.closest('[data-robot-barre], nav, [role=navigation], [role=menu], [role=menubar]');
  return { texte: texte.slice(0, 300), attributs: attributs.join(' | '), refus: refus, propre: propre.slice(0, 300),
           menuNomme: menuNomme, ligatures: motsIcones.slice(0, 5) };
}"""

# Le point où cliquer : le propre texte de l'élément (pas son centre, qui peut tomber sur un sous-menu) ; si le mot
# tapé est un morceau du bloc cliquable (data-robot-mot), c'est sur CE mot, jamais sur le premier texte du bloc.
# Marque le parent de ce texte (data-robot-texte) pour revérifier le même point au dernier moment.
JS_POINT = r"""(c, jeton) => {
  const SOUS = 'ul, ol, [role=menu], [role=group], [role=listbox], table, select';
  const mot = c.querySelector('[data-robot-mot]');
  (mot || c).scrollIntoView({block: 'nearest', inline: 'nearest'});
  const rc = c.getBoundingClientRect();
  let r = null, p = c;
  const marcheur = document.createTreeWalker(mot || c, NodeFilter.SHOW_TEXT);
  for (let n = marcheur.nextNode(); n; n = marcheur.nextNode()) {
    if (!n.textContent.trim()) continue;
    const parent = n.parentElement;
    if (parent && parent !== c && parent.closest(SOUS) && c.contains(parent.closest(SOUS))) continue;
    const plage = document.createRange(); plage.selectNodeContents(n);
    const b = plage.getBoundingClientRect();
    if (b.width > 0 && b.height > 0) { r = b; p = parent || c; break; }
  }
  if (!r) r = rc;
  p.setAttribute('data-robot-texte', jeton);
  return { x: r.left + Math.min(r.width / 2, 40) - rc.left, y: r.top + r.height / 2 - rc.top };
}"""

# Ce qui est sous la souris est-il bien le texte choisi (et pas un sous-menu ouvert par-dessus, un calque...) ?
JS_VERIFIER_POINT = r"""(c, [x, y, jeton]) => {
  const SOUS = 'ul, ol, [role=menu], [role=group], [role=listbox], table, select';
  const p = document.querySelector('[data-robot-texte="' + jeton + '"]');
  const rc = c.getBoundingClientRect();
  const sous = document.elementFromPoint(rc.left + x, rc.top + y);
  if (!p || !sous) return false;
  const liste = sous.closest(SOUS);
  return (sous === p || p.contains(sous)) && !(liste && liste !== p && p.contains(liste));
}"""

# Élément caché dans un menu fermé : l'entrée de menu visible la plus proche, à survoler (comme la souris).
JS_ENTREE_MENU = r"""e => {
  const vis = n => { const s = getComputedStyle(n); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     const r = n.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  if (vis(e)) return null;
  let n = e.parentElement, haut = 1;
  while (n && !vis(n)) { n = n.parentElement; haut++; }
  if (!n || n === document.body || n === document.documentElement || haut > 12) return { haut: 0, enfant: -2 };
  const enfant = Array.from(n.children).findIndex(c => !c.contains(e) && vis(c) && (c.innerText || '').trim() &&
    !c.matches('ul, ol, [role=menu], [role=group], [role=listbox], .dropdown-menu, .submenu, .sub-menu'));
  return { haut: haut, enfant: enfant };
}"""


def _visible(loc):
    try:
        return loc.is_visible()
    except Exception:
        return False


def _encadrer(loc, ms=1200):
    try:
        loc.evaluate(JS_ENCADRER, ms)
        time.sleep(ms / 1000 + 0.2)
    except Exception:
        pass


def _taper_lentement(loc, texte):
    if hasattr(loc, "press_sequentially"):
        loc.press_sequentially(texte, delay=150)
    else:
        loc.type(texte, delay=150)


def _ouvrir_menus(loc):
    """Survole, de haut en bas, les entrées de menu qui cachent l'élément (4 niveaux au plus)."""
    deja = set()
    for _ in range(4):
        try:
            point = loc.evaluate(JS_ENTREE_MENU)
        except Exception:
            return False
        if not point:
            return True
        if point["enfant"] == -2 or (point["haut"], point["enfant"]) in deja:
            return False
        deja.add((point["haut"], point["enfant"]))
        entree = loc.locator("xpath=" + "/".join([".."] * point["haut"]))
        if point["enfant"] >= 0:
            entree = entree.locator(f"xpath=./*[{point['enfant'] + 1}]")
        try:
            entree.hover(timeout=3000)
            loc.wait_for(state="visible", timeout=1000)
            return True
        except Exception:
            continue
    return _visible(loc)


# ---------------------------------------------------------------------------- éléments qui réagissent au clic (Chrome)
# Dans chaque zone de la page (page, cadres, composants web) : l'outil de Chrome « getEventListeners » dit quels
# éléments ont un écouteur de clic ou de survol ; ils sont marqués (data-robot-ecoute / data-robot-survol), et les
# marques devenues fausses sont retirées. Renvoie [éléments marqués, la page ENTIÈRE écoute les clics].
JS_MARQUER_ECOUTEURS = r"""(() => {
  const CLIC = ['click', 'mousedown', 'mouseup', 'pointerdown', 'pointerup', 'touchstart', 'touchend', 'dblclick', 'auxclick'];
  const SURVOL = ['mouseenter', 'mouseover', 'pointerenter', 'pointerover'];
  const types = x => { try { return Object.keys(getEventListeners(x) || {}); } catch (e) { return []; } };
  const poser = (e, nom, val) => { if (val) { if (e.getAttribute(nom) !== val) e.setAttribute(nom, val); }
                                   else if (e.hasAttribute(nom)) e.removeAttribute(nom); };
  let n = 0;
  const fouiller = racine => {
    for (const e of racine.querySelectorAll('*')) {
      const t = types(e);
      const c = t.filter(x => CLIC.includes(x)).sort().join(' '), s = t.filter(x => SURVOL.includes(x)).sort().join(' ');
      if (c || s || e.hasAttribute('data-robot-ecoute') || e.hasAttribute('data-robot-survol')) {
        poser(e, 'data-robot-ecoute', c); poser(e, 'data-robot-survol', s); if (c || s) n++; }
      if (e.shadowRoot) fouiller(e.shadowRoot);
    }
  };
  fouiller(document);
  return [n, types(document).includes('click') || types(window).includes('click')];
})()"""
_CLIC_DELEGUE = {}  # page -> la page entière (document, fenêtre) écoute les clics
_LIAISONS = {}  # (page, cadre) -> liaison avec Chrome gardée ouverte, et ses zones (page, cadres, après navigation)


def _liaison(page, cible):
    """La liaison avec Chrome pour la page (ou un cadre d'un autre site), ouverte une fois et gardée : Chrome dit
    au fur et à mesure quelles zones existent (nouvelle page, cadre ajouté...). None si ce n'est pas possible."""
    cle = (id(page), id(cible))
    if cle in _LIAISONS:
        return _LIAISONS[cle]
    try:
        session = page.context.new_cdp_session(cible)
    except Exception:
        _LIAISONS[cle] = None  # (cadre du même site : déjà dans les zones de la page)
        return None
    zones = {}

    def creee(ev):
        c = ev.get("context") or {}
        if (c.get("auxData") or {}).get("isDefault"):
            zones[c["id"]] = True

    session.on("Runtime.executionContextCreated", creee)
    session.on("Runtime.executionContextDestroyed", lambda ev: zones.pop(ev.get("executionContextId"), None))
    session.on("Runtime.executionContextsCleared", lambda ev: zones.clear())
    try:
        session.send("Runtime.enable")
        page.wait_for_timeout(30)
    except Exception:
        _LIAISONS[cle] = None
        return None
    _LIAISONS[cle] = {"session": session, "zones": zones}
    return _LIAISONS[cle]


def _marquer_ecouteurs(page):
    """Chrome sait quels éléments réagissent au clic ou au survol (ce que montre son outil F12, « Event Listeners ») :
    le robot le lui demande, zone par zone (page, cadres, composants), et les marque. Rien n'est cliqué.
    Renvoie le nombre d'éléments marqués."""
    marques = 0
    for cible in [page] + list(page.frames[1:]):
        liaison = _liaison(page, cible)
        if not liaison:
            continue
        for zone in list(liaison["zones"]) or [None]:
            demande = {"expression": JS_MARQUER_ECOUTEURS, "includeCommandLineAPI": True, "returnByValue": True}
            if zone is not None:
                demande["contextId"] = zone
            try:
                n, delegue = liaison["session"].send("Runtime.evaluate", demande)["result"]["value"]
            except Exception:
                liaison["zones"].pop(zone, None)
                if zone is None:  # la liaison elle-même ne répond plus (page fermée...) : elle sera rouverte
                    _LIAISONS.pop((id(page), id(cible)), None)
                continue
            marques += n
            if delegue:
                _CLIC_DELEGUE[id(page)] = True
    return marques


# ---------------------------------------------------------------------------- menus qui s'ouvrent au survol
# (menu de gauche dont les sous-menus n'apparaissent que sous la souris). Le robot SURVOLE, il ne clique jamais ici.
def _menus_de_la_page(page):
    """Repère et marque les entrées des menus de la page (tous les cadres).
    Renvoie [{"cadre", "numero", "nom", "groupe", "fort" (menu certain), "loc"}]."""
    entrees = []
    for n_cadre, cadre in enumerate(page.frames):
        try:
            for numero, e in enumerate(cadre.evaluate(JS_MENUS)):
                if e.get("compte") or _menu_du_compte([e.get("cache", "")]):
                    nom = "(menu du compte)"
                elif _titre_de_fiches(e.get("titre")) or _titre_de_fiches(e.get("section")):   # liste « Récents »
                    nom = "(donnee masquee)"
                else:
                    nom = e["nom"]
                entrees.append({"cadre": cadre, "numero": numero, "nom": nom, "texte": e.get("texte", ""),
                                "groupe": (n_cadre, e["groupe"]),
                                "fort": e["fort"], "loc": cadre.locator(f'[data-robot-barre="{numero}"]')})
        except Exception:
            continue
    return entrees


def _titre_de_fiches(titre):
    """Le titre d'une liste de FICHES (« Récents », « Favoris », « Mes plans »...), pas celui d'un menu
    (« Mes applications », « Suivi réseau » restent des menus)."""
    t = " ".join(sans_accents(str(titre or "")).lower().split())
    return bool(re.match(r"(recents?|recemment|favoris?|historique|derniers?|dernieres?|epingles?|pinned|history|bookmarks?|"
                         r"signets?|last)\b", t) or
                re.match(r"(mes|my) (favoris|recherches|plans|dossiers|demandes|fiches|documents|clients|contrats|projets|"
                         r"taches|alertes|notifications|messages)\b", t))


def _menu_confirme(entree):
    """Le survol de cette entrée a ouvert un sous-menu : c'est bien un menu (data-robot-fort)."""
    try:
        entree["loc"].evaluate("e => e.setAttribute('data-robot-fort', '1')")
    except Exception:
        pass


def _oublier_groupes(entrees, groupes):
    """Ces groupes ne sont pas des menus (aucun sous-menu au survol) : le robot retire ses marques."""
    for e in entrees:
        if e["groupe"] in groupes:
            try:
                e["loc"].evaluate("e => { e.removeAttribute('data-robot-barre'); e.removeAttribute('data-robot-fort'); }")
            except Exception:
                pass


def _souris_au_repos(page):
    """La souris quitte les menus (ils se referment) : en bas à droite de la page, loin du menu de gauche."""
    try:
        largeur, hauteur = page.evaluate("() => [innerWidth, innerHeight]")
        page.mouse.move(int(largeur * 0.7), int(hauteur * 0.85))
        page.wait_for_timeout(500)
    except Exception:
        pass


def _survoler_et_voir(page, loc, cadre_entree=None, numero=-1, patience=False):
    """Pose la souris sur un élément (jamais de clic) et renvoie ce qui vient d'apparaître À CÔTÉ de lui :
    [{"texte", "sous" (ouvre lui-même un sous-menu), "loc"}]. Un sous-menu qui se charge (« Chargement... ») est
    attendu jusqu'à 6 secondes ; patience : attendre 6 secondes de toute façon."""
    for cadre in page.frames:
        try:
            cadre.evaluate(JS_VUS)
        except Exception:
            pass
    try:
        loc.evaluate("e => e.setAttribute('data-robot-survole', '1')")
    except Exception:
        pass
    loc.hover(timeout=3000)

    def ce_qui_est_apparu():
        _marquer_ecouteurs(page)  # les éléments du sous-menu qui réagissent au clic (signalés par Chrome)
        jeton = secrets.token_hex(4)
        vus = []
        for cadre in page.frames:
            try:
                for n in cadre.evaluate(JS_NOUVEAUX, [numero if cadre == cadre_entree else -1, jeton]):
                    n["loc"] = cadre.locator(f'[data-robot-nouveau="{n["cle"]}"]')
                    if _titre_de_fiches(n.get("section")):  # sous le titre « Derniers plans consultés » : des fiches
                        n["texte"] = "(donnee masquee)"
                    vus.append(n)
            except Exception:
                continue
        return vus

    def en_attente():
        for cadre in page.frames:
            try:
                if cadre.evaluate(JS_EN_ATTENTE):
                    return True
            except Exception:
                continue
        return False

    nouveaux = []
    for attente in (500, 800, 1200):  # un sous-menu peut arriver un peu après (animation, chargement)
        page.wait_for_timeout(attente)
        nouveaux = ce_qui_est_apparu()
        if nouveaux:
            break
    fin = time.time() + 3.5
    while not nouveaux and time.time() < fin and (patience or en_attente()):
        page.wait_for_timeout(500)
        nouveaux = ce_qui_est_apparu()
    for _ in range(3):  # il peut se remplir en plusieurs fois (titre tout de suite, éléments chargés ensuite)
        if not nouveaux:
            break
        page.wait_for_timeout(600)
        encore = ce_qui_est_apparu()  # (nouvelles marques : la liste précédente n'est plus valable)
        if not encore:
            break
        plus, nouveaux = len(encore) > len(nouveaux), encore
        if not plus:
            break
    return nouveaux


def _rang_menu(loc):
    """0 : apparu en survolant un menu ; 1 : entrée d'un menu ; 2 : ailleurs (titre, texte de la page)."""
    try:
        return loc.evaluate("e => e.closest('[data-robot-nouveau]') ? 0 : (e.closest('[data-robot-barre], nav, aside, "
                            "[role=navigation], [role=menu], [role=menubar], [role=menuitem]') ? 1 : 2)")
    except Exception:
        return 2


def _bouton_de_page(loc):
    """Un bouton (ou un élément d'un formulaire ou d'une fenêtre) : une action de la page, pas une entrée de menu."""
    try:
        return loc.evaluate("e => e.matches('button, [role=button], input[type=submit], input[type=button], input[type=image]') || "
                            "!!e.closest('button, [role=button], form, dialog, [role=dialog], [role=alertdialog], [aria-modal=true]')")
    except Exception:
        return True


def _apparu(trouve):
    """Un élément trouvé APRÈS un survol, dans ce qui vient d'apparaître (pas un bouton de la page du même nom)."""
    return _visible(trouve[2]) and _rang_menu(trouve[2]) == 0


def _rouvrir(page, entree, texte):
    """Le sous-menu s'est refermé : on survole à nouveau l'entrée et on retrouve l'élément « texte »."""
    _souris_au_repos(page)
    return next((n["loc"] for n in _survoler_et_voir(page, entree["loc"], entree["cadre"], entree["numero"])
                 if n["texte"] == texte), None)


def _explorer_menus(page, memoire=None):
    """Survole chaque entrée de menu, comme votre souris, et note ce qui apparaît (2 niveaux). JAMAIS de clic.
    Renvoie [(entrée, [(élément, [sous-éléments])])]. memoire : le même menu déjà exploré sur une autre page.
    Un groupe incertain (liste au bord de la page) dont aucune entrée n'ouvre de sous-menu n'est pas un menu :
    il est oublié (le relevé le traite comme le reste de la page)."""
    entrees = _menus_de_la_page(page)
    cle = tuple((e["nom"], e["fort"], e["groupe"]) for e in entrees)
    if memoire is not None and cle in memoire:
        resultat, a_oublier = memoire[cle]
        _oublier_groupes(entrees, a_oublier)
        return resultat
    explores = []
    for e in entrees:
        elements = []
        _souris_au_repos(page)
        try:
            nouveaux = _survoler_et_voir(page, e["loc"], e["cadre"], e["numero"])
        except Exception:
            nouveaux = []
        if nouveaux:
            _menu_confirme(e)
        a_ouvrir = 4
        for n in nouveaux:
            sous = []
            if n["sous"] and a_ouvrir > 0:
                a_ouvrir -= 1
                try:
                    loc = n["loc"] if _visible(n["loc"]) else _rouvrir(page, e, n["texte"])
                    if loc is not None:
                        sous = [x["texte"] for x in _survoler_et_voir(page, loc)]
                except Exception:
                    pass
            elements.append((n["texte"], sous))
        explores.append((e, elements))
    _souris_au_repos(page)
    ouverts = {e["groupe"] for e, elements in explores if elements}
    for groupe in {e["groupe"] for e, _ in explores if not e["fort"]} - ouverts:
        # groupe incertain qui n'a rien ouvert : sa première entrée est revue en attendant plus longtemps
        # (sous-menu chargé lentement la première fois)
        i, (e, _) = next((i, x) for i, x in enumerate(explores) if x[0]["groupe"] == groupe)
        try:
            nouveaux = _survoler_et_voir(page, e["loc"], e["cadre"], e["numero"], patience=True)
        except Exception:
            nouveaux = []
        _souris_au_repos(page)
        if nouveaux:
            _menu_confirme(e)
            explores[i] = (e, [(n["texte"], []) for n in nouveaux])
            ouverts.add(groupe)
    a_oublier = {e["groupe"] for e, _ in explores if not e["fort"] and e["groupe"] not in ouverts}
    _oublier_groupes(entrees, a_oublier)
    resultat = [(e["nom"], elements) for e, elements in explores if e["groupe"] not in a_oublier]
    if memoire is not None:  # (même vide : les mêmes entrées ne sont pas resurvolées sur chaque page)
        memoire[cle] = (resultat, a_oublier)
    return resultat


def _chercher_en_survolant(page, mot):
    """Le texte n'est pas (visible) sur la page : il apparaît peut-être en survolant un menu. Le robot survole
    chaque entrée (et les sous-menus qu'elle ouvre), SANS CLIQUER, jusqu'à le voir apparaître.
    Renvoie (éléments apparus qui portent ce texte, chemin pour rouvrir le menu, ce qui a été vu en survolant)."""
    vu = []
    fin = time.time() + 75  # au plus 1 minute et quart (poste lent, très grand menu)
    for e in _menus_de_la_page(page):
        if time.time() > fin:
            ecrire("   (le robot arrete de survoler : trop long)")
            break
        _souris_au_repos(page)
        try:
            nouveaux = _survoler_et_voir(page, e["loc"], e["cadre"], e["numero"])
        except Exception:
            continue
        if nouveaux:
            _menu_confirme(e)
        textes = [n["texte"] for n in nouveaux]
        if e["fort"] or nouveaux:
            vu.append((e["nom"], textes))
        chemin = [dict(e, textes=textes)]
        trouves = _candidats(page, mot, garder=_apparu)
        if trouves:
            return trouves, chemin, vu
        for n in [n for n in nouveaux if n["sous"]][:4]:
            try:
                loc = n["loc"] if _visible(n["loc"]) else _rouvrir(page, e, n["texte"])
                if loc is None:
                    continue
                _survoler_et_voir(page, loc)
            except Exception:
                continue
            trouves = _candidats(page, mot, garder=_apparu)
            if trouves:
                return trouves, chemin + [{"texte": n["texte"]}], vu
    _souris_au_repos(page)
    return [], [], vu


def _resurvoler(page, chemin, mot):
    """Rouvre le menu où le texte a été trouvé (il s'est refermé pendant la question). Sans cliquer."""
    entree = dict(chemin[0])
    if entree["loc"].count() != 1:  # la page a redessiné son menu : on le repère à nouveau
        pareille = next((e for e in _menus_de_la_page(page)
                         if e["cadre"] == entree["cadre"] and e["texte"] == entree["texte"]), None)
        if pareille is None:
            return []
        entree.update(loc=pareille["loc"], numero=pareille["numero"])
    _souris_au_repos(page)
    nouveaux = _survoler_et_voir(page, entree["loc"], entree["cadre"], entree["numero"])
    if len(chemin) > 1:
        loc2 = next((n["loc"] for n in nouveaux if n["texte"] == chemin[1]["texte"]), None)
        if loc2 is None:
            return []
        _survoler_et_voir(page, loc2)
    return _candidats(page, mot, garder=_apparu)


def _confirmer_entrees(page, trouves):
    """Le mot tapé est une entrée d'un menu incertain (ex : GATES dans le menu de gauche) : le robot la survole ;
    si un sous-menu s'ouvre, c'est bien un menu (il pourra cliquer dessus après votre accord)."""
    for _, cadre, loc in [t for t in trouves if _visible(t[2])][:3]:
        try:
            numero = loc.evaluate("e => { const b = e.closest('[data-robot-barre]');"
                                  " return b && !b.hasAttribute('data-robot-fort') ? Number(b.getAttribute('data-robot-barre')) : -1; }")
        except Exception:
            continue
        if numero < 0:
            continue
        entree = {"cadre": cadre, "numero": numero, "loc": cadre.locator(f'[data-robot-barre="{numero}"]')}
        try:
            if _survoler_et_voir(page, entree["loc"], cadre, numero):
                _menu_confirme(entree)
        except Exception:
            pass
        _souris_au_repos(page)


def _menu_du_compte(textes):
    """Un menu dont les éléments sont « Se déconnecter », « Mon profil »... : c'est le menu de la personne connectée,
    et son nom (le titre du menu) est une donnée."""
    return any(re.search(r"deconnex|deconnect|logout|log out|sign out|signout|mon profil|my profile|mon compte|my account|"
                         r"mes informations|mes preferences|changer de role|mot de passe|password|mon espace|"
                         r"fermer la session|^quitter\b|log ?off",
                         sans_accents(str(t)).lower()) for t in textes)


# une entrée faite d'une icône dont le sous-menu est une liste de fiches ou de messages (historique, notifications,
# favoris, messagerie) ou le menu de la personne connectée : le nom ENTIER de l'icône est comparé (« account_tree »,
# l'organigramme, reste un menu)
ICONE_DE_FICHES = re.compile(r"(?:^|[-_])(history|recent|restore|inbox|e?mail|envelope|notifications?|bell|bookmarks?|"
                             r"star|grade|favorite|heart|chat|forum|messages?|comments?|sms|cart|shopping|contacts?|"
                             r"schedule|access_time|watch_later)(?=$|[-_\d]|s\b)|^clock-rotate-left$", re.I)
# l'icône de la personne connectée (son menu, son profil) : jamais visitée, son sous-menu jamais écrit ; « users-gear »
# (l'administration des utilisateurs) ou « person-digging » (les chantiers) sont des modules
ICONE_DE_COMPTE = re.compile(r"(person|user|account|account_circle|account_box|manage_accounts|circle-user|user-circle|"
                             r"face|id-badge|avatar|user-tie|portrait|users|people)([-_](circle|alt|outline|outlined|solid|fill|"
                             r"filled|round|rounded|o))*", re.I)


def _contenu_a_masquer(nom, textes):
    """Le sous-menu d'une liste de fiches (Récents, Mes plans...), du menu du compte, ou d'une entrée masquée."""
    t = " ".join(str(nom or "").split())
    if t.lower().startswith("icone "):
        icone = re.sub(r" \d+$", "", t[6:]).strip()
        fiches = bool(ICONE_DE_FICHES.search(icone) or ICONE_DE_COMPTE.fullmatch(icone))
    else:
        fiches = _liste_de_donnees(nom)
    return (fiches or _menu_du_compte(textes) or
            _structure(nom) in (None, "(donnee masquee)", "(menu du compte)"))


def _contenu_survole(nom, textes, limite=12):
    """Ce qui est écrit d'un sous-menu : la liste, ou « (contenu masque) » si c'est une liste de fiches
    (Récents, Favoris, Mes plans, Notifications...) ou le menu du compte, ou un nombre si la liste est longue."""
    if not textes:
        return []
    if _contenu_a_masquer(nom, textes):
        return ["(contenu masque)"]
    if len(textes) > limite:
        return [f"({len(textes)} elements)"]
    return _liste(textes, limite)


def _nom_menu(nom, textes=()):
    """Le nom d'un menu tel qu'il est écrit (relevé, écran) : jamais le nom de la personne connectée."""
    return "(menu du compte)" if _menu_du_compte(textes) or nom == "(menu du compte)" else _structure(nom)


def _strategie_profil_impose(canal):
    """Stratégie d'entreprise « UserDataDir » : le navigateur ignorerait le profil vide du robot et
    ouvrirait VOTRE profil. Dans ce cas, le robot n'utilise pas ce navigateur."""
    if not WINDOWS or canal not in ("chrome", "msedge"):
        return False
    import winreg
    cle = r"SOFTWARE\Policies\Google\Chrome" if canal == "chrome" else r"SOFTWARE\Policies\Microsoft\Edge"
    for racine in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        for vue in (winreg.KEY_WOW64_64KEY, 0):
            try:
                with winreg.OpenKey(racine, cle, 0, winreg.KEY_READ | vue) as k:
                    winreg.QueryValueEx(k, "UserDataDir")
                    return True
            except OSError:
                continue
    return False


class Notes:
    """Résultats des étapes 6 à 8, affichés au fil de l'eau et sauvés pour le programme principal."""

    def __init__(self, fichier):
        self.fichier = fichier
        self.liste = []

    def noter(self, lettre, intitule, statut, detail=""):
        detail = masquer(detail, self.url) if hasattr(self, "url") else masquer(detail)
        self.liste = [n for n in self.liste if n[0] != lettre]
        self.liste.append([lettre, intitule, statut, detail])
        ecrire(f"   --> {intitule} : {statut}" + (f"  ({detail})" if detail else ""))
        if statut == "ECHEC":
            ecrire("       Ce n'est pas grave : c'est justement ce qu'on cherche a savoir. On continue.")
        self.sauver()

    def statut(self, lettre):
        return next((n[2] for n in self.liste if n[0] == lettre), None)

    def sauver(self):
        if self.fichier is None:
            return
        try:
            self.fichier.write_text(json.dumps(self.liste, ensure_ascii=False), encoding="utf-8")
        except OSError:
            pass


def _ouvrir_navigateur(p, notes):
    essais = []
    for canal, nom in (("chrome", "Chrome"), ("msedge", "Edge"), (None, "le navigateur integre du robot")):
        if _strategie_profil_impose(canal):
            essais.append(f"{nom} : profil impose par l'entreprise, non utilise")
            ecrire(f"   {nom} : l'entreprise impose son profil ; le robot ne l'utilise pas (pour ne pas toucher au votre).")
            continue
        ecrire(f"   Ouverture du navigateur du robot avec {nom} (jusqu'a 30 secondes). Ne touchez a rien...")
        try:
            options = {"headless": INVISIBLE, "slow_mo": 0 if INVISIBLE else 120, "timeout": 30000}
            if not INVISIBLE:  # fenêtre en grand : le portail s'affiche comme dans votre Chrome (menu de gauche ouvert)
                options["args"] = ["--start-maximized"]
            if canal:
                options["channel"] = canal
            navigateur = p.chromium.launch(**options)
            return navigateur, nom, canal, essais
        except Exception as e:
            raison = expliquer(e)
            essais.append(f"{nom} : {raison}")
            ecrire(f"   {nom} : impossible ({raison}). Le robot essaie autre chose...")
    notes.noter("N", "Navigateur du robot (page d'essai)", "ECHEC", " | ".join(essais))
    return None, "", None, essais


def _verifier_demarrage(p, canal, notes):
    """Chrome avec un profil (comme un robot qui garde votre connexion) : quels onglets s'ouvrent tout
    seuls au démarrage ? Une page Google imposée peut détourner un robot (« il ouvre Google et ne fait rien »)."""
    if canal is None:
        notes.noter("D", "Demarrage avec un profil", "NON DISPONIBLE", "seul le navigateur integre a demarre")
        return
    ecrire("   Controle rapide : une deuxieme fenetre s'ouvre puis se ferme toute seule (5 secondes)...")
    profil = tempfile.mkdtemp(prefix="robot_profil_")
    try:
        try:
            contexte = p.chromium.launch_persistent_context(profil, channel=canal, headless=INVISIBLE, timeout=30000)
        except Exception as e:
            notes.noter("D", "Demarrage avec un profil", "ECHEC", expliquer(e))
            return
        try:
            time.sleep(3)
            sortes = []
            for page in contexte.pages:
                adresse = (page.url or "").lower()
                hote = urlsplit(adresse).hostname or ""
                if adresse in ("", "about:blank") or adresse.startswith(("chrome://newtab", "chrome://new-tab-page",
                                                                           "edge://newtab", "chrome-search://")):
                    sortes.append("vide")
                elif "google." in hote:
                    sortes.append("google")
                elif adresse.startswith(("chrome://", "edge://")):
                    sortes.append("page du navigateur")
                else:
                    sortes.append("autre site")
            notes.noter("D", "Demarrage avec un profil", "OK",
                        f"{len(contexte.pages)} onglet(s) au demarrage : " + (", ".join(sortes) or "aucun"))
        finally:
            try:
                contexte.close()
            except Exception:
                pass
    finally:
        effacer_dossier(profil)


def _page_essai(page, nom, notes):
    page.set_content(PAGE_ESSAI)  # pas de fichier ni d'adresse file:// (souvent bloquées en entreprise)
    page.bring_to_front()
    champ = page.get_by_label("Votre prénom")
    _encadrer(champ, 800)
    champ.click()
    _taper_lentement(champ, "Robot")
    bouton = page.get_by_role("button", name="Valider")
    _encadrer(bouton, 800)
    bouton.click()
    page.get_by_text("Bonjour Robot, le clic a marché !").wait_for(timeout=5000)
    time.sleep(1.5)
    notes.noter("N", "Navigateur du robot (page d'essai)", "OK", f"avec {nom}")


def _analyser_portail(page, url, notes):
    """Étape 7. Renvoie vrai si le portail s'est affiché (l'étape 8 peut alors avoir lieu)."""
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
    except Exception as e:
        notes.noter("R", "Navigateur du robot sur le portail", "ECHEC", expliquer(e))
        return False
    page.bring_to_front()
    ecrire()
    ecrire("   La fenetre du robot (celle qui affichait la « Page d'essai du robot ») montre le portail.")
    ecrire("   SI le portail demande de vous connecter : connectez-vous DANS CETTE FENETRE-LA")
    ecrire("   (pas dans votre Chrome habituel). Le robot ne lit pas votre mot de passe, et cette")
    ecrire("   fenetre est oubliee a la fin.")
    pause("Quand vous voyez la page d'ACCUEIL du portail, revenez ici et appuyez sur Entree...")
    try:
        attendre_portail(page)
    except ErreurTache:
        notes.noter("R", "Navigateur du robot sur le portail", "ECHEC", "toujours sur la page de connexion")
        return False
    _marquer_ecouteurs(page)  # Chrome signale les éléments qui réagissent au clic
    total = {"cliquables": 0, "nouveaux": 0, "champs": 0, "ombres": 0, "dessins": 0}
    cadres_lus, cadres_illisibles, frameset = 0, 0, False
    meilleur, max_cliquables = None, -1
    for cadre in page.frames:
        try:
            s = cadre.evaluate(JS_STRUCTURE)
        except Exception:
            cadres_illisibles += 1
            continue
        cadres_lus += 1
        frameset = frameset or s["frameset"]
        for cle in total:
            total[cle] += s[cle]
        if s["cliquables"] > max_cliquables:
            meilleur, max_cliquables = cadre, s["cliquables"]
    detail = (f"{total['cliquables']} cliquables (dont {total['nouveaux']} trouves grace a Chrome), "
              f"{total['champs']} champs, {cadres_lus} cadre(s)"
              + (", frameset" if frameset else "")
              + (f", {cadres_illisibles} illisible(s)" if cadres_illisibles else "")
              + (f", {total['ombres']} composant(s) web" if total["ombres"] else "")
              + (f", {total['dessins']} zone(s) dessinee(s)" if total["dessins"] else ""))
    ecrire("   Le robot voit : " + detail)
    encadres = 0
    if meilleur is not None and max_cliquables > 0:
        ecrire("   Il encadre en ROUGE quelques elements qu'il saurait cliquer (SANS cliquer). Les cadres")
        ecrire("   restent affiches jusqu'a votre reponse : REGARDEZ LA FENETRE DU ROBOT (cliquez dessus")
        ecrire("   dans la barre des taches si elle est cachee), puis revenez repondre ici.")
        try:
            page.bring_to_front()
            meilleur.evaluate(JS_MARQUER_CLIQUABLES)
            marques = meilleur.locator("[data-robot-essai]")
            for i in range(min(marques.count(), 5)):
                _encadrer(marques.nth(i), 0)  # 0 = reste jusqu'au nettoyage
                encadres += 1
        except Exception:
            pass
    if encadres:
        time.sleep(1.5)
        revenez_ici()
        vu = oui("   Voyez-vous des cadres rouges sur le portail, dans la fenetre du robot ?")
        _nettoyer(page)
        notes.noter("R", "Navigateur du robot sur le portail", "OK" if vu else "PAS VU", detail)
    else:
        _nettoyer(page)
        notes.noter("R", "Navigateur du robot sur le portail", "ECHEC",
                    detail + " ; aucun element cliquable classique (menus faits en JavaScript ?)")
    return True


def _candidats(page, mot, garder=None):
    """Les éléments dont le texte est EXACTEMENT celui tapé, dans tous les cadres : texte complet,
    propre texte d'une entrée de menu (sous-menu à part), puis la même chose sans tenir compte des majuscules.
    Jamais un élément qui contient seulement ce mot (« plans » ne désigne pas « Suppression des plans »).
    garder : ne retenir que certains éléments (ex : ceux qui viennent d'apparaître en survolant un menu)."""
    exact_sans_casse = re.compile(r"^\s*" + re.escape(mot) + r"\s*$", re.IGNORECASE)
    jeton = secrets.token_hex(8)

    def propre_texte_sans_casse(cadre):
        cadre.evaluate(JS_TROUVER_PROPRE_TEXTE, [mot, jeton])
        return cadre.locator(f'[data-robot-trouve="{jeton}"]')

    def tolerant(mode):
        def chercher(cadre):
            cadre.evaluate(JS_TROUVER_TOLERANT, [mot, jeton, mode])
            return cadre.locator(f'[data-robot-trouve="{jeton}"]')
        return chercher

    facons = (
        lambda cadre: cadre.get_by_text(mot, exact=True),
        lambda cadre: cadre.locator("text=" + json.dumps(mot, ensure_ascii=False)),
        lambda cadre: cadre.get_by_text(exact_sans_casse),
        propre_texte_sans_casse,
        lambda cadre: cadre.get_by_title(mot, exact=True),
        lambda cadre: cadre.get_by_alt_text(mot, exact=True),
        tolerant("propre"), tolerant("entier"), tolerant("image"),
    )
    for chercher in facons:
        trouves = []
        for numero, cadre in enumerate(page.frames):
            try:
                loc = chercher(cadre)
                for i in range(min(loc.count(), 10)):
                    trouves.append((numero, cadre, loc.nth(i)))
            except Exception:
                continue
        if garder is not None:
            trouves = [t for t in trouves if garder(t)]
        if trouves:
            return trouves
    return []


def _montrer_ce_que_je_vois(page, vu_en_survolant=()):
    """Les noms des menus, liens et boutons que le robot voit sur la page, et ce qui apparaît en survolant
    les menus (pour retaper le bon texte)."""
    vus = []
    comptes = {nom for nom, textes in vu_en_survolant if _menu_du_compte(textes)}  # le nom de la personne connectée
    caches = {t for nom, textes in vu_en_survolant if _contenu_a_masquer(nom, textes) for t in textes}  # listes de fiches
    for cadre in page.frames:
        try:
            for t, section in cadre.evaluate(JS_CE_QUE_JE_VOIS):
                if t in caches or _titre_de_fiches(section):  # sous « Récents », « Favoris »... : des fiches
                    continue
                t = "(menu du compte)" if t in comptes else _structure(t)
                if t and t not in vus:
                    vus.append(t)
        except Exception:
            continue
    if vus:
        ecrire("   Ce que le robot voit comme menus, liens et boutons : " + " | ".join(vus[:30]))
    else:
        ecrire("   Le robot ne voit aucun menu ni lien sur cette page (page de connexion ? page vide ?).")
    survols = []
    for nom, textes in vu_en_survolant:
        t = _nom_menu(nom, textes)
        if t:
            s = _contenu_survole(nom, textes, 8)
            survols.append(t + " : " + (" | ".join(s) if s else "rien"))
    if survols:
        ecrire("   En survolant les menus, le robot a vu :")
        for ligne in survols[:15]:
            ecrire("     - " + ligne)
    elif not vu_en_survolant:
        nombre = len(_menus_de_la_page(page))
        ecrire(f"   Le robot a survole {nombre} entree(s) de menu : rien ne s'est ouvert a temps." if nombre else
               "   Le robot n'a repere aucun menu a survoler sur cette page.")


def _examiner(cadre, loc, premier=True):
    """Fixe l'élément (il ne peut plus changer entre la vérification et le clic) et dit s'il peut être cliqué.
    premier : premier examen (le mot trouvé est noté : c'est là que la souris se posera)."""
    jeton = secrets.token_hex(8)
    try:
        info = loc.evaluate(JS_EXAMINER, [jeton, premier])
    except Exception:
        return None, "texte illisible"
    cible = cadre.locator(f'[data-robot-cible="{jeton}"]')
    if info["refus"]:
        return cible, info["refus"]
    if not info["texte"]:
        return cible, "texte illisible"
    if interdit(info["texte"]) or interdit(info["attributs"]):
        return cible, "ressemble a une action qui modifie"
    return cible, ""


def _choisir(trouves):
    """Le premier élément que le robot a le droit de cliquer, les visibles d'abord.
    Renvoie ((numéro du cadre, cadre, cible) ou None, raison du refus, nombre de visibles)."""
    visibles = [t for t in trouves if _visible(t[2])]
    refus = ""
    for numero_cadre, cadre, loc in (visibles + [t for t in trouves if t not in visibles])[:10]:
        cible, raison = _examiner(cadre, loc)
        if not raison:
            return (numero_cadre, cadre, cible), "", len(visibles)
        refus = refus or raison
    return None, refus, len(visibles)


def _cliquer_prudemment(page, cadre, cible, chemin=None, mot=""):
    """Rouvre le menu si besoin, revérifie l'élément au dernier moment, pose la souris sur son texte,
    vérifie ce qui est VRAIMENT sous la souris, et seulement alors appuie. Renvoie "" si le clic est fait.
    chemin : le menu survolé où l'élément est apparu (à rouvrir s'il s'est refermé pendant la question)."""
    if chemin and (cible.count() != 1 or not _visible(cible)):
        choix, refus, _ = _choisir(_resurvoler(page, chemin, mot))
        if choix is None:
            return "le menu ne montre plus cet element" + (f" ({refus})" if refus else "")
        _, cadre, cible = choix
    if cible.count() != 1:
        return "l'element a disparu de la page"
    if not _visible(cible) and not _ouvrir_menus(cible):
        return "toujours cache : menu ferme"
    cible2, refus = _examiner(cadre, cible, premier=False)
    if refus or cible2 is None:
        return "refuse au dernier moment : " + (refus or "texte illisible")
    jeton = secrets.token_hex(8)
    point = cible2.evaluate(JS_POINT, jeton)
    cible2.hover(position=point, timeout=5000)  # la souris arrive sur le texte (un menu peut s'ouvrir)
    time.sleep(0.5)
    if not cible2.evaluate(JS_VERIFIER_POINT, [point["x"], point["y"], jeton]):
        return "refuse : autre chose est sous la souris (sous-menu ouvert par-dessus ?)"
    _, refus = _examiner(cadre, cible2, premier=False)
    if refus:
        return "refuse au dernier moment : " + refus
    avant = len(page.context.pages)
    page.mouse.down()  # la souris ne bouge plus : appui et relâchement là où on a vérifié
    page.mouse.up()
    page.wait_for_timeout(1500)
    return "nouvel onglet" if len(page.context.pages) > avant else ""


def _nettoyer(page, tout=True):
    """Retire les cadres rouges (et, avec tout=True, les marques du robot) de toutes les zones de la page."""
    for cadre in page.frames:
        try:
            cadre.evaluate(JS_NETTOYER, tout)
        except Exception:
            pass


def _clic_choisi(page, notes):
    ecrire()
    ecrire("=" * 70)
    ecrire(" ETAPE 8 : un clic que VOUS choisissez sur le portail")
    ecrire("=" * 70)
    ecrire("   Tapez le texte EXACT d'un MENU ou d'un LIEN du portail, tel qu'il est ecrit")
    ecrire("   (par exemple GATES), accents compris. Ca peut etre un element d'un sous-menu qui")
    ecrire("   s'ouvre quand la souris passe dessus (menu de gauche) : le robot survole les menus")
    ecrire("   pour le trouver. Il l'encadre en rouge, et vous demande avant de cliquer.")
    ecrire("   Il ne clique que sur des liens et des menus : jamais")
    ecrire("   sur un bouton de formulaire, une case, ni sur Supprimer, Enregistrer, Valider, OK...")
    ecrire("   Pendant que vous repondez, ne passez pas la souris sur la fenetre du robot.")
    codes = []  # un code court par essai, pour le RESULTAT (jamais le mot tapé)
    while len(codes) < 5:
        mot = demander("   Texte a chercher (Entree sans rien = finir) :")
        if not mot:
            break
        n = len(codes) + 1
        if len(mot) < 3 or sans_accents(mot).lower() in ("oui", "non", "ok", "yes"):
            ecrire("   Tapez le texte d'un menu ou d'un lien (3 lettres au moins), pas une reponse o / n.")
            continue
        if interdit(mot):
            ecrire("   Refuse : ce mot ressemble a une action qui modifie. Choisissez un menu ou un lien.")
            codes.append(f"{n}:refus-mot")
            continue
        _marquer_ecouteurs(page)
        _menus_de_la_page(page)  # le robot repère les entrées des menus (il sait alors ce qui est un menu)
        trouves = _candidats(page, mot)
        _confirmer_entrees(page, trouves)
        choix, refus, nb_visibles = _choisir(trouves) if trouves else (None, "", 0)
        chemin, vu = [], []
        if choix is None or not _visible(choix[2]):
            # pas d'élément visible autorisé : il est peut-être dans un sous-menu qui s'ouvre au survol
            ecrire("   Pas visible sur la page : le robot survole les menus (sans cliquer) pour le chercher.")
            ecrire("   Ne touchez pas a la fenetre du robot (1 minute au plus)...")
            trouves_survol, chemin_survol, vu = _chercher_en_survolant(page, mot)
            choix_survol = None
            if trouves_survol:
                choix_survol, refus_survol, nb_survol = _choisir(trouves_survol)
                if choix_survol is None:
                    refus = refus or refus_survol
            if choix_survol is not None:
                choix, nb_visibles, chemin = choix_survol, nb_survol, chemin_survol
            elif trouves_survol and trouves:   # le premier examen est refait (c'est lui qui guide le clic)
                choix, refus, nb_visibles = _choisir(trouves)
        if choix is None:
            if refus:
                ecrire(f"   Refuse par prudence : {refus}.")
                codes.append(f"{n}:refus({refus.split()[0].rstrip(',')})")
            else:
                ecrire("   Pas trouve, meme en survolant les menus.")
                _montrer_ce_que_je_vois(page, vu)
                codes.append(f"{n}:introuvable")
            _souris_au_repos(page)
            _nettoyer(page)
            continue
        numero_cadre, cadre, cible = choix
        ou = "dans la page" if numero_cadre == 0 else "dans une zone interieure de la page"
        if chemin:
            ou = f"en survolant le menu « {_nom_menu(chemin[0]['nom'], chemin[0].get('textes', ())) or '?'} »"
        cache = not _visible(cible)
        if cache:
            question = (f"   Trouve {ou}, mais cache dans un menu ferme. Le robot va survoler le menu pour l'ouvrir,"
                        " puis cliquer. Tapez o pour cliquer, n pour ne rien faire")
        else:
            if nb_visibles > 1:
                ecrire(f"   {nb_visibles} elements portent ce texte : le robot prend le premier qu'il a le droit de cliquer.")
            _encadrer(cible, 0)  # le cadre rouge reste pendant la question
            question = f"   Trouve {ou}, encadre en rouge. Tapez o pour que le robot clique dessus, n pour ne rien faire"
        revenez_ici()
        accord = oui(question, "n")
        _nettoyer(page, tout=False)  # le cadre rouge part ; la marque de la cible reste jusqu'au clic
        if not accord:
            codes.append(f"{n}:non-clique")
            continue
        try:
            souci = _cliquer_prudemment(page, cadre, cible, chemin, mot)
        except Exception as e:
            souci = "clic en echec (" + expliquer(e) + ")"
        if souci and souci != "nouvel onglet":
            ecrire("   Pas de clic : " + souci + ".")
            codes.append(f"{n}:pas-de-clic")
            continue
        if souci == "nouvel onglet":
            page = page.context.pages[-1]
            try:
                page.bring_to_front()
            except Exception:
                pass
            ecrire("   Le clic a ouvert un nouvel onglet : le robot continue dans ce nouvel onglet.")
        ecrire("   Clic fait. Regardez la fenetre du robot : la page a-t-elle reagi comme avec votre souris ?")
        revenez_ici()
        reagi = oui("   Ca a marche ?")
        codes.append(f"{n}:clic" + ("-survol" if cache or chemin else "") + ("" if reagi else "-sans-effet"))
        _nettoyer(page)
    _nettoyer(page)
    if not codes:
        notes.noter("C", "Clic choisi sur le portail", "SAUTE", "aucun mot donne")
    else:
        bons = [c for c in codes if ":clic" in c and "sans-effet" not in c]
        notes.noter("C", "Clic choisi sur le portail", "OK" if bons else "ECHEC", " ".join(codes))


def sous_programme_navigateur(dossier_echange):
    """Étapes 6 à 8, lancées avec le Python qui a le pilote de navigateur."""
    preparer_console()
    echange = Path(dossier_echange)
    notes = Notes(echange / "resultat.json")
    try:
        url = (echange / "adresse.txt").read_text(encoding="utf-8").strip()
        (echange / "adresse.txt").unlink()
    except OSError:
        url = ""
    notes.url = url
    etape = ["N", "Navigateur du robot (page d'essai)"]
    navigateur = None
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            navigateur, nom, canal, _ = _ouvrir_navigateur(p, notes)
            if navigateur is None:
                for lettre, intitule in (("D", "Demarrage avec un profil"), ("R", "Navigateur du robot sur le portail"),
                                         ("C", "Clic choisi sur le portail")):
                    notes.noter(lettre, intitule, "SAUTE", "navigateur du robot impossible a ouvrir")
                return 1
            contexte = navigateur.new_context(no_viewport=True)
            contexte.add_init_script(script_appris())
            page = contexte.new_page()
            try:
                _page_essai(page, nom, notes)
            except Exception as e:
                notes.noter("N", "Navigateur du robot (page d'essai)", "ECHEC", f"avec {nom} : " + expliquer(e))
            if notes.statut("N") == "OK":
                revenez_ici()
                if not oui("   Avez-vous vu le robot ecrire « Robot » et cliquer sur Valider ?"):
                    notes.noter("N", "Navigateur du robot (page d'essai)", "PAS VU", f"avec {nom}")
            etape[:] = ["D", "Demarrage avec un profil"]
            try:
                _verifier_demarrage(p, canal, notes)
            except Exception as e:  # ce contrôle ne doit jamais empêcher les étapes du portail
                notes.noter("D", "Demarrage avec un profil", "ECHEC", expliquer(e))
            titre_etape(7, "Le navigateur du robot sur votre portail")
            if not url:
                notes.noter("R", "Navigateur du robot sur le portail", "SAUTE", "pas d'adresse donnee")
                notes.noter("C", "Clic choisi sur le portail", "SAUTE", "pas d'adresse donnee")
            else:
                etape[:] = ["R", "Navigateur du robot sur le portail"]
                if _analyser_portail(page, url, notes):
                    etape[:] = ["C", "Clic choisi sur le portail"]
                    _clic_choisi(page, notes)
                    etape[:] = ["V", "Releve des pages (pour Claude)"]
                    titre_etape(9, "Releve des pages du portail (pour Claude)")
                    fichier, nombre = boucle_releve(page)
                    if fichier:
                        notes.noter("V", "Releve des pages (pour Claude)", "OK",
                                    f"{nombre} page(s) : Bureau > Robot > A_ENVOYER_releve.txt")
                    else:
                        notes.noter("V", "Releve des pages (pour Claude)", "SAUTE", "aucune page relevee")
                else:
                    notes.noter("C", "Clic choisi sur le portail", "SAUTE", "portail non affiche")
            pause("Fin des essais du navigateur. Entree pour fermer la fenetre du robot...")
            try:
                navigateur.close()
            except Exception:
                pass
    except KeyboardInterrupt:
        notes.noter(etape[0], etape[1], "ECHEC", "arrete par Ctrl+C")
        return 130
    except Exception as e:
        notes.noter(etape[0], etape[1], "ECHEC", expliquer(e))
    return 0


# ---------------------------------------------------------------------------- outils communs au relevé et aux tâches
def dossier_robot(*sous):
    """Bureau > Robot (et ses sous-dossiers) : tout ce que le robot écrit est là."""
    d = dossier_bureau() / "Robot"
    for s in sous:
        d = d / s
    d.mkdir(parents=True, exist_ok=True)
    return d


def adresse_retenue():
    try:
        url = (dossier_robot() / "adresse_portail.txt").read_text(encoding="utf-8-sig").strip().lstrip("\ufeff")
    except (OSError, UnicodeDecodeError):
        return ""
    try:
        return url if re.match(r"^https?://", url, re.I) and urlsplit(url).hostname else ""
    except ValueError:
        return ""


def retenir_adresse(url):
    if url:
        try:
            (dossier_robot() / "adresse_portail.txt").write_text(url, encoding="utf-8")
        except OSError:
            pass


def adresse_du_portail():
    """L'adresse retenue la dernière fois ; sinon on la demande (et on la retient)."""
    url = adresse_retenue()
    if url:
        ecrire(f"   Adresse du portail : {url}")
        ecrire("   (pour la changer : supprimez Bureau > Robot > adresse_portail.txt, puis relancez)")
        return url
    url = demander_adresse()
    retenir_adresse(url)
    return url


def _champ_mot_de_passe(page):
    for cadre in page.frames:
        try:
            if cadre.locator("input[type=password]").filter(visible=True).count():
                return True
        except Exception:
            try:
                loc = cadre.locator("input[type=password]")
                if any(_visible(loc.nth(i)) for i in range(min(loc.count(), 5))):
                    return True
            except Exception:
                continue
    return False


def attendre_chargement(page, secondes=10):
    for etat in ("domcontentloaded", "networkidle"):
        try:
            page.wait_for_load_state(etat, timeout=int(secondes * 1000))
        except Exception:
            pass
    try:
        page.wait_for_timeout(400)  # (attente qui laisse aussi passer les fenêtres du portail)
    except Exception:
        time.sleep(0.4)


def attendre_portail(page):
    """Si le portail affiche une page de connexion, on attend que vous vous connectiez DANS la fenêtre du robot."""
    attendre_chargement(page)
    attendu = False
    for _ in range(5):
        if not _champ_mot_de_passe(page):
            return attendu
        attendu = True
        ecrire()
        ecrire("   Le portail demande de vous connecter : connectez-vous DANS LA FENETRE DU ROBOT.")
        pause("Quand vous voyez l'accueil du portail, revenez ici et appuyez sur Entree...")
        attendre_chargement(page)
    raise ErreurTache("toujours sur la page de connexion")


def ouvrir_robot(p, lecture=False):
    """Le navigateur du robot (Chrome, sinon Edge, sinon le navigateur intégré) et sa page. Sans profil :
    rien de votre Chrome n'est lu ni gardé. lecture : visite (choix 5 et 6) : aucun téléchargement, et pas de
    « service worker » (ses échanges avec le portail échapperaient à la protection de la visite)."""
    notes = Notes(None)
    navigateur, nom, _, _ = _ouvrir_navigateur(p, notes)
    if navigateur is None:
        raise ErreurTache("le navigateur du robot ne s'ouvre pas sur ce poste (faites le test : choix 1)")
    options = {"no_viewport": True, "accept_downloads": not lecture}
    if lecture:
        options["service_workers"] = "block"
    try:
        contexte = navigateur.new_context(**options)
    except TypeError:  # (pilote de navigateur ancien)
        options.pop("service_workers", None)
        contexte = navigateur.new_context(**options)
    if lecture:  # (un « shared worker » échangerait avec le portail hors de la protection de la visite)
        contexte.add_init_script("try { delete window.SharedWorker; } catch (e) {}")
    contexte.add_init_script(script_appris())
    return navigateur, contexte, contexte.new_page()


# ---------------------------------------------------------------------------- relevé des pages (pour Claude)
# STRUCTURE seulement : noms des menus, onglets, champs, boutons, colonnes de tableaux, libellés d'une fiche.
# Jamais une valeur : ni le contenu des champs, ni les cellules des tableaux, ni les choix des listes.
JS_RELEVE = r"""() => {
  const vis = e => { const s = getComputedStyle(e); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const court = t => (t || '').replace(/\s+/g, ' ').trim();
  const sansChamps = n => { const c = n.cloneNode(true);
    c.querySelectorAll('select, option, input, textarea, button, script, style').forEach(x => x.remove()); return court(c.textContent); };
  // une ligne de DONNÉES d'un tableau : son contenu n'est jamais relevé. Ne sont pas des données : un tableau
  // de mise en page (qui contient d'autres tableaux), la ligne des titres, une ligne de formulaire, une ligne de fiche « Statut : »
  const ligneDeDonnees = e => { if (e.closest('[role=row], [role=gridcell]')) return true;
    const tr = e.closest('tr'); if (!tr) return false; const t = tr.closest('table');
    if (!t || t.querySelector('table') || tr.closest('thead')) return false;
    const premiere = t.querySelector('tr');
    const titres = !!t.querySelector('thead th') || (premiere && premiere.children.length > 1 &&
                   Array.from(premiere.children).every(c => c.tagName === 'TH'));
    if (tr === premiere && (titres || !premiere.querySelector('td'))) return false;     // la ligne des titres
    const corps = Array.from(t.querySelectorAll('tr')).filter(l => l.querySelector('td'));
    if (titres && corps.length >= 2) return true;      // tableau à titres : chaque ligne est un enregistrement
    if (tr.querySelector('input:not([type=hidden]), select, textarea')) return false;   // ligne de formulaire
    const c = tr.querySelector('td, th'); if (c && /:\s*$/.test(c.textContent || '')) return false;   // fiche
    return true; };
  const MENU = 'nav, header, [role=navigation], [role=menu], [role=menubar], [class*=menu i], [id*=menu i], [data-robot-barre]';
  const libelleDe = e => {
    if (e.id) { try { const l = document.querySelector('label[for="' + CSS.escape(e.id) + '"]'); if (l && sansChamps(l)) return sansChamps(l); } catch (x) {} }
    const p = e.closest('label'); if (p && sansChamps(p)) return sansChamps(p);
    const a = e.getAttribute('aria-label'); if (a) return court(a);
    const ph = e.getAttribute('placeholder'); if (ph) return court(ph) + ' (indice)';
    const cell = e.closest('td');
    // la cellule voisine sert de libellé seulement si elle finit par « : », ou si le champ est vide (formulaire de
    // recherche) : à côté d'un champ déjà rempli, c'est souvent une donnée de la fiche
    if (cell && cell.previousElementSibling && (/:\s*$/.test(cell.previousElementSibling.textContent || '') ||
        (cell.parentElement.children.length <= 3 && !(e.value || '').trim()))) {
      const t = sansChamps(cell.previousElementSibling); if (t && t.length < 50) return t; }
    const prev = e.previousElementSibling; if (prev && !['INPUT', 'SELECT', 'TEXTAREA'].includes(prev.tagName)) { const t = sansChamps(prev); if (t && t.length < 50) return t; }
    return e.getAttribute('title') || '(sans nom)';
  };
  const enregistrement = a => { const h = (a && a.getAttribute && a.getAttribute('href')) || '';
    return /^(tel|mailto|callto|sip):/i.test(h) || /[?&][^=&]+=([^&]*\d{2,}|[A-Za-z]+-\d+)/.test(h) || /\/\d{2,}(\/|$|\?)/.test(h); };
  const nomDe = e => enregistrement(e.closest('a') || e.querySelector('a')) ? '(donnee masquee)' : court(e.textContent);
  const r = { autresMenus: [], onglets: [], champs: [], boutons: [], tableaux: [], libelles: [], liens: [],
              cadres: document.querySelectorAll('iframe, frame').length };
  for (const e of document.querySelectorAll('nav a, [role=menuitem], [class*=menu i] a, [id*=menu i] a')) {
    if (!vis(e) || ligneDeDonnees(e) || e.closest('[data-robot-barre], [data-robot-nouveau]')) continue;
    r.autresMenus.push(nomDe(e));
  }
  // un bouton sans rôle dans une barre d'onglets : un onglet, sauf si la barre a de vrais onglets (c'est alors « x », « + »)
  const ongletSansRole = e => e.matches('[role=tablist] > button:not([role])') && !e.parentElement.querySelector(':scope > [role=tab]');
  for (const e of document.querySelectorAll('[role=tab], [role=tablist] > button:not([role]), [class*=onglet i] a, [class*=onglet i] li, [class*=tabs i] a, [class*=tabs i] li'))
    if (vis(e) && !ligneDeDonnees(e) && (!e.matches('[role=tablist] > button:not([role])') || ongletSansRole(e))) r.onglets.push(nomDe(e));
  for (const e of document.querySelectorAll('input, select, textarea')) {
    const type = (e.getAttribute('type') || '').toLowerCase();
    if (['hidden', 'submit', 'button', 'image', 'reset'].includes(type) || !vis(e) || ligneDeDonnees(e)) continue;
    let genre = e.tagName === 'SELECT' ? 'liste de ' + e.options.length + ' choix' : (e.tagName === 'TEXTAREA' ? 'texte long' : (type || 'texte'));
    if (type === 'password') genre = 'mot de passe';
    r.champs.push(libelleDe(e) + ' [' + genre + (e.required ? ', obligatoire' : '') + (e.readOnly || e.disabled ? ', lecture seule' : '') + ']');
  }
  const listeDeDonnees = e => !!e.closest('li') && !e.closest(MENU) && !e.closest('[class*=onglet i], [class*=tabs i], [role=tablist]');
  for (const e of document.querySelectorAll('button, input[type=submit], input[type=button], input[type=reset], [role=button]')) {
    if (!vis(e) || ligneDeDonnees(e) || listeDeDonnees(e) || e.closest('[data-robot-barre], [data-robot-nouveau]') ||
        e.matches('[role=tab]') || e.closest('[role=tab]') || ongletSansRole(e)) continue;   // (onglets : notés avec les onglets)
    const texte = e.tagName === 'INPUT' ? court(e.value) : court(e.textContent);
    const bulle = court(e.getAttribute('title') || e.getAttribute('aria-label') || '');
    r.boutons.push(texte && /[A-Za-zÀ-ÿ]/.test(texte) ? texte : (bulle ? '(icone) ' + bulle : '(icone sans nom)'));
  }
  for (const t of document.querySelectorAll('table, [role=grid]')) {
    if (!vis(t) || t.parentElement.closest('table td')) continue;
    let entetes = Array.from(t.querySelectorAll('thead th, [role=columnheader]'));
    if (!entetes.length) { const l = t.querySelector('tr'); if (l && l.children.length > 1 && Array.from(l.children).every(c => c.tagName === 'TH')) entetes = Array.from(l.children); }
    if (!entetes.length) continue;
    const lignes = Array.from(t.querySelectorAll('tr, [role=row]')).filter(l => l.querySelector('td, [role=gridcell]'));
    const colonnes = entetes.map((h, i) => { let n = court(h.textContent) || '(sans titre)';
      const exemple = lignes.find(l => l.children.length === entetes.length);
      if (exemple && exemple.children[i]) { const c = exemple.children[i];
        if (c.querySelector('a[href]')) n += ' (lien)';
        if (c.querySelector('input:not([type=hidden]), select, textarea')) n += ' (champ a remplir)';
        const actions = Array.from(c.querySelectorAll('button, input[type=button], input[type=submit], [role=button], a[title], img[title]'))
          .map(b => { const x = b.tagName === 'INPUT' ? court(b.value) : court(b.textContent);
                      return x && /[A-Za-zÀ-ÿ]/.test(x) && b.tagName !== 'A' ? 'bouton ' + x.split(' ').slice(0, 2).join(' ')
                        : (b.getAttribute('title') ? 'icone ' + court(b.getAttribute('title')).split(' ')[0] : ''); })
          .filter(Boolean);
        if (actions.length) n += ' (' + actions.join(', ') + ')'; }
      return n; });
    const vide = lignes.length === 1 && lignes[0].children.length === 1 &&
                 /aucun|aucune|vide|pas de|no data|no record|empty/i.test(lignes[0].textContent || '');
    r.tableaux.push({ colonnes: colonnes, lignes: vide ? 0 : lignes.length });
  }
  // libellés d'une fiche (« Statut : » ...) : le libellé seulement, jamais la valeur à côté
  for (const e of document.querySelectorAll('td, th, dt, label, span, b, strong, div, p')) {
    if (!vis(e) || ligneDeDonnees(e) || e.closest(MENU)) continue;
    let propre = ''; for (const n of e.childNodes) if (n.nodeType === 3) propre += n.textContent;
    propre = court(propre);
    if (propre && /:$/.test(propre) && propre.length <= 30) r.libelles.push(propre.replace(/\s*:$/, ''));
  }
  for (const a of document.querySelectorAll('a[href]'))
    if (vis(a) && !a.closest('td, th, [role=gridcell], [role=row], [data-robot-nouveau]') && !listeDeDonnees(a) && !a.closest(MENU) &&
        !a.closest('[class*=onglet i], [class*=tabs i], [role=tablist]'))
      r.liens.push(enregistrement(a) ? '(donnee masquee)' : court(a.textContent || a.getAttribute('title') || ''));
  return r;
}"""


# « Jean DUPONT », « DUPONT Jean », « Mme Durand », « Paul Martin » : un nom de personne n'est jamais de la structure
PRENOMS = set("""jean marie pierre paul jacques michel philippe alain nicolas christophe patrick daniel bernard eric
laurent frederic stephane david olivier sebastien thomas julien francois pascal thierry vincent christian didier
dominique gerard guillaume antoine alexandre maxime romain mathieu julie nathalie isabelle sylvie catherine sophie
christine martine francoise valerie sandrine stephanie anne celine helene veronique monique nicole chantal brigitte
caroline emilie aurelie claire camille laura sarah lea manon chloe emma louise lucas hugo louis gabriel arthur jules
martin marc yves luc andre rene henri georges claude roger kevin jeremy jerome fabien florian cedric arnaud benoit bruno
damien franck gilles herve ludovic mickael michael sylvain xavier yann yannick anthony adrien alexis aurelien baptiste
clement corentin dylan enzo quentin remi simon theo valentin william raphael tom nathan mathis ethan noah samuel victor
charles joseph fabrice serge denis patrice regis lionel joel pascale agnes alice amandine anais audrey carole cecile
charlotte christelle delphine elodie estelle eva fanny florence gaelle ines jade jeanne jessica justine karine laetitia
laurence lucie magali marion mathilde melanie nadia noemie oceane pauline sandra severine sonia virginie zoe juliette
margaux clara elise elisa ambre lina mia rose anna romane lola maeva nina myriam muriel odile josiane danielle jacqueline
mohamed ahmed karim samir nicolas sophie""".split())
NOM_DE_PERSONNE = re.compile(r"\b[A-ZÉÈÀÂÎÔÛÇ][a-zéèêëàâîïôûüç]+(?:-[A-Z][a-z]+)?\s+[A-ZÉÈÀÂÎÔÛÇ]{2,}(?![\w])|"
                             r"\b[A-ZÉÈÀÂÎÔÛÇ]{2,}\s+[A-ZÉÈÀÂÎÔÛÇ][a-zéèêëàâîïôûüç]+\b|"
                             r"\b(?:M\.|Mme|Mlle|Mr|Monsieur|Madame|Mademoiselle)\s+[A-ZÉÈ]|"
                             # « J. Dupont » (au début, ou après « de », « par »... ; jamais un sigle « P.L.U. Zonage »),
                             # « DUPONT J. », identifiant « marie.lefebvre »
                             r"(?:^|\b(?:de|du|par|pour|chez|avec|à)\s+)(?<![A-Z]\.)[A-Z]\.(?![A-Z]\.)\s*[A-ZÉÈ][a-zéèêëàâîïôûüç]+|"
                             r"\b[A-ZÉÈ]{3,}\s+[A-Z]\.$|\b[a-z]{2,}\.[a-z]{2,}\b")
DONNEES_DE_BASE = r"\d{3,}|@|https?://|\\|\b[A-Za-z]{1,6}-\d+|\d{2}[ .]\d{2}[ .]\d{2}|\d{1,2}/\d{1,2}/\d{2,4}"
AUTRES_DONNEES = re.compile(DONNEES_DE_BASE + r"|\b[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,6}\b")  # (+ serveur, j.dupont, fichier.pdf)
# pour décrire la STRUCTURE (classes « div.ag-cell.ag-row ») : les points y séparent des classes, pas un nom de serveur
AUTRES_DONNEES_CSS = re.compile(DONNEES_DE_BASE)


def _prenom_nom(texte):
    mots_ = re.findall(r"[A-Za-zÀ-ÿ'-]+", texte)

    def prenom(m):  # (« Jean-Pierre », « Marie-Claire » : toutes les parties ; jamais « Saint-Denis »)
        parties = [x for x in sans_accents(m).lower().split("-") if x]
        return bool(parties) and all(x in PRENOMS for x in parties)
    return any((prenom(a) and b[:1].isupper() and len(b) > 1) or
               (prenom(b) and a[:1].isupper() and len(a) > 2 and a.lower() not in ("de", "du", "la")) or
               (prenom(a) and a[:1].isupper() and re.fullmatch(r"[A-Z]\.?", b))   # « Agnès K. »
               for a, b in zip(mots_, mots_[1:]))


def _structure(texte):
    """Un nom d'élément du portail, gardé seulement s'il ressemble à de la structure (pas à une donnée)."""
    t = " ".join(str(texte or "").split())
    if not t or len(t) > 45:
        return None
    if t == "(donnee masquee)":
        return t
    if AUTRES_DONNEES.search(t) or NOM_DE_PERSONNE.search(t) or _prenom_nom(t):
        return "(donnee masquee)"
    return t


def _liste_de_donnees(nom):
    """« Récents », « Favoris », « Mes plans », « Notifications »... : ce sous-menu liste des fiches ou des messages
    (des données), pas des menus."""
    return bool(re.search(r"recent|recemm|favori|historique|dernier|consult|epingl|pinned|suivi|watch|\blast\b|^mes |^my |"
                          r"notif|alerte|\balert|message|compte|profil|account|utilisateur|\buser\b|history|bookmark|"
                          r"signet|inbox|boite|corbeille|panier|\btaches?\b|todo", sans_accents(str(nom or "")).lower().strip()))


def _liste(elements, limite=40):
    vus = []
    for e in elements:
        t = _structure(e)
        if t and t not in vus:
            vus.append(t)
    return vus[:limite]


def relever_page(page, memoire=None):
    """La structure de la page affichée dans la fenêtre du robot (tous les cadres). Les menus sont survolés
    (jamais cliqués) pour noter leurs sous-menus. memoire : les menus déjà explorés sur les pages précédentes."""
    return _releve_et_menus(page, memoire)[0]


def _releve_et_menus(page, memoire=None):
    """Le relevé de la page, et ses menus tels que survolés : [(entrée, [(élément, [sous-éléments])])]."""
    _marquer_ecouteurs(page)
    try:
        explores = _explorer_menus(page, memoire)
    except Exception:
        explores = []
    total = {"autresMenus": [], "onglets": [], "champs": [], "boutons": [], "tableaux": [], "libelles": [], "liens": [], "cadres": 0}
    for cadre in page.frames:
        try:
            r = cadre.evaluate(JS_RELEVE)
        except Exception:
            continue
        for cle in total:
            if cle == "cadres":
                total[cle] += r[cle]
            else:
                total[cle] += r[cle]
    _nettoyer(page)
    # « UD & Plans > Recherche de plans ; Mes plans ; Plans récents > (Cette semaine ; Ce mois) »
    menus = []
    for nom, elements in explores:
        textes = [texte for texte, _ in elements]
        t = _nom_menu(nom, textes)
        if not t:
            continue
        parties, deja = [], set()
        if _contenu_a_masquer(nom, textes) or len(elements) > 25:
            parties = _contenu_survole(nom, textes, 25)
        else:
            for texte, sous in elements:
                s = _structure(texte)
                if not s or s in deja:
                    continue
                deja.add(s)
                ss = [x[1:-1] if re.fullmatch(r"\((contenu masque|\d+ elements)\)", x) else x
                      for x in _contenu_survole(texte, sous, 12)]
                parties.append(s + (" > (" + " ; ".join(ss) + ")" if ss else ""))
        menus.append(t + (" > " + " ; ".join(parties) if parties else ""))
    tableaux = []
    for tab in total["tableaux"]:
        cols = [c for c in (_structure(c) for c in tab["colonnes"]) if c]
        if cols:
            tableaux.append(" | ".join(cols) + f"   ({tab['lignes']} ligne(s))")
    return {
        "menus": menus, "autres menus": _liste(total["autresMenus"]), "onglets": _liste(total["onglets"]),
        "champs": _liste(total["champs"]), "boutons": _liste(total["boutons"]), "tableaux": tableaux,
        "libelles de fiche": _liste(total["libelles"]), "liens": _liste(total["liens"], 25),
        "zones interieures (iframes)": total["cadres"],
    }, explores


def _lignes_releve(r, menus_vus):
    """Les lignes d'une page relevée. menus_vus : [menus de la page d'avant] (des menus identiques : « les memes »)."""
    lignes = []
    if r["menus"] and r["menus"] != menus_vus[0]:
        lignes.append("  Menus (survoles par le robot, sans cliquer) :")
        lignes += ["    " + m for m in r["menus"]]
        menus_vus[0] = r["menus"]
    elif r["menus"]:
        lignes.append("  Menus : les memes")
    for cle in ("autres menus", "onglets", "champs", "boutons", "libelles de fiche", "liens"):
        if r[cle]:
            lignes.append(f"  {cle.capitalize()} : " + " ; ".join(r[cle]))
    for t in r["tableaux"]:
        lignes.append("  Tableau : " + t)
    if r["zones interieures (iframes)"]:
        lignes.append(f"  Zones interieures (iframes) : {r['zones interieures (iframes)']}")
    return lignes


def ecrire_releve(pages):
    lignes = [f"RELEVE DU PORTAIL - pour Claude - {datetime.now().strftime('%d/%m/%Y %H:%M')}",
              "Structure seulement (noms des menus, champs, boutons, colonnes). Aucune valeur.",
              "Relisez : une vraie valeur (nom de client, numero de plan...) -> effacez ce mot seulement.", ""]
    menus_vus = [None]
    for numero, (nom, r) in enumerate(pages, 1):
        lignes.append(f"PAGE {numero} : {nom}")
        lignes += _lignes_releve(r, menus_vus)
        lignes.append("")
    for nom in ("A_ENVOYER_releve.txt", "A_ENVOYER_releve_2.txt", "A_ENVOYER_releve_3.txt"):
        fichier = dossier_robot() / nom
        try:
            fichier.write_text("\n".join(lignes), encoding="utf-8")
            return fichier
        except OSError:
            continue
    raise ErreurTache("impossible d'ecrire le releve dans Bureau > Robot")


def boucle_releve(page):
    """Vous allez sur une page dans la fenêtre du robot, vous la nommez, le robot la relève. Entrée vide = fini."""
    pages, memoire = [], {}
    ecrire("   Dans la FENETRE DU ROBOT, allez sur une page du portail (comme d'habitude, a la souris).")
    ecrire("   Puis revenez ici, tapez un nom court pour cette page et Entree.")
    ecrire("   Exemples : recherche plans, liste resultats, fiche plan, onglet references.")
    ecrire("   Entree sans rien = fini.")
    while len(pages) < 30:
        nom = demander(f"   Nom de la page {len(pages) + 1} (Entree sans rien = fini) :")
        if not nom:
            if pages or not oui("   Aucune page relevee pour l'instant. Continuer le releve ?", "o"):
                break
            continue
        try:
            page = page.context.pages[-1] if page.context.pages else page  # l'onglet le plus récent
            page.bring_to_front()
            attendre_chargement(page, 5)
            ecrire("   Releve en cours : le robot survole les menus (sans cliquer).")
            ecrire("   Ne touchez ni a la souris ni a la fenetre du robot (1 minute au plus)...")
            avant = page.url
            r = relever_page(page, memoire)
            if page.url != avant:
                ecrire("   --> la page a change pendant le releve : revenez sur cette page, puis retapez son nom.")
                continue
            pages.append((_structure(nom) or f"page {len(pages) + 1}", r))
            nb = sum(len(v) for k, v in r.items() if isinstance(v, list))
            ecrire(f"   --> page relevee ({nb} elements). Allez sur la page suivante, ou Entree sans rien = fini.")
        except Exception as e:
            ecrire("   --> cette page n'a pas pu etre relevee : " + expliquer(e))
    if not pages:
        return None, 0
    return ecrire_releve(pages), len(pages)


# ---------------------------------------------------------------------------- choix 4 : voir ce que le robot trouve cliquable
# Le robot ENCADRE tout ce qu'il trouve cliquable (rouge : ce qu'il trouvait déjà ; orange : trouvé en plus grâce à
# Chrome, aux applications, ou appris ; vert : ce que vous lui avez montré). Vous posez la souris 3 secondes sur un
# élément cliquable SANS cadre : il le note et l'apprend. Rien n'est cliqué. Rapport : structure seulement.
JS_VOIR = r"""([appris, bandeau]) => {""" + JS_DETECTION + r"""
  if (Array.isArray(appris)) window.__robotAppris = appris;
  const liste = detecter();
  window.__robotDetectes = new WeakSet(liste.map(x => x.e));
  let boite = document.getElementById('robot-cadres');
  if (!boite) {
    boite = document.createElement('div'); boite.id = 'robot-cadres';
    boite.style.cssText = 'position:fixed;left:0;top:0;width:100vw;height:100vh;overflow:hidden;pointer-events:none;z-index:2147483647';
    document.documentElement.appendChild(boite);
  }
  const couleur = { avant: '#dc2626', chrome: '#f59e0b', attribut: '#f59e0b', appris: '#f59e0b' };
  const cadre = (r, c, epais) => { const d = document.createElement('div');
    d.style.cssText = 'position:absolute;box-sizing:border-box;pointer-events:none;border:' + epais + 'px solid ' + c +
      ';left:' + (r.left - 2) + 'px;top:' + (r.top - 2) + 'px;width:' + (r.width + 4) + 'px;height:' + (r.height + 4) + 'px';
    return d; };
  // seulement ce qui se VOIT : dans l'écran, et pas caché sous autre chose (ligne défilée hors d'une grille...)
  const lie = (a, b) => { for (let y = a; y; y = y.parentElement || (y.getRootNode && y.getRootNode() !== document ? y.getRootNode().host : null))
                            if (y === b) return true; return false; };
  const sevoit = x => { const r = x.r; if (r.bottom <= 0 || r.right <= 0 || r.top >= innerHeight || r.left >= innerWidth) return false;
    const h = document.elementFromPoint(Math.min(Math.max(r.left + r.width / 2, 0), innerWidth - 1),
                                        Math.min(Math.max(r.top + r.height / 2, 0), innerHeight - 1));
    return !!h && (lie(h, x.e) || lie(x.e, h)); };
  const visibles = liste.filter(sevoit);
  const morceaux = visibles.slice(0, 1500).map(x => cadre(x.r, couleur[x.s], 3));
  for (const e of (window.__robotMontres || [])) if (e.isConnected) morceaux.push(cadre(e.getBoundingClientRect(), '#16a34a', 5));
  // bande jaune en bas de l'écran : ce que fait le robot (si on la voit, le dessin marche)
  if (bandeau) {
    const b = document.createElement('div');
    b.textContent = bandeau.replace('{n}', String(liste.length));
    b.style.cssText = 'position:absolute;left:50%;bottom:8px;transform:translateX(-50%);max-width:90vw;pointer-events:none;' +
      'background:#fde047;color:#111;font:bold 15px/20px Arial,sans-serif;padding:6px 12px;border:2px solid #111;border-radius:6px;' +
      'box-shadow:0 2px 8px rgba(0,0,0,.4);white-space:nowrap;overflow:hidden;text-overflow:ellipsis';
    morceaux.push(b);
  }
  boite.replaceChildren(...morceaux);
  const n = s => liste.filter(x => x.s === s).length;
  let ombres = 0;
  for (const e of document.querySelectorAll('*')) if (e.shadowRoot) ombres++;
  // la page entière écoute les clics : body, html, ou un bloc qui couvre presque tout l'écran (racine de l'application)
  const ecran = innerWidth * innerHeight;
  const delegation = Array.from(document.querySelectorAll('[data-robot-ecoute~=click]')).some(e => {
    const r = e.getBoundingClientRect(); return e === document.body || e === document.documentElement || r.width * r.height >= 0.9 * ecran; });
  return { total: liste.length, avant: n('avant'), chrome: n('chrome'), attribut: n('attribut'), appris: n('appris'), ombres: ombres,
           delegation: delegation, dessines: Math.min(visibles.length, 1500) };
}"""

JS_EFFACER_CADRES = r"""() => { const b = document.getElementById('robot-cadres'); if (b) b.remove(); window.__robotMontres = []; }"""

# La souris est posée depuis « delai » ms sur un élément SANS cadre : le robot le note (vert) et décrit sa STRUCTURE
# (balise, classes, attributs, curseur, écouteurs, parents). Son texte est renvoyé pour être masqué par le robot.
JS_POINTE = r"""([delai, nouveau]) => {
  // les écouteurs de la souris sont reposés à chaque fois (le navigateur ne pose jamais deux fois le même) : une page
  // qui remplace son contenu les efface
  if (!window.__robotMm) {
    // (Chrome envoie aussi un « mouvement » quand la page change sous la souris immobile, par exemple quand le robot
    //  redessine ses cadres : seul un VRAI déplacement remet le compteur des 3 secondes à zéro)
    window.__robotMm = ev => {
      const e = ev.composedPath ? ev.composedPath()[0] : ev.target, p = window.__robotPointeur;
      if (p && p.e && p.x === ev.clientX && p.y === ev.clientY) p.e = e;
      else window.__robotPointeur = { e: e, t: performance.now(), x: ev.clientX, y: ev.clientY }; };
    window.__robotMo = ev => { if (!ev.relatedTarget) window.__robotPointeur = { e: null, t: performance.now() }; };
  }
  document.addEventListener('mousemove', window.__robotMm, true);
  document.addEventListener('mouseout', window.__robotMo, true);
  window.__robotMontres = window.__robotMontres || [];
  if (nouveau || !window.__robotPointeur) { window.__robotPointeur = { e: null, t: performance.now() }; return null; }
  const p = window.__robotPointeur, e = p.e;
  if (!e || e.nodeType !== 1 || performance.now() - p.t < delai) return null;
  if (e === document.body || e === document.documentElement || (e.closest && e.closest('#robot-cadres'))) return null;
  const det = window.__robotDetectes || new WeakSet();
  for (let x = e; x; x = x.parentElement || (x.getRootNode && x.getRootNode() !== document && x.getRootNode().host) || null) {
    const b = x.getBoundingClientRect();
    if (det.has(x) && b.width * b.height <= innerWidth * innerHeight * 0.12) return null;  // déjà encadré (lui ou son bouton)
  }
  if (window.__robotMontres.some(m => m === e || m.contains(e))) return null;
  const r = e.getBoundingClientRect();
  if (r.width * r.height > innerWidth * innerHeight * 0.3) return null;   // un grand bloc (fond de page) : pas un élément
  window.__robotMontres.push(e);
  window.__robotPointeur = { e: null, t: performance.now() };
  // seulement des classes PARTAGÉES par plusieurs éléments (un genre d'élément), jamais l'identifiant d'un élément :
  // une classe ou un identifiant propre à un seul élément peut contenir une donnée (nom de dossier, numéro...)
  const classes = x => (typeof x.className === 'string' ? x.className : ((x.className && x.className.baseVal) || ''))
                         .split(/\s+/).filter(c => c && !/\d{2,}/.test(c) && document.getElementsByClassName(c).length >= 2);
  const sig = x => x.tagName.toLowerCase() + classes(x).slice(0, 4).map(c => '.' + c).join('');
  const chemin = [];
  for (let x = e.parentElement, n = 0; x && x !== document.body && n < 6; x = x.parentElement, n++) chemin.push(sig(x));
  let parentEcoute = '';
  for (let x = e.parentElement, n = 1; x && n <= 8 && !parentEcoute; x = x.parentElement, n++)
    if (x.hasAttribute('data-robot-ecoute')) parentEcoute = n + ' niveau(x) au-dessus : ' + x.getAttribute('data-robot-ecoute') +
      ' (' + Math.round(x.getBoundingClientRect().width) + 'x' + Math.round(x.getBoundingClientRect().height) + ')';
  // un nom de classe stable pour reconnaître les éléments du même genre (jamais un état : actif, ouvert, survolé...)
  const STABLE = c => !/^(ng-|cdk-|mat-focus|mat-ripple|active|actif|selected|hover|focus|focused|open|opened|ouvert|show|in$|collapsed|expanded|disabled)/i.test(c);
  const base = x => { const c = classes(x).filter(STABLE).slice(0, 2);
                      return c.length ? x.tagName.toLowerCase() + c.map(k => '.' + CSS.escape(k)).join('') : ''; };
  let selecteur = base(e);
  if (!selecteur && e.parentElement && base(e.parentElement)) selecteur = base(e.parentElement) + ' > ' + e.tagName.toLowerCase();
  try { const n = selecteur ? document.querySelectorAll(selecteur).length : 0; if (n < 1 || n > 60) selecteur = ''; }
  catch (x) { selecteur = ''; }
  const s = getComputedStyle(e);
  return { signature: sig(e), chemin: chemin.join(' < '), role: e.getAttribute('role') || '',
           attributs: Array.from(e.attributes).map(a => a.name)
                        .filter(n => !/^(class|style|id)$/.test(n) && !n.startsWith('data-robot') && !n.startsWith('_ng')).slice(0, 12).join(' '),
           curseur: s.cursor, taille: Math.round(r.width) + 'x' + Math.round(r.height),
           ecoute: e.getAttribute('data-robot-ecoute') || '', survol: e.getAttribute('data-robot-survol') || '',
           parentEcoute: parentEcoute, ombre: e.getRootNode() !== document,
           mots: (e.innerText || e.textContent || '').trim().split(/\s+/).filter(Boolean).length, selecteur: selecteur };
}"""


def lire_appris():
    """Les genres d'éléments que vous avez montrés au robot (Bureau > Robot > appris.txt), un par ligne."""
    try:
        lignes = (dossier_robot() / "appris.txt").read_text(encoding="utf-8").splitlines()
    except (OSError, UnicodeDecodeError):
        return []
    return [l.strip() for l in lignes if l.strip() and not l.strip().startswith("#")][:200]


def script_appris():
    """À chaque page du robot : ce qu'il a appris (pour reconnaître ces éléments comme cliquables)."""
    return "window.__robotAppris = " + json.dumps(lire_appris(), ensure_ascii=False) + ";"


def ajouter_appris(selecteurs):
    actuels = lire_appris()
    nouveaux = [s for s in dict.fromkeys(selecteurs) if s and s not in actuels and "\n" not in s and len(s) <= 120]
    if nouveaux:
        try:
            with open(dossier_robot() / "appris.txt", "a", encoding="utf-8") as f:
                f.write("\n".join(nouveaux) + "\n")
        except OSError:
            return actuels
    return actuels + nouveaux


FIN_DU_CLAVIER = [False]


def ligne_tapee(tampon):
    """Le clavier SANS attendre (la fenêtre du robot continue d'être suivie) : la ligne tapée quand Entrée est appuyée,
    sinon None."""
    if WINDOWS:
        import msvcrt
        while msvcrt.kbhit():
            c = msvcrt.getwch()
            if c in ("\r", "\n"):
                ligne = "".join(tampon)
                tampon.clear()
                print(flush=True)
                return ligne
            if c == "\x03":
                raise KeyboardInterrupt
            if c == "\b":
                if tampon:
                    tampon.pop()
                    print("\b \b", end="", flush=True)
                continue
            if c in ("\x00", "\xe0"):  # touche spéciale (flèche...) : ignorée
                msvcrt.getwch()
                continue
            tampon.append(c)
            print(c, end="", flush=True)
        return None
    import select
    if not FIN_DU_CLAVIER[0] and select.select([sys.stdin], [], [], 0)[0]:
        ligne = sys.stdin.readline()
        if not ligne:  # fin de l'entrée (pas de clavier) : ce n'est pas une Entrée tapée
            FIN_DU_CLAVIER[0] = True
            return None
        return ligne.rstrip("\r\n")
    return None


def _dessiner(page, appris, bandeau=""):
    """Cadres (et bande jaune) dans toutes les zones de la page. Renvoie les nombres, et l'erreur s'il y en a une."""
    total = {"total": 0, "avant": 0, "chrome": 0, "attribut": 0, "appris": 0, "ombres": 0, "delegation": False, "cadres": 0,
             "dessines": 0, "erreur": ""}
    for cadre in page.frames:
        try:
            r = cadre.evaluate(JS_VOIR, [appris, bandeau if cadre == page.main_frame else ""])
        except Exception as e:
            total["erreur"] = total["erreur"] or expliquer(e)
            continue
        total["cadres"] += 1
        for cle in ("total", "avant", "chrome", "attribut", "appris", "ombres", "dessines"):
            total[cle] += r[cle]
        total["delegation"] = total["delegation"] or r["delegation"]
    total["delegation"] = total["delegation"] or _CLIC_DELEGUE.get(id(page), False)
    return total


def _effacer_cadres(page):
    for cadre in page.frames:
        try:
            cadre.evaluate(JS_EFFACER_CADRES)
        except Exception:
            pass


def _resume_montre(d):
    return _sans_donnee(d["signature"], css=True)[:60] + f" ({d.get('mots', 0)} mot(s), taille {d['taille']})"


# les mots des classes d'éléments qu'on peut écrire (structure courante des pages et des bibliothèques connues) ;
# un autre mot (le préfixe propre à l'entreprise ou à l'application : « acme-nav ») devient « * »
MOTS_CSS = set("""a abbr app apps area article aside nav navbar navigation navs menu menus menuitem menubar submenu sub item
items list lists listitem link links rail side sidebar sidenav drawer bar bars top bottom left right header footer main
content contents container wrapper wrap inner outer body panel panels pane card cards tile tiles grid row rows col cols
column columns cell cells table tbody thead tr td th tab tabs tablist tabpanel btn button buttons icon icons ico label
labels title titles text texts active selected disabled open opened closed close expanded collapsed collapse hover
focus focused flyout popup popover dropdown dropdowns overlay dialog modal toolbar tool tools group groups section
sections block box area zone region layout page pages view views root shell host component widget module feature home
dashboard user avatar badge badges search input field fields form select option options checkbox radio toggle switch
chip chips tag tags tooltip trigger handle arrow chevron caret divider separator spacer small large sm md lg xl xs full
flex d inline hidden visible show hide fade in out primary secondary tertiary default light dark mini compact normal
height width size wide narrow mobile desktop level parent child children first last odd even start end center middle
align justify auto fixed sticky absolute relative cdk mat mdc ng star inserted ag theme quartz alpine balham material p
k dx e x v el ant q bs ui mud fa fas far mdi bi ripple indicator state layer unbounded leading trailing touch target
outline outlined filled raised flat stroked elevation z mt mb ml mr mx my pt pb pl pr px py m w h font bold color bg
border rounded shadow cursor pointer clickable router current highlight highlighted vertical horizontal scroll
scrollable overflow viewport backdrop portal accordion expansion panel header headers footer label logo brand breadcrumb
breadcrumbs pagination pager paginator step steps stepper progress spinner loader loading empty placeholder
notification notifications toast snackbar alert alerts message messages info warning error success danger help
action actions control controls wrapper body cell normal value values key heading subheading caption legend fieldset
tree node nodes leaf branch expand expander icon toggler hamburger burger more kebab ellipsis overflow content
mdc ng c tns pristine valid invalid touched untouched dirty focused animating animation enter leave done collapsed
mat focus indicator primary list item unscoped content text secondary line lines density outer inner component host
root base core common shared global layout page app view frame frames iframe svg img image images pic thumb thumbnail
mui notranslate css ngx kendo chakra mantine prime primeng pi dx devextreme syncfusion vuetify quasar bootstrap
""".split())


# les balises HTML et les noms d'attributs courants (une balise propre à l'application, « acme-shell », devient « *-shell »)
BALISES = set("""abbr address article aside b body button canvas caption code dd details dialog div dl dt em fieldset
figure footer form frame h1 h2 h3 h4 h5 h6 header html i iframe img input label legend li main mark menu nav ol
optgroup option output p path pre section select small span strong sub summary sup svg table tbody td template textarea
tfoot th thead tr u ul use video g circle rect polygon tabindex role aria labelledby describedby haspopup expanded
controls routerlink router reflect href target rel title type name value tooltip mattooltip click onclick jsaction
jsname jscontroller sref toggle placement original readonly checked for true false pointer default url gridcell
columnheader rowheader navigation menubar listitem treeitem treegrid presentation heading banner complementary region
alertdialog checkbox switch textbox combobox listbox searchbox spinbutton slider tabpanel status none data href id
action pressed owns posinset setsize level current selected disabled hidden live atomic busy invalid required
multiselectable orientation valuemin valuemax valuenow valuetext modal application menuitemcheckbox menuitemradio
radiogroup rowgroup document feed img list math note row separator term definition directory polyline ellipse tspan
text map area defs clippath mask pattern symbol lineargradient marker foreignobject hx x v wire phx routerlinkactive
matmenutriggerfor matbadge cdkdrag src alt download lang dir mouse over out enter leave move key press dbl context
pointer touch start end focus blur change input submit wheel drag drop menutrigger trigger get post put patch delete
swap""".split())
EVENEMENTS = set("""click dblclick mousedown mouseup mouseover mouseout mouseenter mouseleave mousemove keydown keyup keypress
focus blur change input submit contextmenu pointerdown pointerup pointerover pointerenter pointerleave touchstart
touchend wheel load""".split())


def _css_generique(t):
    """Une description d'éléments (« acme-shell > a.acme-nav-item#plans », noms d'attributs) réduite aux mots de
    structure connus : balises, classes, identifiants ; le préfixe propre à l'entreprise (souvent son nom) devient
    « * »."""
    def mot(m):
        ws = mots(m.group(0)).split()
        return m.group(0) if ws and all(w in MOTS_CSS or w in BALISES or w.isdigit() or
                                        (w.startswith("on") and w[2:] in EVENEMENTS) for w in ws) else "*"
    return re.sub(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9]*", mot, t)


def _sans_donnee(t, css=False):
    """Une description de structure, sans rien qui ressemble à une donnée (numéros, codes, dates). css : une description
    d'éléments (« div.classe.autre ») dont les points ne sont pas ceux d'un nom de serveur ; ses classes sont réduites
    aux mots de structure connus (voir _css_generique)."""
    motif = AUTRES_DONNEES_CSS if css else AUTRES_DONNEES
    # (les deux ordres : « PL-10 » entier, pas « PL-# » ; et « 64AB-2 » entier aussi)
    propre = motif.sub("#", re.sub(r"\d{2,}", "#", motif.sub("#", str(t or ""))))
    return _css_generique(propre) if css else propre


def _cle_page(page):
    """Pour reconnaître qu'on a changé de page (jamais écrite : l'adresse peut contenir une donnée)."""
    try:
        morceaux = urlsplit(page.url)
        return morceaux.path + "#" + morceaux.fragment.split("?")[0]
    except Exception:
        return ""


def boucle_voir(page):
    """Choix 4. Dès que le portail est là : cadres EN DIRECT sur tout ce que le robot trouve cliquable, sur chaque page
    où vous allez. Vous posez la souris 3 secondes sur ce qu'il rate : il le note (vert). Entrée ici = fini.
    Rien n'est cliqué par le robot. Il apprend ce que vous avez montré (avec votre accord) et écrit un rapport."""
    appris = lire_appris()
    pages, notes = [], []          # pages vues (dans l'ordre), éléments montrés
    numero_de = {}                 # page -> son numéro (l'adresse elle-même n'est jamais écrite)
    ecrire("   Le robot cherche les elements cliquables de la page...")
    tampon, dernier, deja_dit, erreur_dite = [], 0.0, None, False
    debut = True
    while True:
        if ligne_tapee(tampon) is not None:
            break
        if page.context.pages and page.context.pages[-1] != page:  # un clic a ouvert un nouvel onglet
            page = page.context.pages[-1]
        if time.time() - dernier >= 1.5:
            cle = _cle_page(page)
            if cle not in numero_de:
                numero_de[cle] = len(pages) + 1
                pages.append({"total": 0, "avant": 0, "chrome": 0, "attribut": 0, "appris": 0, "ombres": 0,
                              "delegation": False, "cadres": 0, "dessines": 0})
                for cadre in page.frames:
                    try:
                        cadre.evaluate(JS_POINTE, [3000, True])
                    except Exception:
                        pass
            numero = numero_de[cle]
            _marquer_ecouteurs(page)
            info = _dessiner(page, appris, "ROBOT : {n} elements cliquables encadres. Souris 3 s sur un oubli = VERT. "
                                           f"Notes : {len(notes)}")
            dernier = time.time()
            garde = pages[numero - 1]
            if info["total"] >= garde["total"]:
                garde.update({k: v for k, v in info.items() if k != "erreur"})
            if info["erreur"] and not erreur_dite:
                ecrire(f"   PROBLEME pour dessiner les cadres : {info['erreur']}")
                ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
                erreur_dite = True
            etat = (numero, info["total"], info["dessines"])
            if etat != deja_dit:
                ecrire(f"   Page {numero} : le robot trouve {info['total']} elements cliquables "
                       f"({info['avant']} en rouge, {info['total'] - info['avant']} en orange), "
                       f"{info['dessines']} encadres a l'ecran.")
                if info["total"] == 0:
                    ecrire("   --> Il ne trouve RIEN sur cette page. Montrez-lui quand meme ce qui est cliquable (souris 3 s).")
                deja_dit = etat
            if debut:
                ecrire("   >>> Allez dans la FENETRE DU ROBOT (Chrome, en bas de l'ecran). Vous devez y voir des cadres")
                ecrire("       et une bande JAUNE en bas. Promenez-vous : survolez les menus, ouvrez des pages.")
                ecrire("   >>> Un element cliquable SANS cadre : posez la souris dessus, NE BOUGEZ PLUS 3 secondes.")
                ecrire("       Il devient VERT : c'est note. (Le robot ne clique pas : c'est normal.)")
                ecrire("   >>> Quand vous avez fini : revenez ICI (fenetre noire) et appuyez sur Entree.")
                debut = False
        for cadre in page.frames:
            try:
                d = cadre.evaluate(JS_POINTE, [3000, False])
            except Exception:
                d = None
            if d:
                d["page"] = numero_de.get(_cle_page(page), 0)
                notes.append(d)
                _dessiner(page, appris, f"NOTE ! ({len(notes)} element(s) note(s)). Continuez, ou Entree dans la fenetre noire.")
                ecrire(f"   --> Note n°{len(notes)} : {_resume_montre(d)}.")
        time.sleep(0.4)
    _effacer_cadres(page)
    if notes:
        revenez_ici()
        if oui(f"   Les {len(notes)} element(s) que vous avez montres (en VERT) sont-ils bien cliquables "
               "(le robot va les apprendre) ?"):
            appris = ajouter_appris([n["selecteur"] for n in notes])
            ecrire(f"   --> Appris. Le robot connait maintenant {len(appris)} genre(s) d'elements en plus.")
        else:
            for n in notes:
                n["douteux"] = True
    if not pages:
        return None
    return ecrire_diagnostic([(f"page {i}", info, [n for n in notes if n.get("page") == i])
                              for i, info in enumerate(pages, 1)], appris)


def ecrire_diagnostic(pages, appris):
    def propre(t):
        return _sans_donnee(t, css=True)
    lignes = [f"DIAGNOSTIC - CE QUE LE ROBOT TROUVE CLIQUABLE - pour Claude - {datetime.now().strftime('%d/%m/%Y %H:%M')}",
              "Structure seulement (genres d'elements, classes communes). Aucun texte de la page.",
              "Relisez : un mot qui ressemble a une vraie valeur -> effacez ce mot seulement.", ""]
    for numero, (nom, info, notes) in enumerate(pages, 1):
        lignes.append(f"PAGE {numero} : {nom}")
        lignes.append(f"  Trouves : {info['total']} (rouge {info['avant']}, Chrome {info['chrome']}, applications "
                      f"{info['attribut']}, appris {info['appris']}) ; zones {info['cadres']} ; composants web {info['ombres']}"
                      f" ; clic gere par la page entiere : {'oui' if info['delegation'] else 'non'}")
        if notes:
            lignes.append("  Montres par vous (le robot ne les voyait pas) :")
        for i, d in enumerate(notes, 1):
            lignes.append(f"    {i}. {propre(d['signature'])} ({d.get('mots', 0)} mot(s))" +
                          (" (PAS SUR : vous avez repondu non)" if d.get("douteux") else ""))
            curseur = re.sub(r"url\([^)]*\)", "url(#)", str(d["curseur"] or ""))  # (l'adresse d'un curseur dessiné)
            lignes.append(f"       role : {propre(d['role']) or '-'} | attributs : {propre(d['attributs']) or '-'} | curseur : {curseur}"
                          f" | taille : {d['taille']} | composant web : {'oui' if d['ombre'] else 'non'}")
            lignes.append(f"       ecouteur Chrome : {d['ecoute'] or '-'} | survol : {d['survol'] or '-'}"
                          f" | parent qui ecoute : {d['parentEcoute'] or '-'}")
            lignes.append(f"       parents : {propre(d['chemin'])[:300]}")
            lignes.append(f"       appris comme : {propre(d['selecteur']) or '(trop general : non appris)'}")
        lignes.append("")
    lignes.append("APPRIS (Bureau > Robot > appris.txt) : " + (" ; ".join(propre(a) for a in appris) if appris else "rien"))
    for nom in ("A_ENVOYER_diagnostic.txt", "A_ENVOYER_diagnostic_2.txt", "A_ENVOYER_diagnostic_3.txt"):
        fichier = dossier_robot() / nom
        try:
            fichier.write_text("\n".join(lignes) + "\n", encoding="utf-8")
            return fichier
        except OSError:
            continue
    raise ErreurTache("impossible d'ecrire le diagnostic dans Bureau > Robot")


# ---------------------------------------------------------------------------- tâches (écrites par Claude, lancées par vous)
class ErreurTache(Exception):
    pass


class Passer(Exception):
    """On ne touche à rien de plus pour cet élément de la liste : on passe au suivant."""


class AVerifier(Exception):
    """Une action a peut-être été faite, puis le portail a affiché autre chose : à vérifier sur le portail."""


class Arreter(Exception):
    pass


ACTIONS_TACHE = {
    # en-tête
    "TACHE": "nom de la tache", "DEMANDER": "nom = question", "LISTE": "fichier", "CONFIRMER": "oui / non",
    # navigation et gestes
    "ACCUEIL": "", "ALLER": "adresse", "MENU": "Menu > Sous-menu", "CLIQUER": "texte [DANS LIGNE valeur]",
    "ECRIRE": "Libelle = valeur", "CHOISIR": "Libelle = choix", "COCHER": "Libelle", "DECOCHER": "Libelle",
    "TOUCHE": "Entree / Tab / Echap", "ATTENDRE": "texte, ou 3 secondes", "VERIFIER": "texte",
    "LIRE": "Libelle", "ACCEPTER FENETRE": "texte de la fenetre", "CAPTURE": "nom",
    # décisions
    "SI PRESENT": "texte [DANS COLONNE titre]", "SI ABSENT": "texte [DANS COLONNE titre]", "SINON": "", "FIN SI": "",
    "REPETER TANT QUE PRESENT": "texte [DANS COLONNE titre]", "REPETER TANT QUE ABSENT": "texte [DANS COLONNE titre]",
    "FIN REPETER": "",
    "PASSER": "raison", "PREVENIR": "message", "ARRETER": "raison", "PAUSE": "message",
}
TOUCHES = {"ENTREE": "Enter", "ENTER": "Enter", "TAB": "Tab", "TABULATION": "Tab", "ECHAP": "Escape",
           "ECHAPPE": "Escape", "ESC": "Escape", "BAS": "ArrowDown", "HAUT": "ArrowUp", "ESPACE": "Space"}
GESTES = ("ACCUEIL", "ALLER", "MENU", "CLIQUER", "ECRIRE", "CHOISIR", "COCHER", "DECOCHER", "TOUCHE")
MOTIF_DUREE = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(s|sec\w*|minutes?|min)?$", re.IGNORECASE)


def _nom_action(brut):
    return " ".join(sans_accents(brut).upper().split())


def lire_texte(chemin):
    """Le texte d'un fichier reçu, quel que soit l'enregistrement (UTF-8, Excel « CSV », Unicode...)."""
    donnees = Path(chemin).read_bytes()
    if donnees[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return donnees.decode("utf-16")
    try:
        return donnees.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        texte = donnees.decode("cp1252")
    except UnicodeDecodeError:
        texte = ""
    if not texte or re.search("[‚„…‡ŠŽƒ]", texte):
        raise ErreurTache(f"« {Path(chemin).name} » illisible : dans Excel, Fichier > Enregistrer sous > « CSV UTF-8 »")
    return texte


def lire_tache(chemin):
    """Lit un fichier de tâche. Renvoie (entête, étapes) ; une erreur dit la ligne et quoi corriger."""
    texte = lire_texte(chemin)
    entete = {"nom": Path(chemin).stem, "demander": [], "liste": None, "confirmer": True}
    racine = []
    pile = [racine]
    ouverts = []
    for numero, ligne in enumerate(texte.splitlines(), 1):
        brut = ligne.strip()
        if not brut or brut.startswith(("#", "//")):
            continue
        m = re.match(r"^([A-Za-zÀ-ÿ' ]+?)\s*(?::\s*(.*))?$", brut)
        action = _nom_action(m.group(1)) if m else ""
        arg = (m.group(2) or "").strip() if m else ""
        if action not in ACTIONS_TACHE:
            raise ErreurTache(f"ligne {numero} : action inconnue « {brut[:40]} »")
        if action == "TACHE":
            entete["nom"] = arg or entete["nom"]
        elif action == "DEMANDER":
            if "=" not in arg:
                raise ErreurTache(f"ligne {numero} : ecrire DEMANDER : nom = question")
            nom, question = (x.strip() for x in arg.split("=", 1))
            entete["demander"].append((nom, question or nom))
        elif action == "LISTE":
            entete["liste"] = arg
        elif action == "CONFIRMER":
            entete["confirmer"] = not sans_accents(arg).lower().startswith("non")
        elif action in ("SI PRESENT", "SI ABSENT", "REPETER TANT QUE PRESENT", "REPETER TANT QUE ABSENT"):
            if not arg:
                raise ErreurTache(f"ligne {numero} : il manque le texte apres « {action} : »")
            genre = "SI" if action.startswith("SI") else "REPETER"
            bloc = {"action": genre, "present": action.endswith("PRESENT"), "arg": arg, "ligne": numero,
                    "alors": [], "sinon": [], "genre": genre}
            pile[-1].append(bloc)
            pile.append(bloc["alors"])
            ouverts.append(bloc)
        elif action in ("SINON", "FIN SI", "FIN REPETER"):
            attendu = "REPETER" if action == "FIN REPETER" else "SI"
            if not ouverts:
                raise ErreurTache(f"ligne {numero} : {action} sans {attendu}")
            if ouverts[-1]["genre"] != attendu:
                fin = "FIN REPETER" if ouverts[-1]["genre"] == "REPETER" else "FIN SI"
                raise ErreurTache(f"ligne {numero} : il manque d'abord {fin} (bloc ouvert ligne {ouverts[-1]['ligne']})")
            if action == "SINON":
                if pile[-1] is not ouverts[-1]["alors"]:
                    raise ErreurTache(f"ligne {numero} : deux SINON pour le meme SI")
                pile[-1] = ouverts[-1]["sinon"]
            else:
                pile.pop()
                ouverts.pop()
        else:
            if action in ("ECRIRE", "CHOISIR") and "=" not in arg:
                raise ErreurTache(f"ligne {numero} : ecrire {action} : Libelle = valeur")
            if action in ("ALLER", "MENU", "CLIQUER", "ECRIRE", "CHOISIR", "COCHER", "DECOCHER", "TOUCHE", "ATTENDRE",
                          "VERIFIER", "LIRE", "ACCEPTER FENETRE") and not arg:
                raise ErreurTache(f"ligne {numero} : il manque ce qu'il faut apres « {action} : »")
            if action == "TOUCHE" and _nom_action(arg) not in TOUCHES:
                raise ErreurTache(f"ligne {numero} : touche inconnue « {arg} » (Entree, Tab, Echap, Bas, Haut, Espace)")
            pile[-1].append({"action": action, "arg": arg, "ligne": numero})
    if ouverts:
        fin = "FIN REPETER" if ouverts[-1]["genre"] == "REPETER" else "FIN SI"
        raise ErreurTache(f"ligne {ouverts[-1]['ligne']} : il manque {fin}")
    if not racine:
        raise ErreurTache("la tache ne contient aucune action")
    return entete, racine


def variables_utilisees(etapes):
    for e in etapes:
        for nom in re.findall(r"\{([^{}]+)\}", e.get("arg", "")):
            yield nom.strip(), e["ligne"]
        yield from variables_utilisees(e.get("alors", []))
        yield from variables_utilisees(e.get("sinon", []))


def _versions(base, nom):
    """Le fichier et ses copies « nom (1).csv » laissées par le navigateur, du plus récent au plus ancien."""
    p = Path(nom)
    trouves = [base / nom] + list(base.glob(f"{p.stem} (*){p.suffix}"))
    trouves = [f for f in trouves if f.is_file()]
    return sorted(trouves, key=lambda f: f.stat().st_mtime, reverse=True)


def _chercher_fichier(nom, pres_de):
    for base in (Path(pres_de).parent, Path.home() / "Downloads", Path.home() / "Téléchargements", dossier_robot(),
                 dossier_robot("taches")):
        try:
            versions = _versions(base, nom)
        except OSError:
            continue
        if versions:
            if len(versions) > 1:
                ecrire(f"   Plusieurs « {nom} » : le robot prend la plus recente, {versions[0].name} "
                       f"(recue le {datetime.fromtimestamp(versions[0].stat().st_mtime):%d/%m a %H:%M}).")
            return versions[0]
    raise ErreurTache(f"liste « {nom} » introuvable (mettez-la dans Telechargements)")


def lire_liste(nom, pres_de):
    """Les lignes de la liste (Excel, CSV ou texte) : une ligne = une fois la tâche."""
    import csv
    f = _chercher_fichier(nom, pres_de)
    suffixe = f.suffix.lower()
    if suffixe in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ErreurTache("pour une liste Excel, enregistrez-la en CSV (Fichier > Enregistrer sous > CSV UTF-8)")
        feuille = load_workbook(f, read_only=True, data_only=True).worksheets[0]
        lignes = [["" if v is None else str(v).strip() for v in row] for row in feuille.iter_rows(values_only=True)]
    elif suffixe == ".csv":
        brut = lire_texte(f)
        entete = next((l for l in brut.splitlines() if l.strip()), "")
        separateur = "," if ("," in entete and ";" not in entete) else ";"
        lignes = [[c.strip() for c in r] for r in csv.reader(brut.splitlines(), delimiter=separateur)]
    elif suffixe == ".txt":
        lignes = [["valeur"]] + [[l.strip()] for l in lire_texte(f).splitlines() if l.strip()]
    else:
        raise ErreurTache(f"liste « {nom} » : format inconnu (CSV UTF-8, .xlsx ou .txt)")
    lignes = [l for l in lignes if any(l)]
    if len(lignes) < 2:
        raise ErreurTache(f"la liste « {nom} » est vide")
    entetes = [h or f"colonne{i + 1}" for i, h in enumerate(lignes[0])]
    return [dict(zip(entetes, l + [""] * (len(entetes) - len(l)))) for l in lignes[1:]]


JS_LIGNE = r"""([valeur, jeton]) => {
  document.querySelectorAll('[data-robot-ligne]').forEach(x => x.removeAttribute('data-robot-ligne'));
  const norme = t => (t || '').replace(/\s+/g, ' ').trim().toLowerCase();
  const v = norme(valeur); let n = 0;
  for (const e of document.querySelectorAll('td, th, [role=gridcell], li, a, span, div, b')) {
    let propre = ''; for (const c of e.childNodes) if (c.nodeType === 3) propre += c.textContent;
    if (norme(propre) !== v && norme(e.textContent) !== v) continue;
    const ligne = e.closest('tr, [role=row], li');
    // une « ligne » qui contient d'autres lignes est de la mise en page : ce n'est pas la ligne de ce plan
    if (!ligne || ligne.querySelector('tr, [role=row], li, table')) continue;
    if (!ligne.hasAttribute('data-robot-ligne')) { ligne.setAttribute('data-robot-ligne', jeton); n++; }
  }
  return n;
}"""

# Les lignes qui portent cette valeur (même règle que JS_LIGNE, sans rien marquer) ; avec un élément : est-il dans l'une d'elles ?
JS_LIGNES_DE = r"""(e, valeur) => {
  const norme = t => (t || '').replace(/\s+/g, ' ').trim().toLowerCase();
  const v = norme(valeur), lignes = new Set();
  for (const x of document.querySelectorAll('td, th, [role=gridcell], li, a, span, div, b')) {
    let propre = ''; for (const c of x.childNodes) if (c.nodeType === 3) propre += c.textContent;
    if (norme(propre) !== v && norme(x.textContent) !== v) continue;
    const ligne = x.closest('tr, [role=row], li');
    if (!ligne || ligne.querySelector('tr, [role=row], li, table')) continue;
    lignes.add(ligne);
  }
  return { n: lignes.size, dedans: !!e && e.nodeType === 1 && Array.from(lignes).some(l => l.contains(e)) };
}"""

JS_CHAMP_PAR_LIBELLE = r"""([libelle, jeton]) => {
  const norme = t => (t || '').replace(/\s+/g, ' ').replace(/[\s:*]+$/, '').trim().toLowerCase();
  const v = norme(libelle); let n = 0;
  const champs = Array.from(document.querySelectorAll('input:not([type=hidden]):not([type=submit]):not([type=button]), select, textarea'));
  for (const e of document.querySelectorAll('label, td, th, span, div, b, strong, p, dt')) {
    let propre = ''; for (const c of e.childNodes) if (c.nodeType === 3) propre += c.textContent;
    if (norme(propre) !== v) continue;
    const bloc = e.closest('tr, li, p, dl, fieldset, .form-group, div') || e.parentElement;
    const suivant = champs.find(c => (e.compareDocumentPosition(c) & Node.DOCUMENT_POSITION_FOLLOWING) && bloc.contains(c));
    if (suivant) { suivant.setAttribute('data-robot-champ', jeton); n++; }
  }
  return n;
}"""

# La valeur affichée en face d'un libellé (fiche « Statut : Validé », champ, ou cellule sous un titre de colonne).
JS_LIRE = r"""([libelle]) => {
  const norme = t => (t || '').replace(/\s+/g, ' ').replace(/[\s:*]+$/, '').trim().toLowerCase();
  const court = t => (t || '').replace(/\s+/g, ' ').trim();
  const valeurDe = c => c.tagName === 'SELECT' ? court((c.options[c.selectedIndex] || {}).text) : court(c.value);
  const v = norme(libelle);
  for (const e of document.querySelectorAll('label, th, [role=columnheader], td, span, div, b, strong, p, dt')) {
    let propre = ''; for (const c of e.childNodes) if (c.nodeType === 3) propre += c.textContent;
    const p = court(propre);
    if (norme(p) !== v && !norme(p).startsWith(v + ' :') && !norme(p).startsWith(v + ':')) continue;
    if (norme(p) !== v) return { valeur: court(p.split(':').slice(1).join(':')) };
    if (e.tagName === 'TH' || e.getAttribute('role') === 'columnheader') {      // titre de colonne
      const titres = e.parentElement, table = e.closest('table, [role=grid]');
      const rang = Array.from(titres.children).indexOf(e);
      const lignes = Array.from(table.querySelectorAll('tr, [role=row]'))
        .filter(l => l !== titres && Array.from(l.children).some(c => c.tagName === 'TD' || c.getAttribute('role') === 'gridcell'));
      if (lignes.length !== 1) return { erreur: lignes.length ? 'plusieurs lignes dans le tableau' : 'tableau vide' };
      const c = lignes[0].children[rang];
      return c ? { valeur: court(c.textContent) } : { erreur: 'colonne vide' };
    }
    if (e.tagName === 'LABEL' && e.control) return { valeur: valeurDe(e.control) };
    const champ = e.parentElement && e.parentElement.querySelector('input, select, textarea');
    if (champ && (e.compareDocumentPosition(champ) & Node.DOCUMENT_POSITION_FOLLOWING)) return { valeur: valeurDe(champ) };
    const suivant = e.nextElementSibling;
    if (suivant && suivant.tagName !== 'TH') return { valeur: court(suivant.textContent) };
    const cell = e.closest('td, dt');
    if (cell && cell.nextElementSibling) return { valeur: court(cell.nextElementSibling.textContent) };
  }
  return null;
}"""

# Une ligne du tableau a-t-elle ce texte dans cette colonne ? (cellule égale, ou qui commence par ce texte :
# « Validé » trouve « Validé le 12/03 », pas « Non validé »). Tableaux visibles seulement.
JS_COLONNE = r"""([colonne, texte]) => {
  const norme = t => (t || '').normalize('NFD').replace(/[̀-ͯ]/g, '').replace(/\s+/g, ' ').trim().toLowerCase();
  const vis = e => { const s = getComputedStyle(e); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const egal = (cellule, t) => { const c = norme(cellule), x = norme(t);
    return c === x || (c.startsWith(x) && /^[^a-z0-9]/.test(c.slice(x.length))); };
  let trouve = false;
  for (const t of document.querySelectorAll('table, [role=grid]')) {
    if (!vis(t)) continue;
    let entetes = Array.from(t.querySelectorAll('thead th, [role=columnheader]'));
    if (!entetes.length) { const l = t.querySelector('tr'); if (l && Array.from(l.children).every(c => c.tagName === 'TH')) entetes = Array.from(l.children); }
    let rang = entetes.findIndex(h => norme(h.textContent) === norme(colonne));
    if (rang < 0) { const proches = entetes.map((h, i) => [norme(h.textContent), i]).filter(([h]) => h.includes(norme(colonne)));
                    if (proches.length === 1) rang = proches[0][1]; }
    if (rang < 0) continue;
    trouve = true;
    for (const l of t.querySelectorAll('tr, [role=row]')) {
      const cellules = Array.from(l.children).filter(c => c.tagName === 'TD' || c.getAttribute('role') === 'gridcell');
      if (!cellules.length) continue;
      if (cellules.length === 1 && /aucun|aucune|vide|pas de|no data|empty/i.test(l.textContent || '')) continue;
      if (cellules.length < entetes.length) return 'illisible';
      // colonne sans titre (cases, actions) : on regarde les deux alignements possibles
      const decalage = cellules.length - entetes.length;
      if (egal(cellules[rang].textContent, texte) || (decalage && egal(cellules[rang + decalage].textContent, texte))) return 'oui';
    }
  }
  return trouve ? 'non' : 'colonne absente';
}"""

JS_BANDEAU = r"""(message) => {
  const b = document.createElement('div');
  b.textContent = 'ROBOT : ' + message;
  b.setAttribute('data-robot-bandeau', '1');
  b.style.cssText = 'position:fixed;top:12px;left:50%;transform:translateX(-50%);z-index:2147483647;background:#f59e0b;' +
    'color:#111;font:600 15px Arial,sans-serif;padding:12px 18px;border-radius:8px;box-shadow:0 2px 12px rgba(0,0,0,.4);max-width:80%';
  document.documentElement.appendChild(b);
  setTimeout(() => b.remove(), 5000);
}"""

# Ce que la touche Entrée (ou Espace) déclencherait : l'élément actif, son formulaire et ses boutons d'envoi.
JS_FOCUS = r"""() => {
  const e = document.activeElement;
  if (!e || e === document.body || e.tagName === 'IFRAME' || e.tagName === 'FRAME') return null;
  e.setAttribute('data-robot-focus', '1');   // pour remettre le curseur ici après une question
  const f = e.form || e.closest('form');
  const boutons = f ? Array.from(f.querySelectorAll('button, input[type=submit], input[type=image]'))
    .map(b => [b.innerText, b.value, b.title, b.getAttribute('aria-label')].join(' ')) : [];
  let touchesParents = '';
  for (let n = e.parentElement; n; n = n.parentElement)
    for (const a of ['onkeydown', 'onkeypress', 'onkeyup']) if (n.getAttribute && n.getAttribute(a)) touchesParents += ' ' + n.getAttribute(a);
  const libelle = (e.labels && e.labels[0] ? e.labels[0].innerText : '') + ' ' + (e.name || '') + ' ' + (e.id || '');
  return { texte: [e.innerText, e.value, e.title, e.getAttribute('aria-label'), e.getAttribute('onclick'),
                   e.getAttribute('onkeydown'), e.getAttribute('onkeypress'), libelle, touchesParents].join(' '),
           formulaire: f ? [f.getAttribute('action') || '', f.getAttribute('onsubmit') || '', boutons.join(' ')].join(' ') : '',
           avecFormulaire: !!f, post: !!f && (f.getAttribute('method') || '').toLowerCase() === 'post',
           genre: e.tagName === 'INPUT' ? (e.type || 'text').toLowerCase() : e.tagName.toLowerCase(),
           lien: e.tagName === 'A' || e.tagName === 'BUTTON' || e.getAttribute('role') === 'button',
           auto: !!(e.getAttribute('onchange') || e.getAttribute('onclick') || touchesParents ||
                    (f && (f.getAttribute('onchange') || f.getAttribute('onclick')))) };
}"""


def _correspond(attendu, message):
    """La fenêtre du portail est-elle celle annoncée ? Mots entiers : « PL-1 » ne vaut pas pour « PL-10 »."""
    a, m = mots(attendu), mots(message)
    return bool(a) and re.search(r"(^| )" + re.escape(a) + r"( |$)", m) is not None


class Tache:
    """Exécute une tâche, élément par élément de la liste, en notant chaque résultat dans un fichier."""
    active = None  # la tâche (ou la visite) en cours : c'est elle qui répond aux fenêtres du portail
    _pages_suivies = set()

    def __init__(self, entete, etapes, page, url, lecture_seule=False):
        self.entete, self.etapes = entete, etapes
        self.page = page
        self.principale = page
        self.url = url
        # visite (choix 5 et 6) : le robot regarde seulement ; tout geste qui pourrait modifier est refusé sans question
        self.lecture_seule = lecture_seule
        self.chargement_max, self.calme_max = 10, 15
        self.variables = {}
        self.tout_confirme = not entete["confirmer"]
        self.oui_lignes = set()  # « t » : oui pour cette ligne de la tâche, pour tous les éléments
        self.element = ""
        self.ligne_courante = 0
        self.lectures = {}
        self.attendues = []
        self.en_cours = {}  # requête au portail en cours -> (sa page, quand elle est partie, interrogation régulière ?)
        self.recentes = {}  # (méthode, adresse sans paramètres) -> heures de ses derniers départs
        Tache.active = self
        self.nouvel_element()
        self.suivre_page(page)
        page.context.on("page", self.suivre_page)
        page.context.on("request", self._requete_partie)
        page.context.on("requestfinished", lambda r: self.en_cours.pop(r, None))
        page.context.on("requestfailed", lambda r: self.en_cours.pop(r, None))

    def _requete_partie(self, requete):
        try:
            if requete.resource_type in ("document", "xhr", "fetch"):
                try:
                    page = requete.frame.page
                except Exception:
                    page = None
                maintenant = time.time()
                # la même adresse 3 fois en 10 secondes : le portail s'interroge tout seul (notifications, session)
                cle = (requete.method, requete.url.split("?")[0].split("#")[0])
                heures = [h for h in self.recentes.get(cle, []) if maintenant - h < 10] + [maintenant]
                if len(self.recentes) > 500:
                    self.recentes.clear()
                self.recentes[cle] = heures
                self.en_cours[requete] = (page, maintenant, len(heures) >= 3)
        except Exception:
            pass

    def oublier_requetes(self, page, avant):
        """La page a été remplacée (retour à l'accueil) : ses requêtes parties avant ne répondront jamais."""
        for r, (p, debut, _) in list(self.en_cours.items()):
            if p is page and debut < avant:
                self.en_cours.pop(r, None)

    def en_vol(self):
        """Les requêtes encore attendues. Celles d'une page fermée ne répondront jamais. En visite, une requête ouverte
        depuis plus de 4 secondes est une écoute permanente du portail (elle ne répond qu'à un événement), et une
        interrogation régulière (la même adresse sans cesse) n'est pas la réponse attendue."""
        vieille = 4 if self.lecture_seule else 3600
        maintenant, n = time.time(), 0
        for r, (p, debut, reguliere) in list(self.en_cours.items()):
            try:
                fermee = p is not None and p.is_closed()
            except Exception:
                fermee = True
            if fermee:
                self.en_cours.pop(r, None)
            elif maintenant - debut < vieille and not (self.lecture_seule and reguliere):
                n += 1
        return n

    def calme(self, maximum=None):
        """Attendre que le portail ait répondu (plus aucune requête en cours pendant une demi-seconde) : sinon on
        regarderait la page d'avant la réponse (ligne pas encore retirée, message pas encore affiché)."""
        fin, tranquille = time.time() + (maximum or self.calme_max), None
        while time.time() < fin:
            if self.en_vol() <= 0:
                tranquille = tranquille or time.time()
                if time.time() - tranquille >= 0.5:
                    break
            else:
                tranquille = None
            try:
                self.page.wait_for_timeout(100)
            except Exception:
                break

    def page_ouverte(self):
        """Une fenêtre (popup d'édition) qui s'est fermée toute seule : on revient sur la dernière fenêtre ouverte."""
        if self.page.is_closed():
            ouvertes = [p for p in self.page.context.pages if not p.is_closed()]
            if not ouvertes:
                raise ErreurTache("la fenetre du robot a ete fermee")
            self.page = ouvertes[-1]
            self.page.bring_to_front()

    def nouvel_element(self):
        self.lectures = {}
        self.attendues = []
        self.acceptee = None  # une action acceptée vaut pour tout l'élément : tout message ensuite = A VERIFIER
        self.apres = []
        self.nouveau_geste()

    def nouveau_geste(self):
        self.refusees = []
        self.accord_geste = False

    # --- fenêtres du portail (« Voulez-vous vraiment supprimer ? ») : seule celle annoncée est acceptée
    def suivre_page(self, page):
        # une seule écoute par page, même si une visite (choix 5) est suivie d'un « aller à » (choix 6)
        if id(page) in Tache._pages_suivies:
            return
        Tache._pages_suivies.add(id(page))
        page.on("dialog", lambda fenetre: Tache.active.sur_fenetre(fenetre))

    def sur_fenetre(self, fenetre):
        message = " ".join((fenetre.message or "").split())
        if fenetre.type == "beforeunload" or not message:
            message = "Quitter cette page ? (des modifications pourraient ne pas etre enregistrees)"
        attendue = next((a for a in self.attendues if _correspond(a, message)), None)
        try:
            if self.lecture_seule:  # visite : toujours refusée, et son texte (qui peut contenir une donnée) n'est pas écrit
                self.refusees.append("une fenetre (refusee : visite en lecture seule)")
                ecrire("      fenetre du portail REFUSEE (visite en lecture seule)")
                fenetre.dismiss()
            elif attendue and not self.refusees and (self.accord_geste or not interdit(message)):
                self.attendues.remove(attendue)
                self.acceptee = self.acceptee or message
                ecrire(f"      fenetre du portail acceptee : « {message[:80]} »")
                fenetre.accept()
            elif fenetre.type == "alert" and self.acceptee:
                # après une action acceptée, une alerte n'a qu'un bouton OK : on la ferme, et on le signale
                self.apres.append(message)
                ecrire(f"      message du portail apres l'action : « {message[:80]} »")
                fenetre.accept()
            else:
                self.refusees.append(message)
                ecrire(f"      fenetre du portail REFUSEE (pas prevue) : « {message[:80]} »")
                fenetre.dismiss()
        except Exception:
            pass

    def verifier_fenetres(self):
        if self.acceptee and (self.apres or self.refusees):
            autre = (self.apres + self.refusees)[0]
            raise AVerifier(f"« {self.acceptee[:50]} » accepte, puis le portail a affiche « {autre[:80]} » : "
                            "verifiez cet element sur le portail")
        if self.refusees:
            raise Passer(f"le portail a affiche « {self.refusees[0][:100]} » : rien n'a ete valide, on passe au suivant")

    # --- outils
    def remplir(self, texte):
        def valeur(m):
            nom = m.group(1).strip()
            for cle, v in self.variables.items():
                if mots(cle) == mots(nom):
                    v = str(v).strip()
                    if not v:
                        raise Passer(f"{{{nom}}} est vide pour cet element : rien n'est fait")
                    return v
            raise ErreurTache(f"valeur inconnue {{{nom}}}")
        return re.sub(r"\{([^{}]+)\}", valeur, texte)

    def cadres(self):
        return list(self.page.frames)

    def retirer_bandeaux(self):
        for cadre in self.cadres():
            try:
                cadre.evaluate("() => document.querySelectorAll('[data-robot-bandeau]').forEach(b => b.remove())")
            except Exception:
                pass

    def candidats(self, texte, racines, attente=5, tous=False):
        """Les éléments qui portent EXACTEMENT ce texte (nom de bouton, lien, texte, bulle), les visibles d'abord.
        On réessaie quelques secondes : la page peut être encore en train de s'afficher.
        tous : réunir ce que trouvent toutes les façons de chercher (pour un menu : l'entrée du menu ET le bouton
        de la page du même nom ; le menu choisit ensuite l'entrée du menu)."""
        if not texte.strip():
            raise ErreurTache("texte vide")
        exact_sans_casse = re.compile(r"^\s*" + re.escape(texte) + r"\s*$", re.IGNORECASE)
        jeton = secrets.token_hex(6)

        def propre_texte(r):
            if hasattr(r, "child_frames"):  # un cadre de la page entière
                r.evaluate(JS_TROUVER_PROPRE_TEXTE, [texte, jeton])
            else:
                r.evaluate("(e, [m, j]) => { const norme = t => (t || '').replace(/\\s+/g, ' ').trim().toLowerCase();"
                           " for (const x of [e, ...e.querySelectorAll('*')]) { let p = '';"
                           " for (const n of x.childNodes) if (n.nodeType === 3) p += n.textContent;"
                           " if (norme(p) === norme(m)) x.setAttribute('data-robot-trouve', j); } }", [texte, jeton])
            return r.locator(f'[data-robot-trouve="{jeton}"]')

        def icone(r):
            """Une entrée faite d'une icône seule (barre d'icônes) : son nom, c'est sa bulle ou son icône ; une icône
            sans aucun nom (« icone 3 ») se retrouve par sa place parmi les entrées des menus."""
            if not hasattr(r, "child_frames"):
                return r.locator("css=[data-robot-rien]")
            # « icone folder », « icone 3 », « icone folder 2 » : l'entrée des menus de ce nom (par sa place) ; sinon
            # (bulle, ou icône hors des menus) l'élément qui porte cette bulle ou cette icône
            numero = None
            if texte.startswith("icone "):
                numero = next((i for i, e in enumerate(r.evaluate(JS_MENUS)) if texte in (e["nom"], e.get("texte"))), None)
                if numero is None and re.fullmatch(r"icone (.+ )?\d+", texte):
                    return r.locator("css=[data-robot-rien]")
            r.evaluate(JS_TROUVER_ICONE, [texte, jeton, numero])
            return r.locator(f'[data-robot-trouve="{jeton}"]')

        def tolerant(r):
            """En dernier : accents, majuscules, flèches et symboles ignorés (« Plans récents › »)."""
            if not hasattr(r, "child_frames"):
                return r.locator("css=[data-robot-rien]")
            r.evaluate(JS_TROUVER_TOLERANT, [texte, jeton, "propre"])
            if not r.locator(f'[data-robot-trouve="{jeton}"]').count():
                r.evaluate(JS_TROUVER_TOLERANT, [texte, jeton, "entier"])
            return r.locator(f'[data-robot-trouve="{jeton}"]')

        facons = [lambda r, role=role: r.get_by_role(role, name=texte, exact=True)
                  for role in ("button", "link", "menuitem", "tab", "treeitem")]
        facons += [lambda r: r.get_by_text(texte, exact=True), propre_texte,
                   lambda r: r.get_by_title(texte, exact=True),
                   lambda r: r.locator(f"[aria-label={json.dumps(texte, ensure_ascii=False)}]"),
                   lambda r: r.get_by_text(exact_sans_casse), tolerant, icone]
        fin = time.time() + attente
        while True:
            reunis = []
            for facon in facons:
                trouves = []
                for r in racines:
                    try:
                        loc = facon(r)
                        for i in range(min(loc.count(), 10)):
                            trouves.append(loc.nth(i))
                    except Exception:
                        continue
                if trouves and not tous:
                    visibles = [t for t in trouves if _visible(t)]
                    return visibles + [t for t in trouves if t not in visibles]
                reunis += trouves
            if reunis:
                visibles = [t for t in reunis if _visible(t)]
                return visibles + [t for t in reunis if t not in visibles]
            if time.time() >= fin:
                return []
            self.page.wait_for_timeout(500)

    def racines_ligne(self, valeur):
        jeton = secrets.token_hex(6)
        racines = []
        for cadre in self.cadres():
            try:
                if cadre.evaluate(JS_LIGNE, [valeur, jeton]):
                    racines.append(cadre.locator(f'[data-robot-ligne="{jeton}"]'))
            except Exception:
                continue
        lignes = sum(r.count() for r in racines)
        if lignes == 0:
            raise ErreurTache(f"aucune ligne « {valeur} » sur la page")
        if lignes > 1:
            raise ErreurTache(f"{lignes} lignes « {valeur} » sur la page : le robot ne choisit pas au hasard")
        return racines

    def _sans_focus(self):
        """Personne ne doit pouvoir taper dans la page pendant une question : on retire le curseur des champs."""
        for cadre in self.cadres():
            try:
                cadre.evaluate("() => { const a = document.activeElement; if (a && a !== document.body && a.blur) a.blur(); }")
            except Exception:
                pass

    def confirmer(self, loc, quoi, cle=None):
        """Avant un geste qui modifie : votre accord (o), ou « t » donné avant pour CETTE action précise
        (même ligne de la tâche ET même texte : « t » pour Archiver ne vaut jamais pour Supprimer)."""
        if self.lecture_seule:
            raise Passer(f"visite : le robot ne fait pas « {quoi} » (lecture seule)")
        cle = (self.ligne_courante, mots(cle or quoi))
        if self.tout_confirme or cle in self.oui_lignes:
            self.accord_geste = True
            return
        attendues, self.attendues = self.attendues, []  # rien n'est accepté pendant la question
        self._sans_focus()
        if loc is not None:
            _encadrer(loc, 0)
        revenez_ici()
        reponse = demander(f"   [{self.element}] Le robot va {quoi}. Tapez o (oui), n (non : on passe au suivant), "
                           "t (oui pour cette action, pour toute la liste) :", "n").lower()
        _nettoyer(self.page, tout=False)
        if reponse.startswith("t"):
            self.oui_lignes.add(cle)
        elif not reponse.startswith(("o", "y")):
            raise Passer("refuse par vous")
        self.attendues = attendues
        self.accord_geste = True

    def _texte_de(self, loc):
        """Le texte de l'élément, ses bulles et actions, et celles de ses PARENTS : le clic remonte jusqu'à eux."""
        try:
            return loc.evaluate(r"""(e, [visite, fortes]) => {
              // (en visite, sans le mot d'une icône, « archive » dans « archive Archives », s'il reste un autre texte ;
              // une icône d'action franche, « logout », reste)
              const c = e.cloneNode(true);
              const lig = Array.from(c.querySelectorAll('mat-icon, i, span')).filter(x => !x.children.length &&
                  /^[a-z0-9_]+$/.test((x.textContent || '').trim()) &&
                  (x.tagName === 'MAT-ICON' || /icon|symbol|material|^fa|\bfa-|mdi|glyph/i.test(String(x.className))));
              let texte = e.innerText;
              const autres = lig.filter(x => !fortes.includes((x.textContent || '').trim()));
              if (visite && autres.length) {   // (une icône seule est jugée par son icône dans _refus_visite)
                autres.forEach(x => x.remove()); texte = (c.textContent || '').replace(/\s+/g, ' ').trim(); }
              const parts = [texte, e.value, e.title, e.getAttribute('aria-label'), e.getAttribute('href'), e.getAttribute('onclick')];
              const GEST = /^(on(click|dblclick|mousedown|mouseup|pointerdown|pointerup|touchstart|touchend))$|click|^hx-|^wire:|^x-on|^@|^data-(ajax|turbo|confirm|method|action|url)|^title$|^aria-label$/i;
              for (let n = e.parentElement; n && n !== document.body; n = n.parentElement)
                for (const a of Array.from(n.attributes)) if (GEST.test(a.name)) parts.push(a.value);
              const f = e.form || e.closest('form'); if (f) parts.push(f.getAttribute('action') || '');
              return parts.join(' '); }""", [bool(self.lecture_seule), sorted(LIGATURES_D_ACTION)])
        except Exception:
            return ""

    def cliquer(self, arg):
        nieme = 1
        m = re.match(r"^(.*?)\s*\[(\d+)\]\s*$", arg)
        if m:
            arg, nieme = m.group(1), int(m.group(2))
        racines = self.cadres()
        quoi, ligne = arg, None
        m = re.match(r"^(.*?)\s+DANS\s+LIGNE\s+(.+)$", arg, re.IGNORECASE)
        if m:
            arg, ligne = m.group(1).strip(), m.group(2).strip()
            racines = self.racines_ligne(ligne)
            quoi = f"{arg} (ligne {ligne})"
        trouves = self.candidats(arg, racines)
        if self.lecture_seule and nieme == 1:
            # en visite : le premier élément de ce nom qui est un lien, un menu ou un onglet (pas un titre du même nom)
            bons = [x for x in trouves[:10] if _visible(x) and not self._refus_visite(x)]
            trouves = bons or trouves
        if len(trouves) < nieme:
            raise ErreurTache(f"« {arg} » introuvable sur la page")
        self._cliquer_element(trouves[nieme - 1], arg, quoi, ligne)

    def _refus_visite(self, loc):
        """En visite : pourquoi cet élément ne peut pas être cliqué ("" s'il le peut). Seulement un lien, une entrée de
        menu ou un onglet ; jamais un bouton d'action, un champ, une case, un élément d'une fenêtre ou d'un tableau,
        ni un nom d'action (même contrôle que l'étape 8 du test)."""
        try:
            info = loc.evaluate(JS_EXAMINER, [secrets.token_hex(6), False])
        except Exception:
            return "illisible"
        if info["refus"]:
            return info["refus"]
        nomme = info.get("menuNomme") and not info.get("propre")
        # (une entrée de menu nommée par sa bulle, faite d'une icône : jugée par son nom, pas par le mot ni la classe de
        # son icône, « bi-shield-lock » d'un module Habilitations)
        attributs = info["attributs"]
        if nomme:
            attributs = " | ".join(x for x in attributs.split(" | ") if not x.startswith("class "))
        if (not nomme and action_en_visite(info["texte"])) or action_en_visite(attributs):
            return "ressemble a une action"
        if set(info.get("ligatures") or []) & LIGATURES_D_ACTION:
            return "icone d une action"
        # un élément SANS texte visible à lui (icône seule, même avec une bulle « Favori ») : son icône dit ce qu'il fait ;
        # un lien « Signalements » avec un petit drapeau devant garde son nom, il se visite
        if not info.get("propre") and not nomme and \
                any(_icone_d_action("icone " + n) for n in _icones_de(info["texte"], info["attributs"])):
            return "icone d une action"
        return ""

    def _danger(self, texte, survol=False):
        """Un geste à contrôler : en tâche, tout mot d'action (v4) ; en visite, un VERBE d'action (une rubrique
        « Factures » ou « Archives » se visite, « Archiver » non). survol : une entrée seulement survolée."""
        return action_en_visite(texte, survol) if self.lecture_seule else interdit(texte)

    def _verifier_visite(self, loc, quoi):
        if self.lecture_seule:
            refus = self._refus_visite(loc)
            if refus:
                raise Passer(f"visite : « {quoi} » pas clique ({refus})")

    def _toujours_le_meme(self, loc, texte_avant, ligne):
        """Juste avant le clic (la question a pu durer) : l'élément n'a pas changé, et il est toujours dans LA ligne de
        cet élément de la liste. Une liste qui se rafraîchit peut afficher un AUTRE plan dans la même ligne."""
        if self._texte_de(loc) != texte_avant:
            raise Passer("l'element a change pendant la question : rien n'a ete clique")
        if ligne is None:
            return
        lignes = 0
        for cadre in self.cadres():
            try:
                lignes += cadre.evaluate("v => (" + JS_LIGNES_DE + ")(null, v).n", ligne)
            except Exception:
                continue
        try:
            dedans = loc.evaluate(JS_LIGNES_DE, ligne)["dedans"]
        except Exception:
            dedans = False
        if lignes != 1 or not dedans:
            raise Passer(f"la liste a change pendant la question (ligne « {ligne} ») : rien n'a ete clique")

    def _cliquer_element(self, loc, arg, quoi, ligne=None):
        if not _visible(loc):
            _ouvrir_menus(loc)
        texte_avant = self._texte_de(loc)
        if (self._danger(arg) or self._danger(texte_avant) or any(interdit(a) for a in self.attendues)) \
                and not self.accord_geste:
            self.confirmer(loc, f"cliquer « {quoi} »", cle="cliquer " + arg)
        self._toujours_le_meme(loc, texte_avant, ligne)
        self._verifier_visite(loc, quoi)
        avant = len(self.page.context.pages)
        loc.click(timeout=15000)
        attendre_chargement(self.page, self.chargement_max)
        self.calme()
        self.page_ouverte()
        if len(self.page.context.pages) > avant:
            self.page = self.page.context.pages[-1]
            self.page.bring_to_front()
            attendre_chargement(self.page, self.chargement_max)

    def menu(self, arg):
        parties = [p.strip() for p in arg.split(">") if p.strip()]
        _marquer_ecouteurs(self.page)
        _menus_de_la_page(self.page)  # les entrées des menus passent avant un titre de la page qui porte le même nom
        for i, partie in enumerate(parties):
            dernier = i == len(parties) - 1
            trouves = self.candidats(partie, self.cadres(), tous=True)[:20]
            if not trouves:
                raise ErreurTache(f"menu « {partie} » introuvable")
            trouves.sort(key=lambda t: (not _visible(t), _rang_menu(t), _bouton_de_page(t)))
            loc = trouves[0]
            if i == 0 and any(self._danger(p, survol=j < len(parties) - 1) for j, p in enumerate(parties)) \
                    and not self.accord_geste:
                self.confirmer(loc, "ouvrir le menu « " + " > ".join(parties) + " »", cle="menu " + " > ".join(parties))
            if not _visible(loc):
                _ouvrir_menus(loc)
            if dernier:
                # un bouton de la PAGE (hors menu) qui porte le même nom n'est jamais cliqué à la place de l'entrée du menu
                if _rang_menu(loc) == 2 and _bouton_de_page(loc):
                    raise ErreurTache(f"« {partie} » : le robot ne trouve que un bouton de la page, pas l'entree du menu")
                self._cliquer_element(loc, partie, partie)
                return
            # survol (comme la souris) : le sous-menu apparaît, parfois ailleurs dans la page (menu de gauche)
            ouvert = False
            for essai in [t for t in trouves if _visible(t)][:3] or [loc]:
                try:
                    _survoler_et_voir(self.page, essai)
                except Exception:
                    continue
                suivant = self.candidats(parties[i + 1], self.cadres(), attente=1, tous=True)
                if any(_visible(s) and _rang_menu(s) < 2 for s in suivant[:20]):  # apparu dans un menu, pas sur la page
                    loc, ouvert = essai, True
                    break
            if not ouvert:
                # menu qui s'ouvre au clic : ce clic aussi est contrôlé
                if self._danger(self._texte_de(loc)) and not self.accord_geste:
                    self.confirmer(loc, f"cliquer « {partie} »", cle="cliquer " + partie)
                self._verifier_visite(loc, partie)
                loc.click(timeout=10000)
                self.page.wait_for_timeout(800)

    def champ(self, libelle):
        if libelle.startswith("#"):
            return self.page.locator(libelle)
        jeton = secrets.token_hex(6)
        facons = [lambda r: r.get_by_label(libelle, exact=True),
                  lambda r: r.get_by_label(re.compile(r"^\s*" + re.escape(libelle) + r"\s*[:*]?\s*$", re.IGNORECASE)),
                  lambda r: r.get_by_placeholder(libelle, exact=True)]
        fin = time.time() + 5
        while True:
            for facon in facons + [None]:
                trouves = []
                for cadre in self.cadres():
                    try:
                        if facon is None:
                            cadre.evaluate(JS_CHAMP_PAR_LIBELLE, [libelle, jeton])
                            loc = cadre.locator(f'[data-robot-champ="{jeton}"]')
                        else:
                            loc = facon(cadre)
                        for i in range(min(loc.count(), 5)):
                            trouves.append(loc.nth(i))
                    except Exception:
                        continue
                visibles = [t for t in trouves if _visible(t)]
                if visibles:
                    return visibles[0]
            if time.time() >= fin:
                raise ErreurTache(f"champ « {libelle} » introuvable sur la page")
            self.page.wait_for_timeout(500)

    def _envoi_automatique(self, loc, texte=False):
        """Le champ enregistre-t-il tout seul (au changement, ou quand on le quitte pour un texte) ?"""
        attributs = "['onchange', 'onblur']" if texte else "['onclick', 'onchange']"
        try:
            return bool(loc.evaluate("e => " + attributs + ".some(a => e.getAttribute(a)) || "
                                     "!!(e.form && (e.form.getAttribute('onchange') || e.form.getAttribute('onclick')))"))
        except Exception:
            return False

    def presence(self, arg, attente=3):
        """Le texte est-il sur la page (mots entiers) ? Avec « DANS COLONNE x » : dans une ligne du tableau, sous ce titre."""
        if not arg.strip():
            raise ErreurTache("texte vide")
        self.retirer_bandeaux()  # le message du robot n'est pas du texte du portail
        m = re.match(r"^(.*?)\s+DANS\s+COLONNE\s+(.+)$", arg, re.IGNORECASE)
        motif = re.compile(r"(?<![\w-])" + re.escape(arg.strip()) + r"(?![\w-])", re.IGNORECASE)
        fin = time.time() + attente
        while True:
            if m:
                reponses = []
                for cadre in self.cadres():
                    try:
                        reponses.append(cadre.evaluate(JS_COLONNE, [m.group(2).strip(), m.group(1).strip()]))
                    except Exception:
                        continue
                if "oui" in reponses:
                    return True
                if "illisible" in reponses:
                    raise ErreurTache(f"tableau illisible sous la colonne « {m.group(2).strip()} »")
                if "non" in reponses and time.time() >= fin:
                    return False
                if time.time() >= fin:
                    # jamais « rien trouvé, on continue » quand le tableau a changé : c'est une erreur
                    raise ErreurTache(f"colonne « {m.group(2).strip()} » absente de la page")
            else:
                for cadre in self.cadres():
                    try:
                        loc = cadre.get_by_text(motif)
                        if any(_visible(loc.nth(i)) for i in range(min(loc.count(), 10))):
                            return True
                    except Exception:
                        continue
                if time.time() >= fin:
                    return False
            self.page.wait_for_timeout(500)

    def lire(self, arg):
        for cadre in self.cadres():
            try:
                r = cadre.evaluate(JS_LIRE, [arg])
            except Exception:
                continue
            if r is None:
                continue
            if r.get("erreur"):
                raise ErreurTache(f"« {arg} » : {r['erreur']}")
            return r.get("valeur", "")
        raise ErreurTache(f"« {arg} » introuvable sur la page")

    def touche(self, arg):
        nom = TOUCHES[_nom_action(arg)]
        info = None
        for cadre in self.cadres():
            try:
                info = cadre.evaluate(JS_FOCUS)
            except Exception:
                continue
            if info:
                break
        demander_accord = False
        if info:
            dangereux = interdit(info["texte"]) or interdit(info["formulaire"])
            if nom == "Enter":
                # Entrée envoie : seul un formulaire de recherche ordinaire (sans POST, sans action) passe sans question
                demander_accord = dangereux or info["post"] or info["auto"] or (not info["avecFormulaire"] and not info["lien"])
            elif nom == "Space":
                demander_accord = (info["lien"] or info["genre"] in ("checkbox", "radio")) and (dangereux or info["auto"])
            elif nom in ("ArrowDown", "ArrowUp"):
                demander_accord = info["genre"] in ("select", "radio") and (dangereux or info["auto"])
        if demander_accord and not self.accord_geste:
            self.confirmer(None, f"appuyer sur {arg} (cela peut enregistrer quelque chose)", cle="touche " + arg)
            for cadre in self.cadres():  # la question a retiré le curseur du champ : on l'y remet
                try:
                    cadre.evaluate("() => { const e = document.querySelector('[data-robot-focus]');"
                                   " if (e) { e.removeAttribute('data-robot-focus'); e.focus(); } }")
                except Exception:
                    pass
        try:
            with self.page.expect_navigation(timeout=5000):
                self.page.keyboard.press(nom)
        except Exception:
            pass  # pas de changement de page : rien à attendre
        attendre_chargement(self.page)
        self.calme()

    def aller(self, arg):
        cible = urljoin(self.url or "", arg)
        if (urlsplit(cible).hostname or "") != (urlsplit(self.url or "").hostname or ""):
            raise ErreurTache("adresse hors du portail : refusee")
        morceaux = urlsplit(cible)
        action = adresse_prudente(cible)[1] or interdit(unquote(" ".join((morceaux.path, morceaux.query, morceaux.fragment))))
        if action and not self.accord_geste:
            self.confirmer(None, "ouvrir une adresse qui ressemble a une action : " + (morceaux.path or "/"),
                           cle="aller " + re.sub(r"\d+", "#", morceaux.path))
        self.page.goto(cible, wait_until="domcontentloaded", timeout=60000)
        connexion = attendre_portail(self.page)
        # après une connexion, le portail a pu ramener à l'accueil : on rouvre l'adresse, mais JAMAIS une adresse
        # d'action (elle referait l'action)
        if connexion and not action and urlsplit(self.page.url).path != morceaux.path:
            self.page.goto(cible, wait_until="domcontentloaded", timeout=60000)
            attendre_chargement(self.page)
        self.calme()

    # --- les actions
    def executer(self, etapes):
        for etape in etapes:
            self.page_ouverte()
            self.verifier_fenetres()  # une fenêtre arrivée en retard, après le geste précédent
            self.nouveau_geste()
            self.ligne_courante = etape["ligne"]
            action = etape["action"]
            try:
                arg = self.remplir(etape["arg"])
            except ErreurTache as e:
                raise ErreurTache(f"ligne {etape['ligne']} : {e}")
            if action in ("SI", "REPETER"):
                self.bloc(etape, arg)
                continue
            ecrire(f"   ligne {etape['ligne']:>3} : {action.lower()}" + (f" : {arg}" if arg else ""))
            try:
                self.faire(action, arg)
                self.verifier_fenetres()
            except (Passer, Arreter, AVerifier):
                raise
            except ErreurTache as e:
                raise ErreurTache(f"ligne {etape['ligne']} ({action.lower()}) : {e}")
            except Exception as e:
                raise ErreurTache(f"ligne {etape['ligne']} ({action.lower()}) : {expliquer(e)}")
            if action in GESTES:
                self.attendues = []  # une fenêtre annoncée ne vaut que pour le geste qui la suit

    def bloc(self, etape, arg):
        mot = "si" if etape["action"] == "SI" else "repeter tant que"
        for tour in range(31 if etape["action"] == "REPETER" else 1):
            self.calme()  # la réponse au geste précédent d'abord : jamais de décision sur la page d'avant
            self.verifier_fenetres()
            try:
                vrai = self.presence(arg) == etape["present"]
            except ErreurTache as e:
                raise ErreurTache(f"ligne {etape['ligne']} ({mot}) : {e}")
            ecrire(f"   ligne {etape['ligne']:>3} : {mot} {'present' if etape['present'] else 'absent'} : {arg}"
                   f" -> {'oui' if vrai else 'non'}")
            if etape["action"] == "SI":
                self.executer(etape["alors"] if vrai else etape["sinon"])
                return
            if not vrai:
                return
            if tour == 30:
                raise ErreurTache(f"ligne {etape['ligne']} : repete 30 fois sans fin, le robot s'arrete")
            self.executer(etape["alors"])

    def faire(self, action, arg):
        page = self.page
        if self.lecture_seule and action not in ("ACCUEIL", "ALLER", "MENU", "CLIQUER", "ATTENDRE", "VERIFIER", "LIRE",
                                                 "PREVENIR", "PASSER", "ARRETER", "PAUSE"):
            raise Passer(f"visite : le robot ne fait pas « {action.lower()} » (lecture seule)")
        if action == "ACCUEIL":
            if not self.url:
                raise ErreurTache("adresse du portail inconnue")
            page.goto(self.url, wait_until="domcontentloaded", timeout=60000)
            attendre_portail(page)
        elif action == "ALLER":
            self.aller(arg)
        elif action == "MENU":
            self.menu(arg)
        elif action == "CLIQUER":
            self.cliquer(arg)
        elif action in ("ECRIRE", "CHOISIR"):
            libelle, valeur = (x.strip() for x in arg.split("=", 1))
            loc = self.champ(libelle)
            balise = loc.evaluate("e => e.tagName.toLowerCase()")
            if balise == "select":
                options = loc.evaluate("e => Array.from(e.options).map(o => [o.text.trim(), o.value])")
                rang = next((i for i, (t, v) in enumerate(options) if mots(t) == mots(valeur) or v == valeur), None)
                if rang is None:
                    raise ErreurTache(f"choix « {valeur} » absent de la liste « {libelle} »")
                if (interdit(libelle) or interdit(valeur) or self._envoi_automatique(loc)) and not self.accord_geste:
                    self.confirmer(loc, f"choisir « {valeur} » dans « {libelle} »", cle=f"choisir {libelle} {valeur}")
                loc.select_option(index=rang, timeout=10000)
                attendre_chargement(page)
                self.calme()
            elif action == "CHOISIR":
                raise ErreurTache(f"« {libelle} » n'est pas une liste de choix (utilisez ECRIRE)")
            else:
                if (interdit(libelle) or self._envoi_automatique(loc, texte=True)) and not self.accord_geste:
                    self.confirmer(loc, f"ecrire dans « {libelle} » (le portail peut l'enregistrer tout seul)",
                                   cle=f"ecrire {libelle}")
                loc.fill(valeur, timeout=10000)
        elif action in ("COCHER", "DECOCHER"):
            loc = self.champ(arg)
            if (interdit(arg) or self._envoi_automatique(loc)) and not self.accord_geste:
                self.confirmer(loc, f"{action.lower()} « {arg} »", cle=f"{action} {arg}")
            loc.set_checked(action == "COCHER", timeout=10000)
            attendre_chargement(page)
            self.calme()
        elif action == "TOUCHE":
            self.touche(arg)
        elif action == "ATTENDRE":
            m = MOTIF_DUREE.match(arg)
            if m:
                duree = float(m.group(1).replace(",", "."))
                if (m.group(2) or "").lower().startswith("min"):
                    duree *= 60
                page.wait_for_timeout(int(min(duree, 600) * 1000))
            elif not self.presence(arg, attente=30):
                raise ErreurTache(f"« {arg} » n'est pas apparu en 30 secondes")
        elif action == "VERIFIER":
            if not self.presence(arg, attente=10):
                raise ErreurTache(f"« {arg} » absent de la page")
        elif action == "LIRE":
            valeur = self.lire(arg)
            self.lectures[arg] = valeur
            ecrire(f"      {arg} : lu (dans le fichier de resultats)")
        elif action == "ACCEPTER FENETRE":
            self.attendues.append(arg)
        elif action == "CAPTURE":
            nom = re.sub(r"[^\w.-]+", "_", arg)[:40] or "capture"
            page.screenshot(path=str(dossier_robot("resultats") / f"{nom}_{datetime.now():%H%M%S}.png"))
        elif action == "PREVENIR":
            ecrire()
            ecrire("   " + "!" * 66)
            ecrire(f"   !!  {arg}")
            ecrire("   " + "!" * 66)
            try:
                page.evaluate(JS_BANDEAU, arg)
            except Exception:
                pass
            self.lectures.setdefault("prevenu", arg)
        elif action == "PASSER":
            raise Passer(arg or "passe")
        elif action == "ARRETER":
            raise Arreter(arg or "arret demande par la tache")
        elif action == "PAUSE":
            self._sans_focus()
            revenez_ici()
            pause((arg or "Pause") + " - puis Entree pour continuer...")


def _cellule_sure(valeur):
    texte = str(valeur)
    formule = texte[:1] in ("=", "@", "\t", "\r") or (texte[:1] in ("+", "-") and len(texte) > 1)
    return "'" + texte if formule else texte


def choisir_tache():
    """Les fichiers tache*.txt reçus, à côté de ce programme, dans Téléchargements ou Bureau > Robot (plus récents d'abord)."""
    ici = Path(__file__).resolve().parent
    vus = {}
    bases = [ici, Path.home() / "Downloads", Path.home() / "Téléchargements"]
    try:
        bases += [dossier_robot(), dossier_robot("taches")]
    except OSError:
        pass
    for base in bases:
        try:
            for f in base.glob("*.txt"):
                if f.name.lower().startswith("tache"):
                    vus.setdefault(f.resolve(), f)
        except OSError:
            continue
    fichiers = sorted(vus.values(), key=lambda f: f.stat().st_mtime, reverse=True)
    if not fichiers:
        ecrire("   Aucune tache trouvee. Enregistrez le fichier tache_....txt recu dans Telechargements.")
        return None
    ecrire()
    for i, f in enumerate(fichiers, 1):
        ecrire(f"   {i}. {f.name}   (recue le {datetime.fromtimestamp(f.stat().st_mtime):%d/%m a %H:%M})")
    choix = demander("   Numero de la tache a lancer :")
    if not choix.isdigit() or not 1 <= int(choix) <= len(fichiers):
        ecrire("   Pas de tache choisie.")
        return None
    return fichiers[int(choix) - 1]


def executer_tache(p, chemin, url):
    """Lance une tâche sur chaque ligne de sa liste (ou une fois), et écrit le fichier de résultats."""
    import csv
    entete, etapes = lire_tache(chemin)  # une faute de rédaction est dite tout de suite, avant d'ouvrir quoi que ce soit
    if not entete["confirmer"]:
        ecrire("   ATTENTION : cette tache demande de cliquer Supprimer, Valider... SANS vous redemander a chaque fois.")
        if demander("   Tapez OUI (en entier) pour l'accepter, ou Entree pour que le robot vous demande a chaque fois :") != "OUI":
            entete["confirmer"] = True
    variables_fixes = {}
    for nom, question in entete["demander"]:
        valeur = ""
        while not valeur:
            valeur = demander(f"   {question} :")
            if not valeur:
                ecrire("   Reponse obligatoire (Ctrl+C pour arreter).")
        variables_fixes[nom] = valeur
    lignes = lire_liste(entete["liste"], chemin) if entete["liste"] else [{}]
    connus = [mots(n) for n in list(variables_fixes) + (list(lignes[0]) if lignes and lignes[0] else [])]
    for nom, ligne in variables_utilisees(etapes):
        if mots(nom) not in connus:
            raise ErreurTache(f"ligne {ligne} : {{{nom}}} inconnu (colonnes de la liste : "
                              + ", ".join(lignes[0]) + ")" if lignes[0] else f"ligne {ligne} : {{{nom}}} inconnu")
    if entete["liste"]:
        ecrire(f"   Liste : {len(lignes)} ligne(s), premiere : {next(iter(lignes[0].values()), '')}")
    ecrire(f"   Tache « {entete['nom']} » : {len(lignes)} fois.")
    navigateur, contexte, page = ouvrir_robot(p)
    if url:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.bring_to_front()
        revenez_ici()
        pause("Connectez-vous si besoin dans la fenetre du robot. Quand vous voyez l'accueil du portail, Entree...")
        attendre_portail(page)
    nom_fichier = re.sub(r"[^\w-]+", "_", sans_accents(entete["nom"]))[:40] or "tache"
    resultats = dossier_robot("resultats") / f"{nom_fichier}_{datetime.now():%Y%m%d-%H%M%S}.csv"
    colonnes_lues = []
    lignes_csv = []
    compte = {}
    tache = Tache(entete, etapes, page, url)
    ecrit = [resultats]

    def sauver():
        for essai in range(5):
            cible = ecrit[0] if essai == 0 else resultats.with_name(f"{resultats.stem}_{essai + 1}.csv")
            try:
                with open(cible, "w", encoding="utf-8-sig", newline="") as f:
                    w = csv.writer(f, delimiter=";")
                    w.writerow(["N", "element", "statut", "message"] + colonnes_lues)
                    for l in lignes_csv:
                        w.writerow([_cellule_sure(x) for x in l[:4] + [l[4].get(c, "") for c in colonnes_lues]])
                if cible != ecrit[0]:
                    ecrire(f"   (fichier de resultats ouvert ailleurs : ecrit dans {cible.name})")
                    ecrit[0] = cible
                return
            except OSError:
                continue

    def noter(numero, element, statut, message, lectures):
        compte[statut] = compte.get(statut, 0) + 1
        for c in lectures:
            if c not in colonnes_lues:
                colonnes_lues.append(c)
        lignes_csv.append([numero, element, statut, message, dict(lectures)])
        sauver()

    fait = 0
    erreurs_de_suite = 0
    arret = ""
    try:
        for numero, ligne in enumerate(lignes, 1):
            element = next(iter(ligne.values()), "") if ligne else (" ".join(variables_fixes.values()) or "-")
            ecrire()
            ecrire(f"=== {numero}/{len(lignes)} : {element}")
            tache.variables = dict(variables_fixes, **ligne)
            tache.element = element
            tache.nouvel_element()
            try:
                tache.executer(tache.etapes)
                tache.calme()
                tache.verifier_fenetres()
                statut, message = "OK", ""
            except Passer as e:
                statut, message = "PASSE", str(e)
            except AVerifier as e:
                statut, message = "A VERIFIER", str(e)
            except ErreurTache as e:
                statut, message = "ERREUR", str(e)
                try:
                    tache.page.screenshot(path=str(dossier_robot("resultats") / f"erreur_{numero}.png"))
                except Exception:
                    pass
            except Arreter as e:
                statut, message = "ARRETE", str(e)
                arret = f"tache arretee : {e}"
            except KeyboardInterrupt:
                statut, message = "INTERROMPU", "Ctrl+C pendant cet element : verifiez-le sur le portail"
                arret = "arrete a votre demande"
            fait = numero
            erreurs_de_suite = erreurs_de_suite + 1 if statut == "ERREUR" else 0
            ecrire(f"    --> {statut}" + (f" : {message}" if message else ""))
            noter(numero, element, statut, message, tache.lectures)
            if arret:
                break
            if erreurs_de_suite >= 3:
                arret = "3 erreurs de suite : le robot s'arrete (le portail a peut-etre change)"
                break
    except KeyboardInterrupt:
        arret = "arrete a votre demande"
    for numero in range(fait + 1, len(lignes) + 1):
        ligne = lignes[numero - 1]
        noter(numero, next(iter(ligne.values()), "") if ligne else "-", "NON FAIT", arret or "", {})
    ecrire()
    ecrire("=" * 70)
    if arret:
        ecrire(" " + arret.upper())
    ecrire(" FIN : " + ", ".join(f"{n} {s}" for s, n in compte.items()))
    for l in lignes_csv:
        if l[2] not in ("OK", "NON FAIT"):
            ecrire(f"   {l[0]}. {l[1]} : {l[2]} - {l[3][:110]}")
    ecrire(f" Resultats : Bureau > Robot > resultats > {ecrit[0].name}")
    ecrire("=" * 70)
    try:
        pause("Photo de cet ecran, puis Entree pour fermer la fenetre du robot...")
    except KeyboardInterrupt:
        pass
    try:
        navigateur.close()
    except Exception:
        pass
    return 0 if set(compte) <= {"OK", "PASSE"} else 1


# ---------------------------------------------------------------------------- choix 5 et 6 : la carte du site, « aller à »
# Choix 5 : le robot VISITE le portail tout seul. Depuis l'accueil, il ouvre chaque menu (survol), clique sur chaque
# entrée qui mène à une page, puis sur les onglets de cette page, et note la structure de chaque page (comme le relevé).
# Choix 6 : vous dites où aller (« MK2 », « contrats »...), il cherche dans sa carte et y va tout seul, étape par étape.
# C'est une VISITE, en LECTURE SEULE :
# - il ne clique que sur des liens, des entrées de menu et des onglets dont le nom n'est pas une action (jamais
#   Nouveau, Supprimer, Valider, Déconnexion, Imprimer...), jamais dans un formulaire, un tableau ou une fenêtre ;
# - une fenêtre du portail (« Êtes-vous sûr ? ») est toujours refusée ;
# - PROTECTION : rien ne peut être envoyé au portail. Seules les lectures passent ; un envoi (formulaire, application)
#   ou une adresse d'action (…/supprimer, …?action=valider) est bloqué, un téléchargement est annulé.
# La carte (Bureau > Robot > carte_du_site.json) reste sur ce poste. A_ENVOYER_carte.txt : structure seulement.
FICHIER_CARTE = "carte_du_site.json"
ENTETE_VISITE = {"nom": "visite", "demander": [], "liste": None, "confirmer": True}
LIMITE_PAGES, LIMITE_MINUTES, PROFONDEUR = 120, 30, 3
# en visite, le robot ne va jamais là : quitter le portail, imprimer, télécharger, changer la langue ou l'affichage
A_EVITER_EN_VISITE = re.compile(r"deconnex|deconnect|logout|log ?out|log ?off|sign ?out|quitter|\bexit\b|imprim|\bprint|"
                                r"export|telecharg|download|\blangues?\b|language|english|deutsch|espanol|\btheme\b|"
                                r"mode sombre|dark mode|plein ecran|full ?screen")
EXPRESSIONS_ACTION = re.compile(r"mise a jour|mettre a jour|prendre en charge|prise en charge|remettre a zero|"
                                r"mettre en|mise en (service|production|ligne|place|attente|forme)")
MOTS_ACTION_VISITE = {"nouveau", "nouvelle", "nouvel", "new", "oui", "ok", "yes", "non", "no", "go", "epingler",
                      "desepingler", "marquer", "follow", "unfollow", "unpin", "favori", "lu", "signet"}
# une entrée faite d'une icône seule (« icone trash ») qui fait une ACTION quand on clique dessus : un des MOTS de son
# nom est une action (poubelle, crayon, plus, coche, croix, étoile des favoris, partage, pièce jointe, envoi...), les
# mots de style (circle, square, solid, fill...) mis à part : « pen-to-square », « circle-plus », « trash-can » ; ou
# c'est une bascule (déconnexion, menu replié, langue, thème, plein écran). « pin_drop » (une carte), « fact_check »
# (un module de contrôles), « remove_red_eye » (voir) ne sont pas des actions.
JETONS_D_ACTION = set("""trash delete remove bin recycle garbage broom eraser pencil pen edit create plus add minus
xmark x times close cancel clear check checkmark done tick undo redo sync share send reply forward paperclip attach
attachment star grade heart like favorite favorites thumbs thumb bookmark pin pushpin thumbtack keep ban block scissors cut
copy paste
clone play stop pause power poweroff logout signout save floppy upload publish archive unarchive restore lock unlock
print download mark drafts off form gavel vote plane move unlink merge split toggle radio approve published unpublished
push
markunread rate recommend import export slash""".split())
JETONS_DE_STYLE = set("""circle square solid outline outlined line fill filled round rounded sharp alt o two tone
regular thin bold duotone rect rectangle lg sm""".split())
ICONE_DE_BASCULE = re.compile(
    r"(bars|navicon|hamburger|menu|menu_open|dehaze|reorder|first_page|last_page|keyboard_double_arrow_(left|right)|"
    r"globe|translate|g_translate|language|dark_mode|light_mode|contrast|moon|sun|brightness\w*|circle-half-stroke|"
    r"fullscreen\w*|compress|maximize|minimize|log-?out|logout|sign-?out\w*|exit_to_app|(arrow-)?right-from-bracket|"
    r"door-(open|closed)|box-arrow-[\w-]+|power_settings_new|get_app|how_to_vote|highlight_off|"
    r"(file|cloud|folder)-arrow-(up|down)|arrow-up-from-bracket|arrows?-rotate\w*|rotate(-left|-right)?|\w+_off|[\w-]+-slash)", re.I)
ICONES_SANS_ACTION = re.compile(r"(fact_check|remove_red_eye|pin_drop|location_pin|map-pin|map-marker\w*|"
                                r"push_pin_drop|edit_location_alt|checklist\w*)", re.I)


# le mot d'une icône d'action FRANCHE, même à côté d'un nom (« <mat-icon>logout</mat-icon> Sortir », « add Plan ») ;
# pas « archive Archives », « publish Publications », « lock Sécurité », « edit_road Voirie » (des modules)
LIGATURES_D_ACTION = set("""add add_circle add_box delete delete_forever delete_outline remove remove_circle edit edit_note
create save save_alt send upload close clear cancel logout exit_to_app power_settings_new post_add playlist_add note_add
person_add group_add content_copy file_copy""".split())


def _icone_d_action(nom):
    """« icone trash », « icone pen-to-square », « icone menu » : une icône qui fait une action ou une bascule (et pas
    « icone pin_drop » ni « icone description »)."""
    t = " ".join(str(nom or "").split()).lower()
    if not t.startswith("icone "):
        return False
    icone = re.sub(r" \d+$", "", t[6:]).strip()
    if ICONES_SANS_ACTION.fullmatch(icone):
        return False
    if ICONE_DE_BASCULE.fullmatch(icone):
        return True
    # le PREMIER mot qui n'est pas un mot de style dit ce que fait l'icône (« pen-to-square », « circle-plus », « trash3 ») ;
    # « list-check », « calendar-plus », « user-check » sont des modules (liste, agenda, utilisateurs)
    jetons = [re.sub(r"\d+$", "", j) for j in re.split(r"[-_\s]+", icone) if j and j not in JETONS_DE_STYLE and not j.isdigit()]
    # (ou un DERNIER mot qui ajoute, retire ou envoie : « file-circle-plus », « folder-minus », « paper-plane »)
    return bool(jetons) and (jetons[0] in JETONS_D_ACTION or
                             (len(jetons) > 1 and jetons[-1] in {"plus", "minus", "xmark", "x", "dash", "plane", "export", "import"}))


def _icones_de(texte, attributs):
    """Les noms d'icônes d'un élément examiné (JS_EXAMINER) : mot d'une icône Material (texte d'un seul mot en
    minuscules), classe d'une police d'icônes (fa-trash, bi-pencil...), composant ou symbole d'icône."""
    noms = [texte] if re.fullmatch(r"[a-z][a-z_]{2,30}", texte or "") else []
    noms += re.findall(r"(?<![\w-])(?:fa|fas|far|mdi|bi|icon|ico|glyphicon|pi|k-i|k-svg-i|ti|ri|la|lucide|feather|bx|bxs|"
                       r"octicon|ion|ph)-([a-z][\w-]{1,30})", attributs or "", re.I)
    noms += re.findall(r"(?:data-mat-icon-name|svgicon|data-icon|iconname|data-icon-name) ([A-Za-z][\w-]{1,30})", attributs or "")
    # (symbole d'une planche d'icônes : « xlink:href #trash », « href #icon-trash » ; jamais « href #/plans » d'un lien)
    noms += re.findall(r"xlink:href [^|]*#([A-Za-z][\w-]{1,30})", attributs or "")
    noms += re.findall(r"(?<![\w:-])href [^|#]*#(?:icon|ico|i|svg|sprite|symbol)[-_]([A-Za-z][\w-]{1,30})", attributs or "")
    return noms


# mots d'une question sans importance pour chercher dans la carte
MOTS_VIDES = set("""a au aux avec ce ces cet cette d de des du en et l la le les ou par pour sur un une dans vers aller
va vais veux voudrais voir je j me moi page pages menu menus ecran onglet ici svp""".split())

JS_NOM_PAGE = r"""() => {
  const vis = e => { const s = getComputedStyle(e); if (s.display === 'none' || s.visibility === 'hidden') return false;
                     const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0; };
  const HORS = 'nav, aside, [role=navigation], [role=menu], [role=menubar], [role=dialog], [role=alertdialog], dialog, ' +
               '[data-robot-bandeau], #robot-cadres';
  const texte = e => (e.innerText || e.textContent || '').replace(/\s+/g, ' ').trim();
  // les onglets ouverts (onglet, puis sous-onglet) : des onglets qui ne changent pas l'adresse sont des pages de la carte.
  // Seulement de vrais onglets (rôle tab, ou li / a / bouton d'une barre d'onglets), jamais un élément choisi dans le
  // contenu (ligne de liste, pagination, tableau).
  const jet = x => (typeof x.className === 'string' ? x.className : '').toLowerCase().split(/[^a-z0-9]+/);
  const ONG = ['tab', 'tabs', 'tabset', 'tablist', 'onglet', 'onglets', 'nav-tabs'], ACT = ['active', 'actif', 'selected', 'current'];
  const PAS_ONGLET = '[role=tabpanel], .tab-content, .tab-pane, table, [role=grid], [role=listbox], .list-group, .pagination, ' +
                     '[role=dialog], dialog, [data-robot-bandeau], #robot-cadres';
  let actifs = Array.from(document.querySelectorAll('[role=tab][aria-selected=true]')).filter(e => vis(e) && !e.closest(PAS_ONGLET));
  if (!actifs.length)
    for (const c of document.querySelectorAll('[class]')) {
      if (!jet(c).some(m => ONG.includes(m)) || c.closest(PAS_ONGLET) || c.closest(HORS) || !vis(c)) continue;
      const k = Array.from(c.querySelectorAll('li, a, button')).find(k => (jet(k).some(m => ACT.includes(m)) ||
                  k.getAttribute('aria-current')) && vis(k) && !k.closest(PAS_ONGLET) && texte(k) && texte(k).length <= 40);
      if (k) { actifs = [k]; break; }
    }
  const onglet = actifs.slice(0, 3).map(e => texte(e).slice(0, 40)).filter(Boolean).join(' > ');
  // le titre de la page : un titre hors de l'en-tête d'abord (l'en-tête porte souvent le nom de l'application).
  // 5e valeur : true si ce titre est sans doute le NOM DE L'APPLICATION (en-tête, titre de la fenêtre)
  // (d'abord un titre hors de tout en-tête ; puis l'en-tête d'une PAGE, dans main ou article ; enfin l'en-tête de
  // l'application, « banniere »)
  const genre = e => { const h = e.closest('header, [role=banner]');
    return !h ? 'libre' : (h.parentElement && h.parentElement.closest('main, article, [role=main]') ? 'page' : 'banniere'); };
  for (const [sel, g] of [['h1', 'libre'], ['[role=heading][aria-level="1"]', 'libre'], ['h2', 'libre'], ['h1', 'page'],
                          ['h2', 'page'], ['h1', 'banniere']])
    for (const e of document.querySelectorAll(sel)) {
      if (!vis(e) || e.closest(HORS) || genre(e) !== g) continue;
      if (texte(e)) return [texte(e).slice(0, 200), onglet, true, innerWidth * innerHeight, g === 'banniere'];
    }
  // pas de titre dans la page : le titre de la fenêtre (false : à prendre seulement si aucune zone n'a de vrai titre) ;
  // c'est le nom de l'application, sauf dans un cadre qui a son propre titre
  let appli = true;
  try { appli = window.top === window || document.title === window.top.document.title; } catch (x) { appli = true; }
  return [(document.title || '').replace(/\s+/g, ' ').trim().slice(0, 200), onglet, !document.querySelector('frameset'),
          innerWidth * innerHeight, appli];
}"""

# Pour Claude, quand la visite ne trouve presque rien : les éléments cliquables du HAUT et de la GAUCHE de la page
# (là où sont les menus), décrits par leur structure seulement (balise, classes communes, nombre de mots, icône).
JS_ZONES = r"""() => {""" + JS_DETECTION + JS_OUTILS_MENUS + r"""
  const partagees = x => classe(x).split(/\s+/).filter(c => c && !/\d{2,}/.test(c) && document.getElementsByClassName(c).length >= 2).slice(0, 3);
  const sig = x => x.tagName.toLowerCase() + partagees(x).map(c => '.' + c).join('');
  const liste = detecter().filter(x => x.r.left < innerWidth * 0.3 || x.r.top < 130)
                          .sort((a, b) => (a.r.top - b.r.top) || (a.r.left - b.r.left));
  return liste.slice(0, 40).map(x => { const e = x.e, t = texteDe(e), ic = (compteDe(e) || enregistrement(e)) ? '' : cleIcone(e);
    const parents = [];
    for (let p = e.parentElement, n = 0; p && p !== document.body && n < 3; p = p.parentElement, n++) parents.push(sig(p));
    return { zone: x.r.top < 130 ? 'haut' : 'gauche', sig: sig(e), role: e.getAttribute('role') || '',
             mots: t ? t.split(/\s+/).length : 0, icone: ic.startsWith('icone ') ? ic.slice(6) : (ic ? '(bulle)' : ''),
             source: x.s, ecoute: e.getAttribute('data-robot-ecoute') || '', survol: e.getAttribute('data-robot-survol') || '',
             popup: e.getAttribute('aria-haspopup') || (e.hasAttribute('aria-expanded') ? 'expanded' : ''),
             taille: Math.round(x.r.width) + 'x' + Math.round(x.r.height), x: Math.round(x.r.left), y: Math.round(x.r.top),
             parents: parents.join(' < '), menu: !!e.closest('[data-robot-barre]') }; });
}"""

JS_BANDEAU_VISITE = r"""(message) => {
  let b = document.getElementById('robot-visite');
  if (!message) { if (b) b.remove(); return; }
  if (!b) {
    b = document.createElement('div'); b.id = 'robot-visite'; b.setAttribute('data-robot-bandeau', '1');
    b.style.cssText = 'position:fixed;left:50%;bottom:8px;transform:translateX(-50%);z-index:2147483647;pointer-events:none;' +
      'background:#fde047;color:#111;font:bold 15px/20px Arial,sans-serif;padding:6px 12px;border:2px solid #111;' +
      'border-radius:6px;box-shadow:0 2px 8px rgba(0,0,0,.4);max-width:90vw;white-space:nowrap;overflow:hidden;text-overflow:ellipsis';
    document.documentElement.appendChild(b);
  }
  b.textContent = 'ROBOT : ' + message;
}"""


def _valeur_masquee(v):
    v = unquote(str(v))
    return "#" if (re.search(r"\d{2,}", v) or AUTRES_DONNEES.search(v) or len(v) > 30 or NOM_DE_PERSONNE.search(v)) else v


SECRET = re.compile(r"pass|pwd|mdp|token|jeton|session|sessid|^sid$|auth|apikey|api_key|secret", re.IGNORECASE)


def _cle_url(url):
    """L'adresse d'une page, les valeurs qui ressemblent à des données remplacées par # (« /plan?n=# » : toutes les
    fiches de plans sont la même page). Elle reste sur ce poste (carte), jamais dans un fichier à envoyer."""
    try:
        m = urlsplit(url)
    except ValueError:
        return ""

    def partie(chemin, requete):
        c = "/".join(_valeur_masquee(s) if s else "" for s in chemin.split("/"))
        q = "&".join(sorted(f"{k}=" + ("#" if SECRET.search(k) else _valeur_masquee(v))
                            for k, v in parse_qsl(requete, keep_blank_values=True)))
        return c + ("?" + q if q else "")

    cle = partie(m.path, m.query)
    if m.fragment:
        f, _, q = m.fragment.partition("?")
        cle += "#" + partie(f, q)
    return cle


# un verbe d'action collé à un nom dans une adresse (« deleteplan », « updateStatut » écrit en minuscules)
VERBES_COLLES = ("delete", "remove", "supprimer", "update", "create", "archive", "cancel", "annuler", "valider",
                 "approve", "reject", "submit", "cloturer", "modifier", "upsert", "insert")


# un nom d'action qui suit un numéro (« /commandes/42/annulation ») ; un verbe à l'infinitif aussi (« /taches/3/terminer »)
NOMS_D_ACTION = set("""annulation suppression refus rejet resiliation revocation abandon cloture approbation envoi
diffusion emission supprime supprimee supprimes efface annule detruit retire purge""".split())
# les clés d'ordre d'une adresse lue (« ?action= », « ?cmd= » ; pas « ?event= » d'un journal)
ORDRES_GET = set("action op operation cmd command method verb do activity function fn call exec".split())
# la valeur d'une clé d'ordre (« ?action=consulter ») qui ne fait que montrer quelque chose
ORDRES_DE_LECTURE = set("""view voir afficher affichage consulter consultation detail details fiche show display list
liste lister search recherche rechercher get load init index read lire open ouvrir""".split())


def _adresse_d_action(url):
    """Une adresse (lue par GET) qui fait une ACTION : un verbe d'action (« .../supprimer »), « markAllAsRead »,
    un verbe collé à un nom (« deleteplan »), un ordre à côté d'un numéro (« /demandes/42/refuser », « /plans/cloture/42 »,
    « /commandes/42/annulation »), ou une clé d'ordre qui ne fait pas que montrer (« ?action=cloture »). Une liste
    (« /plans/a-valider », « /api/v1/statuts », « /plans/42/validations ») reste une lecture."""
    try:
        if adresse_prudente(url)[1]:
            return True
        morceaux = urlsplit(url)
        # (un identifiant fait de lettres et de chiffres, « S64uHM50QpNoY0Nb », n'est pas découpé en mots)
        brut = re.sub(r"[A-Za-z0-9_-]{10,}", lambda m: "0" if re.search(r"\d", m.group(0)) and re.search(r"[A-Za-z]", m.group(0))
                      else m.group(0), unquote(morceaux.path + " " + morceaux.fragment))
        ws = mots(brut).split()
        # « markAsRead », « mark-all-read », « toggleFavori » (pas « feature-toggle », « sidenav-toggle »)
        for segment in re.split(r"[/#]", unquote(morceaux.path + "#" + morceaux.fragment)):
            m = mots(segment).split()
            if len(m) > 1 and m[0] in ("mark", "marquer", "toggle"):
                return True
        for w in ws:
            for v in VERBES_COLLES:  # (« deleteplan » ; pas « updated », « cancellations », « submitted »)
                reste = w[len(v):]
                if w.startswith(v) and len(reste) > 2 and \
                        not re.fullmatch(r"(e?d|l?led|ted|s|d?ions?|e?urs?|ings?|t?als?|ables?|ments?|lations?|ations?)", reste):
                    return True
        # « .../notifications/seen », « .../alerts/42/dismiss » : marquer vu, lu, écarté ; « ?query=mutation{...} »
        segments = [x for x in unquote(morceaux.path).split("/") if x]
        if segments and mots(segments[-1]).replace(" ", "") in {"seen", "readall", "markallread", "markread", "markseen",
                                                                "ack", "acknowledge", "dismiss", "lu", "vu"}:
            return True
        if re.search(r"\b(mutation|subscription)\b", unquote(morceaux.query)):
            return True
        numero = lambda x: bool(re.fullmatch(r"\d+|[0-9a-f]{8,}", x))
        for i, w in enumerate(ws):
            voisins = ws[max(0, i - 1):i] + ws[i + 1:i + 2]
            ordre = w in VERBES_ACTION | MOTS_TACHE | NOMS_D_ACTION or (re.search(r"(er|ir|re)$", w) and interdit(w))
            if ordre and any(numero(x) for x in voisins):
                return True
        for cle, valeur in parse_qsl(morceaux.query, keep_blank_values=True):
            mots_cle, mots_valeur = mots(cle).split(), mots(valeur).split()
            if mots_cle and mots_cle[-1] in ORDRES_GET and mots_valeur and \
                    not set(mots_valeur) <= ORDRES_DE_LECTURE | {"0", "1"}:
                return True
        return False
    except Exception:
        return True


# la connexion (identifiant, authentification unique de l'entreprise) passe aussi par des envois : ils ne modifient
# rien sur le portail, et sans eux la session expirerait pendant la visite
MOTS_CONNEXION = set("saml saml2 oauth oauth2 openid oidc token sso signin login logon adfs wsfed".split())


def _envoi_de_connexion(req):
    try:
        chemin = unquote(urlsplit(req.url).path)
        jetons = set(mots(chemin).split())
        if jetons & VERBES_ACTION or interdit(chemin):  # (…/token/revoke, …/login/reset : des actions)
            return False
        if jetons & MOTS_CONNEXION:
            return True
        corps = req.post_data or ""
    except Exception:
        return False
    return bool(re.search(r"(^|&)(SAMLResponse|SAMLRequest|wresult|id_token|grant_type)=", corps))


# un envoi qui ne fait que LIRE : beaucoup d'applications demandent leurs listes par un envoi (recherche, page d'une
# grille, GraphQL « query ») ; sans lui, la page reste vide pendant la visite. Permis seulement dans 3 cas sûrs :
# une demande GraphQL « query », une demande de lignes d'une grille (AG Grid), ou une adresse qui FINIT par un verbe de
# recherche ou de liste (« .../plans/search », « .../getRows ») ; jamais le nom d'une collection (« POST .../rapports »
# crée un rapport), jamais une adresse ou un contenu qui contient un mot d'action ou d'écriture.
VERBES_LECTURE = set("""search recherche rechercher find list liste lister get lookup autocomplete typeahead suggest
fetch""".split())
# (« getRows », « searchPlans » : le verbe de lecture est le début du dernier morceau de l'adresse)
DEBUT_LECTURE = re.compile(r"(get|search|find|list|lookup|fetch)(rows|list|all|data|items|page|results|by(id|ids|name|code|key)?)?")
# après le verbe de lecture : jamais « fetchAndClaim », « getOrCreate », « getNextTask » (prend la tâche suivante)
SUITE_INTERDITE = {"and", "or", "next", "then", "et", "ou", "suivant", "suivante"}
# le dernier morceau d'une adresse de lignes de grille (sinon, un verbe de lecture)
FINS_DE_GRILLE = set("rows row data grid datasource results items records lignes page".split())
# mots d'une ÉCRITURE (en plus des verbes d'action) : marquer comme lu, épingler, suivre, enregistrer un état...
MOTS_ECRITURE = set("""set mark marquer toggle increment decrement patch merge put post change changer modif modifier
store persist write ecrire insert upsert register track tracking log logs logging record follow unfollow subscribe
unsubscribe like unlike favori favoris favorite favorites pin unpin epingler star unstar vote rate rating read lu lue
seen vu ack acknowledge dismiss hide masquer collect beacon telemetry metrics analytics event events""".split())
# les ordres d'un flux de travail (prendre, terminer, accepter une tâche...) : dans l'adresse, dans un nom de clé du
# contenu, ou comme valeur d'une clé d'ordre (« "method": "task.claim" »)
MOTS_TACHE = set("""complete completer claim unclaim accept accepter decline refuse refuser assign reassign start stop
finish terminer cloturer approve reject submit execute run launch lancer trigger declencher notify notifier send
generate generer compute calculer recalculer enable disable link unlink attach detach share partager transfer escalate
resolve reopen rouvrir""".split())
# dans l'ADRESSE : un enregistrement de réglages (état d'une grille, préférences, disposition)
MOTS_ETAT = set("""state etat preference preferences settings setting layout config configuration parametres parametre
personalization personnalisation""".split())
# une page de serveur à l'ancienne, un canal de messages, un lot de demandes : ils peuvent faire n'importe quoi
ADRESSE_SERVEUR = re.compile(r"\.(aspx|ashx|xhtml|jsf|faces|do|action|php|jsp|cfm|pl|cgi)$|^(\$?batch|hubs?|signalr|"
                             r"negotiate|socket\.io|cometd|rpc|jsonrpc|json-rpc|xmlrpc|xml-rpc|soap|livewire|wicket|ajax)$", re.I)
# les clés d'un contenu qui donnent un ORDRE (leur valeur est un nom d'opération)
CLES_D_ORDRE = set("""action op operation cmd command method verb target event handler activity function fn call exec
task""".split())
# les clés d'une demande de lignes d'une grille AG Grid (modèle « côté serveur ») : rien d'autre
CLES_AG_GRID = set("""startrow endrow rowgroupcols valuecols pivotcols pivotmode groupkeys filtermodel sortmodel""".split())
CLES_GRAPHQL = {"query", "variables", "operationname", "extensions"}
# mots techniques d'une adresse qu'on peut écrire dans la carte (jamais un nom, un identifiant ni une valeur)
MOTS_TECHNIQUES = set("""api apis rest rs ws service services svc graphql gql odata batch data datas json rpc app apps
application backend bff gateway public private internal core common shared web ajax handler controller endpoint
resource resources ressources plans plan dossiers dossier documents document fichiers fichier modules module user users
utilisateur utilisateurs me profil profile session sessions log logs logging telemetry track tracking metrics analytics
events event collect beacon notifications notification news actualites messages message items item entities entity
signalr hub hubs negotiate rows row grid search recherche list liste query lookup""".split())


def _chemin_technique(url):
    """La fin d'une adresse, réduite aux mots techniques connus (« .../api/…/search ») : jamais un nom ni une valeur."""
    try:
        morceaux = [x for x in unquote(urlsplit(url).path).split("/") if x][-4:]
    except ValueError:
        return "…"
    connus = MOTS_TECHNIQUES | VERBES_LECTURE | MOTS_ECRITURE | MOTS_CONNEXION | VERBES_ACTION
    garde = []
    for m in morceaux:
        ws = mots(m).split()
        bon = re.fullmatch(r"v\d{1,2}", m) or (ws and re.fullmatch(r"[A-Za-z_.$-]+", m) and all(w in connus for w in ws))
        garde.append(m if bon else "…")
    return ".../" + "/".join(garde)


def _ordre_dans_demande(requete):
    """La partie « ?... » d'une adresse donne un ORDRE : clé d'ordre (« ?action= », « ?pyActivity= », « ?_method= »),
    clé qui est un ordre (« ?accept=5 »), ou demande GraphQL « mutation »."""
    for cle, valeur in parse_qsl(requete or "", keep_blank_values=True):
        mots_cle = mots(cle).split()
        if not mots_cle:
            continue
        if mots_cle[-1] in CLES_D_ORDRE | {"do", "method"} or ("".join(mots_cle) in MOTS_TACHE and mots_cle != ["start"]):
            return True
        if re.search(r"\b(mutation|subscription)\b", valeur):
            return True
    return False


def _envoi_de_lecture(req):
    """Un envoi (POST) de l'application (jamais d'un formulaire de page) qui ne fait que LIRE : demande GraphQL
    « query » (jamais « mutation »), demande de lignes d'une grille, ou adresse qui finit par un verbe de recherche ou
    de liste ; contenu vide ou JSON ; et aucun mot d'action ou d'écriture, ni dans l'adresse ni dans le contenu."""
    try:
        if req.resource_type not in ("xhr", "fetch"):
            return False
        entetes = {str(k).lower(): str(v).lower() for k, v in (req.headers or {}).items()}
        # (« POST » qui annonce un autre ordre : X-HTTP-Method: DELETE, MERGE... ; formulaire ou fichier envoyé)
        if {"x-http-method", "x-http-method-override", "x-method-override"} & set(entetes) or \
                re.search(r"multipart|form-urlencoded", entetes.get("content-type", "")):
            return False
        morceaux = urlsplit(req.url)
        chemin, demande = unquote(morceaux.path), unquote(morceaux.query)
        # (« list.do;jsessionid=... » : les paramètres d'un morceau d'adresse sont ôtés)
        segments = [x.split(";")[0] for x in chemin.split("/") if x.split(";")[0]]
        mots_adresse = set(mots(chemin + " " + demande).split())
        if (action_en_visite(chemin + " " + demande) or mots_adresse & (VERBES_ACTION | MOTS_ECRITURE) or
                set(mots(chemin).split()) & (MOTS_ETAT | MOTS_TACHE) or
                _ordre_dans_demande(morceaux.query) or
                any(ADRESSE_SERVEUR.search(x) for x in segments)):
            return False
        corps = req.post_data or ""
    except Exception:
        return False
    if len(corps) > 20000:
        return False
    try:
        objet = json.loads(corps) if corps.strip() else None
    except ValueError:
        return False  # ni vide ni JSON (formulaire, texte, XML) : jamais une lecture sûre
    demandes = objet if isinstance(objet, list) else [objet]
    # GraphQL : par son adresse (« graphql », « gql », « .../query ») ou par son contenu ; seul ce contrôle décide
    chemin_gql = bool(segments) and ({"graphql", "gql"} & {s.lower() for s in segments} or segments[-1].lower() == "query")
    contenu_gql = any(isinstance(d, dict) and ("operationName" in d or (isinstance(d.get("query"), str) and
                      re.match(r"\s*(query|mutation|subscription|fragment)\b|\s*\{", d["query"]))) for d in demandes)
    if chemin_gql or contenu_gql:
        return bool(objet is not None and demandes and chemin_gql and
                    all(isinstance(d, dict) and _graphql_lecture(d) for d in demandes))
    if objet is not None and not isinstance(objet, (dict, list)):
        return False
    dernier = mots(segments[-1]).split() if segments else []
    verbe = bool(dernier) and (dernier[0] in VERBES_LECTURE or bool(DEBUT_LECTURE.fullmatch(dernier[0]))) and \
        not SUITE_INTERDITE & set(dernier[1:])
    if isinstance(objet, dict):
        cles = {mots(k).replace(" ", "") for k in objet}
        if {"startrow", "endrow"} <= cles <= CLES_AG_GRID:
            # lignes d'une grille (de telle ligne à telle ligne, tri, filtre) : rien d'autre dans la demande ; les valeurs
            # (« dateModification », filtre « set ») sont des noms de colonnes, pas des ordres
            return verbe or bool(dernier) and dernier[-1] in FINS_DE_GRILLE
    if (VERBES_ACTION | MOTS_ECRITURE) & set(mots(corps).split()):
        return False
    # les noms des clés (« "complete": true ») et la valeur d'une clé d'ordre (« "method": "task.claim" »)
    for cle, valeur in re.findall(r'"([^"\\]{1,60})"\s*:\s*("(?:[^"\\]|\\.){0,200}"|[^,}\]]{0,40})', corps):
        mots_cle = mots(cle).split()
        if set(mots_cle) & VERBES_ACTION or ("".join(mots_cle) in MOTS_TACHE and mots_cle != ["start"]):
            return False
        if "".join(mots_cle) in CLES_D_ORDRE and (set(mots(valeur).split()) & (VERBES_ACTION | MOTS_ECRITURE | MOTS_TACHE) or
                                                 action_en_visite(valeur)):
            return False
    return verbe


def _graphql_lecture(d):
    """Une demande GraphQL qui ne fait que LIRE : un document « query » (jamais « mutation » ni « subscription »), sans
    champ qui donne un ordre (« completeTask(...) », « markAllAsRead », « logout »). Les champs d'état (« canEdit »,
    « isRead », « dateValidation ») sont des lectures."""
    q = d.get("query")
    if not isinstance(q, str) or not q.strip() or not {sans_accents(k).lower() for k in d} <= CLES_GRAPHQL:
        return False
    if re.search(r"\b(mutation|subscription)\b", q):
        return False
    texte = re.sub(r"\b(can|is|has|date|last|nb|count|total|nombre)[A-Z_]\w*", " ", q)
    ws = set(mots(texte).split())
    if ws & (VERBES_ACTION | MOTS_ECRITURE) or ws & {"logout", "signout", "deconnexion", "deconnecter"}:
        return False
    champs = [mots(c).split() for c in re.findall(r"([A-Za-z_]\w*)\s*[({]", texte)]
    return not any(c and c[0] in MOTS_TACHE for c in champs)


def proteger_en_visite(contexte):
    """Lecture seule pour tout le navigateur du robot (à mettre APRÈS votre connexion) : chaque demande au portail
    passe par ce contrôle. Seules les lectures (GET) passent, et jamais vers une adresse d'action ; tout envoi (POST,
    PUT, DELETE... : formulaire, application) est bloqué, sauf ce qui sert à la connexion.
    (Ce contrôle est fait par le robot lui-même : pendant qu'il attend une réponse dans la fenêtre noire, il doit
    continuer à faire vivre le navigateur, sinon la page attendrait ; voir demander_en_vivant et _garde_coupee.)"""
    garde = {"bloquees": 0, "exemples": [], "liaison": False, "contexte": contexte, "liaison_bloquee": False,
             "lectures": 0, "exemples_lectures": []}

    def filtre(route):
        req, methode = None, "?"
        try:
            req = route.request
            methode = req.method.upper()
            lecture = methode == "POST" and not _envoi_de_connexion(req) and _envoi_de_lecture(req)
            bloquer = ((methode not in ("GET", "HEAD", "OPTIONS") and not (methode == "POST" and _envoi_de_connexion(req))
                        and not lecture)
                       or (req.resource_type in ("document", "xhr", "fetch", "ping", "eventsource", "other")
                           and _adresse_d_action(req.url)))
            if lecture and not bloquer:
                garde["lectures"] += 1
                exemple = f"POST {_chemin_technique(req.url)}"
                if exemple not in garde["exemples_lectures"] and len(garde["exemples_lectures"]) < 8:
                    garde["exemples_lectures"].append(exemple)
        except Exception:
            bloquer = True
        try:
            if bloquer:
                garde["bloquees"] += 1
                if req is not None:
                    # pour Claude : le genre d'envoi et le mot d'action de son adresse, rien d'autre de l'adresse
                    # (un morceau d'adresse peut contenir un identifiant ou un nom)
                    genre = "page" if req.resource_type == "document" else "application"
                    exemple = f"{methode} {_chemin_technique(req.url)} ({genre})"
                    if exemple not in garde["exemples"] and len(garde["exemples"]) < 10:
                        garde["exemples"].append(exemple)
                route.abort("blockedbyclient")
            else:
                route.continue_()
        except Exception:
            pass

    garde["filtre"] = filtre
    contexte.route("**/*", filtre)
    # liaison directe (websocket : SignalR, Blazor, socket.io...) : ses messages ne passent pas par le contrôle
    # ci-dessus. Elle est remplacée par une liaison factice, jamais reliée au portail (une page ouverte avant garde la
    # sienne jusqu'à son rechargement : programme_enfant recharge la page).
    if hasattr(contexte, "route_web_socket"):
        try:
            contexte.route_web_socket("**/*", lambda ws: None)
            garde["liaison_bloquee"] = True
        except Exception:
            pass

    def liaison(_):
        garde["liaison"] = True  # (bloquée si liaison_bloquee ; sinon la visite s'arrête : voir _liaison_libre)

    for p in contexte.pages:
        p.on("websocket", liaison)
    contexte.on("page", lambda p: p.on("websocket", liaison))
    return garde


def _liaison_libre(garde):
    """Une liaison directe avec le portail que le pilote de navigateur (trop ancien) ne sait pas bloquer : la lecture
    seule n'est plus garantie, la visite s'arrête."""
    return garde["liaison"] and not garde["liaison_bloquee"]


def _garde_coupee(garde, faire):
    """Le portail redemande de se connecter : la protection est retirée le temps que VOUS vous connectiez (le robot
    attend votre Entrée, il ne contrôlerait plus rien et la page attendrait), puis remise."""
    garde["contexte"].unroute("**/*", garde["filtre"])
    try:
        return faire()
    finally:
        garde["contexte"].route("**/*", garde["filtre"])


def demander_en_vivant(page, question, defaut=""):
    """Comme demander(), mais le navigateur du robot continue de vivre pendant que vous répondez (la protection de la
    visite est faite par le robot : sans cela, la page du portail attendrait)."""
    vider_clavier()
    print(question + " ", end="", flush=True)
    tampon = []
    while True:
        ligne = ligne_tapee(tampon)
        if ligne is not None:
            return ligne.strip() or defaut
        if FIN_DU_CLAVIER[0]:
            return defaut
        try:
            page.wait_for_timeout(150)
        except Exception:
            time.sleep(0.15)


def oui_en_vivant(page, question, defaut=""):
    while True:
        reponse = demander_en_vivant(page, f"{question} (o/n)" + (f" [{defaut}]" if defaut else "") + " :", defaut).lower()
        if reponse[:1] in ("o", "y", "n"):
            return reponse[:1] != "n"
        ecrire("   Repondez o ou n.")


def _nom_de_page(texte):
    """Le titre d'une page tel qu'il est écrit dans la carte : sans numéro, code, adresse, nom de serveur ni nom de
    personne. Le texte ENTIER est contrôlé avant d'être coupé (« ... pour Jean Martin » ne sort pas coupé en « Jean »)."""
    t = " ".join(str(texte or "").split())
    if not t:
        return ""
    t = _sans_donnee(t)  # « Plan PL-10 » -> « Plan # » (puis le texte entier est contrôlé, avant d'être coupé)
    # « Dossier : Boulangerie Kowalski », « Client : ... » : après les deux-points, c'est souvent la fiche elle-même
    t = re.sub(r"\s*:\s*\S.*$", " : #", t)
    if NOM_DE_PERSONNE.search(t) or _prenom_nom(t):
        return "(donnee masquee)"
    return t[:60]


def _identite(page):
    """(identité de la page, son titre) : son adresse sans les valeurs (et celles de ses cadres) + son titre + l'onglet
    ouvert. Deux pages de même identité sont la même page (deux fiches de plans, par exemple)."""
    return _identite_complete(page)[:2]


def _identite_complete(page):
    """(identité, titre, titre de l'application ?) : le 3e est vrai quand le titre vient de l'en-tête ou du titre de la
    fenêtre (le nom de l'application, souvent celui de l'entreprise : il n'est pas écrit dans la carte)."""
    cles, titres, autres, onglets = [_cle_url(page.url)], [], [], []
    for cadre in page.frames:
        if cadre != page.main_frame and (cadre.url or "").startswith("http"):
            cles.append(_cle_url(cadre.url))
        try:
            n, o, vrai, aire, appli = (list(cadre.evaluate(JS_NOM_PAGE)) + [True])[:5]
        except Exception:
            n, o, vrai, aire, appli = "", "", False, 0, True
        if n:
            (titres if vrai else autres).append((aire, n, appli))
        if o:
            onglets.append(o)
    # un vrai titre de page ; s'il y a plusieurs zones (frameset, cadres), celui de la plus grande : le contenu, pas le
    # bandeau du haut ni le menu
    titres.sort(key=lambda x: -x[0])
    choix = (titres + autres + [(0, "", False)])[0]
    nom = _nom_de_page(choix[1])
    # l'onglet : seulement si c'est de la structure (jamais le nom d'une fiche choisie)
    onglet = " > ".join(o for o in (x.strip() for x in (onglets[0] if onglets else "").split(" > ")) if _structure(o) == o)
    if onglet and onglet != nom:
        nom = f"{nom} - onglet {onglet}"
    return " | ".join(cles) + " || " + re.sub(r"\d+", "#", nom), nom, bool(choix[1]) and bool(choix[2])


def _zones_de_la_page(page):
    """La description (structure seulement) du haut et de la gauche de la page, pour Claude."""
    _marquer_ecouteurs(page)
    _menus_de_la_page(page)
    lignes = []
    for cadre in page.frames:
        try:
            zones = cadre.evaluate(JS_ZONES)
        except Exception:
            continue
        for z in zones:
            icone = _structure(z["icone"]) or "-" if z["icone"] else "-"
            sig, parents = _sans_donnee(z["sig"][:60], css=True), _sans_donnee(z["parents"][:120], css=True)
            role = _sans_donnee(z["role"] or "-", css=True)
            lignes.append(f"{z['zone']:6} {sig} | role {role} | {z['mots']} mot(s) | icone {icone} | "
                          f"{z['source']} | ecoute {z['ecoute'] or '-'} | survol {z['survol'] or '-'} | "
                          f"popup {_sans_donnee(z['popup'] or '-', css=True)} | {z['taille']} a {z['x']},{z['y']} | "
                          f"menu reconnu {'oui' if z['menu'] else 'non'} | parents {parents}")
    _nettoyer(page)
    return lignes[:40]


def _nombre_cliquables(page):
    n = 0
    for cadre in page.frames:
        try:
            n += cadre.evaluate(JS_STRUCTURE)["cliquables"]
        except Exception:
            continue
    return n


def _bandeau_visite(page, message=""):
    try:
        page.evaluate(JS_BANDEAU_VISITE, message)
    except Exception:
        pass


def _nom_sur(texte):
    """Un nom d'élément que le robot peut écrire dans un chemin : de la structure (pas une donnée), sans les signes
    que le langage des tâches interprète (« > », « [2] », « DANS LIGNE »)."""
    t = " ".join(str(texte or "").split())
    return bool(t) and _structure(t) == t and ">" not in t and not re.search(r"\[\d+\]\s*$|\bdans\s+ligne\b", t, re.I)


def action_en_visite(texte, survol=False):
    """En visite : ce nom (ou cet attribut) désigne-t-il une ACTION ? Un VERBE d'action (Valider, Archiver, Créer,
    Facturer, Validez, Delete, Nouveau...), Déconnexion, Imprimer, Exporter... Un NOM de rubrique (Factures, Archives,
    Validations en attente, État des stocks, Paiements) se visite : la visite ne fait que regarder, et la protection
    réseau et le contrôle de l'élément cliqué empêchent toujours tout envoi. survol : l'élément est seulement survolé
    (une icône de module dont le sous-menu s'ouvre au survol), pas cliqué."""
    t = mots(texte)
    if not t:
        return False
    if A_EVITER_EN_VISITE.search(t) or EXPRESSIONS_ACTION.search(t):
        return True
    if survol and t.startswith("icone "):
        # (une icône de module seulement SURVOLÉE : son sous-menu s'ouvre, rien n'est cliqué ; « icone edit_road » se
        # survole. Le nom de ce qui sera cliqué dans le sous-menu est contrôlé à son tour.)
        return bool(A_EVITER_EN_VISITE.search(t))
    brut = " ".join(str(texte or "").split())
    if re.fullmatch(r"(ne plus )?suivre", t):  # (le bouton « Suivre », pas « Dossiers à suivre »)
        return True
    if not survol and (_icone_d_action(brut) or
                       # (le nom d'une icône en minuscules, bulle d'une icône Ant Design : « poweroff », « pushpin »)
                       (re.fullmatch(r"[a-z][a-z_-]{2,30}", brut) and _icone_d_action("icone " + brut))):
        return True
    liste = t.split()
    for i, m in enumerate(liste):
        if m in MOTS_ACTION_VISITE or m in VERBES_ACTION:
            return True
        if re.search(r"(er|ir|re|ez)$", m) and interdit(m):  # infinitif ou impératif d'un verbe d'action
            return True
        if i and interdit(m) and (re.search(r"\d", liste[i - 1]) or
                                  liste[i - 1] in ("action", "op", "operation", "cmd", "do", "faire", "traitement")):
            # « /commandes/42/annulation », « ?action=cloture » : une action sur une fiche, même dite par un nom
            return True
    return False


def _a_eviter(texte, survol=False):
    return action_en_visite(texte, survol)


def _lisible(chemin):
    """« GATES > Données MK2 puis Références » (sans les mots du langage des tâches)."""
    return " puis ".join(e.partition(" : ")[2] for e in chemin) or "accueil"


def _chemins_de_la_page(chemin, r, menus, deja, par_onglet, par_lien, communs=()):
    """Les chemins à essayer depuis cette page : entrées des menus (et éléments de leurs sous-menus), autres menus,
    onglets, liens. deja : ce qui est déjà prévu (un menu, un lien de l'accueil : une seule fois pour tout le portail ;
    un onglet, un lien du contenu : une fois par page). par_onglet : chemin qui finit par un clic sur un onglet -> les
    onglets de la page d'avant (déjà prévus ; seuls les sous-onglets restent à voir). Une longue liste (plus de 25
    éléments, ou plus de 12 dans un sous-menu) est une liste de fiches : ni visitée, ni nommée. par_lien : chemins qui
    finissent par un lien du contenu (la page peut être une fiche : son titre ne sera pas écrit). communs : les liens
    de l'accueil (une colonne de liens présente sur toutes les pages : chaque lien n'est essayé qu'une fois).
    Renvoie (chemins, évités par prudence)."""
    nouveaux, evites = [], []

    def ajouter(etape, noms, propre=False, onglets=None):
        cle = " ++ ".join(chemin + [etape]) if propre else etape
        if cle in deja or (propre and onglets is None and etape in deja):  # (lien déjà prévu depuis l'accueil)
            return False
        deja.add(cle)
        if any(_a_eviter(n, survol=i < len(noms) - 1) for i, n in enumerate(noms)):  # (les premiers : survolés)
            evites.append(" > ".join(noms))
            return False
        nouveaux.append(chemin + [etape])
        if onglets is not None:
            par_onglet[cle] = list(onglets)
        return True

    for nom, elements in menus:
        textes = [t for t, _ in elements]
        if not _nom_sur(nom):
            continue
        if _contenu_a_masquer(nom, textes):  # liste de fiches (Récents...), menu du compte
            # (une icône de liste de fiches ou de la personne, « icone history » : ni visitée, ni son sous-menu écrit ;
            # elle est notée comme évitée)
            if str(nom).startswith("icone ") and _structure(nom) == nom:
                evites.append(f"{nom} (liste de fiches ou compte, pas visite)")
            continue
        if len(elements) > 25:
            evites.append(f"{nom} > ({len(elements)} elements : liste de fiches, pas visitee)")
            continue
        if not elements:
            ajouter("MENU : " + nom, [nom])
        for texte, sous in elements:
            if not _nom_sur(texte):
                continue
            if not sous:
                ajouter(f"MENU : {nom} > {texte}", [nom, texte])
            elif len(sous) > 12:
                evites.append(f"{nom} > {texte} > ({len(sous)} elements : liste de fiches, pas visitee)")
            elif not _contenu_a_masquer(texte, sous):
                for s in sous:
                    if _nom_sur(s):
                        ajouter(f"MENU : {nom} > {texte} > {s}", [nom, texte, s])
    onglets_d_avant = par_onglet.get(" ++ ".join(chemin), [])
    for cle_r, limite in (("autres menus", 30), ("onglets", 15), ("liens", 25 if not chemin else 10)):
        ajoutes = 0
        for texte in r[cle_r]:  # (la limite compte les chemins NOUVEAUX : la liste n'est pas coupée avant)
            if ajoutes >= limite:
                break
            if not _nom_sur(texte) or (cle_r == "onglets" and texte in onglets_d_avant):
                continue
            if ajouter(f"CLIQUER : {texte}", [texte],
                       propre=cle_r == "onglets" or (cle_r == "liens" and texte not in communs),
                       onglets=r["onglets"] if cle_r == "onglets" else None):
                ajoutes += 1
                if cle_r == "liens":
                    par_lien.add(" ++ ".join(chemin + [f"CLIQUER : {texte}"]))
    return nouveaux, evites


def _empreinte(r):
    """Ce qui distingue un écran d'un autre quand l'adresse et le titre ne changent pas (applications à adresse
    unique) : ses champs, boutons, colonnes, onglets et libellés (sans le nombre de lignes des tableaux)."""
    tableaux = [re.sub(r"\s+\(\d+ ligne\(s\)\)$", "", x) for x in r["tableaux"]]
    return json.dumps([r["champs"], r["boutons"], tableaux, r["onglets"], r["libelles de fiche"]], ensure_ascii=False)


def _adresse_carte(url):
    """L'adresse de l'accueil pour reconnaître la carte (deux applications d'un même serveur ont deux cartes)."""
    try:
        m = urlsplit(url)
        return f"{m.scheme}://{m.netloc}{_cle_url(url)}"
    except ValueError:
        return url


def lire_carte(url=None):
    """La carte du portail (None si elle n'existe pas, est abîmée, ou est celle d'une autre adresse)."""
    try:
        carte = json.loads((dossier_robot() / FICHIER_CARTE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if not (isinstance(carte, dict) and isinstance(carte.get("pages"), list) and isinstance(carte.get("essais", {}), dict)):
        return None
    if url is not None and carte.get("adresse") != _adresse_carte(url):
        return None
    if not isinstance(carte.get("par_onglet"), dict):
        carte["par_onglet"] = {}
    return carte


def sauver_carte(carte):
    try:
        f = dossier_robot() / FICHIER_CARTE
        temporaire = f.with_suffix(".tmp")
        temporaire.write_text(json.dumps(carte, ensure_ascii=False, indent=1), encoding="utf-8")
        temporaire.replace(f)
    except OSError:
        pass


def _vider_entree():
    """Une Entrée tapée en avance (double appui, impatience pendant le relevé de l'accueil) ne doit pas arrêter la
    visite dès son début."""
    if WINDOWS:
        vider_clavier()
        return
    import select
    try:
        while not FIN_DU_CLAVIER[0] and select.select([sys.stdin], [], [], 0)[0]:
            if not sys.stdin.readline():
                FIN_DU_CLAVIER[0] = True
                break
    except Exception:
        pass


def _accueil(t, url, garde):
    """Retour à l'accueil du portail dans la fenêtre principale du robot (les onglets ouverts en route sont fermés).
    La page est RECHARGÉE (une application à adresses « #/ » garderait sinon ses panneaux ouverts). Si le portail
    redemande de se connecter, la protection est coupée le temps que VOUS vous connectiez."""
    principale = t.principale
    if principale.is_closed():
        raise ErreurTache("la fenetre du robot a ete fermee")
    for p in list(principale.context.pages):
        if p != principale and not p.is_closed():
            try:
                p.close()
            except Exception:
                pass
    t.page = principale
    depart = time.time()
    try:
        meme_document = urlsplit(principale.url)[:4] == urlsplit(url)[:4]
    except ValueError:
        meme_document = False
    for essai in range(3):
        try:
            if principale.url != url:
                principale.goto(url, wait_until="domcontentloaded", timeout=60000)
            if meme_document:  # (seul le « # » changeait, ou rien : rien n'a été rechargé)
                principale.reload(wait_until="domcontentloaded", timeout=60000)
            break
        except Exception as e:
            # (une page bloquée par la protection affiche une page d'erreur qui arrive un peu après : on réessaie)
            if essai == 2 or "interrupted" not in str(e):
                raise
            principale.wait_for_timeout(1000)
    t.oublier_requetes(principale, depart)
    attendre_chargement(principale, t.chargement_max)
    if _champ_mot_de_passe(principale):
        revenez_ici()
        _garde_coupee(garde, lambda: attendre_portail(principale))
        try:
            principale.bring_to_front()
        except Exception:
            pass
    t.calme()


def _suivre_chemin(t, url, garde, chemin, dire=False):
    """Depuis l'accueil, refait le chemin (menus survolés puis cliqués, onglets), comme votre souris. Lecture seule."""
    _accueil(t, url, garde)
    t.nouvel_element()
    for numero, etape in enumerate(chemin, 1):
        action, _, arg = etape.partition(" : ")
        if dire:
            ecrire(f"   Etape {numero}/{len(chemin)} : " + ("menu " if action == "MENU" else "clic sur ") + arg)
        t.page_ouverte()
        t.nouveau_geste()
        t.ligne_courante = numero
        t.faire(action.strip(), arg.strip())
        t.verifier_fenetres()
    t.calme()
    t.verifier_fenetres()


def _raison_courte(e):
    """Pour la carte (fichier à envoyer) : le genre d'échec, jamais le texte du portail."""
    m = sans_accents(str(e)).lower()
    if isinstance(e, Passer):
        return "fenetre du portail refusee" if "le portail a affiche" in m else "refuse par prudence"
    if "introuvable" in m:
        return "introuvable"
    if "connexion" in m:
        return "connexion demandee"
    return "echec"


PAGE_DE_CONNEXION = re.compile(r"session.{0,20}expir|expir.{0,20}session|deconnecte|reconnect|connectez-vous|"
                               r"identifiez-vous|sign in|log in|choisir un compte|pick an account")
# ce qui a été noté pour un chemin et ne changera pas en recommençant (les autres sont réessayés en complétant la carte)
STATUTS_DEFINITIFS = ("page", "deja vue", "reste sur l'accueil", "autre site", "bloque", "fenetre du portail refusee",
                      "refuse par prudence")


def explorer_portail(page, url, garde):
    """Choix 5. Le robot visite le portail tout seul (lecture seule) et fait la carte du site. Entrée = arrêter.
    Renvoie (fichier à envoyer, raison de l'arrêt)."""
    hote, adresse = urlsplit(url).hostname or "", _adresse_carte(url)
    carte = lire_carte(url)
    if carte and carte["pages"]:
        revenez_ici()
        if not oui_en_vivant(page, f"   Le robot a deja une carte de ce portail ({len(carte['pages'])} pages). Tapez o pour la "
                                   "COMPLETER, n pour tout recommencer", "o"):
            carte = None
    elif lire_carte() is not None:
        ecrire("   (La carte deja sur ce poste est celle d'une autre adresse : le robot en fait une nouvelle.)")
        carte = None
    if carte is None:
        carte = {"version": 2, "hote": hote, "adresse": adresse, "pages": [], "essais": {}, "evites": [], "par_onglet": {}}
    par_onglet = carte["par_onglet"]
    par_lien = set(carte.get("par_lien", []))
    carte.update(date=datetime.now().strftime("%d/%m/%Y %H:%M"), bloquees=0, exemples=[], liaison=False)
    # en complétant : les chemins en échec la fois d'avant sont réessayés (session expirée, portail lent...)
    carte["essais"] = {c: st for c, st in carte["essais"].items() if st.startswith(STATUTS_DEFINITIFS)}
    index = {p["id"]: i for i, p in enumerate(carte["pages"])}
    t = Tache(dict(ENTETE_VISITE), [], page, url, lecture_seule=True)
    # (« plus aucun échange » n'arrive jamais sur un portail qui s'interroge sans cesse : on attend peu, et calme()
    #  attend la vraie réponse)
    t.chargement_max, t.calme_max = 1, 6
    memoire, deja, file = {}, set(), deque()
    for p in carte["pages"]:  # carte complétée : les chemins pas encore essayés la fois d'avant
        for c in p.get("a_essayer", []):
            deja.update((c[-1], " ++ ".join(c)))
            file.append(c)
    pages_au_depart, debut, essais, arret, interrompu = len(carte["pages"]), time.time(), 0, "", False
    accueil = [0]

    def noter_page(chemin):
        """La page affichée : (son numéro dans la carte, nouvelle ?). Une nouvelle page est relevée, et ses chemins
        (sous-menus, onglets, liens) sont ajoutés à ceux à essayer."""
        t.retirer_bandeaux()
        ident, nom, appli = _identite_complete(t.page)
        if appli:
            # le titre est celui de l'APPLICATION (en-tête, fenêtre) : la page porte le nom de ce qui y mène
            nom = (chemin[-1].partition(" : ")[2].split(">")[-1].strip() if chemin else "Accueil") + " (titre non ecrit)"
        r = menus = None
        if ident in index:
            connue = carte["pages"][index[ident]]
            if not chemin or "empreinte" not in connue:
                return index[ident], False
            # même adresse et même titre : le même écran, ou un autre écran d'une application à adresse unique ?
            r, menus = _releve_et_menus(t.page, memoire)
            empreinte = _empreinte(r)
            if empreinte == connue["empreinte"]:
                return index[ident], False
            ident += " || ecran " + hashlib.sha1(empreinte.encode("utf-8")).hexdigest()[:12]
            if ident in index:
                return index[ident], False
            nom = chemin[-1].partition(" : ")[2].split(">")[-1].strip()  # (le titre est celui de l'autre écran)
        elif chemin and " ++ ".join(chemin) in par_lien:
            # page ouverte par un lien du contenu (« Voir », « Ma fiche ») : c'est peut-être une fiche, dont le titre
            # EST une donnée (nom d'un client...) : elle porte le nom du lien
            nom = chemin[-1].partition(" : ")[2].strip() + " (titre non ecrit)"
        if r is None:
            r, menus = _releve_et_menus(t.page, memoire)
        if not chemin:  # (les liens de l'accueil : une navigation commune à toutes les pages)
            carte["liens_accueil"] = list(r["liens"])
        nouveaux, evites = (_chemins_de_la_page(chemin, r, menus, deja, par_onglet, par_lien,
                                                set(carte.get("liens_accueil", [])))
                            if len(chemin) < PROFONDEUR else ([], []))
        carte["par_lien"] = sorted(par_lien)
        carte["pages"].append({"nom": nom, "id": ident, "chemin": chemin, "releve": r, "a_essayer": nouveaux,
                               "empreinte": _empreinte(r), "cliquables": _nombre_cliquables(t.page),
                               "autres_chemins": []})
        index[ident] = len(carte["pages"]) - 1
        file.extend(nouveaux)
        for e in evites:
            if e not in carte["evites"]:
                carte["evites"].append(e)
        sauver_carte(carte)
        return index[ident], True

    def noter_arrivee(chemin, bloquees_avant):
        """Le chemin a été suivi : où est-on ? Renvoie ce qui est noté pour ce chemin."""
        ici = t.page
        if ici.url.startswith("chrome-error") or (garde["bloquees"] > bloquees_avant and not ici.url.startswith("http")):
            statut = "bloque par la protection (adresse d'action)"
        elif (urlsplit(ici.url).hostname or "") != hote:
            statut = "autre site (pas visite)"
        elif PAGE_DE_CONNEXION.search(sans_accents(_identite(ici)[1]).lower()) or _champ_mot_de_passe(ici):
            # (« Votre session a expiré », page de connexion : pas une page du portail ; compte comme un échec)
            statut = "connexion demandee"
        else:
            numero, nouvelle = noter_page(chemin)
            p = carte["pages"][numero]
            if nouvelle:
                statut = f"page {numero + 1}"
                ecrire(f"      -> NOUVELLE PAGE {numero + 1} : « {p['nom'] or '(sans titre)'} » "
                       f"({p['cliquables']} elements cliquables)")
            elif numero == accueil[0]:
                statut = "reste sur l'accueil"
            else:
                statut = f"deja vue (page {numero + 1})"
                if chemin != p["chemin"] and chemin not in p["autres_chemins"] and len(p["autres_chemins"]) < 3:
                    p["autres_chemins"].append(chemin)
        if ici != t.principale:
            statut += " (nouvel onglet)"
        if garde["bloquees"] > bloquees_avant and not statut.startswith("bloque"):
            statut += " (envoi bloque par la protection)"
        if not statut.startswith("page"):
            ecrire(f"      -> {statut}")
        return statut

    def toujours_connecte():
        """Après 3 échecs de suite : l'accueil est-il toujours l'accueil (pas une page « session expirée », « choisir un
        compte »...) ? Sinon, vous vous reconnectez. Renvoie (l'accueil est là, il a fallu se reconnecter)."""
        try:
            _accueil(t, url, garde)
            if _identite(t.principale)[0] == carte["pages"][accueil[0]]["id"]:
                return True, False
        except ErreurTache:
            raise
        except Exception:
            pass
        ecrire()
        ecrire("   Le portail ne montre plus son accueil habituel : il vous a peut-etre deconnecte.")
        revenez_ici()
        _garde_coupee(garde, lambda: pause("Reconnectez-vous DANS LA FENETRE DU ROBOT si besoin, revenez a l'accueil du "
                                           "portail, puis Entree ici..."))
        _vider_entree()
        try:
            t.principale.bring_to_front()
            _accueil(t, url, garde)
            return _identite(t.principale)[0] == carte["pages"][accueil[0]]["id"], True
        except ErreurTache:
            raise
        except Exception:
            return False, True

    ecrire("   Patientez : le robot regarde l'accueil et survole ses menus (sans cliquer), 1 a 2 minutes...")
    _accueil(t, url, garde)
    connu = next((p for p in carte["pages"] if not p["chemin"]), None)
    if connu is not None and _identite(t.principale)[0] != connu["id"]:
        # en complétant : une page « session expirée » ne doit pas devenir l'accueil de la carte
        ecrire("   L'accueil ne ressemble pas a celui de la carte : le portail vous a peut-etre deconnecte.")
        revenez_ici()
        _garde_coupee(garde, lambda: pause("Reconnectez-vous DANS LA FENETRE DU ROBOT si besoin, revenez a l'accueil du "
                                           "portail, puis Entree ici..."))
        _vider_entree()
        _accueil(t, url, garde)
    accueil[0], _ = noter_page([])
    if _liaison_libre(garde):
        arret = "liaison directe avec le portail que ce robot ne sait pas bloquer : visite arretee par prudence"
        file.clear()
    elif garde["liaison"]:
        ecrire("   (Le portail utilise une liaison directe : elle est BLOQUEE pendant la visite, par prudence. Une partie")
        ecrire("    du portail peut ne pas marcher pendant la visite : c'est normal.)")
    if len(file) < 5:
        # presque rien à visiter depuis l'accueil : la carte décrira le haut et la gauche de l'accueil, pour Claude
        ecrire("   Le robot ne trouve presque rien a visiter depuis l'accueil. Le fichier pour Claude contiendra une")
        ecrire("   description de l'accueil (sans aucun texte de la page) pour qu'il corrige le robot.")
        try:
            carte["diagnostic_accueil"] = _zones_de_la_page(t.principale)
        except Exception as e:
            carte["diagnostic_accueil"] = ["(description impossible : " + _sans_donnee(masquer(premiere_ligne(e)))[:60] + ")"]
    else:
        carte.pop("diagnostic_accueil", None)
    evites = [e for e in carte["evites"] if "pas visitee" not in e]
    if evites:
        ecrire("   Par prudence, le robot n'ira PAS dans : " + " | ".join(evites[:15]))
    ecrire(f"   {len(file)} chemin(s) a essayer.")
    ecrire("   >>> NE TOUCHEZ PLUS A LA SOURIS : le robot visite le portail tout seul. Regardez sa fenetre (Chrome,")
    ecrire("       en bas de l'ecran si elle est cachee) : bande JAUNE en bas.")
    ecrire("   >>> Pour l'arreter avant la fin : cliquez dans CETTE fenetre noire, puis Entree.")
    try:
        t.principale.bring_to_front()
    except Exception:
        pass
    _vider_entree()
    tampon, echecs_de_suite, en_echec, verifie = [], 0, [], False
    try:
        while file:
            if ligne_tapee(tampon) is not None:
                arret = "visite arretee par vous"
                break
            if t.principale.is_closed():
                arret = "fenetre du robot fermee"
                break
            if _liaison_libre(garde):
                arret = "liaison directe avec le portail que ce robot ne sait pas bloquer : visite arretee par prudence"
                break
            if time.time() - debut > LIMITE_MINUTES * 60:
                arret = f"visite arretee au bout de {LIMITE_MINUTES} minutes (relancez le choix 5 pour la completer)"
                break
            if len(carte["pages"]) - pages_au_depart >= LIMITE_PAGES:
                arret = f"visite arretee a {LIMITE_PAGES} nouvelles pages (relancez le choix 5 pour la completer)"
                break
            if echecs_de_suite >= 3 and not verifie:
                verifie = True
                accueil_la, reconnecte = toujours_connecte()
                if not accueil_la:
                    arret = "le portail ne montre plus son accueil habituel (deconnexion ?)"
                    break
                if reconnecte:  # les chemins en échec pendant la déconnexion sont réessayés
                    for c in en_echec:
                        carte["essais"].pop(" ++ ".join(c), None)
                        file.append(c)
                    echecs_de_suite, en_echec, verifie = 0, [], False
            if echecs_de_suite >= 8:
                arret = "trop d'echecs de suite : le portail ne repond pas comme prevu"
                break
            chemin = file.popleft()
            cle = " ++ ".join(chemin)
            if cle in carte["essais"]:  # (déjà essayé : la fois d'avant, ou par un autre chemin)
                continue
            essais += 1
            ecrire(f"   [{essais}] {_lisible(chemin)}")
            avant = garde["bloquees"]
            try:
                _bandeau_visite(t.principale, f"VISITE, lecture seule ({len(carte['pages'])} pages) : "
                                              f"{_lisible(chemin)}  -  ne touchez pas a la souris")
                _suivre_chemin(t, url, garde, chemin)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                statut = _raison_courte(e)
                if statut in ("echec", "introuvable"):  # (pour Claude : la cause, sans adresse ni chemin)
                    statut += " : " + _sans_donnee(masquer(premiere_ligne(e)))[:80]
                ecrire(f"      -> pas de nouvelle page : {statut}")
            else:
                try:
                    statut = noter_arrivee(chemin, avant)
                except KeyboardInterrupt:
                    raise
                except Exception as e:
                    statut = "echec : page pas relevee"
                    ecrire(f"      -> page pas relevee ({masquer(premiere_ligne(e))[:90]})")
            carte["essais"][cle] = statut
            if statut.startswith(("echec", "introuvable", "connexion")):
                echecs_de_suite += 1
                en_echec.append(chemin)
            else:
                echecs_de_suite, en_echec, verifie = 0, [], False
    except KeyboardInterrupt:
        arret, interrompu = "visite arretee par Ctrl+C", True
    finally:
        carte.update(bloquees=garde["bloquees"], exemples=list(garde["exemples"]), liaison=garde["liaison"],
                     liaison_bloquee=garde["liaison_bloquee"], duree=round((time.time() - debut) / 60, 1), arret=arret,
                     lectures=garde.get("lectures", 0), exemples_lectures=list(garde.get("exemples_lectures", [])),
                     comparaison=garde.get("comparaison"))
        sauver_carte(carte)
    fichier = ecrire_carte(carte)  # (avant tout appel au navigateur : après un Ctrl+C, il peut ne plus répondre)
    if not interrompu:
        try:
            if not t.principale.is_closed():
                _bandeau_visite(t.principale, "")
        except Exception:
            pass
    ecrire()
    ecrire(f"   CARTE : {len(carte['pages'])} page(s) connue(s) ({len(carte['pages']) - pages_au_depart} nouvelle(s)), "
           f"{essais} chemin(s) essaye(s) en {carte['duree']} minute(s)." + (f" ({arret})" if arret else ""))
    if garde["bloquees"]:
        ecrire(f"   La protection a bloque {garde['bloquees']} envoi(s) au portail (rien n'a ete modifie).")
    return fichier, arret


def ecrire_carte(carte):
    """A_ENVOYER_carte.txt : les pages du portail et comment y aller. Structure seulement (aucune adresse, aucune valeur)."""
    if carte.get("liaison"):
        liaison = "oui, bloquee pendant la visite" if carte.get("liaison_bloquee") else "oui, NON BLOQUEE (visite arretee)"
    else:
        liaison = "non"
    lignes = [f"CARTE DU PORTAIL - pour Claude - {carte.get('date', '')}",
              "Faite par le robot (choix 5) en VISITE : lecture seule (aucun envoi au portail).",
              "Structure seulement (menus, pages, champs, boutons, colonnes). Aucune valeur, aucune adresse.",
              "Relisez : une vraie valeur (nom de client, numero de plan...) -> effacez ce mot seulement.", "",
              f"Pages connues : {len(carte['pages'])} ; duree : {carte.get('duree', '?')} min"
              + (f" ; {carte['arret']}" if carte.get("arret") else " ; visite complete"),
              f"Envois bloques par la protection : {carte.get('bloquees', 0)}"
              + (" (" + " ; ".join(carte.get("exemples", [])) + ")" if carte.get("exemples") else ""),
              f"Envois de lecture permis (recherches, listes) : {carte.get('lectures', 0)}"
              + (" (" + " ; ".join(carte.get("exemples_lectures", [])) + ")" if carte.get("exemples_lectures") else ""),
              f"Liaison directe avec le portail (websocket) : {liaison}"]
    if carte.get("comparaison"):
        avant, apres = carte["comparaison"]
        lignes.append(f"Elements cliquables a l'accueil : {avant} sans la protection, {apres} avec"
                      + (" (LA PROTECTION EMPECHE UNE PARTIE DU PORTAIL DE S'AFFICHER)" if apres < 0.7 * avant else ""))
    lignes.append("")
    menus_vus = [None]
    for numero, p in enumerate(carte["pages"], 1):
        lignes.append(f"PAGE {numero} : « {p['nom'] or '(sans titre)'} »  ({p['cliquables']} elements cliquables)")
        lignes.append("  Chemin : " + _lisible(p["chemin"]))
        for autre in p.get("autres_chemins", []):
            lignes.append("  Autre chemin : " + _lisible(autre))
        lignes += _lignes_releve(p["releve"], menus_vus)
        lignes.append("")
    if carte.get("evites"):
        lignes.append("PAS VISITES PAR PRUDENCE : " + " ; ".join(carte["evites"]))
    if carte.get("diagnostic_accueil"):
        lignes.append("")
        lignes.append("DESCRIPTION DE L'ACCUEIL pour Claude (le robot n'y trouve presque rien a visiter) : elements cliquables")
        lignes.append("du haut et de la gauche, structure seulement (balise.classes | role | nombre de mots | icone | ...) :")
        lignes += ["  " + l for l in carte["diagnostic_accueil"]]
        lignes.append("")
    autres = [(c, s) for c, s in carte.get("essais", {}).items() if not s.startswith(("page", "deja vue"))]
    if autres:
        lignes.append("CHEMINS SANS NOUVELLE PAGE :")
        for c, s in autres:
            lignes.append(f"  {_lisible(c.split(' ++ '))} : {s}")
    for nom in ("A_ENVOYER_carte.txt", "A_ENVOYER_carte_2.txt", "A_ENVOYER_carte_3.txt"):
        fichier = dossier_robot() / nom
        try:
            fichier.write_text("\n".join(lignes) + "\n", encoding="utf-8")
            return fichier
        except OSError:
            continue
    raise ErreurTache("impossible d'ecrire la carte dans Bureau > Robot")


def _elements_de_page(p):
    """Ce qu'on peut chercher dans une page de la carte : (genre, texte à montrer, texte à trouver sur la page, poids)."""
    r = p["releve"]
    items = [("page", re.sub(r" \(titre non ecrit\)$", "", p["nom"] or ""), "", 4)]
    for etape in p["chemin"]:
        for x in etape.partition(" : ")[2].split(">"):
            items.append(("chemin", x.strip(), "", 3))
    items += [("onglet", x, x, 2) for x in r["onglets"]]
    for x in r["champs"]:
        libelle = re.sub(r"\s*\(indice\)$", "", x.split(" [")[0])
        items.append(("champ", x, libelle, 2))
    items += [("bouton", x, re.sub(r"^\(icone\)\s*", "", x), 2) for x in r["boutons"]]
    for tableau in r["tableaux"]:
        for c in re.sub(r"\s+\(\d+ ligne\(s\)\)$", "", tableau).split(" | "):
            items.append(("colonne", c, c.split(" (")[0], 1))
    items += [("libelle", x, x, 1) for x in r["libelles de fiche"]]
    items += [("lien", x, x, 1) for x in r["liens"] + r["autres menus"]]
    return [i for i in items if i[1] and not i[1].startswith("(")]


def chercher_dans_carte(carte, question):
    """Les endroits de la carte qui correspondent à la question (tous ses mots, au singulier comme au pluriel) :
    [(page, genre, texte à montrer, texte à trouver)], les meilleurs d'abord."""
    tous = mots(question).split()
    cherches = [m for m in tous if m not in MOTS_VIDES] or tous
    if not cherches:
        return []

    def touche(m, cible):
        return any(c == m or (len(m) >= 3 and c.startswith(m)) or (len(c) >= 3 and m.startswith(c) and len(m) - len(c) <= 2)
                   for c in cible)

    def singulier(texte):
        return " ".join(re.sub(r"(?<=\w{3})[sx]$", "", m) for m in mots(texte).split())

    trouves = {}
    for i, p in enumerate(carte["pages"]):
        for genre, montre, a_trouver, poids in _elements_de_page(p):
            cible = mots(montre).split()
            n = sum(1 for m in cherches if touche(m, cible))
            if n:
                exact = singulier(a_trouver or montre) == singulier(" ".join(cherches))
                cle = (i, genre, montre)
                rang = (n, exact, poids, -len(p["chemin"]), -i)
                if cle not in trouves or trouves[cle][0] < rang:
                    trouves[cle] = (rang, a_trouver)
        # la page entière : ses mots dans son titre et son chemin ensemble (« base MK2 » + « GATES »)
        ensemble = mots(re.sub(r" \(titre non ecrit\)$", "", p["nom"] or "") + " " +
                        _lisible(p["chemin"]).replace(" puis ", " ")).split()
        n = sum(1 for m in cherches if touche(m, ensemble))
        if n == len(cherches) and (i, "page", p["nom"]) not in trouves:
            trouves[(i, "page", p["nom"])] = ((n, False, 4, -len(p["chemin"]), -i), "")
    if not trouves:
        return []
    meilleur = max(rang[0] for rang, _ in trouves.values())
    garde = sorted(((rang, cle, a) for cle, (rang, a) in trouves.items() if rang[0] == meilleur), reverse=True)
    pages = any(cle[1] in ("page", "chemin") for _, cle, _ in garde)
    vus, resultats = set(), []
    for _, (i, genre, montre), a_trouver in garde:
        if genre in ("page", "chemin"):
            if i in vus:  # une seule ligne par page quand c'est la page elle-même qui correspond
                continue
            vus.add(i)
            genre, montre, a_trouver = "page", carte["pages"][i]["nom"], ""
        elif genre in ("onglet", "bouton", "lien"):
            # un onglet, un lien du même nom qu'une page trouvée : c'est le chemin vers elle ; et le même onglet sur
            # chaque onglet de la page : une seule fois
            if pages or (genre, montre) in vus:
                continue
            vus.add((genre, montre))
        resultats.append((i, genre, montre, a_trouver))
    return resultats


def _decrire_resultat(carte, i, genre, montre):
    p = carte["pages"][i]
    ou = f"page « {p['nom'] or '(sans titre)'} » (chemin : {_lisible(p['chemin'])})"
    return ou if genre == "page" else f"{genre} « {montre} » sur la {ou}"


def _montrer_sur_la_page(t, genre, a_trouver):
    """Encadre en rouge, sur la page, l'élément cherché (champ, bouton, onglet, colonne...). Sans cliquer."""
    try:
        if genre == "champ":
            loc = t.champ(a_trouver)
        else:
            trouves = [x for x in t.candidats(a_trouver, t.cadres(), attente=3) if _visible(x)]
            loc = trouves[0] if trouves else None
        if loc is not None and _visible(loc):
            _encadrer(loc, 0)
            return True
    except Exception:
        pass
    return False


def _aller_vers(t, url, garde, carte, i, genre="page", a_trouver=""):
    p = carte["pages"][i]
    ecrire()
    ecrire(f"   Le robot va sur « {p['nom'] or '(sans titre)'} » : {_lisible(p['chemin'])}")
    ecrire("   Ne touchez pas a la souris : regardez la fenetre du robot.")
    try:
        t.principale.bring_to_front()
    except Exception:
        pass
    try:
        _suivre_chemin(t, url, garde, p["chemin"], dire=True)
    except KeyboardInterrupt:
        raise
    except Exception as e:
        ecrire(f"   PAS PU Y ALLER ({_raison_courte(e)}).")  # (jamais le texte d'une fenêtre du portail)
        ecrire("   Le portail a peut-etre change : refaites le choix 5 pour mettre la carte a jour.")
        return False
    ident, nom = _identite(t.page)
    if ident == p["id"]:
        ecrire(f"   ARRIVE : « {p['nom'] or '(sans titre)'} ».")
    else:
        ecrire(f"   Arrive sur « {nom or '(sans titre)'} » (la carte disait « {p['nom']} ») : le portail a peut-etre change.")
    _bandeau_visite(t.page, f"ARRIVE : {p['nom'] or '(sans titre)'}")
    if genre != "page" and a_trouver:
        if _montrer_sur_la_page(t, genre, a_trouver):
            ecrire(f"   Le robot vous montre {genre} « {a_trouver} » (cadre ROUGE). Il ne clique pas dessus : en visite,")
            ecrire("   il ne fait que regarder.")
        else:
            ecrire(f"   ({genre} « {a_trouver} » pas visible sur la page pour l'instant)")
    ecrire("   Regardez la fenetre du robot. Puis revenez ici (fenetre noire) pour aller ailleurs.")
    return True


def _numero(reponse, maximum):
    """Le numéro tapé (la rangée du haut d'un clavier AZERTY sans Maj : & é " ' ( - è _ ç à compte aussi)."""
    reponse = reponse.strip().translate(str.maketrans("&é\"'(-è_çà", "1234567890"))
    if not reponse.isdecimal():
        return None
    try:
        n = int(reponse)
    except ValueError:
        return None
    return n if 1 <= n <= maximum else None


def boucle_gps(page, url, garde):
    """Choix 6. Vous dites où aller ; le robot cherche dans sa carte et y va tout seul, étape par étape (lecture seule)."""
    carte = lire_carte(url)
    if not carte or not carte["pages"]:
        if lire_carte() is not None:
            ecrire("   La carte sur ce poste est celle d'une autre adresse : faites d'abord le choix 5 pour ce portail.")
        else:
            ecrire("   Le robot n'a pas encore la carte de ce portail : faites d'abord le choix 5 (il explore le portail).")
        return
    t = Tache(dict(ENTETE_VISITE), [], page, url, lecture_seule=True)
    t.chargement_max, t.calme_max = 2, 10
    ecrire(f"   Le robot connait {len(carte['pages'])} page(s) de ce portail (carte du {carte.get('date', '?')}).")
    while True:
        ecrire()
        if _liaison_libre(garde):
            ecrire("   Le portail utilise une liaison directe que ce robot ne sait pas bloquer : arret par prudence.")
            break
        question = demander_en_vivant(t.page, "   Ou voulez-vous aller ? (un nom de page, de menu, de champ ou de bouton ; « liste » = "
                            "toutes les pages ; Entree sans rien = fini) :")
        if not question:
            break
        if mots(question) in ("liste", "carte", "pages", "tout"):
            for numero, p in enumerate(carte["pages"], 1):
                ecrire(f"   {numero:>3}. « {p['nom'] or '(sans titre)'} »  (chemin : {_lisible(p['chemin'])})")
            ecrire("   Tapez le numero d'une page pour y aller.")
            continue
        if question.strip().translate(str.maketrans("&é\"'(-è_çà", "1234567890")).isdecimal():
            n = _numero(question, len(carte["pages"]))
            if n:
                _aller_vers(t, url, garde, carte, n - 1)
            else:
                ecrire(f"   Tapez un numero entre 1 et {len(carte['pages'])} (tapez « liste » pour les voir).")
            continue
        resultats = chercher_dans_carte(carte, question)
        if not resultats:
            evites = [e for e in carte.get("evites", []) if set(mots(question).split()) & set(mots(e).split())]
            if evites:
                ecrire(f"   « {question} » : le robot ne l'a PAS visite, par prudence (le nom ressemble a une action,")
                ecrire("   ou c'est une liste de fiches) : " + " | ".join(evites[:3]))
                continue
            ecrire(f"   Rien dans la carte pour « {question} ». Essayez un autre mot, ou « liste ».")
            ecrire("   (Si c'est une page derriere un bouton ou une recherche, le robot ne l'a pas visitee : en visite, il ne")
            ecrire("    clique que sur les menus, les liens et les onglets.)")
            continue
        resultats = resultats[:9]
        if len(resultats) == 1:
            ecrire("   Trouve : " + _decrire_resultat(carte, *resultats[0][:3]))
            choix = 1
        else:
            for numero, (i, genre, montre, _) in enumerate(resultats, 1):
                ecrire(f"   {numero}. " + _decrire_resultat(carte, i, genre, montre))
            reponse = demander_en_vivant(t.page, "   Numero de votre choix (Entree = 1, n = aucun) :", "1")
            choix = _numero(reponse, len(resultats))
            if not choix:
                if not reponse.lower().startswith("n"):
                    ecrire(f"   Pas compris : tapez un numero entre 1 et {len(resultats)}. Reposez la question.")
                continue
        i, genre, montre, a_trouver = resultats[choix - 1]
        _aller_vers(t, url, garde, carte, i, genre, a_trouver)
    _bandeau_visite(t.page, "")
    if garde["bloquees"]:
        ecrire(f"   La protection a bloque {garde['bloquees']} envoi(s) au portail (rien n'a ete modifie).")


# ---------------------------------------------------------------------------- programmes du navigateur (processus enfant)
def programme_enfant(mode, dossier_echange):
    """Relevé (choix 2), tâche (choix 3), « voir » (choix 4), visite (choix 5) ou « aller à » (choix 6), lancés avec
    le Python qui a le pilote de navigateur."""
    preparer_console()
    fichier_demande = Path(dossier_echange) / "demande.json"
    echange = json.loads(fichier_demande.read_text(encoding="utf-8"))
    try:
        fichier_demande.unlink()  # l'adresse du portail ne reste pas dans le dossier temporaire
    except OSError:
        pass
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            if mode == "tache":
                return executer_tache(p, echange["fichier"], echange["url"])
            visite = mode in ("explorer", "gps")
            navigateur, contexte, page = ouvrir_robot(p, lecture=visite)
            page.goto(echange["url"], wait_until="domcontentloaded", timeout=60000)
            page.bring_to_front()
            revenez_ici()
            pause("Connectez-vous si besoin dans la fenetre du robot. Quand vous voyez l'accueil du portail, Entree...")
            attendre_portail(page)
            garde = None
            if visite:
                # ce que montre l'accueil SANS la protection (pour voir si la protection empêche le portail de s'afficher)
                try:
                    _marquer_ecouteurs(page)
                    avant = _nombre_cliquables(page)
                except Exception:
                    avant = 0
                garde = proteger_en_visite(contexte)  # à partir d'ici : lecture seule, rien ne peut être envoyé au portail
                ecrire("   VISITE EN LECTURE SEULE : le robot ne peut rien envoyer au portail (rien ne sera modifie).")
                # (une liaison directe ouverte avant la protection resterait ouverte : la page est rechargée)
                page.reload(wait_until="domcontentloaded", timeout=60000)
                attendre_chargement(page)
                if _champ_mot_de_passe(page):
                    _garde_coupee(garde, lambda: attendre_portail(page))
                try:
                    page.wait_for_timeout(1500)
                    _marquer_ecouteurs(page)
                    apres = _nombre_cliquables(page)
                except Exception:
                    apres = 0
                if avant:
                    garde["comparaison"] = (avant, apres)
                    if apres < 0.7 * avant:
                        ecrire(f"   Attention : avec la protection, l'accueil montre moins de choses ({apres} elements au lieu "
                               f"de {avant}) : le portail a sans doute besoin d'envois que le robot bloque. C'est note pour Claude.")
            if mode == "voir":
                fichier = boucle_voir(page)
            elif mode == "explorer":
                fichier, arret = explorer_portail(page, echange["url"], garde)
                if "Ctrl+C" not in arret and "liaison directe" not in arret and not page.is_closed():
                    revenez_ici()
                    if oui_en_vivant(page, "   Voulez-vous essayer tout de suite « aller a une page » (le robot y va tout "
                                           "seul) ?", "n"):
                        boucle_gps(page, echange["url"], garde)
            elif mode == "gps":
                boucle_gps(page, echange["url"], garde)
                fichier = False
            else:
                fichier, _ = boucle_releve(page)
            ecrire()
            if fichier is False:
                pass
            elif fichier:
                ecrire(f"   Fichier cree : Bureau > Robot > {fichier.name}")
                ecrire("   Ouvrez-le, relisez-le, puis envoyez-le-moi (ou une photo).")
            else:
                ecrire("   Aucune page : rien a envoyer.")
            try:
                navigateur.close()
            except Exception:
                pass
            return 0
    except KeyboardInterrupt:
        ecrire("   Arrete a votre demande.")
        return 130
    except ErreurTache as e:
        ecrire(f"   PROBLEME : {e}")
        ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
        return 1
    except Exception as e:
        ecrire(f"   PROBLEME : {expliquer(e)}")
        ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
        return 1


def lancer_programme_enfant(mode, demande):
    python = obtenir_python_playwright()
    if not python:
        ecrire(f"   Le robot ne peut pas ouvrir son navigateur ({RAISON_PILOTE[0]}).")
        ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
        return 1
    echange = Path(tempfile.mkdtemp(prefix="robot_"))
    try:
        (echange / "demande.json").write_text(json.dumps(demande, ensure_ascii=False), encoding="utf-8")
        enfant = subprocess.Popen([python, str(Path(__file__).resolve()), "--" + mode, str(echange)])
        try:
            return enfant.wait()
        except KeyboardInterrupt:
            try:
                enfant.wait(timeout=5)
            except Exception:
                enfant.kill()
            return 130
    finally:
        effacer_dossier(echange)


# ---------------------------------------------------------------------------- programme principal
def robot_plus_recent():
    """« robot (1).txt » : le navigateur a gardé l'ancien robot.txt et rangé le nouveau à côté."""
    try:
        moi = Path(__file__).resolve()
        recents = [f for f in moi.parent.glob(f"{moi.stem} (*){moi.suffix}") if f.stat().st_mtime > moi.stat().st_mtime]
        return max(recents, key=lambda f: f.stat().st_mtime) if recents else None
    except OSError:
        return None


def principal():
    preparer_console()
    rendre_net()
    ancien_mode = edition_rapide_coupee()
    try:
        ecrire("=" * 70)
        ecrire(" ROBOT (version %s)" % VERSION)
        ecrire("=" * 70)
        recent = robot_plus_recent()
        if recent:
            ecrire(f"   ATTENTION : un robot plus recent est la : « {recent.name} ».")
            ecrire(f"   Dans Telechargements : supprimez « {Path(__file__).stem} », renommez « {recent.stem} » en "
                   f"{Path(__file__).stem} (sans .txt), puis relancez.")
        ecrire("   1 = Test de prise en main")
        ecrire("   2 = Releve des pages du portail (pour Claude)")
        ecrire("   3 = Lancer une tache")
        ecrire("   4 = Voir ce que le robot trouve cliquable (et lui montrer ce qu'il rate)")
        ecrire("   5 = Le robot explore le portail tout seul et fait la carte du site (sans rien modifier)")
        ecrire("   6 = Aller a une page : vous dites ou, le robot y va tout seul (avec sa carte)")
        choix = demander(" Votre choix (1 a 6) :")
        if choix == "1":
            test_prise_en_main()
            return 0
        if choix in ("4", "5", "6"):
            url = adresse_du_portail()
            if url:
                lancer_programme_enfant({"4": "voir", "5": "explorer", "6": "gps"}[choix], {"url": url})
            else:
                ecrire("   Pas d'adresse du portail : rien a faire.")
        elif choix in ("2", "3"):
            fichier = None
            if choix == "3":
                fichier = choisir_tache()
                if fichier:
                    try:
                        lire_tache(fichier)
                    except ErreurTache as e:
                        ecrire(f"   La tache contient une erreur : {e}")
                        ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
                        fichier = None
            if choix == "2" or fichier:
                url = adresse_du_portail()
                if not url:
                    ecrire("   Pas d'adresse du portail : rien a faire.")
                else:
                    demande = {"url": url} if choix == "2" else {"fichier": str(fichier), "url": url}
                    lancer_programme_enfant("releve" if choix == "2" else "tache", demande)
        else:
            ecrire("   Rien a faire.")
    except KeyboardInterrupt:
        ecrire("\n   Arrete a votre demande.")
    except Exception as e:
        ecrire("   PROBLEME : " + (erreur_courte(e) if isinstance(e, OSError) else expliquer(e)))
        ecrire("   Prenez cet ecran en photo et envoyez-la-moi.")
    finally:
        remettre_console(ancien_mode)
    try:
        demander("\n Entree pour terminer.")
    except KeyboardInterrupt:
        pass
    return 0


def test_prise_en_main():
    ancien_mode = None
    ecrire("=" * 70)
    ecrire(" TEST DE PRISE EN MAIN")
    ecrire("=" * 70)
    ecrire(" Le robot va montrer, une chose a la fois, qu'il sait se servir de votre")
    ecrire(" ordinateur : fichiers, clavier, souris, navigateur. Il ne modifie RIEN dans vos")
    ecrire(" outils. Avant chaque etape, il explique et attend votre Entree.")
    ecrire(" Quand il pose une question : cliquez d'abord dans CETTE fenetre noire, puis repondez.")
    ecrire(" Pour arreter a tout moment : Ctrl + C dans cette fenetre.")
    bilan = Bilan()
    etat = {"dossier": None, "hwnd": None, "url": ""}

    def etape(lettre, intitule, faire):
        """Une étape qui plante inopinément est notée ECHEC ; les suivantes ont quand même lieu."""
        try:
            faire()
        except KeyboardInterrupt:
            raise
        except Exception as e:
            bilan.noter(lettre, intitule, "ECHEC", "erreur imprevue : " + (erreur_courte(e) if isinstance(e, OSError)
                                                                             else masquer(premiere_ligne(e))))

    try:
        etape("P", "Python fonctionne", lambda: etape_python(bilan))
        etape("F", "Fichiers et dossiers", lambda: etat.update(dossier=etape_fichiers(bilan)))
        etape("K", "Clavier (Bloc-notes)", lambda: etat.update(hwnd=etape_clavier(bilan, etat["dossier"])))
        etape("S", "Souris", lambda: etape_souris(bilan, etat["hwnd"]))
        etape("W", "Portail dans votre navigateur", lambda: etat.update(url=demander_adresse()))
        etape("W", "Portail dans votre navigateur", lambda: etape_navigateur_habituel(bilan, etat["url"]))
        etape("N", "Navigateur du robot (page d'essai)", lambda: etape_navigateur_robot(bilan, etat["url"]))
    except KeyboardInterrupt:
        ecrire("\n   Arrete a votre demande.")
    dossier = etat["dossier"]
    texte = bilan.texte()
    # la photo ne doit montrer que le résultat (ni l'adresse, ni les mots tapés plus haut) : des lignes
    # vides d'abord (marche partout), puis l'effacement de l'écran si la fenêtre le permet
    ecrire("\n" * 60)
    if WINDOWS:
        try:
            os.system("cls")
        except Exception:
            pass
    ecrire("=" * 70)
    ecrire(texte)
    ecrire("=" * 70)
    ecrire(" PRENEZ CE RESULTAT EN PHOTO et envoyez-le-moi.")
    if dossier is not None:
        try:
            (dossier / "resultat_prise_en_main.txt").write_text(texte + "\n", encoding="utf-8")
            ecrire(" Il est aussi ecrit dans : Bureau > Robot > resultat_prise_en_main.txt")
        except OSError:
            pass
    ecrire(" Vous pouvez fermer le Bloc-notes (essai_clavier) ; enregistrer ou non, peu importe.")
    if bilan.lignes.get("V", ("", ""))[1] == "OK":
        ecrire(" Envoyez aussi le releve : Bureau > Robot > A_ENVOYER_releve.txt (relisez-le d'abord).")
    remettre_console(ancien_mode)
    demander("\n Entree pour terminer.")
    return 0


if __name__ == "__main__":
    if sys.argv[1:2] == ["--navigateur"] and len(sys.argv) >= 3:
        sys.exit(sous_programme_navigateur(sys.argv[2]))
    if sys.argv[1:2] in (["--releve"], ["--tache"], ["--voir"], ["--explorer"], ["--gps"]) and len(sys.argv) >= 3:
        sys.exit(programme_enfant(sys.argv[1][2:], sys.argv[2]))
    sys.exit(principal())
