"""Tests de bout en bout : le scénario de démo contre le formulaire local, navigateur invisible."""

import datetime as dt
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from autoweb.cli import main
from autoweb.erreurs import ErreurAutoweb
from autoweb.excel import creer_classeur
from autoweb.runner import Options, analyser_lignes, lancer
from autoweb.scenario import charger


def _preparer_demo(tmp_path: Path, dossier_modeles: Path) -> Path:
    for nom in ("formulaire_demo.html", "demo.yaml"):
        shutil.copyfile(dossier_modeles / nom, tmp_path / nom)
    date = dt.datetime(2026, 9, 21)
    creer_classeur(
        tmp_path / "demo_suivi.xlsx",
        ["Numéro plan", "Titre", "Type", "Date", "Urgent", "Commentaire"],
        [
            ["PL-2026-001", "Rez-de-chaussée", "Architecture", date, "oui", "Première saisie"],
            ["PL-2026-002", "Charpente", "STR", date, "non", None],
            [None, "Plan sans numéro", "ELE", None, None, None],
            ["PL-2026-004", "Réseau", "Électricité", date, "x", "Vérifier"],
        ],
        feuille="Plans",
    )
    return tmp_path / "demo.yaml"


def _statuts(chemin_excel: Path):
    ws = load_workbook(chemin_excel)["Plans"]
    entetes = [c.value for c in ws[1]]
    col = {nom: i + 1 for i, nom in enumerate(entetes)}
    lignes = {}
    for r in range(2, ws.max_row + 1):
        lignes[r] = {nom: ws.cell(row=r, column=i).value for nom, i in col.items()}
    return lignes


def test_simulation_sans_navigateur(tmp_path, dossier_modeles, caplog):
    scenario = charger(_preparer_demo(tmp_path, dossier_modeles))
    with caplog.at_level("INFO", logger="autoweb"):
        bilan = lancer(scenario, Options(simuler=True))
    assert bilan.simule and bilan.total == 4 and bilan.message == ""
    texte = caplog.text
    assert "PL-2026-001" in texte and "21/09/2026" in texte and "Saisi automatiquement" in texte
    # rien n'a été écrit dans l'Excel
    assert all(l["Statut"] is None for l in _statuts(tmp_path / "demo_suivi.xlsx").values())


def test_analyser_lignes():
    assert analyser_lignes("5, 8,10-12") == {5, 8, 10, 11, 12}
    with pytest.raises(ErreurAutoweb):
        analyser_lignes("a")


def test_colonne_manquante_detectee_avant_de_lancer(tmp_path, dossier_modeles):
    chemin = _preparer_demo(tmp_path, dossier_modeles)
    texte = chemin.read_text(encoding="utf-8").replace("{{Titre}}", "{{Titre du plan}}")
    chemin.write_text(texte, encoding="utf-8")
    with pytest.raises(ErreurAutoweb) as exc:
        lancer(charger(chemin), Options(visible=False))
    assert "Titre du plan" in str(exc.value)


def test_bout_en_bout(tmp_path, dossier_modeles, navigateur_ok):
    chemin = _preparer_demo(tmp_path, dossier_modeles)
    scenario = charger(chemin)
    bilan = lancer(scenario, Options(visible=False, interactif=False))
    assert (bilan.ok, bilan.erreurs, bilan.ignorees, bilan.interrompu) == (3, 1, 0, False)

    lignes = _statuts(tmp_path / "demo_suivi.xlsx")
    assert lignes[2]["Statut"] == "OK" and lignes[2]["Référence outil"] == "REF-0001"
    assert lignes[3]["Statut"] == "OK" and lignes[3]["Référence outil"] == "REF-0002"
    assert lignes[4]["Statut"] == "ERREUR" and "verifier" in lignes[4]["Message"].lower()
    assert "capture" in lignes[4]["Message"]
    assert lignes[5]["Statut"] == "OK" and lignes[5]["Référence outil"] == "REF-0003"
    assert lignes[2]["Horodatage"] and lignes[2]["Message"] in (None, "")
    assert (tmp_path / "captures" / "PL-2026-001.png").exists()
    assert list((tmp_path / "captures" / "erreurs").glob("ligne-4-*.png"))
    assert list((tmp_path / "sauvegardes").glob("demo_suivi.*.xlsx"))

    # relance : plus rien à faire (les OK et l'ERREUR ne sont pas retraités)
    bilan2 = lancer(scenario, Options(visible=False, interactif=False))
    assert bilan2.total == 0

    # on corrige la ligne en erreur puis --reprendre-erreurs
    from openpyxl import load_workbook as lw
    wb = lw(tmp_path / "demo_suivi.xlsx")
    wb["Plans"]["A4"] = "PL-2026-003"
    wb.save(tmp_path / "demo_suivi.xlsx")
    bilan3 = lancer(scenario, Options(visible=False, interactif=False, reprendre_erreurs=True))
    assert (bilan3.total, bilan3.ok) == (1, 1)
    lignes = _statuts(tmp_path / "demo_suivi.xlsx")
    assert lignes[4]["Statut"] == "OK" and lignes[4]["Référence outil"].startswith("REF-")


def test_limite_lignes_et_arret_premiere_erreur(tmp_path, dossier_modeles, navigateur_ok):
    scenario = charger(_preparer_demo(tmp_path, dossier_modeles))
    bilan = lancer(scenario, Options(visible=False, interactif=False, lignes={4, 5}, arret_premiere_erreur=True))
    assert bilan.total == 2 and bilan.erreurs == 1 and bilan.ok == 0 and bilan.interrompu
    lignes = _statuts(tmp_path / "demo_suivi.xlsx")
    assert lignes[2]["Statut"] is None and lignes[4]["Statut"] == "ERREUR" and lignes[5]["Statut"] is None

    bilan = lancer(scenario, Options(visible=False, interactif=False, limite=1))
    assert bilan.total == 1 and bilan.ok == 1
    assert _statuts(tmp_path / "demo_suivi.xlsx")[2]["Statut"] == "OK"


def test_etapes_avancees(tmp_path, dossier_modeles, navigateur_ok):
    """si/valeur, ignorer, echouer, optionnel, executer_js, verifier absent, lire attribut."""
    _preparer_demo(tmp_path, dossier_modeles)
    (tmp_path / "avance.yaml").write_text("""
nom: avancé
navigateur: {canal: auto, visible: false}
excel: {fichier: demo_suivi.xlsx, feuille: Plans}
etapes:
  - aller: {fichier: formulaire_demo.html}
  - si:
      valeur: "{{Numéro plan}}"
      vide: true
      alors:
        - ignorer: "pas de numéro"
  - si:
      valeur: "{{Urgent}}"
      vrai: true
      alors:
        - journal: "urgent !"
      sinon:
        - si:
            valeur: "{{Titre}}"
            contient: "charpente"
            alors:
              - echouer: "charpente refusée"
  - cliquer: {selecteur: "#element-inexistant", delai: 300}
    optionnel: true
  - verifier: {selecteur: "#application", cache: true}
  - executer_js: {script: "document.title", vers: "Titre page"}
  - lire: {selecteur: "#utilisateur", vers: "Nom champ", attribut: name}
  - lire: {selecteur: "#etat-connexion", vers: "Etat", regex: "Non (\\\\w+)"}
""", encoding="utf-8")
    bilan = lancer(charger(tmp_path / "avance.yaml"), Options(interactif=False))
    assert (bilan.ok, bilan.erreurs, bilan.ignorees) == (2, 1, 1)
    lignes = _statuts(tmp_path / "demo_suivi.xlsx")
    assert lignes[2]["Statut"] == "OK" and lignes[2]["Titre page"].startswith("Outil démo")
    assert lignes[2]["Nom champ"] == "utilisateur" and lignes[2]["Etat"] == "connecté"
    assert lignes[3]["Statut"] == "ERREUR" and "charpente refusée" in lignes[3]["Message"]
    assert lignes[4]["Statut"] == "IGNORE" and lignes[4]["Message"] == "pas de numéro"
    assert lignes[5]["Statut"] == "OK"


def test_cli_demo_et_verifier(tmp_path, dossier_modeles, navigateur_ok, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    code = main(["demo", "--cache", "--sans-pause", "--dossier", "d"])
    sortie = capsys.readouterr().out
    assert code == 0, sortie
    assert "3 OK, 1 ERREUR" in sortie
    assert (tmp_path / "d" / "demo_suivi.xlsx").exists()

    code = main(["verifier", "d/demo.yaml"])
    sortie = capsys.readouterr().out
    assert code == 0 and "valide" in sortie and "Numéro plan" in sortie

    code = main(["initialiser", "nouveau", "--nom", "Mon outil"])
    sortie = capsys.readouterr().out
    assert code == 0 and (tmp_path / "nouveau" / "mon_outil.yaml").exists() and (tmp_path / "nouveau" / "suivi.xlsx").exists()
    code = main(["verifier", "nouveau/mon_outil.yaml"])
    sortie = capsys.readouterr().out
    assert code == 0, sortie
